"""Motion 05 — acceleration and new velocity readings (PW 05)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Motion_05_Acceleration.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


_GREEK_SUP = {"α": "ᵅ", "β": "ᵝ", "γ": "ᵞ", "δ": "ᵟ", "ε": "ᵋ", "θ": "ᶿ"}


_SUB_CHARS = {
    "0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄",
    "5": "₅", "6": "₆", "7": "₇", "8": "₈", "9": "₉",
    "a": "ₐ", "e": "ₑ", "h": "ₕ", "i": "ᵢ", "j": "ⱼ",
    "k": "ₖ", "l": "ₗ", "m": "ₘ", "n": "ₙ", "o": "ₒ",
    "p": "ₚ", "r": "ᵣ", "s": "ₛ", "t": "ₜ", "u": "ᵤ",
    "v": "ᵥ", "x": "ₓ", "y": "ᵧ",
    # no reliable Unicode subscript f — use PPT native subscript run instead
}


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(-1/2)", "⁻½").replace("^(-3/2)", "⁻³⁄²")
    text = text.replace("^(-5/2)", "⁻⁵⁄²").replace("^(-7/2)", "⁻⁷⁄²")
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³⁄²")
    text = text.replace("^(5/2)", "⁵⁄²").replace("^(7/2)", "⁷⁄²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([αβγδεθ])", lambda m: _GREEK_SUP[m.group(1)], text)
    return text


def _format_subs_unicode(text: str) -> str:
    """Convert _i, _{fi}, _x into Unicode subscripts; unknown letters → PPT subscript markers."""
    def brace(m):
        out = []
        for ch in m.group(1):
            if ch in _SUB_CHARS:
                out.append(_SUB_CHARS[ch])
            elif ch.isalnum():
                out.append(f"\0{ch}\0")
            else:
                out.append(ch)
        return "".join(out)

    text = re.sub(r"_\{([^}]+)\}", brace, text)
    text = re.sub(
        r"_([A-Za-z0-9])",
        lambda m: _SUB_CHARS[m.group(1)] if m.group(1) in _SUB_CHARS else f"\0{m.group(1)}\0",
        text,
    )
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    # Subscripts: prefer Unicode; letters without a Unicode form use a PPT subscript run.
    text = _format_subs_unicode(text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    pt_size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))

    def _add_run(chunk: str, as_sub: bool = False):
        if not chunk:
            return
        run = p.add_run()
        run.text = chunk
        run.font.name = font
        run.font.size = pt_size
        run.font.bold = bold
        run.font.color.rgb = color
        if as_sub:
            # python-pptx Font.subscript is unreliable — set OOXML baseline directly
            rPr = run._r.get_or_add_rPr()
            for old in rPr.findall(qn("a:baseline")):
                rPr.remove(old)
            bl = OxmlElement("a:baseline")
            bl.set("val", "-25000")
            rPr.append(bl)

    # \0X\0 markers = letters that need PPT native subscript (e.g. f)
    parts = re.split(r"\0([A-Za-z0-9])\0", text)
    if len(parts) == 1:
        _add_run(text)
    else:
        for i, part in enumerate(parts):
            _add_run(part, as_sub=(i % 2 == 1))
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    step = 0.52
    for i, item in enumerate(items):
        cy = y + i * step
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.06), Inches(0.32), Inches(0.32))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.08, 0.32, 0.28, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.46, cy + 0.04, w - 0.46, 0.42, 17, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# ===== Figure helpers (Motion 01) =====

def stick(slide, x, y, color=WHITE, h=0.55, label=None, lab_color=None):
    """Simple standing person: head + body."""
    r = 0.09
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - h), Inches(2 * r), Inches(2 * r))
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    draw_line(slide, x, y - h + 2 * r, x, y, color, 2.0)
    draw_line(slide, x, y - h * 0.45, x - 0.16, y - h * 0.15, color, 1.6)
    draw_line(slide, x, y - h * 0.45, x + 0.16, y - h * 0.15, color, 1.6)
    draw_line(slide, x, y, x - 0.12, y + 0.22, color, 1.6)
    draw_line(slide, x, y, x + 0.12, y + 0.22, color, 1.6)
    if label:
        tx(slide, label, x - 0.35, y + 0.26, 0.70, 0.28, 13, lab_color or color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)


def ground_line(slide, x1, x2, y, color=MUTED, label="ground"):
    draw_line(slide, x1, y, x2, y, color, 2.0)
    for i in range(8):
        gx = x1 + 0.25 + i * ((x2 - x1 - 0.4) / 7)
        draw_line(slide, gx, y, gx - 0.12, y + 0.14, color, 1.2)
    tx(slide, label, (x1 + x2) / 2 - 0.5, y + 0.16, 1.0, 0.26, 12, color, True,
       PP_ALIGN.CENTER, spa=0, scaled=False)


def rounded_box(slide, x, y, w, h, fill, border, pt=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(pt)
    return s


def number_line(slide, x, y, w, xmin, xmax, ticks=None):
    """Horizontal number line; returns function pos -> screen x."""
    draw_line(slide, x, y, x + w, y, MUTED, 2.0)
    arrow_line(slide, x + w - 0.15, y, x + w + 0.05, y, MUTED, 1.6)

    def sx(v):
        return x + (v - xmin) * w / (xmax - xmin)

    ticks = ticks if ticks is not None else list(range(xmin, xmax + 1))
    for v in ticks:
        px = sx(v)
        draw_line(slide, px, y - 0.08, px, y + 0.08, MUTED, 1.2)
        tx(slide, str(v), px - 0.18, y + 0.10, 0.36, 0.22, 11, MUTED, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    return sx


def mark_point(slide, px, y, label, color=GOLD, above=True):
    small_dot(slide, px, y, color)
    dy = -0.32 if above else 0.18
    tx(slide, label, px - 0.35, y + dy, 0.70, 0.28, 13, color, True,
       PP_ALIGN.CENTER, spa=0, scaled=False)


def path_polyline(slide, points, color=GOLD, pt=2.6):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        arrow_line(slide, x1, y1, x2, y2, color, pt)


# ===== COVER =====
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx_3d(slide, "EXCELLENCE", XL, 0.52, CW, 1.10, 62, GOLD, RGBColor(210, 210, 220))
tx_3d(slide, "CLASS 11   -   2026-27", XL, 1.64, CW, 0.58, 30, WHITE, RGBColor(90, 90, 100), font="Aptos")
_solid(slide, 4.15, 2.28, 5.00, 0.028, WHITE)
bx, by, bw = 2.05, 2.70, 9.20
_card_bg(slide, bx, by, bw, 2.35, WHITE, WHITE, 0)
_card_bg(slide, bx, by + 0.72, bw, 1.63, BLACK, RGBColor(220, 220, 220), 1.25)
_solid(slide, bx + bw / 2, by + 0.16, 0.018, 0.42, RGBColor(40, 40, 40))
tx(slide, "Physics", bx, by + 0.14, bw / 2 - 0.10, 0.50, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Motion 05", bx + bw / 2, by + 0.14, bw / 2, 0.50, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Motion in a Straight Line", bx, by + 1.02, bw, 0.55, 28, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False, font="Georgia")
tx(slide, "Acceleration · later intervals · when a stays constant", bx, by + 1.62, bw, 0.40, 14, GOLD, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "By - Excellence Programme", XL, 5.28, CW, 0.50, 25, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "22 October 2026", XL, 5.80, CW, 0.32, 18, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, 1)
add_notes(slide, "PW Motion 05. Skips speed, U-turn, half-circle, and endpoint-average recipes already in Motion 03 and 04.")
after_content_slide()

slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx(slide, "What is new today", XL, 0.35, CW, 0.64, 32, GOLD, True, spa=0, scaled=False)
_solid(slide, XL, 1.00, 3.40, 0.035, GOLD)
topics = [
    "Average velocity on an interval that starts after t = 0",
    "Which of two x–t curves is faster, early and late",
    "Acceleration: a = dv/dt and a = v dv/dx",
    "Speed can rise while a falls; v = 0 can sit with a ≠ 0",
    "Opposite velocity and acceleration",
    "A later velocity from two readings, when a is constant",
    "Which formulas hide a constant acceleration",
]
for i, t in enumerate(topics):
    y = 1.22 + i * 0.72
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(XL), Inches(y), Inches(0.42), Inches(0.42))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.fill.background()
    tx(slide, str(i + 1), XL, y + 0.04, 0.42, 0.34, 16, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
    tx(slide, t, XL + 0.58, y + 0.02, CW - 0.58, 0.42, 18, TEXT, False, spa=0, scaled=False)
_footer(slide, len(prs.slides))
after_content_slide()

# ===== INTERVAL NOT FROM ZERO =====
slide = new_slide("Example", "Average velocity between t = 2 s and t = 5 s", TEAL)
para(slide, "Position is x = 5 + 2t^2. The interval asked for starts at t = 2 s, so the average uses those two clock readings.", XL, Y0, CW, 0.85, 16)
fml(slide, "x(5) = 5 + 2(25) = 55 m", XL, Y0 + 1.00, CW, LH, 18)
fml(slide, "x(2) = 5 + 2(4) = 13 m", XL, Y0 + 1.00 + LH + NG, CW, LH, 18)
fml(slide, "v_{avg} = (55 − 13) / (5 − 2) = 42 / 3 = 14 m/s", XL, Y0 + 1.00 + 2 * (LH + NG), CW, LH, 18)
para(slide, "v = 4t stays positive from 2 s to 5 s, so the particle keeps the same direction. Distance equals the 42 m displacement, and the average speed is 14 m/s as well.", XL, Y0 + 2.85, CW, 0.95, 16)
result_card(slide, "14 m/s", XL, Y0 + 4.00, 3.2, 0.48, 20)

slide = new_slide("Check", "The integral agrees, with one careful step", ORANGE)
para(slide, "Average velocity is also the time-average of v. Here v = dx/dt = 4t.", XL, Y0, CW, 0.70, 16)
fml(slide, "v_{avg} = (1 / 3) ∫_2^5 4t dt = [2t^2]_2^5 / 3", XL, Y0 + 0.85, CW, LH, 18)
fml(slide, "= 2(5^2 − 2^2) / 3 = 2(25 − 4) / 3 = 14 m/s", XL, Y0 + 0.85 + LH + NG, CW, LH, 18)
trap_card(slide, "2(5 − 2) = 6 uses the times in place of their squares. The antiderivative of 4t is 2t^2, so the squares stay.", XL, Y0 + 2.40, CW, 0.85)
result_card(slide, "Endpoints and the integral both give 14 m/s", XL, Y0 + 3.50, CW, 0.48, 18)

# ===== TWO CURVES =====
slide = question_slide("Which particle is faster?")
qtx(slide, "A and B move on the x–t graph. Which statement is correct?", XL, Y0, 6.7, 0.85, 18)
choices(slide, [
    "B is faster at the start and slower at the end",
    "B is slower at the start and faster at the end",
    "B is faster at the start and at the end",
    "B is slower at the start and at the end",
], XL, Y0 + 1.05, 6.7)
sx, sy = graph_axes(slide, 7.7, 1.55, 4.9, 4.9, -0.3, 5.4, -0.3, 5.2, "t", "x")
plot_function(slide, lambda t: 0.72 * t, sx, sy, 0, 5, -0.3, 5.2, GOLD, 40, 2.4)
plot_function(slide, lambda t: 0.11 * t * t + 1.55, sx, sy, 0, 5, -0.3, 5.2, TEAL, 50, 2.4)
tx(slide, "A", sx(5) - 0.35, sy(0.72 * 5) - 0.32, 0.40, 0.24, 14, GOLD, True, spa=0, scaled=False)
tx(slide, "B", sx(5) + 0.08, sy(0.11 * 25 + 1.55) - 0.28, 0.40, 0.24, 14, TEAL, True, spa=0, scaled=False)

slide = answer_slide()
para(slide, "The slope of an x–t graph is the velocity. A is a straight line, so its speed stays fixed. B opens with a horizontal tangent, so B’s initial speed is zero, below A’s. At the right-hand end B’s tangent is steeper than A’s line, so B is faster there.", XL, Y0, 7.3, 2.15, 16)
result_card(slide, "Answer: B is slower at first, faster at the end", XL, Y0 + 2.35, 7.3, 0.70, 16)
sx, sy = graph_axes(slide, 8.3, 1.45, 4.3, 4.6, -0.3, 5.4, -0.3, 5.2, "t", "x")
plot_function(slide, lambda t: 0.72 * t, sx, sy, 0, 5, -0.3, 5.2, GOLD, 40, 2.2)
plot_function(slide, lambda t: 0.11 * t * t + 1.55, sx, sy, 0, 5, -0.3, 5.2, TEAL, 50, 2.2)
draw_line(slide, sx(0), sy(1.55), sx(1.4), sy(1.55), WHITE, 1.8)
# tangent at the right-hand end of B, extended past the curve
draw_line(slide, sx(4.5), sy(4.3 + 1.1 * (4.5 - 5)),
          sx(5.35), sy(4.3 + 1.1 * (5.35 - 5)), WHITE, 2.0)

# ===== ACCELERATION =====
slide = new_slide("Idea", "Acceleration is how velocity changes", ORANGE)
para(slide, "Velocity tells how fast the position is changing, and in which direction. Acceleration tells how fast that velocity itself is changing. Its direction is the direction of the change in velocity.", XL, Y0, CW, 1.15, 16)
ncert(slide, "a_{avg} = (v_f − v_i) / Δt", XL, Y0 + 1.30, CW, 0.72, 22)
fml(slide, "a = dv/dt = d^2x/dt^2", XL, Y0 + 2.20, CW, LH, 20)
fml(slide, "a = v dv/dx", XL, Y0 + 2.20 + LH + NG, CW, LH, 20)
result_card(slide, "Constant velocity means a = 0", XL, Y0 + 3.55, CW, 0.48, 18)
add_notes(slide, "Chain rule: a = dv/dt = (dv/dx)(dx/dt) = v dv/dx.")

slide = new_slide("Table", "Same clock, five different accelerations", BLUE)
para(slide, "Each row gives the velocity at t = 0, 1, 2 and 3 s. The acceleration column is the change from one second to the next.", XL, Y0, CW, 0.80, 16)
grid_table(
    slide,
    ["Story", "Velocities (m/s)", "Accelerations (m/s^2)"],
    [
        ["Steady", "10, 10, 10, 10", "0, 0, 0"],
        ["Even speed-up", "10, 20, 30, 40", "10, 10, 10"],
        ["Gentler speed-up", "10, 20, 28, 34", "10, 8, 6"],
        ["Harder speed-up", "10, 20, 35, 55", "10, 15, 20"],
        ["Slowing down", "20, 10, 6, 4", "−10, −4, −2"],
    ],
    XL, Y0 + 0.95, [3.3, 4.3, 4.45], 0.48, 14, 15,
)
para(slide, "On the gentler row the velocity keeps rising while the acceleration falls from 10 to 6. Rising velocity and falling acceleration can happen together.", XL, Y0 + 4.00, CW, 0.85, 16)

slide = new_slide("Example", "Slowing down, with a growing acceleration", TEAL)
para(slide, "Velocity readings: 20, 18, 14, 8 m/s at t = 0, 1, 2, 3 s. The object is still moving forward, and it is moving more slowly each second.", XL, Y0, CW, 1.00, 16)
fml(slide, "a = (18 − 20) / 1 = −2 m/s^2", XL, Y0 + 1.15, CW, LH, 18)
fml(slide, "a = (14 − 18) / 1 = −4 m/s^2", XL, Y0 + 1.15 + LH + NG, CW, LH, 18)
fml(slide, "a = (8 − 14) / 1 = −6 m/s^2", XL, Y0 + 1.15 + 2 * (LH + NG), CW, LH, 18)
result_card(slide, "Velocity falls. The size of a grows from 2 to 6.", XL, Y0 + 3.15, CW, 0.50, 18)

slide = new_slide("Figure", "v = 0 while acceleration is still there", GREEN)
para(slide, "At the top of a vertical throw the velocity is zero for one instant. The acceleration is still g, downward. That acceleration carries the velocity from upward, through zero, and on to downward, so the object turns around.", XL, Y0, 6.5, 2.30, 16)
result_card(slide, "v = 0 and a = g downward, together", XL, Y0 + 2.50, 6.4, 0.70, 16)
sx, sy = graph_axes(slide, 7.6, 1.50, 5.0, 4.8, -0.4, 4.4, -0.6, 4.6, "t", "x")
plot_function(slide, lambda t: 4 * t - t * t, sx, sy, 0, 4, -0.6, 4.6, GOLD, 80, 2.4)
arrow_line(slide, sx(2), sy(3.55), sx(2), sy(1.15), RED, 2.4)
tx(slide, "a", sx(2) + 0.14, sy(2.3), 0.40, 0.28, 16, RED, True, spa=0, scaled=False)
tx(slide, "v = 0", sx(2) - 0.85, sy(4) - 0.36, 0.80, 0.28, 14, WHITE, True, spa=0, scaled=False)

slide = question_slide("Which statements can be true?")
qtx(slide, "Which option is correct?", XL, Y0, CW, 0.50, 18)
choices(slide, [
    "Velocity may increase while acceleration decreases",
    "Velocity may decrease while the size of the acceleration increases",
    "Acceleration may be non-zero at the instant velocity is zero",
    "All of these",
], XL, Y0 + 0.70)

slide = answer_slide()
bullet(slide, "Gentler speed-up: 10, 20, 28, 34 m/s with a = 10, 8, 6 m/s^2.", XL, Y0, CW, 0.70, 16)
bullet(slide, "Harder slowing: 20, 18, 14, 8 m/s with a = −2, −4, −6 m/s^2.", XL, Y0 + 0.80, CW, 0.70, 16)
bullet(slide, "Top of a throw: v = 0 while a = g downward.", XL, Y0 + 1.60, CW, 0.60, 16)
result_card(slide, "Answer: all of these", XL, Y0 + 2.50, CW, 0.48, 20)

slide = new_slide("Rule", "Velocity and acceleration in opposite directions", BLUE)
para(slide, "Opposite directions mean the acceleration points against the velocity. That part of the acceleration reduces the speed. A perpendicular acceleration, as in uniform circular motion, leaves the speed unchanged. The opposite case is the one in which the speed falls.", XL, Y0, CW, 1.55, 16)
vec(slide, 4.4, 3.35, 1.8, 3.35, "", GOLD, 3.0)
tx(slide, "v", 4.55, 3.05, 0.40, 0.30, 16, GOLD, True, spa=0, scaled=False)
vec(slide, 1.8, 4.20, 4.4, 4.20, "", RED, 3.0)
tx(slide, "a", 1.25, 3.95, 0.40, 0.30, 16, RED, True, spa=0, scaled=False)
result_card(slide, "Speed must be decreasing", XL, Y0 + 3.70, CW, 0.48, 20)

# ===== CONSTANT a ARITHMETIC =====
slide = new_slide("Example", "Two readings with constant acceleration", GREEN)
para(slide, "Acceleration is 2 m/s^2 and the velocity at t = 0 is 10 m/s. Find the velocity at t = 4 s.", XL, Y0, CW, 0.75, 16)
fml(slide, "v = 10 + 2 × 4 = 18 m/s", XL, Y0 + 0.85, CW, LH, 20)
para(slide, "A second object goes from 10 m/s to 40 m/s in 6 s. The acceleration over that interval is the change in velocity divided by the time.", XL, Y0 + 1.55, CW, 0.90, 16)
fml(slide, "a_{avg} = (40 − 10) / 6 = 5 m/s^2", XL, Y0 + 2.55, CW, LH, 20)
result_card(slide, "18 m/s    and    5 m/s^2", XL, Y0 + 3.40, CW, 0.48, 18)

slide = new_slide("Example", "A later velocity, once a is constant", TEAL)
para(slide, "Velocity is 20 m/s at t = 2 s and 32 m/s at t = 5 s. These two readings determine the velocity at t = 7 s when the acceleration is constant.", XL, Y0, CW, 1.05, 16)
fml(slide, "a = (32 − 20) / (5 − 2) = 12 / 3 = 4 m/s^2", XL, Y0 + 1.20, CW, LH, 18)
fml(slide, "From t = 5 s to t = 7 s is another 2 s", XL, Y0 + 1.20 + LH + NG, CW, LH, 18)
fml(slide, "v(7) = 32 + 4 × 2 = 40 m/s", XL, Y0 + 1.20 + 2 * (LH + NG), CW, LH, 18)
result_card(slide, "40 m/s", XL, Y0 + 3.30, CW, 0.48, 20)

# ===== DIFFERENTIATE =====
slide = new_slide("Method", "Acceleration from x(t)", ORANGE)
para(slide, "Differentiate once for v, and once more for a.", XL, Y0, CW, 0.50, 16)
fml(slide, "x = 4t^2 + 6   →   v = 8t   →   a = 8 m/s^2", XL, Y0 + 0.65, CW, LH, 16)
fml(slide, "x = 3t^2 + 4t + 6   →   v = 6t + 4   →   a = 6 m/s^2", XL, Y0 + 0.65 + LH + NG, CW, LH, 16)
fml(slide, "x = 2t^3 + 5t   →   v = 6t^2 + 5   →   a = 12t m/s^2", XL, Y0 + 0.65 + 2 * (LH + NG), CW, LH, 16)
fml(slide, "x = t^4 + 4t   →   v = 4t^3 + 4   →   a = 12t^2 m/s^2", XL, Y0 + 0.65 + 3 * (LH + NG), CW, LH, 16)

slide = new_slide("Method", "Acceleration from v, including v(x)", BLUE)
fml(slide, "v = 3t^2 + 4   →   a = 6t m/s^2", XL, Y0, CW, LH, 18)
fml(slide, "v = t^3 + 4   →   a = 3t^2 m/s^2", XL, Y0 + LH + NG, CW, LH, 18)
para(slide, "When velocity is given against position, differentiate with respect to x and multiply by v.", XL, Y0 + 1.55, CW, 0.80, 16)
fml(slide, "v = 2x + 4    dv/dx = 2", XL, Y0 + 2.45, CW, LH, 18)
fml(slide, "a = v dv/dx = (2x + 4) × 2 = 4x + 8 m/s^2", XL, Y0 + 2.45 + LH + NG, CW, LH, 18)

slide = new_slide("Example", "One instant: position, velocity, acceleration", GREEN)
para(slide, "x = 2t^3 − 4t^2 + 4. Read all three at t = 2 s.", XL, Y0, CW, 0.60, 16)
fml(slide, "x(2) = 2(8) − 4(4) + 4 = 4 m", XL, Y0 + 0.75, CW, LH, 18)
fml(slide, "v = 6t^2 − 8t    →    v(2) = 24 − 16 = 8 m/s", XL, Y0 + 0.75 + LH + NG, CW, LH, 18)
fml(slide, "a = 12t − 8    →    a(2) = 24 − 8 = 16 m/s^2", XL, Y0 + 0.75 + 2 * (LH + NG), CW, LH, 18)
result_card(slide, "4 m ,  8 m/s ,  16 m/s^2", XL, Y0 + 3.00, CW, 0.50, 20)

slide = new_slide("Example", "The instant when acceleration is zero", TEAL)
para(slide, "For x = t^3 − 6t^2 + 10, set a = 0 and then read the velocity there.", XL, Y0, CW, 0.70, 16)
fml(slide, "v = 3t^2 − 12t", XL, Y0 + 0.85, CW, LH, 18)
fml(slide, "a = 6t − 12 = 0    →    t = 2 s", XL, Y0 + 0.85 + LH + NG, CW, LH, 18)
fml(slide, "v(2) = 3(4) − 12(2) = −12 m/s", XL, Y0 + 0.85 + 2 * (LH + NG), CW, LH, 18)
para(slide, "A second position, x = t^3 − 9t^2 + 4, has a = 6t − 18. That is zero at t = 3 s, and v(3) = 3(9) − 18(3) = −27 m/s.", XL, Y0 + 2.85, CW, 1.05, 16)
result_card(slide, "t = 2 s, v = −12 m/s     and     t = 3 s, v = −27 m/s", XL, Y0 + 4.05, CW, 0.50, 16)

slide = new_slide("Example", "Acceleration at the instant velocity is zero", ORANGE)
para(slide, "x = 2t^3 − 24t + 4. Find a at every time where v = 0.", XL, Y0, CW, 0.65, 16)
fml(slide, "v = 6t^2 − 24 = 6(t^2 − 4) = 0", XL, Y0 + 0.80, CW, LH, 18)
fml(slide, "t = 2 s    and    t = −2 s", XL, Y0 + 0.80 + LH + NG, CW, LH, 18)
fml(slide, "a = 12t", XL, Y0 + 0.80 + 2 * (LH + NG), CW, LH, 18)
fml(slide, "a(2) = 24 m/s^2      a(−2) = −24 m/s^2", XL, Y0 + 0.80 + 3 * (LH + NG), CW, LH, 18)
result_card(slide, "24 m/s^2 at t = 2 s,   and   −24 m/s^2 at t = −2 s", XL, Y0 + 3.55, CW, 0.55, 16)

# ===== WHICH a IS CONSTANT =====
slide = new_slide("Test", "Six formulas whose acceleration is a pure number", BLUE)
para(slide, "In this list y is a constant. Differentiate until a appears. These six come out as numbers.", XL, Y0, CW, 0.80, 16)
bullet(slide, "√x = t − 3    →    x = (t − 3)^2    →    a = 2 m/s^2", XL, Y0 + 0.95, CW, 0.48, 16)
bullet(slide, "x = t^2 + y    →    a = 2 m/s^2", XL, Y0 + 1.50, CW, 0.45, 16)
bullet(slide, "x = 4t^2 + 6    →    a = 8 m/s^2", XL, Y0 + 2.00, CW, 0.45, 16)
bullet(slide, "v = 6t − y    →    a = 6 m/s^2", XL, Y0 + 2.50, CW, 0.45, 16)
bullet(slide, "v = √x    →    a = 1/2 m/s^2", XL, Y0 + 3.00, CW, 0.45, 16)
bullet(slide, "a = 5 m/s^2 is already constant", XL, Y0 + 3.50, CW, 0.45, 16)

slide = new_slide("Why", "v = √x really is constant acceleration", GREEN)
fml(slide, "v = √x", XL, Y0, CW, LH, 20)
fml(slide, "dv/dx = 1 / (2√x)", XL, Y0 + LH + NG, CW, LH, 20)
fml(slide, "a = v dv/dx = √x × 1/(2√x) = 1/2 m/s^2", XL, Y0 + 2 * (LH + NG), CW, LH, 18)
para(slide, "Three neighbours still depend on t or on x, so their acceleration keeps changing: v = 3t^2 + 4 gives a = 6t; v = 2x + 3 gives a = 4x + 6; x = A sin(ωt) gives a = −Aω^2 sin(ωt).", XL, Y0 + 2.30, CW, 1.35, 16)

# ===== RETURN GRAPH =====
slide = question_slide("Which graph can have average velocity zero?")
qtx(slide, "Average velocity from t = 0 to some later time can be zero for only one of these graphs. Which one?", XL, Y0, CW, 0.70, 17)


def _mini_xt(slide, x, y, w, h, pts, color, tag):
    ox, oy = x + 0.38, y + h - 0.38
    draw_line(slide, ox, oy, x + w - 0.12, oy, MUTED, 1.3)
    draw_line(slide, ox, oy, ox, y + 0.12, MUTED, 1.3)
    tx(slide, "t", x + w - 0.28, oy + 0.02, 0.24, 0.20, 11, MUTED, True, spa=0, scaled=False)
    tx(slide, "x", ox + 0.06, y + 0.02, 0.24, 0.20, 11, MUTED, True, spa=0, scaled=False)
    screen = []
    for px, py in pts:
        screen.append((ox + px * (w - 0.55), oy - py * (h - 0.55)))
    curve(slide, screen, color, 2.3)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.02), Inches(y + 0.02), Inches(0.32), Inches(0.32))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.fill.background()
    tx(slide, tag, x + 0.02, y + 0.04, 0.32, 0.28, 13, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)


_mini_xt(slide, 0.55, 2.35, 3.00, 2.55, [(0.0, 0.45), (0.35, 0.85), (0.70, 1.00), (1.0, 1.02)], GOLD, "1")
_mini_xt(slide, 3.65, 2.35, 3.00, 2.55, [(0.0, 0.55), (0.30, 0.95), (0.55, 0.70), (0.75, 0.55), (1.0, 0.18)], TEAL, "2")
_mini_xt(slide, 6.75, 2.35, 3.00, 2.55, [(0.0, 0.90), (0.40, 0.45), (1.0, 0.15)], ORANGE, "3")
_mini_xt(slide, 9.85, 2.35, 3.00, 2.55, [(0.0, 0.0), (0.40, 0.45), (1.0, 0.95)], BLUE, "4")

slide = answer_slide()
para(slide, "Average velocity from t = 0 to a later time is zero when the position has returned to its starting value. Graph 2 climbs and then crosses that same height again, so there is a later time with x equal to the start. Graphs 1, 3 and 4 leave their starting height and remain on one side of it.", XL, Y0, CW, 1.70, 16)
result_card(slide, "Answer: graph 2", XL, Y0 + 1.90, CW, 0.48, 20)

slide = new_slide("Lock lines", "Only today’s lines", GREEN)
bullet(slide, "Between two times, average velocity is the change in position divided by the change in time.", XL, Y0, CW, 0.62, 16)
bullet(slide, "a = dv/dt = d^2x/dt^2, and a = v dv/dx.", XL, Y0 + 0.70, CW, 0.50, 16)
bullet(slide, "Constant velocity means a = 0. At the top of a throw, v = 0 while a is still g downward.", XL, Y0 + 1.30, CW, 0.62, 16)
bullet(slide, "When a points with v the speed rises. When a points against v the speed falls.", XL, Y0 + 2.00, CW, 0.55, 16)
bullet(slide, "Two velocity readings fix a later velocity once a is constant: v = u + at.", XL, Y0 + 2.65, CW, 0.55, 16)
bullet(slide, "a is constant when it comes out as a pure number. For v = √x, that number is 1/2.", XL, Y0 + 3.30, CW, 0.55, 16)

slide = prs.slides.add_slide(BLANK)
_bg(slide)
_bar(slide, "End", GOLD)
tx_3d(slide, "Thank you", XL, 2.4, CW, 1.0, 48, GOLD, RGBColor(120, 120, 130))
tx(slide, "Excellence  ·  Motion 05  ·  Acceleration", XL, 3.7, CW, 0.45, 15, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, len(prs.slides))
after_content_slide()

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(OUTPUT)
print("slides", len(prs.slides))
