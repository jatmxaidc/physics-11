"""Vector 02 — components, magnitude, and vector addition."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Vector_02_Components_and_Resultant.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(1/2)", "½"),
        ("^(-12)", "⁻¹²"),
        ("^(-7)", "⁻⁷"),
        ("^(-4)", "⁻⁴"),
        ("^(-2)", "⁻²"),
        ("^(-n)", "⁻ⁿ"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("cos inverse", "cos⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "F_z"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)




# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vector 02",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vectors",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Components and Resultant",bx,by+1.62,bw,0.40,20,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"29 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Vector class notes 02 and 03 in one lecture."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Types of vectors, part 2",
"Multiplying a vector by a scalar",
"How to read east of south",
"Rectangular components and magnitude",
"Triangle law and the resultant formula",
"Subtraction of vectors",
"Addition by components",
"Dot product, angle, and unit vector",
]
for i,item in enumerate(topics):
    yy=1.10+i*0.70
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.46,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.50,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"One class for Vector 02 and Vector 03 notes."); after_content_slide()

slide=new_slide("Copy","Types of vector — equality and lines",TEAL)
body(slide,"Equal: same magnitude, same direction, same kind of quantity.",XL,Y0,CW,0.50)
body(slide,"Parallel: same direction. Magnitudes may differ. Every equal pair is parallel. A parallel pair need not be equal.",XL,Y0+0.70,CW,0.70)
body(slide,"Negative: same magnitude, exactly opposite direction.",XL,Y0+1.55,CW,0.50)
body(slide,"Anti-parallel: opposite direction, magnitudes may differ. Every negative pair is anti-parallel. An anti-parallel pair need not be a negative pair.",XL,Y0+2.20,CW,0.75)

slide=new_slide("Copy","Types of vector — size and plane",TEAL)
body(slide,"Unit vector: magnitude 1, no unit, direction only.",XL,Y0,CW,0.50)
body(slide,"Zero vector: magnitude 0. Example: velocity of a car at rest.",XL,Y0+0.65,CW,0.50)
body(slide,"Collinear: all lie on the same straight line.",XL,Y0+1.30,CW,0.45)
body(slide,"Coplanar: all lie in the same plane.",XL,Y0+1.90,CW,0.45)
body(slide,"Orthogonal: angle between them is 90 degrees. Magnitudes may differ.",XL,Y0+2.50,CW,0.50)
body(slide,"Concurrent: arrows act at the same point, even if magnitudes and directions differ.",XL,Y0+3.15,CW,0.55)
body(slide,"Axial: direction is along the axis of rotation. Examples: torque, angular velocity, angular acceleration, angular momentum.",XL,Y0+3.80,CW,0.70)

# TYPES
slide=question_slide("Direction does not make every quantity a vector")
qtx(slide,"A: a quantity with no direction must be a scalar.   B: a quantity with a direction must be a vector.",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"A is true. A scalar has no direction.",XL,Y0,CW,LH,18)
fml(slide,"B is false. Current has a direction but does not obey the triangle law.",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"A true, B false.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=new_slide("Copy","Types of vectors — lock lines",TEAL)
grid_table(slide,["Type","Magnitude","Angle between them"],[
    ["Equal, same kind","same","0 degrees"],
    ["Parallel","may differ","0 degrees"],
    ["Opposite (negative)","same","180 degrees"],
    ["Anti-parallel","may differ","180 degrees"],
    ["Orthogonal","may differ","90 degrees"],
],XL,Y0,[4.0,3.6,4.4], row_h=0.62, hsize=14, csize=14)

slide=new_slide("Copy","Unit, null, collinear, coplanar",TEAL)
body(slide,"Unit vector: magnitude 1, no unit. A-hat = A / |A|.",XL,Y0,CW,0.48)
body(slide,"Null vector: magnitude 0. Direction is not defined. Example: two equal opposite forces, net force 0.",XL,Y0+0.62,CW,0.55)
body(slide,"Collinear: along one straight line, so the angle is 0 degrees or 180 degrees.",XL,Y0+1.30,CW,0.50)
body(slide,"Coplanar: every arrow lies in the same plane.",XL,Y0+1.95,CW,0.45)
body(slide,"Concurrent: the arrows pass through one common point.",XL,Y0+2.55,CW,0.45)
body(slide,"Axial: along the axis of rotation. Examples: angular velocity, angular momentum, torque. Moment of inertia is not a vector.",XL,Y0+3.15,CW,0.65)

slide=new_slide("Visual idea","Equal, parallel, opposite, anti-parallel",TEAL)
body(slide,"Do not remember these by words only. Compare three things: kind, length and direction.",XL,Y0,CW,0.45)
vec(slide,1.20,2.20,3.10,2.20,"20 N",GOLD); vec(slide,1.20,2.75,3.10,2.75,"20 N",GOLD)
tx(slide,"Equal vectors",1.20,3.12,2.4,0.28,15,GREEN,True,spa=0,scaled=False)
tx(slide,"same kind + same length + same direction",1.20,3.45,3.4,0.40,12,TEXT,False,spa=0,scaled=False)
vec(slide,5.00,2.20,7.00,2.20,"A",GOLD); vec(slide,5.00,2.75,6.25,2.75,"B",GREEN)
tx(slide,"Parallel",5.00,3.12,1.8,0.28,15,GREEN,True,spa=0,scaled=False)
tx(slide,"same direction, length may differ",5.00,3.45,3.2,0.40,12,TEXT,False,spa=0,scaled=False)
vec(slide,8.65,2.20,10.55,2.20,"A",GOLD); vec(slide,10.55,2.75,8.65,2.75,"-A",ORANGE)
tx(slide,"Opposite / negative",8.65,3.12,2.6,0.28,15,GREEN,True,spa=0,scaled=False)
tx(slide,"same length, exactly 180° direction",8.65,3.45,3.4,0.40,12,TEXT,False,spa=0,scaled=False)
vec(slide,1.20,4.70,3.50,4.70,"A",GOLD); vec(slide,3.10,5.25,1.20,5.25,"B",GREEN)
tx(slide,"Anti-parallel",1.20,5.62,2.2,0.28,15,GREEN,True,spa=0,scaled=False)
tx(slide,"opposite direction, length may differ",1.20,5.95,3.4,0.40,12,TEXT,False,spa=0,scaled=False)
trap_card(slide,"Exam trap: equal is stronger than parallel. Opposite is stronger than anti-parallel.",5.00,5.10,6.80,0.60)

slide=new_slide("Visual idea","Line, plane, point and axis",TEAL)
body(slide,"These names tell you where the arrows live geometrically.",XL,Y0,CW,0.45)
draw_line(slide,1.15,2.15,4.35,2.15,MUTED,1.2)
vec(slide,1.40,2.15,2.60,2.15,"A",GOLD); vec(slide,3.95,2.15,2.95,2.15,"B",GREEN)
tx(slide,"Collinear",1.25,2.55,2.0,0.30,15,GREEN,True,spa=0,scaled=False)
_card_bg(slide,5.00,1.70,3.15,1.55,RGBColor(10,18,25),STROKE,1.0,False)
vec(slide,5.35,2.70,7.45,2.05,"A",GOLD); vec(slide,5.55,2.15,6.95,2.85,"B",WORK)
tx(slide,"Coplanar",5.15,3.40,2.0,0.30,15,GREEN,True,spa=0,scaled=False)
small_dot(slide,10.00,2.55,GOLD)
vec(slide,10.00,2.55,8.95,1.85,"A",GOLD); vec(slide,10.00,2.55,11.20,1.95,"B",WORK); vec(slide,10.00,2.55,10.25,3.70,"C",GREEN)
tx(slide,"Concurrent",9.35,3.95,2.1,0.30,15,GREEN,True,spa=0,scaled=False)
arrow_line(slide,2.35,5.75,2.35,4.35,MUTED,1.2); vec(slide,2.35,5.05,3.45,4.35,"τ",ORANGE); tx(slide,"Axial vector: along rotation axis",1.25,6.05,3.8,0.30,15,GREEN,True,spa=0,scaled=False)
vec(slide,7.00,5.70,8.55,5.70,"A",GOLD); vec(slide,7.00,5.70,7.00,4.25,"B",GREEN); angle_mark(slide,7.00,5.70,0.52,0,90,"90°")
tx(slide,"Orthogonal",6.50,6.05,2.0,0.30,15,GREEN,True,spa=0,scaled=False)

slide=question_slide("Which statements are true?")
qtx(slide,"1. Every parallel pair is equal.   2. Every equal pair is parallel.   3. Every anti-parallel pair is opposite.   4. Every opposite pair is anti-parallel.",XL,Y0,CW,1.20,18)
slide=answer_slide()
fml(slide,"1 is false. Parallel vectors may have different lengths.",XL,Y0,CW,LH,17)
fml(slide,"2 is true. Equal vectors have the same direction.",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"3 is false. Anti-parallel lengths may differ.",XL,Y0+2*LH+2*NG,CW,LH,17)
fml(slide,"4 is true. Opposite means same length and 180 degrees, so they are anti-parallel.",XL,Y0+3*LH+3*NG,CW,LH,16)

slide=question_slide("Same number and same direction")
qtx(slide,"F = 20 N east and v = 20 m/s east. Are F and v equal vectors?",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"Equal vectors must be the same kind of quantity",XL,Y0,CW,LH,18)
fml(slide,"Force and velocity are different quantities, so they cannot be equal",XL,Y0+LH+NG,CW,LH,17)
result_card(slide,"Not equal. Two 20 N forces both east are equal.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=new_slide("How it works","Thrown up and thrown down",TEAL)
body(slide,"Weight mg is always downward.",XL,Y0,CW,0.45)
body(slide,"Object projected upward: velocity up, weight down. They are anti-parallel.",XL,Y0+0.65,CW,0.55)
body(slide,"Object thrown downward: velocity and weight are both down. They are parallel.",XL,Y0+1.35,CW,0.55)
trap_card(slide,"Parallel means the same direction. It does not demand equal magnitude.",XL,Y0+2.15,CW,0.55)
vec(slide,8.4,4.35,8.4,3.15,"v",GREEN)
vec(slide,9.05,3.15,9.05,4.35,"mg",ORANGE)
tx(slide,"upward throw",7.65,4.60,2.2,0.28,13,MUTED,True,spa=0,scaled=False)
vec(slide,10.65,3.15,10.65,4.35,"v",GREEN)
vec(slide,11.30,3.15,11.30,4.35,"mg",ORANGE)
tx(slide,"downward throw",10.05,4.60,2.4,0.28,13,MUTED,True,spa=0,scaled=False)

# SCALAR
slide=new_slide("New topic","Multiplying by a scalar",GREEN)
ncert(slide,"Positive scalar: magnitude changes, direction stays.",XL,Y0,CW,0.90,20)
ncert(slide,"Negative scalar: magnitude changes, direction reverses.",XL,Y0+1.10,CW,0.90,20)

slide=new_slide("Visual idea","What a scalar does to an arrow",GREEN)
body(slide,"A scalar does not create a new direction unless the scalar is negative.",XL,Y0,CW,0.45)
vec(slide,1.35,2.30,3.05,2.30,"A",GOLD)
vec(slide,4.35,2.30,7.75,2.30,"2A",GREEN)
tx(slide,"positive scalar: same direction",4.35,2.72,3.2,0.30,14,TEXT,False,spa=0,scaled=False)
vec(slide,1.35,4.20,3.05,4.20,"A",GOLD)
vec(slide,6.20,4.20,3.00,4.20,"-2A",ORANGE)
tx(slide,"negative scalar: reversed direction",3.25,4.62,3.6,0.30,14,TEXT,False,spa=0,scaled=False)
mini_card(slide,"Board line","kA keeps A's line. k > 0 keeps direction. k < 0 reverses direction.",8.25,2.25,3.85,1.55,GREEN)

slide=question_slide("3F and minus 2F")
qtx(slide,"F = 10 N i-hat. Find 3F, F/2 and -2F.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"3F = 30 N i-hat. Same direction, longer.",XL,Y0,CW,LH,18)
fml(slide,"F/2 = 5 N i-hat. Same direction, shorter.",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"-2F = -20 N i-hat. Opposite direction.",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"The minus sign flips the arrow. It does not change the kind of quantity.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("A = 30 i-hat + 20 j-hat")
qtx(slide,"Find 2A, A/2 and -5A.",XL,Y0,CW,0.70,22)
slide=answer_slide()
fml(slide,"2A = 60 i-hat + 40 j-hat",XL,Y0,CW,LH,18)
fml(slide,"A/2 = 15 i-hat + 10 j-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"-5A = -150 i-hat - 100 j-hat",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"Multiply every component. A minus outside flips both.",XL,Y0+3*LH+3*NG,CW,0.52)

slide=question_slide("A = 30 m i-hat")
qtx(slide,"Find 3A, A/3, -5A, the magnitude of A, and the unit vector.",XL,Y0,CW,0.90,19)
slide=answer_slide()
fml(slide,"3A = 90 m i-hat and A/3 = 10 m i-hat",XL,Y0,CW,LH,18)
fml(slide,"-5A = -150 m i-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"|A| = 30 m. The unit vector is i-hat, with no unit.",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"Unit vector = vector / magnitude.",XL,Y0+3*LH+3*NG,CW,0.50)

# DIRECTION LANGUAGE
slide=new_slide("New topic","How to read a compass phrase",GREEN)
ncert(slide,"The last word is the reference line. The first word is the side you turn toward.",XL,Y0,CW,1.10,18)
body(slide,"30 degrees east of south: start facing south, turn 30 degrees toward east.",XL,Y0+1.30,CW,0.55)
body(slide,"60 degrees south of east: start facing east, turn 60 degrees toward south.",XL,Y0+1.95,CW,0.55)
body(slide,"Those two phrases describe the same arrow.",XL,Y0+2.60,CW,0.45)
trap_card(slide,"East of south is not the same wording as south of east, even when the angles add to 90 degrees.",XL,Y0+3.20,CW,0.60)
ox,oy=axis_diagram(slide,8.4,2.25,3.2,2.8,coord=False)
tx(slide,"N",ox+0.08,2.18,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
tx(slide,"S",ox+0.08,4.95,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
tx(slide,"E",11.55,oy+0.05,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
tx(slide,"W",8.15,oy+0.05,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
vec(slide,ox,oy,ox+1.15,oy+1.95,"30° E of S",GOLD)
tx(slide,"also 60° S of E",9.85,5.20,2.4,0.30,13,WORK,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Compass phrases: one ray can have two names",GREEN)
body(slide,"Start from the last word. Turn toward the first word. The same slanted ray can be described from either nearby axis.",XL,Y0,CW,0.55)
ox,oy=axis_diagram(slide,1.15,2.10,4.4,3.7,coord=False)
tx(slide,"N",ox+0.05,2.02,0.35,0.25,12,MUTED,True,spa=0,scaled=False); tx(slide,"S",ox+0.05,5.78,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
tx(slide,"E",5.48,oy+0.05,0.35,0.25,12,MUTED,True,spa=0,scaled=False); tx(slide,"W",1.02,oy+0.05,0.35,0.25,12,MUTED,True,spa=0,scaled=False)
vec(slide,ox,oy,ox+1.75,oy+2.95,"F",GOLD)
angle_mark(slide,ox,oy,0.72,-90,-60,"30°")
angle_mark(slide,ox,oy,1.10,0,-60,"60°",WORK)
mini_card(slide,"30° east of south","Reference = south. Turn 30° toward east.",6.55,2.10,4.80,1.25,GOLD)
mini_card(slide,"60° south of east","Reference = east. Turn 60° toward south.",6.55,3.65,4.80,1.25,WORK)
result_card(slide,"Both names point to the same ray.",6.55,5.35,4.80,0.50)

slide=question_slide("Angle between the equal vectors changes when one is reversed")
qtx(slide,"In an equilateral layout, angle between A and B is 120 degrees. Find the angle between A and -B, and between A and -C if angle A with C is 60 degrees.",XL,Y0,CW,1.15,17)
slide=answer_slide()
fml(slide,"Reversing one arrow changes 120 degrees into 180 - 120 = 60",XL,Y0,CW,LH,16)
fml(slide,"If angle between A and C is 60, angle between A and -C is 120",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"Flip one arrow, and the angle becomes its supplement.",XL,Y0+2*LH+2*NG,CW,0.55)
vec(slide,8.2,4.45,10.2,4.45,"A",GOLD)
vec(slide,8.2,4.45,7.2,2.75,"B",GREEN)
vec(slide,8.2,4.45,9.2,2.75,"C",WORK)
vec(slide,10.75,4.45,11.75,2.75,"-C",ORANGE)
tx(slide,"same tail first",8.15,4.80,2.2,0.25,13,MUTED,True,spa=0,scaled=False)

slide=question_slide("Angle 30 degrees, then multiply")
qtx(slide,"Angle between A and B is 30 degrees. Find the angles between A and -B, between -A and -B, and between 2A and -3B.",XL,Y0,CW,1.10,17)
slide=answer_slide()
fml(slide,"A positive multiplier does not change direction, so 2A stays with A",XL,Y0,CW,LH,16)
fml(slide,"A minus flips the arrow, so angle with -B is 150 degrees",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"Two minuses bring both arrows back, so -A and -B stay at 30 degrees",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"2A and -3B are 150 degrees apart. 2A and B stay at 30 degrees.",XL,Y0+3*LH+3*NG,CW,0.55)

# COMPONENTS
slide=new_slide("New topic","Rectangular components",GREEN)
ncert(slide,"If theta is measured from the +x axis: Fx = F cos theta, Fy = F sin theta",XL,Y0,CW,1.05,18)
ncert(slide,"F = sqrt(Fx^2 + Fy^2)     and     tan theta = Fy / Fx",XL,Y0+1.25,CW,0.95,20)
body(slide,"The sign of a component is the sign of that axis. Left of the y-axis means a negative x-component. Below the x-axis means a negative y-component.",XL,Y0+2.45,7.5,0.75)
trap_card(slide,"Cos goes with the axis you measured the angle from. Do not swap sin and cos.",XL,Y0+3.40,7.5,0.55)
ox,oy=axis_diagram(slide,8.75,3.55,3.3,2.6)
vec(slide,ox,oy,ox+2.0,oy-1.25,"F",GOLD)
draw_line(slide,ox+2.0,oy-1.25,ox+2.0,oy,WORK,1.5)
draw_line(slide,ox,oy,ox+2.0,oy,WORK,1.5)
tx(slide,"Fₓ",ox+0.85,oy+0.10,0.45,0.25,13,WORK,True,spa=0,scaled=False)
tx(slide,"Fᵧ",ox+2.07,oy-0.75,0.45,0.25,13,WORK,True,spa=0,scaled=False)
tx(slide,"θ",ox+0.48,oy-0.27,0.35,0.25,15,ORANGE,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Resolving a vector means making a right triangle",GREEN)
body(slide,"A slanted vector can be replaced by two perpendicular vectors: one along x and one along y.",XL,Y0,CW,0.55)
ox,oy=axis_diagram(slide,1.15,2.00,5.0,3.8)
vec(slide,ox,oy,ox+2.85,oy-1.85,"F",GOLD,3.0)
arrow_line(slide,ox,oy,ox+2.85,oy,WORK,2.2); tx(slide,"Fₓ",ox+1.25,oy+0.15,0.65,0.28,15,WORK,True,spa=0,scaled=False)
arrow_line(slide,ox+2.85,oy,ox+2.85,oy-1.85,WORK,2.2); tx(slide,"Fᵧ",ox+3.02,oy-1.05,0.65,0.28,15,WORK,True,spa=0,scaled=False)
draw_line(slide,ox+2.62,oy,ox+2.62,oy-0.23,MUTED,1.0); draw_line(slide,ox+2.62,oy-0.23,ox+2.85,oy-0.23,MUTED,1.0)
angle_mark(slide,ox,oy,0.70,0,33,"θ")
mini_card(slide,"If θ is from x-axis","Adjacent side = Fₓ = F cos θ. Opposite side = Fᵧ = F sin θ.",7.15,2.20,4.75,1.45,GREEN)
mini_card(slide,"Sign comes later","First find size. Then attach + or - by quadrant.",7.15,4.05,4.75,1.20,ORANGE)

slide=new_slide("Visual idea","Quadrant decides the signs",GREEN)
body(slide,"Magnitude is always positive. Components carry signs because axes have signs.",XL,Y0,CW,0.45)
ox,oy=axis_diagram(slide,4.65,1.60,4.0,4.3)
labels=[("+, +",ox+1.08,oy-1.05,GREEN),("-, +",ox-1.45,oy-1.05,WORK),("-, -",ox-0.95,oy+0.42,ORANGE),("+, -",ox+0.72,oy+0.42,GOLD)]
for text,x,y,c in labels:
    tx(slide,text,x,y,1.05,0.30,17,c,True,spa=0,scaled=False)
vec(slide,ox,oy,ox+1.25,oy-0.85,"Q1",GREEN)
vec(slide,ox,oy,ox-1.25,oy-0.85,"Q2",WORK)
vec(slide,ox,oy,ox-1.25,oy+0.85,"Q3",ORANGE)
vec(slide,ox,oy,ox+1.25,oy+0.85,"Q4",GOLD)
mini_card(slide,"Memory line","Right means Fₓ positive. Left means Fₓ negative. Up means Fᵧ positive. Down means Fᵧ negative.",0.90,4.95,4.0,1.20,TEAL)
mini_card(slide,"Do not change sin/cos","Only the sign changes with quadrant.",8.75,4.95,3.6,1.20,ORANGE)

slide=new_slide("Copy","37 degrees and 53 degrees",TEAL)
grid_table(slide,["Angle","sin","cos"],[
    ["37 degrees","3/5","4/5"],
    ["53 degrees","4/5","3/5"],
    ["30 degrees","1/2","sqrt(3)/2"],
    ["45 degrees","1/sqrt(2)","1/sqrt(2)"],
    ["60 degrees","sqrt(3)/2","1/2"],
],XL,Y0,[4.0,4.0,4.0], row_h=0.62, hsize=16, csize=16)

slide=question_slide("10 N at 45 degrees")
qtx(slide,"A 10 N force is at 45 degrees above the +x axis. Find both components.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"Fx = 10 cos 45 = 10 / sqrt(2) = 5 sqrt(2)",XL,Y0,CW,LH,18)
fml(slide,"Fy = 10 sin 45 = 5 sqrt(2)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"Both components are 5 sqrt(2) N, and both are positive.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("10 N at 60 degrees")
qtx(slide,"A 10 N force is at 60 degrees above the +x axis. Find Fx and Fy.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"Fx = 10 cos 60 = 10 times 1/2 = 5",XL,Y0,CW,LH,18)
fml(slide,"Fy = 10 sin 60 = 10 times sqrt(3)/2 = 5 sqrt(3)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"Fx = 5 N, Fy = 5 sqrt(3) N.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("53 degrees in the second quadrant")
qtx(slide,"A 30 N force is in the second quadrant, 53 degrees from the +y axis toward the -x axis. Find the components.",XL,Y0,CW,1.10,18)
slide=answer_slide()
fml(slide,"The angle from +y is 53 degrees, so the y-component uses cos 53",XL,Y0,CW,LH,16)
fml(slide,"Fy = 30 cos 53 = 30 times 3/5 = 18, so +18 j-hat",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"Fx = 30 sin 53 = 30 times 4/5 = 24, but toward -x",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"Fx = -24 i-hat, Fy = +18 j-hat.",XL,Y0+3*LH+3*NG,CW,0.50)
ox,oy=axis_diagram(slide,8.2,3.20,3.7,2.8)
vec(slide,ox,oy,ox-1.55,oy-1.15,"30 N",GOLD)
draw_line(slide,ox-1.55,oy-1.15,ox-1.55,oy,WORK,1.4)
draw_line(slide,ox,oy,ox-1.55,oy,WORK,1.4)
tx(slide,"-24 î",ox-1.20,oy+0.10,0.80,0.25,13,WORK,True,spa=0,scaled=False)
tx(slide,"+18 ĵ",ox-1.95,oy-0.65,0.80,0.25,13,WORK,True,spa=0,scaled=False)

slide=question_slide("50 N at 37 degrees below +x")
qtx(slide,"F = 50 N in the fourth quadrant, 37 degrees below the +x axis. Find the components.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"Fx = 50 cos 37 = 50 times 4/5 = 40",XL,Y0,CW,LH,18)
fml(slide,"Fy = -50 sin 37 = -50 times 3/5 = -30",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"+40 i-hat - 30 j-hat.",XL,Y0+2*LH+2*NG,CW,0.50)
ox,oy=axis_diagram(slide,8.3,3.10,3.6,2.8)
vec(slide,ox,oy,ox+1.85,oy+1.25,"50 N",GOLD)
draw_line(slide,ox+1.85,oy+1.25,ox+1.85,oy,WORK,1.4)
draw_line(slide,ox,oy,ox+1.85,oy,WORK,1.4)
tx(slide,"+40 î",ox+0.78,oy+0.10,0.85,0.25,13,WORK,True,spa=0,scaled=False)
tx(slide,"-30 ĵ",ox+1.92,oy+0.65,0.85,0.25,13,WORK,True,spa=0,scaled=False)

slide=new_slide("Copy","Magnitude from components",TEAL)
ncert(slide,"|F| = sqrt(Fx^2 + Fy^2 + Fz^2)",XL,Y0,CW,0.90,24)

slide=new_slide("Visual idea","Magnitude is the length of the diagonal",TEAL)
body(slide,"Components form a right triangle. The vector is the diagonal, so Pythagoras gives its length.",XL,Y0,CW,0.55)
ox,oy=axis_diagram(slide,1.20,2.10,5.2,3.7)
arrow_line(slide,ox,oy,ox+2.70,oy,WORK,2.4); tx(slide,"Fₓ",ox+1.20,oy+0.17,0.60,0.30,16,WORK,True,spa=0,scaled=False)
arrow_line(slide,ox+2.70,oy,ox+2.70,oy-1.85,WORK,2.4); tx(slide,"Fᵧ",ox+2.85,oy-1.05,0.60,0.30,16,WORK,True,spa=0,scaled=False)
vec(slide,ox,oy,ox+2.70,oy-1.85,"F",GOLD,3.0)
draw_line(slide,ox+2.45,oy,ox+2.45,oy-0.25,MUTED,1.0); draw_line(slide,ox+2.45,oy-0.25,ox+2.70,oy-0.25,MUTED,1.0)
ncert(slide,"F² = Fₓ² + Fᵧ²     so     |F| = √(Fₓ² + Fᵧ²)",6.75,2.20,5.30,1.05,19)
mini_card(slide,"With z-component","In 3D, add one more perpendicular square: |F| = √(Fₓ² + Fᵧ² + F_z²).",6.75,3.75,5.30,1.25,TEAL)

slide=question_slide("Magnitude of 30 i-hat + 40 j-hat")
qtx(slide,"Find the magnitude of F = 30 i-hat + 40 j-hat.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"|F| = sqrt(30^2 + 40^2) = sqrt(900 + 1600)",XL,Y0,CW,LH,18)
fml(slide,"sqrt(2500) = 50",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"50",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("Magnitude when a component is negative")
qtx(slide,"Find |2 i-hat + 2 j-hat|, |10 i-hat + 10 j-hat| and |3 i-hat - 4 j-hat|.",XL,Y0,CW,0.90,18)
slide=answer_slide()
fml(slide,"sqrt(4 + 4) = sqrt(8) = 2 sqrt(2)",XL,Y0,CW,LH,18)
fml(slide,"sqrt(100 + 100) = 10 sqrt(2)",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"sqrt(9 + 16) = sqrt(25) = 5. The minus does not enter the square.",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"2 sqrt(2), 10 sqrt(2), and 5.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=question_slide("Three components")
qtx(slide,"Find |3 i-hat + 4 j-hat + 5 k-hat| and |6 i-hat - 8 j-hat + 10 k-hat|.",XL,Y0,CW,0.90,18)
slide=answer_slide()
fml(slide,"sqrt(9 + 16 + 25) = sqrt(50) = 5 sqrt(2)",XL,Y0,CW,LH,18)
fml(slide,"sqrt(36 + 64 + 100) = sqrt(200) = 10 sqrt(2)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"5 sqrt(2) and 10 sqrt(2).",XL,Y0+2*LH+2*NG,CW,0.50)

# ADDITION
slide=new_slide("New topic","Triangle law of addition",GREEN)
body(slide,"To find an angle, put tails together, or heads together.",XL,Y0,CW,0.50)
body(slide,"To add, put the tail of the second arrow on the head of the first.",XL,Y0+0.65,CW,0.55)
ncert(slide,"Resultant R runs from the first tail to the last head.",XL,Y0+1.35,CW,0.90,20)
body(slide,"A + B = B + A. For three arrows the drawing order does not change R.",XL,Y0+2.45,CW,0.55)
trap_card(slide,"The resultant of two arrows lies between them, closer to the longer one.",XL,Y0+3.15,CW,0.55)
vec(slide,8.1,4.95,9.75,4.95,"A",GOLD)
vec(slide,9.75,4.95,10.85,3.65,"B",GREEN)
vec(slide,8.1,4.95,10.85,3.65,"R = A + B",WORK)
tx(slide,"head-to-tail",9.25,5.35,1.8,0.25,13,MUTED,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Why drawing order does not change R",GREEN)
body(slide,"A + B and B + A reach the same final point. The path changes, but the start and finish do not.",XL,Y0,CW,0.55)
vec(slide,1.35,3.20,3.10,3.20,"A",GOLD); vec(slide,3.10,3.20,4.30,2.25,"B",GREEN); vec(slide,1.35,3.20,4.30,2.25,"R",WORK)
tx(slide,"A then B",2.35,3.78,1.6,0.30,14,MUTED,True,spa=0,scaled=False)
vec(slide,7.00,3.20,8.20,2.25,"B",GREEN); vec(slide,8.20,2.25,9.95,2.25,"A",GOLD); vec(slide,7.00,3.20,9.95,2.25,"R",WORK)
tx(slide,"B then A",8.00,3.78,1.6,0.30,14,MUTED,True,spa=0,scaled=False)
result_card(slide,"Same starting point + same ending point = same resultant.",XL,5.35,CW,0.50)

slide=new_slide("Visual idea","Resultant range: minimum to maximum",GREEN)
body(slide,"The largest resultant happens when both arrows pull in the same direction. The smallest happens when they oppose.",XL,Y0,CW,0.55)
vec(slide,1.25,2.45,3.05,2.45,"A",GOLD); vec(slide,3.05,2.45,4.35,2.45,"B",GREEN); vec(slide,1.25,2.95,4.35,2.95,"Rmax = A + B",WORK)
vec(slide,7.45,2.45,9.25,2.45,"A",GOLD); vec(slide,9.25,2.45,8.00,2.45,"B",GREEN); vec(slide,7.45,2.95,8.00,2.95,"Rmin = |A - B|",WORK)
mini_card(slide,"Allowed range","|A - B| ≤ R ≤ A + B",4.25,4.65,4.8,0.95,TEAL)
trap_card(slide,"A resultant outside this range is impossible.",4.25,5.85,4.8,0.50)

slide=new_slide("Copy","Magnitude of the resultant",TEAL)
ncert(slide,"R = sqrt(A^2 + B^2 + 2 A B cos theta)",XL,Y0,CW,0.90,22)
ncert(slide,"tan alpha = (B sin theta) / (A + B cos theta)",XL,Y0+1.10,CW,0.85,20)
body(slide,"alpha is the angle between A and the resultant.",XL,Y0+2.15,CW,0.45)
body(slide,"theta = 0: R max = A + B.   theta = 180: R min = |A - B|.   theta = 90: R = sqrt(A^2 + B^2).",XL,Y0+2.75,CW,0.70)
trap_card(slide,"R is allowed only from |A - B| to A + B. Both ends are included.",XL,Y0+3.60,CW,0.55)
vec(slide,8.3,5.65,10.35,5.65,"A",GOLD)
vec(slide,8.3,5.65,9.40,4.40,"B",GREEN)
vec(slide,8.3,5.65,10.10,4.75,"R",WORK)
tx(slide,"θ",8.85,5.20,0.35,0.25,14,ORANGE,True,spa=0,scaled=False)
tx(slide,"α",9.40,5.45,0.35,0.25,14,ORANGE,True,spa=0,scaled=False)

slide=question_slide("Two 4 N forces at 60 degrees")
qtx(slide,"|A| = 4 N, |B| = 4 N, angle = 60 degrees. Find R.",XL,Y0,CW,0.85,20)
slide=answer_slide()
fml(slide,"R = sqrt(16 + 16 + 32 times cos 60)",XL,Y0,CW,LH,18)
fml(slide,"cos 60 = 1/2, so the middle term is 16",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"R = sqrt(48) = 4 sqrt(3)",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"R = 4 sqrt(3) N.",XL,Y0+3*LH+3*NG,CW,0.48)
vec(slide,8.4,4.95,10.0,4.95,"4 N",GOLD)
vec(slide,8.4,4.95,9.25,3.55,"4 N",GREEN)
vec(slide,8.4,4.95,10.25,4.05,"R",WORK)
tx(slide,"60°",8.95,4.48,0.60,0.25,14,ORANGE,True,spa=0,scaled=False)

slide=question_slide("Magnitudes 3 and 4")
qtx(slide,"Two vectors have magnitudes 3 and 4. Find R at 0 degrees, 90 degrees and 180 degrees.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"At 0 degrees, R = 3 + 4 = 7",XL,Y0,CW,LH,18)
fml(slide,"At 90 degrees, R = sqrt(9 + 16) = 5",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"At 180 degrees, R = 4 - 3 = 1",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"7, 5 and 1.",XL,Y0+3*LH+3*NG,CW,0.48)

slide=question_slide("3 and 4 at 60 degrees and 120 degrees")
qtx(slide,"Same magnitudes 3 and 4. Find R at 60 degrees and at 120 degrees.",XL,Y0,CW,0.90,19)
slide=answer_slide()
fml(slide,"At 60 degrees: sqrt(9 + 16 + 12) = sqrt(37)",XL,Y0,CW,LH,18)
fml(slide,"At 120 degrees cos = -1/2, so sqrt(9 + 16 - 12) = sqrt(13)",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"sqrt(37) and sqrt(13).",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("5 N east and 10 N north")
qtx(slide,"Two forces, 5 N due east and 10 N due north. Find the resultant.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"The angle is 90 degrees, so the cross term is 0",XL,Y0,CW,LH,18)
fml(slide,"R = sqrt(25 + 100) = sqrt(125) = sqrt(25*5) = 5 sqrt(5)",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"5 sqrt(5) N.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("Two 5 N forces at 120 degrees")
qtx(slide,"Two forces of 5 N have 120 degrees between them. Find the net force.",XL,Y0,CW,0.90,19)
slide=answer_slide()
fml(slide,"Use the angle between the arrows, which is 120 degrees, not a nearby 60 degree mark",XL,Y0,CW,LH,15)
fml(slide,"R = sqrt(25 + 25 + 50*(-1/2)) = sqrt(50 - 25) = sqrt(25)",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"Net force = 5 N.",XL,Y0+2*LH+2*NG,CW,0.48)
vec(slide,8.5,4.65,10.0,4.65,"5 N",GOLD)
vec(slide,8.5,4.65,7.75,3.35,"5 N",GREEN)
vec(slide,8.5,4.65,9.05,3.95,"R",WORK)
tx(slide,"120°",8.15,4.08,0.65,0.25,14,ORANGE,True,spa=0,scaled=False)

slide=question_slide("Which pair cannot give a resultant of 4?")
qtx(slide,"Pairs: (2, 4), (4, 4), (4, 9), (1, 5). Which cannot give resultant 4?",XL,Y0,CW,0.90,18)
slide=answer_slide()
fml(slide,"(2, 4) allows 2 to 6. 4 is inside.",XL,Y0,CW,LH,17)
fml(slide,"(4, 4) allows 0 to 8. (1, 5) allows 4 to 6.",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"(4, 9) allows 5 to 13. 4 is smaller than the minimum.",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"(4, 9) cannot give 4.",XL,Y0+3*LH+3*NG,CW,0.48)

slide=question_slide("Three forces, zero resultant")
qtx(slide,"Which set can have zero resultant: (2, 4, 7), (3, 1, 5), (2, 8, 11), or (3, 4, 2)?",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"The third force must lie between |first - second| and first + second",XL,Y0,CW,LH,16)
fml(slide,"For 3 and 4, the range is 1 to 7. The force 2 lies inside it.",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"7 is outside 2 to 6. 5 is outside 2 to 4. 11 is outside 6 to 10.",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"Only (3, 4, 2) can give zero.",XL,Y0+3*LH+3*NG,CW,0.48)

slide=new_slide("Copy","Equal magnitudes",TEAL)
ncert(slide,"If |A| = |B| = A, then R = 2 A cos(theta / 2)",XL,Y0,CW,0.90,20)
ncert(slide,"Their difference has magnitude D = 2 A sin(theta / 2)",XL,Y0+1.10,CW,0.90,20)
body(slide,"The resultant then bisects the angle. Angle between A and R is theta / 2.",XL,Y0+2.20,CW,0.60)
body(slide,"theta = 0: R = 2A, D = 0.   theta = 90: both are sqrt(2) A.   theta = 120: R = A and D = sqrt(3) A.   theta = 180: R = 0 and D = 2A.",XL,Y0+2.95,CW,0.85)

slide=new_slide("Visual idea","Equal vectors: resultant bisects the angle",TEAL)
body(slide,"When |A| = |B|, the triangle is symmetric. The resultant must lie exactly in the middle.",XL,Y0,CW,0.55)
cx,cy=3.50,4.55
vec(slide,cx,cy,cx+2.25,cy,"A",GOLD)
vec(slide,cx,cy,cx+1.12,cy-1.95,"B",GREEN)
vec(slide,cx,cy,cx+2.00,cy-1.15,"R",WORK,3.0)
angle_mark(slide,cx,cy,0.70,0,60,"θ")
angle_mark(slide,cx,cy,1.05,0,30,"θ/2",ORANGE)
ncert(slide,"R = 2A cos(θ / 2)",7.00,2.35,4.80,0.90,22)
mini_card(slide,"Fast memory","Equal vectors make an isosceles triangle, so R is the angle bisector.",7.00,3.80,4.80,1.15,TEAL)

slide=question_slide("Sum and difference have the same magnitude")
qtx(slide,"If |A + B| = |A - B|, what is the angle between A and B?",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"Equal lengths of sum and difference means the cross term must vanish",XL,Y0,CW,LH,16)
fml(slide,"2 A B cos theta = 0, so cos theta = 0",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"90 degrees.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("Two unit vectors add to a unit vector")
qtx(slide,"The sum of two unit vectors is itself a unit vector. Find the magnitude of their difference.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"R = 1 and each magnitude is 1, so 1 = 2 cos(theta / 2), so theta = 120 degrees",XL,Y0,CW,LH,15)
fml(slide,"D = 2 sin(60) = 2 times sqrt(3)/2 = sqrt(3)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"Magnitude of the difference = sqrt(3).",XL,Y0+2*LH+2*NG,CW,0.50)

# SUBTRACTION
slide=new_slide("New topic","Vector subtraction",GREEN)
ncert(slide,"A - B = A + (-B). Reverse B, then add.",XL,Y0,CW,0.90,22)
ncert(slide,"D = sqrt(A^2 + B^2 - 2 A B cos theta)",XL,Y0+1.10,CW,0.90,22)
body(slide,"theta = 0: D = |A - B|, the minimum.   theta = 180: D = A + B, the maximum.   theta = 90: D = sqrt(A^2 + B^2).",XL,Y0+2.20,CW,0.75)
trap_card(slide,"A + B equals B + A. A - B does not equal B - A.",XL,Y0+3.15,6.80,0.60)
vec(slide,8.15,4.90,9.80,4.90,"A",GOLD)
vec(slide,9.80,4.90,8.75,3.75,"-B",ORANGE)
vec(slide,8.15,4.90,8.75,3.75,"D = A - B",WORK)
tx(slide,"reverse B first",9.00,5.25,2.0,0.25,13,MUTED,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Subtraction is not a new operation",GREEN)
body(slide,"Every subtraction question is an addition question after reversing the second vector.",XL,Y0,CW,0.50)
vec(slide,1.20,2.55,3.05,2.55,"A",GOLD); vec(slide,1.20,3.25,2.35,2.30,"B",GREEN)
tx(slide,"Step 1: given A and B",1.20,3.80,3.6,0.30,14,TEXT,False,spa=0,scaled=False)
vec(slide,5.00,2.55,6.85,2.55,"A",GOLD); vec(slide,5.00,3.25,3.85,4.20,"-B",ORANGE)
tx(slide,"Step 2: reverse B",5.00,4.48,3.0,0.30,14,TEXT,False,spa=0,scaled=False)
vec(slide,8.85,3.65,10.70,3.65,"A",GOLD); vec(slide,10.70,3.65,9.55,4.60,"-B",ORANGE); vec(slide,8.85,3.65,9.55,4.60,"A - B",WORK)
tx(slide,"Step 3: add A + (-B)",8.85,4.95,3.2,0.30,14,TEXT,False,spa=0,scaled=False)
result_card(slide,"A - B and B - A have the same length but opposite directions.",XL,5.75,CW,0.48)

slide=new_slide("Visual idea","Adding by components is adding columns",GREEN)
body(slide,"The x-parts combine with x-parts. The y-parts combine with y-parts. This is the cleanest method for many vectors.",XL,Y0,CW,0.55)
ox,oy=axis_diagram(slide,1.10,2.05,4.7,3.7)
vec(slide,ox,oy,ox+1.80,oy-1.15,"A",GOLD); draw_line(slide,ox+1.80,oy-1.15,ox+1.80,oy,WORK,1.2); draw_line(slide,ox,oy,ox+1.80,oy,WORK,1.2)
vec(slide,ox,oy,ox+1.25,oy+1.45,"B",GREEN); draw_line(slide,ox+1.25,oy+1.45,ox+1.25,oy,ORANGE,1.2); draw_line(slide,ox,oy,ox+1.25,oy,ORANGE,1.2)
ncert(slide,"A + B = (Aₓ + Bₓ) î + (Aᵧ + Bᵧ) ĵ",6.60,2.10,5.45,0.90,19)
mini_card(slide,"Column idea","x-column decides left-right. y-column decides up-down.",6.60,3.55,5.45,1.15,TEAL)
mini_card(slide,"3D same rule","A + B = (Aₓ+Bₓ)î + (Aᵧ+Bᵧ)ĵ + (A_z+B_z)k̂",6.60,5.00,5.45,1.15,GREEN)

slide=question_slide("Add by components")
qtx(slide,"A = 4 i-hat + 3 j-hat and B = 3 i-hat - 4 j-hat. Find A + B, A - B, and |A| + |B|.",XL,Y0,CW,1.00,17)
slide=answer_slide()
fml(slide,"A + B = 7 i-hat - 1 j-hat",XL,Y0,CW,LH,18)
fml(slide,"A - B = 1 i-hat + 7 j-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"|A| = 5 and |B| = 5, so |A| + |B| = 10",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"Sum of magnitudes is 10. Magnitude of the sum is sqrt(50).",XL,Y0+3*LH+3*NG,CW,0.55)
ox,oy=axis_diagram(slide,8.25,3.05,3.6,2.9)
vec(slide,ox,oy,ox+1.45,oy-1.05,"A",GOLD)
vec(slide,ox,oy,ox+1.10,oy+1.35,"B",GREEN)
vec(slide,ox,oy,ox+2.05,oy+0.30,"A+B",WORK)

slide=question_slide("Change in velocity")
qtx(slide,"Velocity changes from 10 m/s east to 10 m/s north. Find the change in velocity.",XL,Y0,CW,0.90,18)
slide=answer_slide()
fml(slide,"Change = final - initial = 10 j-hat - 10 i-hat",XL,Y0,CW,LH,18)
fml(slide,"Magnitude = sqrt(100 + 100) = 10 sqrt(2)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"10 sqrt(2) m/s, pointing north of west.",XL,Y0+2*LH+2*NG,CW,0.52)
ox,oy=axis_diagram(slide,8.4,3.15,3.4,2.8)
vec(slide,ox,oy,ox+1.55,oy,"vᵢ",GOLD)
vec(slide,ox,oy,ox,oy-1.55,"vᶠ",GREEN)
vec(slide,ox+1.55,oy,ox,oy-1.55,"Δv",WORK)

slide=question_slide("Rotate 2 i-hat")
qtx(slide,"A = 2 i-hat is rotated by 90 degrees, 180 degrees, 270 degrees and 360 degrees anticlockwise. Write each new vector.",XL,Y0,CW,0.95,17)
slide=answer_slide()
fml(slide,"90 degrees: 2 j-hat",XL,Y0,CW,LH,18)
fml(slide,"180 degrees: -2 i-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"270 degrees: -2 j-hat.   360 degrees: 2 i-hat again",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"Rotation keeps the magnitude and changes only the direction.",XL,Y0+3*LH+3*NG,CW,0.52)
ox,oy=axis_diagram(slide,8.4,3.05,3.4,2.9)
vec(slide,ox,oy,ox+1.25,oy,"0°",GOLD)
vec(slide,ox,oy,ox,oy-1.25,"90°",GREEN)
vec(slide,ox,oy,ox-1.25,oy,"180°",WORK)
vec(slide,ox,oy,ox,oy+1.25,"270°",ORANGE)

# DOT
slide=new_slide("New topic","Dot product",GREEN)
ncert(slide,"A dot B = A B cos theta = Aₓ Bₓ + Aᵧ Bᵧ + A_z B_z",XL,Y0,CW,1.00,20)
body(slide,"If the dot product is 0, the vectors are perpendicular.",XL,Y0+1.20,CW,0.50)
body(slide,"cos theta = (A dot B) / (A B). That is how the angle is found from components.",XL,Y0+1.80,CW,0.55)
body(slide,"Component of B along A = (A dot B) / A.",XL,Y0+2.45,CW,0.50)
body(slide,"i-hat dot i-hat = 1. i-hat dot j-hat = 0. The same rule holds for j-hat and k-hat.",XL,Y0+3.10,CW,0.55)
trap_card(slide,"A unit vector dotted with itself is 1.",XL,Y0+3.80,CW,0.48)
vec(slide,8.15,5.10,10.50,5.10,"A",GOLD)
vec(slide,8.15,5.10,9.65,3.70,"B",GREEN)
draw_line(slide,9.65,3.70,9.65,5.10,WORK,1.2)
arrow_line(slide,8.15,5.10,9.65,5.10,WORK,2.0)
tx(slide,"B cos θ",8.55,5.22,1.25,0.28,13,WORK,True,spa=0,scaled=False)
tx(slide,"projection of B along A",9.30,5.42,2.5,0.25,13,MUTED,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Dot product tells alignment",GREEN)
body(slide,"Dot product is positive when vectors face mostly the same way, zero at 90°, and negative when they face opposite way.",XL,Y0,CW,0.55)
vec(slide,1.25,3.00,3.10,3.00,"A",GOLD); vec(slide,1.25,3.00,2.85,2.25,"B",GREEN); angle_mark(slide,1.25,3.00,0.60,0,25,"θ<90°")
tx(slide,"A · B > 0",1.35,3.55,1.8,0.30,17,GREEN,True,spa=0,scaled=False)
vec(slide,5.35,3.25,7.05,3.25,"A",GOLD); vec(slide,5.35,3.25,5.35,1.60,"B",GREEN); angle_mark(slide,5.35,3.25,0.60,0,90,"90°")
tx(slide,"A · B = 0",5.35,3.75,1.8,0.30,17,WORK,True,spa=0,scaled=False)
vec(slide,9.15,3.00,10.95,3.00,"A",GOLD); vec(slide,9.15,3.00,7.85,2.20,"B",GREEN); angle_mark(slide,9.15,3.00,0.65,140,180,"θ>90°")
tx(slide,"A · B < 0",9.05,3.55,1.8,0.30,17,ORANGE,True,spa=0,scaled=False)
mini_card(slide,"Why this matters","It quickly tests perpendicular vectors and helps find the angle between two vectors.",3.90,5.20,5.50,1.00,TEAL)

slide=question_slide("Are these two vectors perpendicular?")
qtx(slide,"A = 3 i-hat + 4 j-hat + 5 k-hat and B = 3 i-hat + 4 j-hat - 5 k-hat. What is the angle between them?",XL,Y0,CW,1.05,17)
slide=answer_slide()
fml(slide,"A dot B = 9 + 16 - 25 = 0",XL,Y0,CW,LH,18)
result_card(slide,"The angle is 90 degrees.",XL,Y0+LH+NG,CW,0.50)

slide=question_slide("Angle with 4 i-hat")
qtx(slide,"A = 2 i-hat + 6 j-hat + 3 k-hat and B = 4 i-hat. Find cos theta.",XL,Y0,CW,0.90,19)
slide=answer_slide()
fml(slide,"A dot B = 8",XL,Y0,CW,LH,18)
fml(slide,"|A| = sqrt(4 + 36 + 9) = 7 and |B| = 4",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"cos theta = 8 / 28 = 2/7",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"theta = cos inverse of 2/7.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=new_slide("Copy","Unit vectors from components",TEAL)
body(slide,"i-hat itself is already a unit vector.",XL,Y0,CW,0.45)
body(slide,"Unit vector of i-hat + j-hat is (i-hat + j-hat) / sqrt(2).",XL,Y0+0.60,CW,0.50)
body(slide,"Unit vector of i-hat + j-hat - k-hat is (i-hat + j-hat - k-hat) / sqrt(3).",XL,Y0+1.25,CW,0.55)
body(slide,"sin theta i-hat - cos theta j-hat already has magnitude 1, because sin^2 + cos^2 = 1.",XL,Y0+1.95,CW,0.60)
body(slide,"Unit vector of 3 i-hat + 4 j-hat is (3 i-hat + 4 j-hat) / 5.",XL,Y0+2.70,CW,0.55)

slide=new_slide("Visual idea","Unit vector keeps direction and removes size",TEAL)
body(slide,"A unit vector is only the direction of a vector. It has length 1 and no physical unit.",XL,Y0,CW,0.50)
ox,oy=axis_diagram(slide,1.20,2.20,4.8,3.5)
vec(slide,ox,oy,ox+2.40,oy-1.60,"A",GOLD,3.0)
vec(slide,ox,oy,ox+1.20,oy-0.80,"Â",GREEN,2.4)
draw_line(slide,ox+1.20,oy-0.80,ox+2.40,oy-1.60,MUTED,1.0)
ncert(slide,"Â = A / |A|",6.75,2.30,4.60,0.90,26)
mini_card(slide,"Example","For A = 3î + 4ĵ, |A| = 5, so Â = (3î + 4ĵ)/5.",6.75,3.80,4.60,1.25,GREEN)
trap_card(slide,"Do not attach N, m/s or kg to a unit vector.",6.75,5.35,4.60,0.55)

slide=question_slide("Perpendicular forces A and A/2")
qtx(slide,"Two forces of magnitudes A and A/2 are perpendicular. Find the magnitude of their resultant.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"cos 90 = 0, so R = sqrt(A^2 + A^2/4)",XL,Y0,CW,LH,18)
fml(slide,"R = sqrt(5 A^2 / 4) = (A/2) sqrt(5)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"(sqrt(5) / 2) times A.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Resultant is half the larger vector")
qtx(slide,"The resultant of two vectors is half the larger one, and it is perpendicular to the smaller one. Find the angle between the vectors.",XL,Y0,CW,1.10,17)
slide=answer_slide()
fml(slide,"R perpendicular to the smaller vector means (A + B) dot B = 0",XL,Y0,CW,LH,15)
fml(slide,"cos theta = -B/A",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"R = A/2 then gives B = (sqrt(3)/2) A",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"cos theta = -sqrt(3)/2, so theta = 150 degrees.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=question_slide("Acceleration of a 4 kg body")
qtx(slide,"F1 = 5 i-hat + 8 j-hat + 7 k-hat and F2 = 3 i-hat - 4 j-hat - 3 k-hat act on 4 kg. Find the acceleration.",XL,Y0,CW,1.05,16)
slide=answer_slide()
fml(slide,"Net force = 8 i-hat + 4 j-hat + 4 k-hat",XL,Y0,CW,LH,18)
fml(slide,"a = net force / 4",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"a = 2 i-hat + 1 j-hat + 1 k-hat.",XL,Y0+2*LH+2*NG,CW,0.52)

slide=new_slide("Class summary","Vector 02 lock lines",TEAL)
fml(slide,"Components: cos along the axis you measured from, then keep the sign",XL,Y0,CW,LH,16)
fml(slide,"|F| = sqrt(Fx^2 + Fy^2 + Fz^2)",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"R = sqrt(A^2 + B^2 + 2AB cos theta), and |A-B| to A+B is the allowed range",XL,Y0+2*LH+2*NG,CW,LH,15)
fml(slide,"A dot B = 0 means 90 degrees. Unit vector = vector divided by magnitude.",XL,Y0+3*LH+3*NG,CW,LH,15)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Draw 30 degrees east of south and 60 degrees south of east on one compass.",XL,Y0,CW,0.55)
body(slide,"2. Resolve 10 N at 30 degrees and at 53 degrees, with signs.",XL,Y0+0.70,CW,0.50)
body(slide,"3. Redo magnitudes 3 and 4 at 0 degrees, 60 degrees, 90 degrees, 120 degrees and 180 degrees.",XL,Y0+1.35,CW,0.50)
body(slide,"4. Redo the unit-vector list and the 4 kg acceleration.",XL,Y0+2.00,CW,0.50)
trap_card(slide,"Do not add magnitudes when the arrows are not along one line.",XL,Y0+2.70,CW,0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
