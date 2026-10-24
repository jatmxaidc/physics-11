"""Units 04 — Significant figures in calculations, accuracy, precision, errors (PW 05)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Units_04_Significant_Figures_Accuracy_and_Errors.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


_GREEK_SUP = {"α": "ᵅ", "β": "ᵝ", "γ": "ᵞ", "δ": "ᵟ", "ε": "ᵋ", "θ": "ᶿ"}


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(-1/2)", "⁻½").replace("^(-3/2)", "⁻³⁄²")
    text = text.replace("^(-5/2)", "⁻⁵⁄²").replace("^(-7/2)", "⁻⁷⁄²")
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³⁄²")
    text = text.replace("^(5/2)", "⁵⁄²").replace("^(7/2)", "⁷⁄²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([αβγδεθ])", lambda m: _GREEK_SUP[m.group(1)], text)
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    step = 0.52
    for i, item in enumerate(items):
        cy = y + i * step
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.06), Inches(0.32), Inches(0.32))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.08, 0.32, 0.28, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.46, cy + 0.04, w - 0.46, 0.42, 17, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units 04",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units and Measurement",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Significant Figures, Accuracy and Errors",bx,by+1.62,bw,0.40,16,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"15 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"From PW Units and Measurement Lecture 05."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Significant figures in a calculated answer",
"Rounding off",
"Accuracy and precision",
"Absolute, relative and percentage error",
"Mean error from repeated readings",
"Dimensional questions with a new set of quantities",
"When a formula cannot be built by dimensions",
]
for i,t in enumerate(topics):
    y=1.22+i*0.72
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(XL),Inches(y),Inches(0.42),Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.fill.background()
    tx(slide,f"{i+1:02d}",XL,y+0.05,0.42,0.32,14,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
    tx(slide,t,XL+0.60,y+0.05,10.5,0.35,20,TEXT,False,spa=0,scaled=False)
_footer(slide,2); after_content_slide()

slide=new_slide("Class note","Where we left Units 03",TEAL)
para(slide,"Last class we used dimensions to find powers, and we learned how to count significant digits in one measured number. A trailing zero after the decimal counts. A leading zero does not. The power of ten is not a digit.",XL,Y0,CW,1.35,16)
para(slide,"Today we use those digits in a sum, a difference, a product and a quotient. Then we separate accuracy from precision, and we start errors: absolute, relative, percentage, and the mean of repeated readings.",XL,Y0+1.50,CW,1.35,16)
ncert(slide,"A calculated answer keeps only the digits the measurements can support.",XL,Y0+3.05,CW,0.90,16)
trap_card(slide,"Class notes will be saved after teaching. This deck is for the board.",XL,Y0+4.15,CW,0.45)

# SIGNIFICANT FIGURES
slide=new_slide("Idea","What a significant figure is",GREEN)
para(slide,"A significant figure is a digit we trust in a measurement. All the certain digits are kept, and one uncertain digit is kept with them. That last digit is the one the instrument is guessing between two marks.",XL,Y0,CW,1.40,16)
para(slide,"Example: a mass written as 256703 kg has six significant digits. Five of them are certain. The last one, 3, is the uncertain digit. More significant digits means a finer measurement. That is precision.",XL,Y0+1.55,CW,1.35,16)
ncert(slide,"Significant figures = all certain digits + one uncertain digit",XL,Y0+3.10,CW,0.85,16)

slide=new_slide("Rule","Changing the unit does not change the count",TEAL)
para(slide,"2.5 kg and 2.5 × 10^3 g are the same mass. Both have two significant digits. Writing 2500 g would look like a different count, so we keep scientific notation when we change the unit.",XL,Y0,CW,1.35,16)
ncert(slide,"2.5 kg  =  2.5 × 10^3 g     both have 2 significant digits",XL,Y0+1.55,CW,0.90,16)
trap_card(slide,"The power 10^3 is only a place-value label. It is not an extra digit.",XL,Y0+2.70,CW,0.50)

slide=new_slide("Rules","How to count — part 1",GREEN)
bullet(slide,"Every non-zero digit is significant.  2427 m has 4.",XL,Y0,CW,0.48,16)
bullet(slide,"A zero between non-zero digits is significant.  24507 has 5.",XL,Y0+0.55,CW,0.48,16)
bullet(slide,"A trailing zero with no decimal point is not significant.  4500 m has 2.",XL,Y0+1.15,CW,0.55,16)
bullet(slide,"Zeros after the decimal point are significant.  45.00 cm has 4.",XL,Y0+1.80,CW,0.55,16)
trap_card(slide,"4500 m and 45.00 cm are not the same kind of zero.",XL,Y0+2.60,CW,0.48)

slide=new_slide("Rules","How to count — part 2",GREEN)
bullet(slide,"In a number smaller than 1, zeros on the left of the first non-zero digit are not significant.  0.02 m has 1.  0.0307 m has 3.",XL,Y0,CW,0.85,16)
bullet(slide,"An exact count has unlimited significant digits. Five pens means exactly 5, not a measured 5.",XL,Y0+1.05,CW,0.70,16)
bullet(slide,"In scientific notation, only the front number is counted.  2.303 × 10^23 m has 4 significant digits.",XL,Y0+1.90,CW,0.75,16)
ncert(slide,"Exact number → unlimited digits.    Power of ten → not a digit.",XL,Y0+2.85,CW,0.85,15)

slide=question_slide("Count these measurements")
qtx(slide,"How many significant figures are in each?",XL,Y0,CW,0.40,18)
body(slide,"1.  36.7 m",XL,Y0+0.55,CW,0.40,18)
body(slide,"2.  0.006606 s",XL,Y0+1.05,CW,0.40,18)
body(slide,"3.  2,002 kg",XL,Y0+1.55,CW,0.40,18)
body(slide,"4.  306,490,000 people, counted exactly",XL,Y0+2.05,CW,0.45,18)
slide=answer_slide()
body(slide,"Apply one rule to each number.",XL,Y0,CW,0.40)
fml(slide,"36.7 m — three non-zero digits → 3.",XL,Y0+0.50,CW,LH,16)
fml(slide,"0.006606 s — leading zeros drop. Digits 6, 6, 0, 6 → 4.",XL,Y0+0.50+LH+NG,CW,LH,15)
fml(slide,"2,002 kg — zeros sit between non-zeros → 4.",XL,Y0+0.50+2*LH+2*NG,CW,LH,16)
fml(slide,"An exact people-count is not a measured guess → unlimited significant figures.",XL,Y0+0.50+3*LH+3*NG,CW,0.70,15)
result_card(slide,"3 ,   4 ,   4 ,   unlimited",XL,Y0+0.50+4*LH+4*NG,CW,0.48)

slide=question_slide("Three numbers together")
qtx(slide,"P = 0.0030 m,  Q = 2.40 m,  R = 3000 m.",XL,Y0,CW,0.55,18)
qtx(slide,"Significant figures in P, Q, R are",XL,Y0+0.60,CW,0.40,18)
choices(slide,["1, 2, 1","2, 3, 1","4, 2, 1","4, 2, 4"], XL, Y0+1.20, 10.5)
slide=answer_slide()
body(slide,"Read each number with its own zero rule.",XL,Y0,CW,0.40)
fml(slide,"P = 0.0030. Leading zeros drop. The last 0 is after the decimal → 2.",XL,Y0+0.50,CW,LH,15)
fml(slide,"Q = 2.40. The zero after the decimal counts → 3.",XL,Y0+0.50+LH+NG,CW,LH,16)
fml(slide,"R = 3000. No decimal point, so trailing zeros drop → 1.",XL,Y0+0.50+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: 2, 3, 1   (option 2)",XL,Y0+0.50+3*LH+3*NG,CW,0.48)

slide=new_slide("Arithmetic","Addition and subtraction",TEAL)
para(slide,"When you add or subtract, the answer is limited by the least precise decimal place. The calculator may print extra digits. Those extra digits are not in the measurements, so they are cut off.",XL,Y0,CW,1.25,16)
fml(slide,"4.71 + 1.3 = 6.01 on a calculator.  1.3 has one decimal place, so the sum is 6.0.",XL,Y0+1.40,CW,0.70,15)
fml(slide,"3.813 + 2.21 = 6.023.  2.21 has two decimal places, so the sum is 6.02.",XL,Y0+2.20,CW,0.70,15)
ncert(slide,"Sum or difference → keep the least number of decimal places.",XL,Y0+3.10,CW,0.80,15)

slide=new_slide("Arithmetic","Multiplication and division",TEAL)
para(slide,"When you multiply or divide, the answer keeps the smaller count of significant figures. Decimal places are not the rule here. The count of trusted digits is the rule.",XL,Y0,CW,1.25,16)
fml(slide,"3.01 × 3.1 = 9.331 on a calculator.",XL,Y0+1.40,CW,LH,16)
fml(slide,"3.01 has 3 significant figures.  3.1 has 2. Keep 2 → 9.3.",XL,Y0+1.40+LH+NG,CW,LH,16)
ncert(slide,"Product or quotient → keep the least number of significant figures.",XL,Y0+1.40+2*LH+2*NG,CW,0.85,15)

slide=question_slide("Area of a rectangle")
qtx(slide,"Length 55.3 m, breadth 25 m. The area, with the proper significant figures, is",XL,Y0,CW,0.85,18)
choices(slide,["1382 m^2","1382.5 m^2","14 × 10^2 m^2","138 × 10^1 m^2"], XL, Y0+1.05, 10.8)
slide=answer_slide()
body(slide,"Area is a product, so count significant figures, not decimal places.",XL,Y0,CW,0.48)
fml(slide,"A = 55.3 × 25 = 1382.5 m^2 on the calculator.",XL,Y0+0.58,CW,LH,16)
fml(slide,"55.3 has 3 significant figures.  25 has 2. Keep 2.",XL,Y0+0.58+LH+NG,CW,LH,16)
fml(slide,"1382.5 to 2 significant figures is 1.4 × 10^3, which is the same as 14 × 10^2.",XL,Y0+0.58+2*LH+2*NG,CW,0.70,15)
result_card(slide,"Final answer: 14 × 10^2 m^2   (option 3)",XL,Y0+0.58+3*LH+3*NG,CW,0.48)

slide=question_slide("A small subtraction")
qtx(slide,"With significant figures, the value of 9.99 m − 0.0099 m is",XL,Y0,CW,0.70,18)
choices(slide,["9.98 m","9.980 m","9.9 m","9.9801 m"], XL, Y0+0.95, 10.5)
slide=answer_slide()
body(slide,"This is a difference, so match decimal places.",XL,Y0,CW,0.42)
fml(slide,"9.99 − 0.0099 = 9.9801 on the calculator.",XL,Y0+0.52,CW,LH,16)
fml(slide,"9.99 has two decimal places.  0.0099 has four. Keep two.",XL,Y0+0.52+LH+NG,CW,LH,15)
fml(slide,"9.9801 written to two decimal places is 9.98 m.",XL,Y0+0.52+2*LH+2*NG,CW,LH,16)
result_card(slide,"Final answer: 9.98 m   (option 1)",XL,Y0+0.52+3*LH+3*NG,CW,0.48)

slide=question_slide("Add three lengths")
qtx(slide,"Add 0.307 m, 0.52 m and 0.4 m.",XL,Y0,CW,0.50,18)
choices(slide,["1.227 m","1.2 m","1.3 m","1.23 m"], XL, Y0+0.75, 10.5)
slide=answer_slide()
body(slide,"The roughest length decides the decimal place.",XL,Y0,CW,0.42)
fml(slide,"0.307 + 0.52 + 0.4 = 1.227 m.",XL,Y0+0.52,CW,LH,16)
fml(slide,"Decimal places available: 3, 2 and 1. The least is 1, from 0.4 m.",XL,Y0+0.52+LH+NG,CW,LH,15)
fml(slide,"1.227 to one decimal place is 1.2 m.",XL,Y0+0.52+2*LH+2*NG,CW,LH,16)
result_card(slide,"Final answer: 1.2 m   (option 2)",XL,Y0+0.52+3*LH+3*NG,CW,0.48)

slide=new_slide("Rounding","The three rounding rules",AMBER)
bullet(slide,"If the first digit you drop is less than 5, leave the last kept digit as it is.",XL,Y0,CW,0.70,16)
bullet(slide,"If the first digit you drop is more than 5, increase the last kept digit by 1.",XL,Y0+0.85,CW,0.70,16)
bullet(slide,"If the first digit you drop is exactly 5, leave an even last digit unchanged, and increase an odd last digit by 1.",XL,Y0+1.70,CW,0.90,16)
ncert(slide,"…4 stays.    …6 goes up.    …5 goes to the even digit.",XL,Y0+2.80,CW,0.80,16)

slide=question_slide("Round to 4 significant figures")
qtx(slide,"(i)  Round 4.7475 to 4 significant figures.",XL,Y0,CW,0.50,18)
qtx(slide,"(ii) Round 4.7465 to 4 significant figures.",XL,Y0+0.70,CW,0.50,18)
slide=answer_slide()
body(slide,"Four significant figures means we keep 4.74x and look at the next digit, which is 5.",XL,Y0,CW,0.55)
fml(slide,"4.7475 — the kept digit before the 5 is 7, which is odd, so it becomes 8.",XL,Y0+0.70,CW,0.70,15)
fml(slide,"4.7465 — the kept digit before the 5 is 6, which is even, so it stays 6.",XL,Y0+1.50,CW,0.70,15)
result_card(slide,"(i) 4.748      (ii) 4.746",XL,Y0+2.40,CW,0.50)

slide=question_slide("Proper final answers")
qtx(slide,"(i)  101.2 + 18.702 = ?",XL,Y0,CW,0.50,18)
qtx(slide,"(ii) 202.88 − 1.013 = ?",XL,Y0+0.65,CW,0.50,18)
slide=answer_slide()
body(slide,"Both are add or subtract, so use decimal places.",XL,Y0,CW,0.42)
fml(slide,"101.2 + 18.702 = 119.902. Least decimal places = 1 → 119.9.",XL,Y0+0.55,CW,0.70,15)
fml(slide,"202.88 − 1.013 = 201.867. Least decimal places = 2 → 201.87.",XL,Y0+1.40,CW,0.70,15)
result_card(slide,"(i) 119.9      (ii) 201.87",XL,Y0+2.30,CW,0.50)

slide=question_slide("One sum and one quotient")
qtx(slide,"(i)  A = 2.413 and B = 1.2. Find A + B.",XL,Y0,CW,0.50,18)
qtx(slide,"(ii) A = 25.5 and B = 5, both measured. Find A/B.",XL,Y0+0.65,CW,0.55,18)
slide=answer_slide()
body(slide,"The sum follows decimal places. The quotient follows significant figures.",XL,Y0,CW,0.50)
fml(slide,"2.413 + 1.2 = 3.613. B has one decimal place → 3.6.",XL,Y0+0.60,CW,0.65,16)
fml(slide,"25.5 / 5 = 5.1 on the calculator. Measured 5 has 1 significant figure.",XL,Y0+1.35,CW,0.65,15)
fml(slide,"So A/B is written as 5. If the 5 were an exact count, we would keep 5.10.",XL,Y0+2.10,CW,0.65,15)
result_card(slide,"(i) 3.6      (ii) 5",XL,Y0+2.90,CW,0.48)

# ACCURACY AND PRECISION
slide=new_slide("Idea","Accuracy and precision",GREEN)
para(slide,"Accuracy tells how close a reading is to the true value. A small error from the true value means a more accurate reading.",XL,Y0,CW,1.15,16)
para(slide,"Precision tells how fine the instrument is, and how close repeated readings are to each other. More decimal places, or more significant figures, means a more precise reading. A precise reading can still sit far from the true value.",XL,Y0+1.30,CW,1.40,16)
ncert(slide,"Accuracy → close to the true value.    Precision → a fine, repeatable reading.",XL,Y0+2.90,CW,0.90,15)

slide=question_slide("Most accurate, most precise")
qtx(slide,"True length = 6.57 m. Which reading is most accurate, and which is most precise?",XL,Y0,CW,0.80,17)
choices(slide,["6.52 m","6.61 m","6.513 m","6.68 m"], XL, Y0+1.00, 10.5)
slide=answer_slide()
body(slide,"Accuracy uses the gap from 6.57 m. Precision uses how fine the number is.",XL,Y0,CW,0.50)
fml(slide,"Gaps: 6.52 → 0.05,   6.61 → 0.04,   6.513 → 0.057,   6.68 → 0.11.",XL,Y0+0.60,CW,0.65,15)
fml(slide,"Smallest gap is 0.04, so 6.61 m is the most accurate.",XL,Y0+1.35,CW,LH,16)
fml(slide,"6.513 m is written to a thousandth of a metre, so it is the most precise.",XL,Y0+1.35+LH+NG,CW,0.65,15)
result_card(slide,"Most accurate: 6.61 m.    Most precise: 6.513 m.",XL,Y0+2.55,CW,0.50)

slide=question_slide("Finest reading")
qtx(slide,"No true value is given. Which reading is the finest?",XL,Y0,CW,0.55,18)
choices(slide,["2.4 m","2.41 m","2 m","2.413 m"], XL, Y0+0.80, 10.5)
slide=answer_slide()
body(slide,"Accuracy needs a true value. Here we only have the readings, so we judge precision.",XL,Y0,CW,0.55)
fml(slide,"2 m has 1 decimal-free digit.  2.4 m has 2 significant figures.",XL,Y0+0.70,CW,LH,15)
fml(slide,"2.41 m has 3.  2.413 m has 4, down to a millimetre.",XL,Y0+0.70+LH+NG,CW,LH,16)
result_card(slide,"Most precise reading: 2.413 m   (option 4)",XL,Y0+0.70+2*LH+2*NG,CW,0.50)
trap_card(slide,"Do not call it more accurate until a true value is given.",XL,Y0+2.70,CW,0.45)

# ERRORS
slide=new_slide("Definitions","Three ways to write an error",ORANGE)
para(slide,"Absolute error is the size of the mistake, in the same unit as the quantity. If the true value is a_0 and the measured value is a, the absolute error is |a − a_0|.",XL,Y0,CW,1.20,16)
para(slide,"Relative error is that absolute error divided by the true value. It has no unit. Percentage error is the relative error multiplied by 100.",XL,Y0+1.35,CW,1.15,16)
ncert(slide,"Relative error = absolute error / true value.    Percentage error = relative error × 100.",XL,Y0+2.70,CW,0.95,15)

slide=question_slide("Relative error")
qtx(slide,"Absolute error = 5 and true value = 15. The relative error is",XL,Y0,CW,0.60,18)
slide=answer_slide()
body(slide,"Relative error is a ratio, so the unit cancels.",XL,Y0,CW,0.42)
fml(slide,"Relative error = 5 / 15 = 1/3.",XL,Y0+0.55,CW,LH,18)
fml(slide,"Percentage error = (1/3) × 100 = 33.3%.",XL,Y0+0.55+LH+NG,CW,LH,16)
result_card(slide,"Relative error = 1/3",XL,Y0+0.55+2*LH+2*NG,CW,0.48)

slide=question_slide("Error in a sum")
qtx(slide,"l_1 = (10 ± 2) cm and l_2 = (20 ± 1) cm. Find l = l_1 + l_2 with its error.",XL,Y0,CW,0.80,17)
slide=answer_slide()
body(slide,"For a sum, the absolute errors are added. They are not averaged.",XL,Y0,CW,0.48)
fml(slide,"l = 10 + 20 = 30 cm.",XL,Y0+0.58,CW,LH,18)
fml(slide,"Δl = 2 + 1 = 3 cm.",XL,Y0+0.58+LH+NG,CW,LH,18)
result_card(slide,"l = (30 ± 3) cm",XL,Y0+0.58+2*LH+2*NG,CW,0.50)

slide=new_slide("Repeated readings","Mean and mean absolute error",TEAL)
para(slide,"When the same quantity is measured several times, the mean is the best single value. Each reading then has its own absolute error, equal to how far it sits from the mean.",XL,Y0,CW,1.25,16)
para(slide,"The mean absolute error is the average of those gaps. The result is written as mean ± mean absolute error. Relative error is that mean absolute error divided by the mean.",XL,Y0+1.40,CW,1.25,16)
ncert(slide,"Result = mean  ±  mean absolute error",XL,Y0+2.85,CW,0.80,16)

slide=question_slide("Pendulum period")
qtx(slide,"Five readings of the time period, in seconds:",XL,Y0,CW,0.42,17)
qtx(slide,"2.63,  2.56,  2.42,  2.71,  2.80",XL,Y0+0.50,CW,0.45,20)
body(slide,"Find the mean, the mean absolute error, the relative error and the percentage error.",XL,Y0+1.15,CW,0.70,16)
slide=answer_slide("Solution: mean of the period")
body(slide,"Add the five readings, then divide by 5.",XL,Y0,CW,0.42)
fml(slide,"Sum = 2.63 + 2.56 + 2.42 + 2.71 + 2.80 = 13.12 s.",XL,Y0+0.52,CW,LH,16)
fml(slide,"Mean T = 13.12 / 5 = 2.624 s.",XL,Y0+0.52+LH+NG,CW,LH,17)
fml(slide,"We will report this as 2.62 s once the error is known to the hundredth.",XL,Y0+0.52+2*LH+2*NG,CW,0.60,15)
slide=answer_slide("Solution: errors in the period")
body(slide,"Gap of each reading from 2.624 s.",XL,Y0,CW,0.40)
fml(slide,"0.006,  0.064,  0.204,  0.086,  0.176.   Sum of gaps = 0.536 s.",XL,Y0+0.50,CW,0.65,15)
fml(slide,"Mean absolute error = 0.536 / 5 = 0.107 s ≈ 0.11 s.",XL,Y0+1.25,CW,LH,16)
fml(slide,"Relative error = 0.11 / 2.62 ≈ 0.04.   Percentage error ≈ 4%.",XL,Y0+1.25+LH+NG,CW,0.65,15)
result_card(slide,"T = (2.62 ± 0.11) s,    relative error 0.04,    percentage error 4%",XL,Y0+2.55,CW,0.70)

# DIMENSIONAL PRACTICE
slide=new_slide("Recall","Matching powers still works",BLUE)
para(slide,"Write the unknown as a product of powers. Put dimensions on both sides. Equate the powers of M, of L and of T separately. A pure number such as 2π or 6π never comes out of this method.",XL,Y0,CW,1.40,16)
para(slide,"If two quantities in the list have the same dimensions, they cannot both be used as separate base quantities. If a quantity is an angle, it is a pure number and cannot build a force or a time by itself.",XL,Y0+1.55,CW,1.30,16)
ncert(slide,"Same dimensions, or a pure angle, means the method stops.",XL,Y0+3.05,CW,0.80,16)

slide=question_slide("Electric field and magnetic field")
qtx(slide,"The relation between [E] and [B] is",XL,Y0,CW,0.45,18)
choices(slide,[
"[E] = [B] [L] [T]",
"[E] = [B] [L]^(-1) [T]",
"[E] = [B] [L] [T]^(-1)",
"[E] = [B] [L]^(-1) [T]^(-1)",
], XL, Y0+0.65, 10.8)
slide=answer_slide()
body(slide,"In an electromagnetic wave, E = c B. Speed c carries the extra dimensions.",XL,Y0,CW,0.50)
fml(slide,"[c] = L T^(-1), so [E] = [B] × L T^(-1).",XL,Y0+0.60,CW,LH,16)
fml(slide,"That is [E] = [B] [L] [T]^(-1).",XL,Y0+0.60+LH+NG,CW,LH,17)
result_card(slide,"Final answer: option 3",XL,Y0+0.60+2*LH+2*NG,CW,0.48)

slide=question_slide("Permittivity and permeability")
qtx(slide,"The relation between [ε_0] and [μ_0] is",XL,Y0,CW,0.45,18)
choices(slide,[
"[μ_0] = [ε_0] [L]^2 [T]^(-2)",
"[μ_0] = [ε_0] [L]^(-2) [T]^2",
"[μ_0] = [ε_0]^(-1) [L]^2 [T]^(-2)",
"[μ_0] = [ε_0]^(-1) [L]^(-2) [T]^2",
], XL, Y0+0.60, 11.0)
slide=answer_slide()
body(slide,"Use c = 1 / √(μ_0 ε_0). Square both sides.",XL,Y0,CW,0.45)
fml(slide,"μ_0 ε_0 = 1 / c^2, so [μ_0] [ε_0] = L^(-2) T^2.",XL,Y0+0.55,CW,LH,16)
fml(slide,"[μ_0] = [ε_0]^(-1) [L]^(-2) [T]^2.",XL,Y0+0.55+LH+NG,CW,LH,16)
result_card(slide,"Final answer: option 4",XL,Y0+0.55+2*LH+2*NG,CW,0.48)

slide=question_slide("Same dimensions as S")
qtx(slide,"S = (E × B) / μ_0. Its dimensions match which of these?",XL,Y0,CW,0.55,17)
choices(slide,[
"Energy / (charge × current)",
"Force / (length × time)",
"Energy / volume",
"Power / area",
], XL, Y0+0.75, 10.8)
slide=answer_slide()
body(slide,"Build [S] from the known formulas, then compare.",XL,Y0,CW,0.42)
fml(slide,"[E] = M L T^(-3) I^(-1),  [B] = M T^(-2) I^(-1),  [μ_0] = M L T^(-2) I^(-2).",XL,Y0+0.50,CW,0.65,14)
fml(slide,"[E B / μ_0] = M T^(-3).",XL,Y0+1.25,CW,LH,17)
fml(slide,"Force / (length × time) = M T^(-3).  Power / area = M T^(-3).",XL,Y0+1.25+LH+NG,CW,0.60,15)
result_card(slide,"Final answer: options 2 and 4",XL,Y0+2.50,CW,0.48)

slide=question_slide("Mass and angular momentum are pure numbers")
qtx(slide,"In a new system, mass and angular momentum are dimensionless, and length has dimension L. Which statements are correct?",XL,Y0,CW,0.85,16)
choices(slide,[
"Dimension of force is L^(-3)",
"Dimension of energy is L^(-2)",
"Dimension of power is L^(-5)",
"Dimension of linear momentum is L^(-1)",
], XL, Y0+1.05, 10.8)
slide=answer_slide()
body(slide,"Angular momentum m v r is a pure number, and m is a pure number, so velocity must cancel the length.",XL,Y0,CW,0.60)
fml(slide,"[v] = L^(-1).  Then [time] = length / speed = L^2.",XL,Y0+0.70,CW,LH,15)
fml(slide,"[force] = [mass × acceleration] = L^(-3).  [energy] = force × length = L^(-2).",XL,Y0+0.70+LH+NG,CW,0.60,14)
fml(slide,"[power] = energy / time = L^(-4), not L^(-5).  [momentum] = m v = L^(-1).",XL,Y0+1.45+LH+NG,CW,0.60,14)
result_card(slide,"Correct: options 1, 2 and 4.  Option 3 is wrong.",XL,Y0+2.70,CW,0.48)

slide=question_slide("Young’s modulus in F, A and V")
qtx(slide,"Speed V, area A and force F are the base quantities. [Young’s modulus] is",XL,Y0,CW,0.70,17)
choices(slide,["F A^(-1) V^0","F A^2 V^(-1)","F A^2 V^(-2)","F A^2 V^(-3)"], XL, Y0+0.90, 10.5)
slide=answer_slide()
body(slide,"Young’s modulus is stress / strain, and strain is a pure number.",XL,Y0,CW,0.48)
fml(slide,"Y = F / A.",XL,Y0+0.58,CW,LH,18)
fml(slide,"Speed is not needed, so its power is 0.",XL,Y0+0.58+LH+NG,CW,LH,16)
result_card(slide,"Final answer: F A^(-1) V^0   (option 1)",XL,Y0+0.58+2*LH+2*NG,CW,0.48)

slide=question_slide("Density from force, velocity and time")
qtx(slide,"F, V and T are base quantities. The dimensional formula of density is",XL,Y0,CW,0.65,17)
choices(slide,["F V^4 T^(-6)","F V^(-4) T^(-2)","F^2 V^(-2) T^6","F V^(-2) T^2"], XL, Y0+0.85, 10.5)
slide=answer_slide()
body(slide,"Assume density = F^a V^b T^c and match M, L, T.",XL,Y0,CW,0.45)
fml(slide,"Mass: a = 1.  Length: a + b = -3, so b = -4.",XL,Y0+0.55,CW,LH,16)
fml(slide,"Time: -2a - b + c = 0 → -2 - (-4) + c = 0 → c = -2.",XL,Y0+0.55+LH+NG,CW,LH,15)
result_card(slide,"Final answer: F V^(-4) T^(-2)   (option 2)",XL,Y0+0.55+2*LH+2*NG,CW,0.48)

slide=question_slide("Viscosity from momentum, area and time")
qtx(slide,"P, A and T are base quantities. The dimensional formula of viscosity is",XL,Y0,CW,0.65,17)
choices(slide,["[P A^(-1)]","[P A T^(-1)]","[P A^(-1) T]","[P A^(-1) T^(-1)]"], XL, Y0+0.85, 10.5)
slide=answer_slide()
body(slide,"From F = 6π η r v, the dimensions of η are M L^(-1) T^(-1).",XL,Y0,CW,0.50)
fml(slide,"η = P^a A^b T^c.  Mass gives a = 1.",XL,Y0+0.60,CW,LH,16)
fml(slide,"Length: 1 + 2b = -1, so b = -1.  Time: -1 + c = -1, so c = 0.",XL,Y0+0.60+LH+NG,CW,LH,15)
result_card(slide,"Final answer: [P A^(-1)]   (option 1)",XL,Y0+0.60+2*LH+2*NG,CW,0.48)

slide=question_slide("Mass from time, velocity and angular momentum")
qtx(slide,"Express mass using time t, velocity v and angular momentum J.",XL,Y0,CW,0.60,17)
slide=answer_slide()
body(slide,"[J] = M L^2 T^(-1). Assume M = J^a v^b t^c.",XL,Y0,CW,0.48)
fml(slide,"Mass: a = 1.  Length: 2 + b = 0, so b = -2.",XL,Y0+0.58,CW,LH,16)
fml(slide,"Time: -1 - (-2) + c = 0, so c = -1.",XL,Y0+0.58+LH+NG,CW,LH,16)
result_card(slide,"Mass = J / (v^2 t)",XL,Y0+0.58+2*LH+2*NG,CW,0.50)

slide=question_slide("Young’s modulus from c, h and G")
qtx(slide,"Young’s modulus is a product of powers of c, h and G. Which powers are correct?",XL,Y0,CW,0.55,17)
choices(slide,[
"α = 7, β = -1, γ = -2",
"α = -7, β = -1, γ = -2",
"α = 7, β = -1, γ = 2",
"α = -7, β = 1, γ = -2",
], XL, Y0+0.70, 10.5)
slide=answer_slide("Solution: equations for Y")
body(slide,"[Y] = M L^(-1) T^(-2),  [c] = L T^(-1),  [h] = M L^2 T^(-1),  [G] = M^(-1) L^3 T^(-2).",XL,Y0,CW,0.70)
fml(slide,"Mass: β - γ = 1.",XL,Y0+0.85,CW,LH,17)
fml(slide,"Length: α + 2β + 3γ = -1.",XL,Y0+0.85+LH+NG,CW,LH,17)
fml(slide,"Time: -α - β - 2γ = -2, so α + β + 2γ = 2.",XL,Y0+0.85+2*LH+2*NG,CW,LH,15)
slide=answer_slide("Solution: the three powers")
body(slide,"From mass, β = 1 + γ. Put that into the time equation.",XL,Y0,CW,0.48)
fml(slide,"α + (1 + γ) + 2γ = 2  →  α + 3γ = 1.",XL,Y0+0.58,CW,LH,16)
fml(slide,"Length then gives α + 5γ = -3. Subtract: 2γ = -4, so γ = -2.",XL,Y0+0.58+LH+NG,CW,0.60,15)
fml(slide,"α = 1 - 3(-2) = 7.  β = 1 + (-2) = -1.",XL,Y0+1.30+LH+NG,CW,LH,16)
result_card(slide,"α = 7, β = -1, γ = -2   (option 1)",XL,Y0+2.40,CW,0.48)

slide=question_slide("A leaking balloon")
qtx(slide,"Gas speed depends on the radius through a power a. The emptying time T depends on surface tension S, outlet area A, density ρ and starting radius R through powers α, β, γ and δ.",XL,Y0,CW,1.05,16)
choices(slide,[
"a = 1/2, α = 1/2, β = -1, γ = 1, δ = 3/2",
"a = -1/2, α = -1/2, β = -1, γ = -1/2, δ = 5/2",
"a = -1/2, α = -1/2, β = -1, γ = 1/2, δ = 7/2",
"a = 1/2, α = 1/2, β = -1/2, γ = 1/2, δ = 7/2",
], XL, Y0+1.20, 12.0)
slide=answer_slide("Solution: speed of the gas")
body(slide,"The speed of the leaving gas is fixed by S, ρ and the present radius r.",XL,Y0,CW,0.50)
fml(slide,"v ∝ S^(1/2) ρ^(-1/2) r^(-1/2).",XL,Y0+0.60,CW,LH,16)
fml(slide,"Mass cancels. What remains is a speed, and the power of r is -1/2. So a = -1/2.",XL,Y0+0.60+LH+NG,CW,0.70,15)
slide=answer_slide("Solution: time to empty")
body(slide,"Volume leaves through area A, so the radius falls as dr/dt ∝ A v / r^2.",XL,Y0,CW,0.55)
fml(slide,"With v ∝ r^(-1/2), we get dr/dt ∝ A S^(1/2) ρ^(-1/2) r^(-5/2).",XL,Y0+0.65,CW,0.65,14)
fml(slide,"Time to fall from R to 0 is T ∝ R^(7/2) / (A S^(1/2) ρ^(-1/2)).",XL,Y0+1.40,CW,0.60,15)
fml(slide,"So α = -1/2, β = -1, γ = 1/2, δ = 7/2.",XL,Y0+2.10,CW,LH,16)
result_card(slide,"Final answer: option 3",XL,Y0+2.65,CW,0.45)

slide=new_slide("Stop","Three cases where dimensions refuse",ORANGE)
bullet(slide,"Torque, energy and work have the same dimensions. Using all three as base quantities cannot produce a new formula for time.",XL,Y0,CW,0.85,16)
bullet(slide,"If g is said to depend on force, energy, mass and velocity, that is four quantities and only three equations (M, L, T). The powers are not unique.",XL,Y0+1.05,CW,0.90,16)
bullet(slide,"If force is said to depend on an angle, a velocity and an energy, the angle is a pure number. It cannot supply the dimensions of force.",XL,Y0+2.15,CW,0.90,16)
trap_card(slide,"Write “not possible” and name the reason. That is a full answer.",XL,Y0+3.20,CW,0.48)

slide=question_slide("Which equation cannot be fully derived")
qtx(slide,"Which equation cannot be obtained from dimensions alone?",XL,Y0,CW,0.50,17)
choices(slide,[
"F = 6π η r v",
"θ = ω t",
"P = ρ g h",
"v = d / t",
], XL, Y0+0.70, 10.5)
slide=answer_slide()
body(slide,"Dimensions can give the powers. They cannot give a pure number such as 6π.",XL,Y0,CW,0.55)
fml(slide,"F, η, r and v can be matched, but 6π stays unknown.",XL,Y0+0.65,CW,LH,16)
fml(slide,"θ = ω t, P = ρ g h and v = d / t are products of powers. Their number in front is 1.",XL,Y0+0.65+LH+NG,CW,0.70,15)
result_card(slide,"Final answer: F = 6π η r v   (option 1)",XL,Y0+2.15,CW,0.50)

# CLOSE
slide=new_slide("Class summary","Units 04 lock lines",TEAL)
fml(slide,"Add or subtract → least decimal places. Multiply or divide → least significant figures.",XL,Y0,CW,0.70,15)
fml(slide,"A dropped 5 rounds the last kept digit to the even choice.",XL,Y0+0.85,CW,LH,16)
fml(slide,"Accuracy is closeness to the true value. Precision is how fine the reading is.",XL,Y0+0.85+LH+NG,CW,0.65,15)
fml(slide,"Relative error = absolute error / true value. For a sum, add the absolute errors.",XL,Y0+1.65+LH+NG,CW,0.65,15)
result_card(slide,"Repeated readings: report mean ± mean absolute error.",XL,Y0+2.50+LH,CW,0.50)

slide=new_slide("CBSE write","Ready theory lines",GREEN)
body(slide,"1. Significant figures are the certain digits plus one uncertain digit.",XL,Y0,CW,0.50)
body(slide,"2. Changing the unit does not change the number of significant figures.",XL,Y0+0.55,CW,0.50)
body(slide,"3. In addition and subtraction the result has the least number of decimal places.",XL,Y0+1.10,CW,0.55)
body(slide,"4. In multiplication and division the result has the least number of significant figures.",XL,Y0+1.70,CW,0.55)
body(slide,"5. Absolute error has a unit. Relative error and percentage error do not add a new kind of unit.",XL,Y0+2.30,CW,0.60)
trap_card(slide,"Percentage error = (absolute error / true value) × 100.",XL,Y0+3.05,CW,0.48)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Round 2.745 and 2.735 to three significant figures.",XL,Y0,CW,0.48)
body(slide,"2. A = 12.5 cm, B = 0.30 cm. Find A × B and A − B with the proper digits.",XL,Y0+0.55,CW,0.55)
body(slide,"3. True value 9.80. Readings 9.79, 9.85, 9.77. Which is most accurate?",XL,Y0+1.20,CW,0.55)
body(slide,"4. x = (5.0 ± 0.2) and y = (3.0 ± 0.1). Find x + y with the error.",XL,Y0+1.85,CW,0.55)
body(slide,"5. If resistance R depends on V and I only, show why a pure number in R = k V/I cannot be found by dimensions.",XL,Y0+2.50,CW,0.70)
trap_card(slide,"Next class: how errors combine in a product, a quotient and a power.",XL,Y0+3.30,CW,0.48)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
