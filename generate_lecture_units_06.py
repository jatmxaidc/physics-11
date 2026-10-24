"""Units 06 — measuring instruments: metre scale, vernier, screw gauge (PW 08)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Units_06_Measuring_Instruments.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


_GREEK_SUP = {"α": "ᵅ", "β": "ᵝ", "γ": "ᵞ", "δ": "ᵟ", "ε": "ᵋ", "θ": "ᶿ"}


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(-1/2)", "⁻½").replace("^(-3/2)", "⁻³⁄²")
    text = text.replace("^(-5/2)", "⁻⁵⁄²").replace("^(-7/2)", "⁻⁷⁄²")
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³⁄²")
    text = text.replace("^(5/2)", "⁵⁄²").replace("^(7/2)", "⁷⁄²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([αβγδεθ])", lambda m: _GREEK_SUP[m.group(1)], text)
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    step = 0.52
    for i, item in enumerate(items):
        cy = y + i * step
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.06), Inches(0.32), Inches(0.32))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.08, 0.32, 0.28, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.46, cy + 0.04, w - 0.46, 0.42, 17, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units 06",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units and Measurement",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Easy Instruments: Scale, Vernier, Screw Gauge",bx,by+1.62,bw,0.40,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"19 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"From PW Units and Measurement Lecture 08. Final Units lecture."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"What we learn today",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Metre scale — the simple ruler",
"Vernier callipers — how to find least count",
"How to read a vernier",
"Zero error — when the instrument is wrong at zero",
"Screw gauge — for very small lengths",
]
for i,t in enumerate(topics):
    y=1.22+i*0.72
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(XL),Inches(y),Inches(0.42),Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.fill.background()
    tx(slide,f"{i+1:02d}",XL,y+0.05,0.42,0.32,14,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
    tx(slide,t,XL+0.60,y+0.05,10.5,0.35,20,TEXT,False,spa=0,scaled=False)
_footer(slide,2); after_content_slide()

slide=new_slide("Class note","Link from last class",TEAL)
para(slide,"Last class we learned how error grows when we add, multiply or raise a number to a power. That part is done.",XL,Y0,CW,1.15,16)
para(slide,"Today we learn the tools that measure length: a metre scale, vernier callipers and a screw gauge. The key word is least count — the smallest length the tool can measure.",XL,Y0+1.35,CW,1.40,16)
ncert(slide,"Least count = smallest length the instrument can show",XL,Y0+3.00,CW,0.90,16)
trap_card(slide,"Class notes come after teaching. These slides are for the board.",XL,Y0+4.15,CW,0.45)

slide=new_slide("Words","Short names we will use",GREEN)
bullet(slide,"L.C. = least count (smallest step the tool can show)",XL,Y0,CW,0.60,17)
bullet(slide,"M.S.D. = one small mark gap on the main scale",XL,Y0+0.70,CW,0.60,17)
bullet(slide,"V.S.D. = one small mark gap on the vernier scale",XL,Y0+1.40,CW,0.60,17)
bullet(slide,"M.S.R. = main scale reading   |   V.S.R. = vernier scale reading",XL,Y0+2.10,CW,0.70,17)
ncert(slide,"Learn these four short names. Questions use them a lot.",XL,Y0+3.00,CW,0.90,16)

# METRE SCALE
slide=new_slide("Idea","Three tools for length",GREEN)
bullet(slide,"Metre scale — a normal ruler with millimetre marks.",XL,Y0,CW,0.55,17)
bullet(slide,"Vernier callipers — a ruler plus a short sliding scale.",XL,Y0+0.70,CW,0.60,17)
bullet(slide,"Screw gauge — a screw tool for very thin objects.",XL,Y0+1.40,CW,0.60,17)
ncert(slide,"All three measure length. Least count tells how fine the reading can be.",XL,Y0+2.25,CW,0.95,16)

slide=new_slide("Rule","Metre scale",TEAL)
para(slide,"On a school metre scale, marks are 1 mm apart. So the smallest length you can read is 1 mm. That 1 mm is the least count.",XL,Y0,CW,1.25,16)
fml(slide,"Least count of metre scale = 1 mm",XL,Y0+1.40,CW,LH,18)
fml(slide,"Write the answer as (reading ± 1 mm)",XL,Y0+1.40+LH+NG,CW,LH,16)
trap_card(slide,"If a rod ends between 4 mm and 5 mm, write 4 mm ± 1 mm. Do not guess 4.2 mm.",XL,Y0+2.90,CW,0.70)

# VERNIER THEORY
slide=new_slide("Idea","Why we need a vernier",BLUE)
para(slide,"A metre scale cannot show lengths smaller than 1 mm. Vernier callipers fix that. They have a fixed main scale and a short scale that slides. That short scale is the vernier.",XL,Y0,CW,1.40,16)
para(slide,"The vernier marks are a little closer than the main marks. So one vernier mark can line up exactly with a main mark. That helps us read a small part of 1 mm.",XL,Y0+1.60,CW,1.25,16)
ncert(slide,"L.C. = 1 M.S.D. − 1 V.S.D.",XL,Y0+3.10,CW,0.80,18)

slide=new_slide("Rule","Easy school vernier",TEAL)
para(slide,"In school, 10 vernier marks match 9 main marks. Each main mark gap is 1 mm. So one vernier gap is a bit less than 1 mm.",XL,Y0,CW,1.20,16)
fml(slide,"10 V.S.D. = 9 mm  →  1 V.S.D. = 0.9 mm",XL,Y0+1.35,CW,LH,17)
fml(slide,"L.C. = 1 mm − 0.9 mm = 0.1 mm",XL,Y0+1.35+LH+NG,CW,LH,18)
ncert(slide,"If n vernier marks match (n − 1) main marks: L.C. = (1 M.S.D.) / n",XL,Y0+2.90,CW,1.00,15)

slide=new_slide("Method","Only two jobs",ORANGE)
bullet(slide,"Job 1: Find the least count from how the two scales match.",XL,Y0,CW,0.65,17)
bullet(slide,"Job 2: Find the reading = main reading + (vernier mark × least count).",XL,Y0+0.80,CW,0.80,17)
fml(slide,"Reading = M.S.R. + (V.S.R. × L.C.)",XL,Y0+1.85,CW,LH,18)
ncert(slide,"Write the final answer as (reading ± least count).",XL,Y0+2.80,CW,0.85,16)

slide=question_slide("Least count: 9 and 12")
qtx(slide,"One main mark gap is 1 cm. 9 main gaps match 12 vernier gaps. Find the least count.",XL,Y0,CW,0.75,17)
slide=answer_slide()
body(slide,"Idea: find one vernier gap, then subtract it from one main gap.",XL,Y0,CW,0.55)
fml(slide,"1 V.S.D. = (9/12) × 1 cm = 0.75 cm",XL,Y0+0.70,CW,LH,16)
fml(slide,"L.C. = 1 − 0.75 = 0.25 cm",XL,Y0+0.70+LH+NG,CW,LH,17)
result_card(slide,"Final answer: 0.25 cm",XL,Y0+0.70+2*LH+2*NG,CW,0.48)

slide=question_slide("25 vernier marks")
qtx(slide,"25 vernier marks match 24 main marks. One main mark gap is 1 mm. Least count is",XL,Y0,CW,0.75,17)
choices(slide,["0.04 mm","0.01 mm","0.02 mm","0.08 mm"], XL, Y0+0.95, 10.5)
slide=answer_slide()
body(slide,"Idea: use L.C. = 1 M.S.D. − 1 V.S.D.",XL,Y0,CW,0.45)
fml(slide,"1 V.S.D. = 24/25 mm = 0.96 mm",XL,Y0+0.60,CW,LH,17)
fml(slide,"L.C. = 1 − 0.96 = 0.04 mm",XL,Y0+0.60+LH+NG,CW,LH,17)
result_card(slide,"Final answer: 0.04 mm   (option 1)",XL,Y0+0.60+2*LH+2*NG,CW,0.48)

slide=question_slide("Formula with x and n")
qtx(slide,"One main mark gap is x cm. n vernier marks match (n − 1) main marks. Least count in cm is",XL,Y0,CW,0.80,16)
choices(slide,[
"(n − 1)x / n",
"n x / (n − 1)",
"x / n",
"x / (n − 1)",
], XL, Y0+1.00, 10.8)
slide=answer_slide()
body(slide,"Idea: this is the same school rule, written with letters.",XL,Y0,CW,0.50)
fml(slide,"1 V.S.D. = ((n − 1)/n) × x",XL,Y0+0.65,CW,LH,16)
fml(slide,"L.C. = x − ((n − 1)/n)x = x/n",XL,Y0+0.65+LH+NG,CW,LH,17)
result_card(slide,"Final answer: x / n   (option 3)",XL,Y0+0.65+2*LH+2*NG,CW,0.48)

slide=question_slide("20 vernier and 8 main")
qtx(slide,"1 cm on the main scale is cut into 10 equal parts. 20 vernier marks match 8 of those small parts. Find least count.",XL,Y0,CW,0.90,16)
slide=answer_slide()
body(slide,"Idea: find 1 M.S.D., then 1 V.S.D., then subtract.",XL,Y0,CW,0.50)
fml(slide,"1 M.S.D. = 1/10 cm = 0.1 cm",XL,Y0+0.65,CW,LH,17)
fml(slide,"1 V.S.D. = (8/20) × 0.1 = 0.04 cm",XL,Y0+0.65+LH+NG,CW,LH,16)
fml(slide,"L.C. = 0.1 − 0.04 = 0.06 cm",XL,Y0+0.65+2*LH+2*NG,CW,LH,17)
result_card(slide,"Final answer: 0.06 cm",XL,Y0+2.75,CW,0.48)

# READING
slide=new_slide("Method","How to read a vernier",GREEN)
para(slide,"Look at the zero mark of the vernier. The main mark just before it is the main scale reading. Then look for the vernier mark that lines up with any main mark. That number is the vernier reading.",XL,Y0,CW,1.40,16)
fml(slide,"Reading = M.S.R. + V.S.R. × L.C.",XL,Y0+1.55,CW,LH,18)
fml(slide,"Example: 3 mm + 4 × 0.1 mm = 3.4 mm",XL,Y0+1.55+LH+NG,CW,LH,16)
ncert(slide,"Write it as (3.4 ± 0.1) mm",XL,Y0+3.00,CW,0.75,17)

# ZERO ERROR
slide=new_slide("Idea","What is zero error?",ORANGE)
para(slide,"Close the jaws with nothing inside. The two zero marks should meet. If they do not meet, the tool already shows a small wrong reading. That wrong reading is called zero error.",XL,Y0,CW,1.40,16)
para(slide,"Always subtract the zero error from the reading you took. Keep the + or − sign. Zero correction is just the opposite of zero error.",XL,Y0+1.60,CW,1.25,16)
ncert(slide,"True reading = reading you took − zero error",XL,Y0+3.10,CW,0.90,17)

slide=new_slide("Rule","Positive zero error",TEAL)
para(slide,"Jaws closed. If vernier zero is to the right of main zero, zero error is positive. Size = (matching vernier mark) × least count.",XL,Y0,CW,1.25,16)
fml(slide,"Zero error = + (matching vernier mark) × L.C.",XL,Y0+1.40,CW,LH,16)
fml(slide,"Example: 4th mark matches, L.C. = 0.1 mm → zero error = +0.4 mm",XL,Y0+1.40+LH+NG,CW,0.70,15)
trap_card(slide,"Zero correction here is −0.4 mm. Error and correction have opposite signs.",XL,Y0+3.05,CW,0.65)

slide=new_slide("Rule","Negative zero error",BLUE)
para(slide,"Jaws closed. If vernier zero is to the left of main zero, zero error is negative. For a 10-mark vernier, use (10 − matching mark) × least count, and put a minus sign.",XL,Y0,CW,1.35,16)
fml(slide,"Zero error = − (N − matching mark) × L.C.",XL,Y0+1.50,CW,LH,16)
fml(slide,"Example: N = 10, 6th mark matches, L.C. = 0.1 mm → zero error = −0.4 mm",XL,Y0+1.50+LH+NG,CW,0.70,14)
ncert(slide,"Minus a negative error means you add. True = reading + 0.4 mm",XL,Y0+3.10,CW,0.90,15)

slide=question_slide("Find zero correction")
qtx(slide,"Jaws closed. Vernier zero is to the right of main zero. 5th vernier mark matches. L.C. = 0.1 mm. Zero correction is",XL,Y0,CW,0.90,16)
choices(slide,["0.4 mm","+0.5 mm","−0.5 mm","−0.4 mm"], XL, Y0+1.10, 10.5)
slide=answer_slide()
body(slide,"Idea: right side means + error. Correction flips the sign.",XL,Y0,CW,0.55)
fml(slide,"Zero error = +5 × 0.1 = +0.5 mm",XL,Y0+0.70,CW,LH,17)
fml(slide,"Zero correction = −0.5 mm",XL,Y0+0.70+LH+NG,CW,LH,17)
result_card(slide,"Final answer: −0.5 mm   (option 3)",XL,Y0+0.70+2*LH+2*NG,CW,0.48)

slide=question_slide("Fix a vernier reading")
qtx(slide,"Zero error = +0.2 mm. Vernier zero is between 8.5 cm and 8.6 cm. Matching vernier mark = 6. L.C. = 0.01 cm. Correct length is",XL,Y0,CW,0.95,16)
choices(slide,["8.58 cm","8.54 cm","8.56 cm","8.36 cm"], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Idea: use one unit. Change 0.2 mm to 0.02 cm.",XL,Y0,CW,0.50)
fml(slide,"Reading taken = 8.5 + 6 × 0.01 = 8.56 cm",XL,Y0+0.65,CW,LH,16)
fml(slide,"True length = 8.56 − 0.02 = 8.54 cm",XL,Y0+0.65+LH+NG,CW,LH,16)
result_card(slide,"Final answer: 8.54 cm   (option 2)",XL,Y0+0.65+2*LH+2*NG,CW,0.48)

slide=question_slide("Sphere with zero error")
qtx(slide,"1 M.S.D. = 0.1 cm. 10 vernier marks match 9 main marks. Jaws closed: 6th vernier mark matches (vernier zero is left). With sphere: M.S.R. = 3.1 cm, 1st vernier mark matches. Diameter is",XL,Y0,CW,1.15,15)
choices(slide,["3.07 cm","3.11 cm","3.15 cm","3.17 cm"], XL, Y0+1.35, 10.5)
slide=answer_slide()
body(slide,"Idea: least count → zero error → sphere reading → true value.",XL,Y0,CW,0.55)
fml(slide,"L.C. = 0.1 / 10 = 0.01 cm",XL,Y0+0.70,CW,LH,17)
fml(slide,"Zero error = −(10 − 6) × 0.01 = −0.04 cm",XL,Y0+0.70+LH+NG,CW,LH,15)
fml(slide,"Reading taken = 3.1 + 1 × 0.01 = 3.11 cm",XL,Y0+0.70+2*LH+2*NG,CW,LH,15)
fml(slide,"True = 3.11 − (−0.04) = 3.15 cm",XL,Y0+0.70+3*LH+3*NG,CW,LH,16)
result_card(slide,"Final answer: 3.15 cm   (option 3)",XL,Y0+3.55,CW,0.48)

slide=question_slide("Diameter of a cylinder")
qtx(slide,"No zero error. Vernier zero is between 5.10 cm and 5.15 cm. 50 vernier marks = 2.45 cm. 24th vernier mark matches. Diameter is",XL,Y0,CW,0.95,16)
choices(slide,["5.112 cm","5.124 cm","5.136 cm","5.148 cm"], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Idea: gap 5.10 to 5.15 means 1 M.S.D. = 0.05 cm.",XL,Y0,CW,0.55)
fml(slide,"1 V.S.D. = 2.45 / 50 = 0.049 cm",XL,Y0+0.70,CW,LH,16)
fml(slide,"L.C. = 0.05 − 0.049 = 0.001 cm",XL,Y0+0.70+LH+NG,CW,LH,16)
fml(slide,"Diameter = 5.10 + 24 × 0.001 = 5.124 cm",XL,Y0+0.70+2*LH+2*NG,CW,LH,16)
result_card(slide,"Final answer: 5.124 cm   (option 2)",XL,Y0+2.85,CW,0.48)

slide=question_slide("20 vernier match 16 main")
qtx(slide,"Main scale marks are 1 mm apart. 20 vernier marks match 16 main marks. Least count is",XL,Y0,CW,0.75,17)
choices(slide,["0.02 mm","0.05 mm","0.1 mm","0.2 mm"], XL, Y0+0.95, 10.5)
slide=answer_slide()
body(slide,"Idea: same subtract rule again.",XL,Y0,CW,0.45)
fml(slide,"1 V.S.D. = 16/20 = 0.8 mm",XL,Y0+0.60,CW,LH,17)
fml(slide,"L.C. = 1 − 0.8 = 0.2 mm",XL,Y0+0.60+LH+NG,CW,LH,18)
result_card(slide,"Final answer: 0.2 mm   (option 4)",XL,Y0+0.60+2*LH+2*NG,CW,0.48)

slide=question_slide("n marks in one centimetre")
qtx(slide,"Main scale has n marks in 1 cm. n vernier marks match (n − 1) main marks. Least count is",XL,Y0,CW,0.80,16)
choices(slide,[
"1 / ((n + 1)(n − 1)) cm",
"1 / n cm",
"1 / n^2 cm",
"1 / (n(n + 1)) cm",
], XL, Y0+1.00, 11.0)
slide=answer_slide()
body(slide,"Idea: one main gap is already 1/n cm.",XL,Y0,CW,0.50)
fml(slide,"1 M.S.D. = 1/n cm",XL,Y0+0.65,CW,LH,17)
fml(slide,"L.C. = (1 M.S.D.) / n = 1/n^2 cm",XL,Y0+0.65+LH+NG,CW,LH,17)
result_card(slide,"Final answer: 1 / n^2 cm   (option 3)",XL,Y0+0.65+2*LH+2*NG,CW,0.48)

slide=question_slide("Fill the blank")
qtx(slide,"Vernier constant = 0.1 mm. Zero error = (−0.05) cm. M.S.R. = 1.7 cm. Matching vernier mark = 5. Correct diameter = ____ × 10^(−2) cm.",XL,Y0,CW,1.00,16)
slide=answer_slide()
body(slide,"Idea: change 0.1 mm to cm first, then find true length.",XL,Y0,CW,0.55)
fml(slide,"L.C. = 0.1 mm = 0.01 cm",XL,Y0+0.70,CW,LH,17)
fml(slide,"Reading taken = 1.7 + 5 × 0.01 = 1.75 cm",XL,Y0+0.70+LH+NG,CW,LH,15)
fml(slide,"True = 1.75 − (−0.05) = 1.80 cm = 180 × 10^(−2) cm",XL,Y0+0.70+2*LH+2*NG,CW,0.65,15)
result_card(slide,"Final answer: 180",XL,Y0+2.90,CW,0.48)

slide=question_slide("Least count in mm")
qtx(slide,"One main mark gap is a cm. The nth vernier mark matches the (n − 1)th main mark. Least count in mm is",XL,Y0,CW,0.85,16)
choices(slide,[
"10 n a / (n − 1)",
"10 a / (n − 1)",
"10 a / n",
"(n − 1)a / (10 n)",
], XL, Y0+1.05, 10.8)
slide=answer_slide()
body(slide,"Idea: in cm, L.C. = a/n. Times 10 gives mm.",XL,Y0,CW,0.50)
fml(slide,"L.C. = a/n cm = 10a/n mm",XL,Y0+0.65,CW,LH,18)
result_card(slide,"Final answer: 10 a / n   (option 3)",XL,Y0+0.65+LH+NG,CW,0.48)

# SCREW GAUGE
slide=new_slide("Idea","Screw gauge in plain words",GREEN)
para(slide,"A screw gauge is used for very small lengths, like the diameter of a thin wire. When you turn the screw once, it moves forward by a fixed length. That fixed length is called pitch.",XL,Y0,CW,1.40,16)
fml(slide,"Pitch = how far the screw moves in one full turn",XL,Y0+1.55,CW,LH,16)
fml(slide,"L.C. = Pitch / (number of marks on the round scale)",XL,Y0+1.55+LH+NG,CW,LH,15)
ncert(slide,"Reading = linear reading + (round-scale mark × L.C.)",XL,Y0+3.05,CW,0.95,15)

slide=new_slide("Rule","Zero error of screw gauge",TEAL)
para(slide,"Close the ends with nothing between them. The zero of the round scale should meet the line. If it does not, there is a zero error — just like in the vernier.",XL,Y0,CW,1.30,16)
fml(slide,"True reading = reading you took − zero error",XL,Y0+1.50,CW,LH,17)
trap_card(slide,"Screw gauge least count is smaller than vernier. That is why we use it for thin wires.",XL,Y0+2.45,CW,0.70)

slide=question_slide("Find the pitch")
qtx(slide,"A screw gauge has least count 0.01 mm and 50 marks on the round scale. Pitch is",XL,Y0,CW,0.75,17)
choices(slide,["0.01 mm","0.25 mm","0.5 mm","1.0 mm"], XL, Y0+0.95, 10.5)
slide=answer_slide()
body(slide,"Idea: Pitch = least count × number of marks.",XL,Y0,CW,0.50)
fml(slide,"Pitch = 0.01 × 50 = 0.5 mm",XL,Y0+0.65,CW,LH,18)
result_card(slide,"Final answer: 0.5 mm   (option 3)",XL,Y0+0.65+LH+NG,CW,0.48)

slide=question_slide("Wire diameter")
qtx(slide,"Pitch = 0.5 mm. Round scale has 50 marks. Linear reading = 2.5 mm. Round reading = 20. Zero error = +0.02 mm. True diameter is",XL,Y0,CW,0.90,16)
slide=answer_slide()
body(slide,"Idea: least count → reading taken → subtract zero error.",XL,Y0,CW,0.55)
fml(slide,"L.C. = 0.5 / 50 = 0.01 mm",XL,Y0+0.70,CW,LH,17)
fml(slide,"Reading taken = 2.5 + 20 × 0.01 = 2.70 mm",XL,Y0+0.70+LH+NG,CW,LH,15)
fml(slide,"True = 2.70 − 0.02 = 2.68 mm",XL,Y0+0.70+2*LH+2*NG,CW,LH,17)
result_card(slide,"Final answer: 2.68 mm",XL,Y0+2.80,CW,0.48)

# CLOSE
slide=new_slide("Class summary","Remember these lines",TEAL)
fml(slide,"Metre scale least count = 1 mm",XL,Y0,CW,LH,17)
fml(slide,"Vernier: L.C. = 1 M.S.D. − 1 V.S.D.   (or 1 M.S.D. / n)",XL,Y0+LH+NG,CW,0.65,16)
fml(slide,"Reading = M.S.R. + V.S.R. × L.C.",XL,Y0+2*LH+2*NG,CW,LH,17)
fml(slide,"True reading = reading taken − zero error",XL,Y0+2.40,CW,LH,17)
ncert(slide,"Screw gauge: L.C. = Pitch / N. Same zero-error rule.",XL,Y0+3.25,CW,0.85,16)

slide=new_slide("For the exam","Four lines to write",GREEN)
body(slide,"1. Vernier least count = 1 M.S.D. − 1 V.S.D.",XL,Y0,CW,0.50)
body(slide,"2. Reading taken = M.S.R. + (matching vernier mark × L.C.)",XL,Y0+0.70,CW,0.60)
body(slide,"3. True reading = reading taken − zero error",XL,Y0+1.40,CW,0.55)
body(slide,"4. Screw gauge least count = Pitch / (round marks)",XL,Y0+2.10,CW,0.55)
ncert(slide,"These four lines are enough for most board questions.",XL,Y0+2.90,CW,0.85,16)

slide=new_slide("Practice","Try these at home",ORANGE)
body(slide,"1. 50 vernier marks match 49 main marks. 1 M.S.D. = 1 mm. Find least count.",XL,Y0,CW,0.65)
body(slide,"2. M.S.R. = 4.2 cm, V.S.R. = 7, L.C. = 0.01 cm, zero error = +0.03 cm. Find true length.",XL,Y0+0.80,CW,0.70)
body(slide,"3. Pitch = 1 mm, round scale has 100 marks. Find least count in mm and in cm.",XL,Y0+1.65,CW,0.70)
body(slide,"4. Ends closed. Round reading = −3, L.C. = 0.01 mm. Write zero error and zero correction.",XL,Y0+2.50,CW,0.75)
trap_card(slide,"Units chapter is finished. Next chapter: Motion in a Straight Line.",XL,Y0+3.50,CW,0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
