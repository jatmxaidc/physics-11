"""Lecture 11 — higher derivatives, chain rule and partial differentiation."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_11_Higher_Derivatives_Chain_Rule_and_Partial_Differentiation.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)


# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 11",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Higher Derivatives, Chain Rule and Partial Differentiation",bx,by+1.62,bw,0.40,16,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"16 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Short recap, then continue with higher derivatives and the chain rule."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["Product and quotient rule continuation","5-minute differentiation recap","Higher-order derivatives","Chain rule: function inside a function","S.H.M. application","Partial differentiation","Parametric and reciprocal derivatives"]
for i,item in enumerate(topics):
    yy=1.18+i*0.78
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.70,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Keep recap to five minutes."); after_content_slide()

# UNFINISHED LECTURE 10 CONTENT — CONTINUED FIRST
slide=new_slide("Rule 6","Product and quotient rules",TEAL)
tx(slide,"PRODUCT",XL,Y0,5.75,0.35,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
ncert(slide,"d(AB)/dx = A dB/dx + B dA/dx",XL,Y0+0.45,5.75,1.08,18)
body(slide,"Differentiate one factor at a time.",XL,Y0+1.75,5.75,0.50)
tx(slide,"QUOTIENT",6.95,Y0,5.75,0.35,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
ncert(slide,"d(A/B)/dx = (B dA/dx - A dB/dx)/B²",6.95,Y0+0.45,5.75,1.08,17)
body(slide,"Keep the denominator squared.",6.95,Y0+1.75,5.75,0.50)
trap_card(slide,"For a quotient, remember: denominator × derivative of numerator minus numerator × derivative of denominator.",XL,Y0+2.62,CW,0.78)

slide=question_slide("Differentiate a simple polynomial")
qtx(slide,"For y=4x³-5x²+2x-7, find dy/dx.",XL,Y0,CW,0.72,22)
slide=answer_slide()
fml(slide,"d(4x³)/dx = 12x²",XL,Y0,CW,LH,19)
fml(slide,"d(-5x²)/dx = -10x",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"d(2x)/dx = 2 and d(-7)/dx = 0",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = 12x²-10x+2.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Product-rule practice")
qtx(slide,"Differentiate y=tan x · logₑx.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Use d(AB)/dx = A dB/dx + B dA/dx",XL,Y0,CW,LH,18)
fml(slide,"d(tan x)/dx=sec²x and d(logₑx)/dx=1/x",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"dy/dx = tan x/x + logₑx·sec²x.",XL,Y0+2*LH+2*NG,CW,0.60)

slide=question_slide("Stationary-point practice")
qtx(slide,"If y=A/x³-B/x⁴, find x for which dy/dx=0. Assume A and B are non-zero constants.",XL,Y0,CW,1.00,21)
slide=answer_slide()
fml(slide,"Write y=Ax⁻³-Bx⁻⁴",XL,Y0,CW,LH,19)
fml(slide,"dy/dx=-3Ax⁻⁴+4Bx⁻⁵",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"Set zero and multiply by x⁵: -3Ax+4B=0",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"x = 4B/(3A).",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quotient-rule practice: exponential")
qtx(slide,"Differentiate y=e^x/x², where x≠0.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Take A=e^x and B=x²",XL,Y0,CW,LH,19)
fml(slide,"dy/dx = [x²e^x-e^x(2x)]/x⁴",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"Factor e^x x from the numerator",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"dy/dx = e^x(x-2)/x³.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quotient-rule practice: algebraic")
qtx(slide,"Differentiate y=(x²-1)/(x²+1).",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Use the quotient rule with A=x²-1 and B=x²+1",XL,Y0,CW,LH,18)
fml(slide,"dy/dx = [(x²+1)(2x)-(x²-1)(2x)]/(x²+1)²",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"The numerator becomes 2x[(x²+1)-(x²-1)]=4x",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"dy/dx = 4x/(x²+1)².",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Mixed trigonometric-algebraic practice")
qtx(slide,"Differentiate y=cos x+2/x³, where x≠0.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Write 2/x³ as 2x⁻³",XL,Y0,CW,LH,19)
fml(slide,"d(cos x)/dx=-sin x",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"d(2x⁻³)/dx=-6x⁻⁴",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = -sin x-6/x⁴.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Simplify before differentiating")
qtx(slide,"Differentiate y=(sqrt(x)+1/sqrt(x))² for x>0.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"Expand first: y=x+2+1/x",XL,Y0,CW,LH,19)
fml(slide,"Write 1/x as x⁻¹",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"dy/dx=1+0-x⁻²",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = 1-1/x².",XL,Y0+3*LH+3*NG,CW,0.55)

# FIVE-MINUTE RECAP
slide=new_slide("5-minute recap","Differentiation formula map",BLUE)
grid_table(slide,["Function","Derivative"],[
    ["C","0"],
    ["xⁿ","nxⁿ⁻¹"],
    ["sin x","cos x"],
    ["cos x","-sin x"],
    ["eˣ","eˣ"],
    ["logₑx","1/x"],
],1.80,Y0,[4.35,5.35],0.62,17,18)
body(slide,"Use radians for trigonometric derivatives.",XL,5.82,CW,0.45)

slide=new_slide("5-minute recap","Rules used in the last class",BLUE)
fml(slide,"constant multiple: d[kf]/dx = kf'",XL,Y0,CW,LH,18)
fml(slide,"sum or difference: differentiate each term separately",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"product: d(AB)/dx = A B' + B A'",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"quotient: d(A/B)/dx = (B A' - A B')/B²",XL,Y0+3*LH+3*NG,CW,LH,18)
trap_card(slide,"Before differentiating, simplify powers and fractions whenever possible.",XL,Y0+4*LH+4*NG,CW,0.60)

slide=question_slide("Recap check: slope at a point")
qtx(slide,"For y=x³-5x+x², find dy/dx at x=1.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"dy/dx = 3x²-5+2x",XL,Y0,CW,LH,20)
fml(slide,"At x=1: dy/dx = 3-5+2",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Slope at x=1 is 0.",XL,Y0+2*LH+2*NG,CW,0.55)

# HIGHER DERIVATIVES
slide=new_slide("New topic","Higher-order derivatives",GREEN)
body(slide,"Differentiate y once to get the first derivative.",XL,Y0,CW,0.52)
ncert(slide,"first derivative: dy/dx",XL,Y0+0.70,CW,0.86,23)
body(slide,"Differentiate the first derivative again to get the second derivative.",XL,Y0+1.78,CW,0.58)
ncert(slide,"second derivative: d²y/dx² = d/dx(dy/dx)",XL,Y0+2.54,CW,0.95,22)
body(slide,"Continue in the same way for d³y/dx³, d⁴y/dx⁴ and higher orders.",XL,Y0+3.73,CW,0.52)

slide=new_slide("How it works","Meaning of the second derivative",TEAL)
body(slide,"The first derivative measures how y changes with x.",XL,Y0,CW,0.52)
body(slide,"The second derivative measures how that first derivative changes.",XL,Y0+0.68,CW,0.52)
fml(slide,"position x(t) → velocity v=dx/dt",XL,Y0+1.52,CW,LH,19)
fml(slide,"velocity v(t) → acceleration a=dv/dt=d²x/dt²",XL,Y0+2.16,CW,LH,19)
body(slide,"For a graph, the second derivative helps describe its curvature.",XL,Y0+3.02,CW,0.52)

slide=new_slide("Worked example","Repeated derivatives of sin x",TEAL)
fml(slide,"y=sin x",XL,Y0,CW,LH,20)
fml(slide,"dy/dx=cos x",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"d²y/dx²=-sin x",XL,Y0+2*LH+2*NG,CW,LH,20)
fml(slide,"d³y/dx³=-cos x",XL,Y0+3*LH+3*NG,CW,LH,20)
fml(slide,"d⁴y/dx⁴=sin x",XL,Y0+4*LH+4*NG,CW,LH,20)
result_card(slide,"The derivatives repeat after every four differentiations.",XL,Y0+5*LH+5*NG,CW,0.55)

slide=question_slide("Higher derivatives: mixed function")
qtx(slide,"For y=4x³+3sin x-4e^x, find the first four derivatives.",XL,Y0,CW,0.85,22)
slide=answer_slide()
fml(slide,"dy/dx=12x²+3cos x-4e^x",XL,Y0,CW,LH,18)
fml(slide,"d²y/dx²=24x-3sin x-4e^x",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"d³y/dx³=24-3cos x-4e^x",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"d⁴y/dx⁴=3sin x-4e^x",XL,Y0+3*LH+3*NG,CW,LH,18)

slide=question_slide("Higher derivatives: polynomial")
qtx(slide,"For y=x⁵+2x⁴+3x³+4x²+5x+6, find derivatives up to d⁴y/dx⁴.",XL,Y0,CW,0.95,21)
slide=answer_slide()
fml(slide,"dy/dx=5x⁴+8x³+9x²+8x+5",XL,Y0,CW,LH,18)
fml(slide,"d²y/dx²=20x³+24x²+18x+8",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"d³y/dx³=60x²+48x+18",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"d⁴y/dx⁴=120x+48",XL,Y0+3*LH+3*NG,CW,LH,18)

# CHAIN RULE
slide=new_slide("New topic","Chain rule: a function inside a function",GREEN)
body(slide,"Use the chain rule when one function is placed inside another.",XL,Y0,CW,0.55)
fml(slide,"sin(x³): sine outside, x³ inside",XL,Y0+0.78,CW,LH,19)
fml(slide,"eˢⁱⁿˣ: exponential outside, sin x inside",XL,Y0+1.42,CW,LH,19)
fml(slide,"(2x+3)⁵: fifth power outside, 2x+3 inside",XL,Y0+2.06,CW,LH,19)
trap_card(slide,"Product rule is for multiplied functions. Chain rule is for nested functions.",XL,Y0+2.94,CW,0.62)

slide=new_slide("How it works","Outside–inside rule",TEAL)
body(slide,"Let y=f(u) and u=g(x). Differentiate in two stages.",XL,Y0,CW,0.55)
ncert(slide,"dy/dx = (dy/du)(du/dx)",XL,Y0+0.78,CW,1.00,24)
body(slide,"Step 1: Differentiate the outer function, keeping the inner expression unchanged.",XL,Y0+2.04,CW,0.58)
body(slide,"Step 2: Multiply by the derivative of the inner expression.",XL,Y0+2.80,CW,0.55)
fml(slide,"d[f(g(x))]/dx = f'(g(x))g'(x)",XL,Y0+3.62,CW,LH,19)

slide=new_slide("Worked examples","Recognising the inner function",TEAL)
tx(slide,"y=sin(x³)",XL,Y0,5.70,0.40,21,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
fml(slide,"dy/dx=cos(x³)·3x²",XL,Y0+0.62,5.70,LH,18)
tx(slide,"y=cos(e^x)",6.95,Y0,5.70,0.40,21,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
fml(slide,"dy/dx=-sin(e^x)·e^x",6.95,Y0+0.62,5.70,LH,18)
tx(slide,"Outer derivative keeps the inner expression as it is.",XL,Y0+1.62,CW,0.48,17,WORK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
ncert(slide,"outer derivative × inner derivative",XL,Y0+2.36,CW,0.95,22)

slide=question_slide("Chain rule: trigonometric function")
qtx(slide,"Differentiate y=sin(2x³+4).",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Outer derivative: cos(2x³+4)",XL,Y0,CW,LH,19)
fml(slide,"Inner derivative: d(2x³+4)/dx=6x²",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=6x²cos(2x³+4).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: exponential of x²")
qtx(slide,"Differentiate y=eˣ².",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Outer function eᵘ differentiates to eᵘ",XL,Y0,CW,LH,19)
fml(slide,"Inner function u=x² has du/dx=2x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=2x eˣ².",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: exponential of sin x")
qtx(slide,"Differentiate y=eˢⁱⁿˣ.",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Outer derivative: eˢⁱⁿˣ",XL,Y0,CW,LH,20)
fml(slide,"Inner derivative: d(sin x)/dx=cos x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=eˢⁱⁿˣ cos x.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: power of a function")
qtx(slide,"Differentiate y=(x²+2x)³.",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Keep the inner expression: 3(x²+2x)²",XL,Y0,CW,LH,19)
fml(slide,"Differentiate inside: d(x²+2x)/dx=2x+2",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=3(x²+2x)²(2x+2).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: linear expression")
qtx(slide,"Differentiate y=(2x+3)⁵.",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Outer derivative: 5(2x+3)⁴",XL,Y0,CW,LH,20)
fml(slide,"Inner derivative: d(2x+3)/dx=2",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=10(2x+3)⁴.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: logarithm")
qtx(slide,"Differentiate y=logₑ(sin x).",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"For logₑu: derivative=(1/u)(du/dx)",XL,Y0,CW,LH,19)
fml(slide,"u=sin x and du/dx=cos x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=cos x/sin x=cot x.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: sine of a logarithm")
qtx(slide,"Differentiate y=sin(logₑx), where x>0.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"Outer derivative: cos(logₑx)",XL,Y0,CW,LH,19)
fml(slide,"Inner derivative: d(logₑx)/dx=1/x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=cos(logₑx)/x.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: square root")
qtx(slide,"Differentiate y=sqrt(x+4), where x>-4.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"Write y=(x+4)¹ᐟ²",XL,Y0,CW,LH,20)
fml(slide,"dy/dx=(1/2)(x+4)⁻¹ᐟ²·1",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=1/[2sqrt(x+4)].",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Chain rule: square of sine")
qtx(slide,"Differentiate y=sin²x.",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Treat y=(sin x)²",XL,Y0,CW,LH,20)
fml(slide,"dy/dx=2sin x·cos x",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dx=sin(2x).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Nested chain rule")
qtx(slide,"Differentiate y=eˢⁱⁿ⁽ˣ³⁾.",XL,Y0,CW,0.75,24)
slide=answer_slide()
fml(slide,"Differentiate eᵘ: eᵘ·du/dx",XL,Y0,CW,LH,19)
fml(slide,"For u=sin(x³): du/dx=cos(x³)·3x²",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=3x² eˢⁱⁿ⁽ˣ³⁾ cos(x³).",XL,Y0+2*LH+2*NG,CW,0.55)

# PHYSICS APPLICATION
slide=new_slide("Physics application","S.H.M. differential equation",GREEN)
body(slide,"For displacement y=A sin(ωt), differentiate with respect to time.",XL,Y0,CW,0.55)
fml(slide,"dy/dt=Aωcos(ωt)",XL,Y0+0.76,CW,LH,20)
fml(slide,"d²y/dt²=-Aω²sin(ωt)",XL,Y0+1.40,CW,LH,20)
body(slide,"Since y=A sin(ωt), replace A sin(ωt) by y.",XL,Y0+2.18,CW,0.52)
ncert(slide,"d²y/dt² = -ω²y     or     d²y/dt²+ω²y=0",XL,Y0+2.92,CW,1.05,22)
body(slide,"Acceleration is opposite to displacement: a=-ω²y.",XL,Y0+4.18,CW,0.52)

# PARTIAL DIFFERENTIATION
slide=new_slide("New topic","Partial differentiation",GREEN)
body(slide,"Some physical quantities depend on more than one independent variable.",XL,Y0,CW,0.55)
fml(slide,"V=V(x,y,z)",XL,Y0+0.78,CW,LH,21)
body(slide,"A partial derivative measures change with respect to one variable while the others are held constant.",XL,Y0+1.56,CW,0.72)
ncert(slide,"∂V/∂x,     ∂V/∂y,     ∂V/∂z",XL,Y0+2.52,CW,1.00,24)

slide=new_slide("How it works","Hold the other variables constant",TEAL)
body(slide,"To find ∂V/∂x, treat y and z as constants.",XL,Y0,CW,0.52)
body(slide,"To find ∂V/∂y, treat x and z as constants.",XL,Y0+0.66,CW,0.52)
body(slide,"To find ∂V/∂z, treat x and y as constants.",XL,Y0+1.32,CW,0.52)
fml(slide,"If V=x²y, then ∂V/∂x=2xy",XL,Y0+2.18,CW,LH,20)
fml(slide,"For the same V, ∂V/∂y=x²",XL,Y0+2.82,CW,LH,20)
trap_card(slide,"The symbol ∂ is used because only one variable changes at a time.",XL,Y0+3.66,CW,0.60)

slide=question_slide("Partial derivatives of three variables")
qtx(slide,"For V=x²y²z⁴, find ∂V/∂x, ∂V/∂y and ∂V/∂z.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"∂V/∂x=2xy²z⁴",XL,Y0,CW,LH,20)
fml(slide,"∂V/∂y=2x²yz⁴",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"∂V/∂z=4x²y²z³",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"Differentiate one variable; keep the other factors unchanged.",XL,Y0+3*LH+3*NG,CW,0.58)

# PARAMETRIC AND RECIPROCAL DERIVATIVES
slide=new_slide("New topic","Parametric differentiation",GREEN)
body(slide,"Sometimes x and y are both written in terms of a third variable t.",XL,Y0,CW,0.55)
fml(slide,"x=f(t),     y=g(t)",XL,Y0+0.78,CW,LH,21)
body(slide,"Differentiate each equation with respect to t, then divide.",XL,Y0+1.56,CW,0.55)
ncert(slide,"dy/dx = (dy/dt)/(dx/dt), provided dx/dt≠0",XL,Y0+2.34,CW,1.05,22)

slide=new_slide("How it works","Reciprocal derivative",TEAL)
body(slide,"If dy/dx is known and non-zero, reverse it to obtain dx/dy.",XL,Y0,CW,0.58)
ncert(slide,"dx/dy = 1/(dy/dx)",XL,Y0+0.84,CW,1.05,24)
body(slide,"This is useful when the question asks for x changing with y instead of y changing with x.",XL,Y0+2.14,CW,0.66)
trap_card(slide,"The reciprocal rule needs dy/dx≠0.",XL,Y0+3.04,CW,0.58)

slide=question_slide("Parametric derivative")
qtx(slide,"If y=2t³+1 and x=t², find dy/dx.",XL,Y0,CW,0.82,23)
slide=answer_slide()
fml(slide,"dy/dt=6t²",XL,Y0,CW,LH,20)
fml(slide,"dx/dt=2t",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"dy/dx=(6t²)/(2t)",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"dy/dx=3t, for t≠0.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Reciprocal derivative")
qtx(slide,"If y=sin u and u=3x, find dx/dy.",XL,Y0,CW,0.82,23)
slide=answer_slide()
fml(slide,"dy/du=cos u and du/dx=3",XL,Y0,CW,LH,19)
fml(slide,"dy/dx=3cos u=3cos(3x)",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"Take the reciprocal",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dx/dy=1/[3cos(3x)], where cos(3x)≠0.",XL,Y0+3*LH+3*NG,CW,0.58)

slide=question_slide("Slope using differentiation")
qtx(slide,"Find the slope of y=3x²+5x-4 at x=1.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"dy/dx=6x+5",XL,Y0,CW,LH,20)
fml(slide,"At x=1: slope=6(1)+5",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Slope = 11.",XL,Y0+2*LH+2*NG,CW,0.55)

# HOMEWORK SET
slide=question_slide("Homework: trigonometric chain rule")
qtx(slide,"Differentiate y=A cos(kx), where A and k are constants.",XL,Y0,CW,0.85,22)
slide=answer_slide()
fml(slide,"Outer derivative: -A sin(kx)",XL,Y0,CW,LH,19)
fml(slide,"Inner derivative: d(kx)/dx=k",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx=-Ak sin(kx).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Homework: exponential chain rule")
qtx(slide,"Differentiate y=5e³ˣ.",XL,Y0,CW,0.80,24)
slide=answer_slide()
fml(slide,"Keep 5 outside",XL,Y0,CW,LH,20)
fml(slide,"d(e³ˣ)/dx=e³ˣ·3",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dx=15e³ˣ.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Homework: negative exponent")
qtx(slide,"Differentiate y=e^(-x).",XL,Y0,CW,0.80,24)
slide=answer_slide()
fml(slide,"Outer derivative remains e⁻ˣ",XL,Y0,CW,LH,20)
fml(slide,"Inner derivative: d(-x)/dx=-1",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dx=-e⁻ˣ.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Homework: nested sine")
qtx(slide,"Differentiate y=3sin(x²+2).",XL,Y0,CW,0.80,24)
slide=answer_slide()
fml(slide,"Keep 3 outside",XL,Y0,CW,LH,20)
fml(slide,"Outer derivative: cos(x²+2)",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"Inner derivative: d(x²+2)/dx=2x",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx=6x cos(x²+2).",XL,Y0+3*LH+3*NG,CW,0.55)

slide=new_slide("Class summary","Higher differentiation lock lines",TEAL)
fml(slide,"differentiate again to obtain d²y/dx², d³y/dx³, ...",XL,Y0,CW,LH,18)
fml(slide,"chain rule: outer derivative × inner derivative",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"S.H.M.: d²y/dt²=-ω²y",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"partial derivative: change one variable, hold the others constant",XL,Y0+3*LH+3*NG,CW,LH,17)
fml(slide,"parametric: dy/dx=(dy/dt)/(dx/dt)",XL,Y0+4*LH+4*NG,CW,LH,18)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Redo all chain-rule questions without seeing the solutions.",XL,Y0,CW,0.52)
body(slide,"2. Write the first four derivatives of sin x and cos x.",XL,Y0+0.70,CW,0.52)
body(slide,"3. Prove the S.H.M. differential equation once.",XL,Y0+1.40,CW,0.52)
body(slide,"4. Practise one partial and one parametric derivative.",XL,Y0+2.10,CW,0.52)
trap_card(slide,"For nested functions, never forget the derivative of the inside.",XL,Y0+3.02,CW,0.62)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)


