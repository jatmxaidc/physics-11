"""Lecture 09 — circles, ellipses, exponential graphs and differentiation."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_09_Circle_Ellipse_Exponential_Differentiation.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)


# ── COVER ─────────────────────────────────────────────────────────────────────
# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 09",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Circle, Ellipse, Exponential and Differentiation",bx,by+1.62,bw,0.40,17,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"14 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Five-minute recap, then circle and new topics."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["5-minute graph recap","Circle: equation, centre, radius and exact graph","Ellipse: standard equation, axes and exact graph","S.H.M. speed-position ellipse","Exponential functions and their graphs","Differentiation, turning points and quadratic practice"]
for i,item in enumerate(topics):
 yy=1.18+i*0.78; tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False); tx(slide,item,XL+0.90,yy,CW-0.90,0.70,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Keep recap to five minutes."); after_content_slide()

# FIVE-MINUTE RECAP
slide=new_slide("5-minute recap","Graph-shape formula map",BLUE)
fml(slide,"Modulus y=|x|: V-shape     inverse graph y=1/x: rectangular hyperbola",XL,Y0,CW,LH,17)
fml(slide,"Parabola y=x^2     square-root graph y=sqrt(x)",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"Kinetic energy K proportional to v^2     pendulum period T proportional to sqrt(L)",XL,Y0+2*LH+2*NG,CW,LH,17)
fml(slide,"Orbital period T proportional to R^(3/2)     at fixed p, kinetic energy K proportional to 1/m",XL,Y0+3*LH+3*NG,CW,LH,16)

slide=new_slide("5-minute recap","How to draw any graph",BLUE)
body(slide,"1. Put the dependent quantity on the y-axis and independent quantity on the x-axis.",XL,Y0,CW)
body(slide,"2. Check the domain (allowed values), intercepts, symmetry and turning points.",XL,Y0+LH+NG,CW)
body(slide,"3. Mark enough exact points; draw the appropriate smooth curve or straight segments.",XL,Y0+2*LH+2*NG,CW)
body(slide,"4. Label both axes and important points.",XL,Y0+3*LH+3*NG,CW)

# CIRCLE
slide=new_slide("New topic","Circle: exact graph",GREEN)
body(slide,"A circle is the locus of points in a plane at a fixed distance R from a fixed point called the centre.",XL,Y0,5.8,0.70)
ncert(slide,"x² + y² = R²",XL,Y0+0.85,5.8,0.95,26)
body(slide,"Shown here: R = 3 and centre = (0, 0).",XL,Y0+2.00,5.8)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-4,4,-4,4)
axis_ticks(slide,sx,sy,[-3,3],[-3,3])
plot_circle(slide,sx,sy,0,0,3,GOLD)
mark_point(slide,sx,sy,0,0,"O",GREEN,0.10,0.10)
draw_line(slide,sx(0),sy(0),sx(3),sy(0),TEAL,1.6)
tx(slide,"R=3",sx(1.1),sy(0)-0.32,0.70,0.22,12,TEAL,True,spa=0,scaled=False)

slide=new_slide("How it works","Equation of a circle: how it is formed",TEAL)
body(slide,"If the centre is C(h, k) and P(x, y) lies on the circle, then CP = R.",XL,Y0,CW,0.48)
fml(slide,"Distance formula:  √[(x-h)²+(y-k)²] = R",XL,Y0+0.62,CW,0.48,19)
ncert(slide,"(x-h)² + (y-k)² = R²",XL,Y0+1.28,CW,1.10,26)
fml(slide,"centre = (h, k)          radius = R",XL,Y0+2.62,CW,0.48,22)
trap_card(slide,"Signs inside brackets are opposite: (x+3)² means h = -3.",XL,Y0+3.30,CW,0.48)

slide=new_slide("Learn","Important points on a circle",TEAL)
body(slide,"For (x-h)²+(y-k)²=R²:",XL,Y0,CW)
fml(slide,"right = (h+R, k)          left = (h-R, k)",XL,Y0+0.62,CW,0.48,20)
fml(slide,"top = (h, k+R)          bottom = (h, k-R)",XL,Y0+1.30,CW,0.48,20)
body(slide,"These four points fix the exact width and height before drawing the circle.",XL,Y0+2.00,CW,0.52)

slide=question_slide("Circle: x^2+y^2=49")
qtx(slide,"Find the centre and radius, then draw the exact graph.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"Compare with x^2+y^2=R^2",XL,Y0,5.8,LH,19)
fml(slide,"R^2=49, so R=7",XL,Y0+LH+NG,5.8,LH,19)
result_card(slide,"Centre = (0, 0), radius = 7.",XL,Y0+2*LH+2*NG,5.8,0.42,20)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-8,8,-8,8)
axis_ticks(slide,sx,sy,[-7,-4,4,7])
plot_circle(slide,sx,sy,0,0,7,GOLD)
draw_line(slide,sx(0),sy(0),sx(7),sy(0),TEAL,1.6)
tx(slide,"R=7",sx(2.7),sy(0)-0.32,0.70,0.22,12,TEAL,True,spa=0,scaled=False)

slide=question_slide("Circle: shifted centre")
qtx(slide,"For (x-4)^2+(y-3)^2=25, find the centre and radius.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"Compare with (x-h)^2+(y-k)^2=R^2",XL,Y0,5.8,LH,18)
fml(slide,"h=4, k=3 and R=sqrt(25)=5",XL,Y0+LH+NG,5.8,LH,18)
result_card(slide,"Centre = (4, 3), radius = 5.",XL,Y0+2*LH+2*NG,5.8,0.42,20)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-2,10,-3,9)
axis_ticks(slide,sx,sy,[-1,4,9],[-2,3,8])
plot_circle(slide,sx,sy,4,3,5,TEAL)
mark_point(slide,sx,sy,4,3,"C(4,3)",GREEN)
draw_line(slide,sx(4),sy(3),sx(9),sy(3),GOLD,1.5)
tx(slide,"R=5",sx(5.8),sy(3)-0.30,0.70,0.22,12,GOLD,True,spa=0,scaled=False)

slide=question_slide("Circle: negative centre coordinate")
qtx(slide,"For (x+3)^2+(y-4)^2=36, find the centre and radius.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"Write x+3 as x-(-3)",XL,Y0,5.8,LH,19)
fml(slide,"h=-3, k=4 and R=sqrt(36)=6",XL,Y0+LH+NG,5.8,LH,19)
result_card(slide,"Centre = (-3, 4), radius = 6.",XL,Y0+2*LH+2*NG,5.8,0.42,20)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-10,4,-3,11)
axis_ticks(slide,sx,sy,[-9,-3,3],[-2,4,10])
plot_circle(slide,sx,sy,-3,4,6,GREEN)
mark_point(slide,sx,sy,-3,4,"C(-3,4)",TEAL)

slide=new_slide("Extension","General equation of a circle",TEAL)
body(slide,"After expanding, a circle may appear as:",XL,Y0,CW)
ncert(slide,"x^2+y^2+2gx+2fy+c=0",XL,Y0+LH+NG,CW,0.95,22)
fml(slide,"centre=(-g,-f)",XL,Y0+2*LH+2*NG+0.88,CW,LH,21)
fml(slide,"radius=sqrt(g^2+f^2-c)",XL,Y0+3*LH+3*NG+0.88,CW,LH,21)

slide=question_slide("General circle equation")
qtx(slide,"Find the centre and radius of x^2+y^2-6x+4y-12=0.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"2g=-6 gives g=-3; 2f=4 gives f=2",XL,Y0,CW,LH,18)
fml(slide,"Centre=(-g,-f)=(3,-2)",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"R=sqrt[(-3)^2+2^2-(-12)]=sqrt(25)=5",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"Centre=(3,-2), radius=5.",XL,Y0+3*LH+3*NG,CW,0.40,21)

# ELLIPSE
slide=new_slide("New topic","Ellipse: exact graph",GREEN)
body(slide,"An ellipse is the locus of points whose total distance from two fixed points is constant.",XL,Y0,5.8,0.70)
ncert(slide,"x²/a² + y²/b² = 1",XL,Y0+0.85,5.8,0.95,24)
body(slide,"Shown here: a = 4 horizontally and b = 2 vertically.",XL,Y0+2.00,5.8,0.48)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-5,5,-5,5)
axis_ticks(slide,sx,sy,[-4,4],[-2,2])
plot_ellipse(slide,sx,sy,0,0,4,2,GOLD)
draw_line(slide,sx(0),sy(0),sx(4),sy(0),TEAL,1.5)
tx(slide,"a=4",sx(1.4),sy(0)-0.30,0.75,0.22,12,TEAL,True,spa=0,scaled=False)
draw_line(slide,sx(0),sy(0),sx(0),sy(2),GREEN,1.5)
tx(slide,"b=2",sx(0)+0.10,sy(1.2)-0.12,0.75,0.22,12,GREEN,True,spa=0,scaled=False)

slide=new_slide("How it works","Standard ellipse: how to read it",TEAL)
body(slide,"For an ellipse centred at the origin:",XL,Y0,CW)
ncert(slide,"x²/a² + y²/b² = 1",XL,Y0+0.62,CW,1.05,24)
fml(slide,"x-intercepts = (±a, 0)",XL,Y0+1.90,CW,0.48,20)
fml(slide,"y-intercepts = (0, ±b)",XL,Y0+2.56,CW,0.48,20)
body(slide,"The larger denominator gives the longer semi-axis.",XL,Y0+3.24,CW,0.48)

slide=new_slide("Learn","Horizontal and vertical ellipses",TEAL)
tx(slide,"a²>b²: horizontal major axis (a,b>0)",0.85,Y0,5.4,0.38,17,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
sx1,sy1=graph_axes(slide,1.35,1.78,4.35,4.35,-5,5,-5,5)
axis_ticks(slide,sx1,sy1,[-4,4],[-2,2])
plot_ellipse(slide,sx1,sy1,0,0,4,2,GOLD)
tx(slide,"b²>a²: vertical major axis (a,b>0)",7.05,Y0,5.4,0.38,17,TEAL,True,PP_ALIGN.CENTER,spa=0,scaled=False)
sx2,sy2=graph_axes(slide,7.55,1.78,4.35,4.35,-5,5,-5,5)
axis_ticks(slide,sx2,sy2,[-2,2],[-4,4])
plot_ellipse(slide,sx2,sy2,0,0,2,4,TEAL)

slide=question_slide("Draw x^2/4 + y^2/9 = 1")
qtx(slide,"Find the intercepts and draw the exact ellipse.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"x^2/2^2 + y^2/3^2 = 1",XL,Y0,5.8,LH,19)
fml(slide,"x-intercepts=(+2,0),(-2,0)",XL,Y0+LH+NG,5.8,LH,18)
fml(slide,"y-intercepts=(0,+3),(0,-3)",XL,Y0+2*LH+2*NG,5.8,LH,18)
result_card(slide,"Vertical ellipse: semi-major axis=3.",XL,Y0+3*LH+3*NG,5.8,0.40,18)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-4,4,-4,4)
axis_ticks(slide,sx,sy,[-2,2],[-3,3])
plot_ellipse(slide,sx,sy,0,0,2,3,TEAL)

slide=question_slide("Shifted ellipse")
qtx(slide,"For (x-2)^2/16 + (y+1)^2/4 = 1, find centre and semi-axes.",XL,Y0,CW,0.48,18)
slide=answer_slide()
fml(slide,"Centre=(2,-1)",XL,Y0,5.8,LH,20)
fml(slide,"Horizontal semi-axis=sqrt(16)=4",XL,Y0+LH+NG,5.8,LH,18)
fml(slide,"Vertical semi-axis=sqrt(4)=2",XL,Y0+2*LH+2*NG,5.8,LH,18)
result_card(slide,"Horizontal ellipse centred at (2,-1).",XL,Y0+3*LH+3*NG,5.8,0.40,19)
sx,sy=graph_axes(slide,7.15,1.40,4.8,4.8,-3,7,-6,4)
axis_ticks(slide,sx,sy,[-2,2,6],[-3,-1,1])
plot_ellipse(slide,sx,sy,2,-1,4,2,GREEN)
mark_point(slide,sx,sy,2,-1,"C(2,-1)",TEAL)

# SHM ELLIPSE
slide=new_slide("Physics application","S.H.M. speed-position graph",GREEN)
body(slide,"For displacement x in simple harmonic motion:",XL,Y0,5.8)
ncert(slide,"|v| = omega sqrt(A^2-x^2)",XL,Y0+LH+NG,5.8,0.95,20)
body(slide,"The speed-position graph is the upper half of an ellipse.",XL,Y0+2*LH+2*NG+0.88,5.8)
sx,sy=graph_axes(slide,7.0,1.40,5.1,4.8,-1.2,1.2,0,1.2,"x","v")
plot_parametric(slide,lambda t:math.cos(t),lambda t:math.sin(t),sx,sy,0,math.pi,TEAL)
draw_line(slide,sx(-1),sy(0)-0.06,sx(-1),sy(0)+0.06,MUTED,1.1)
draw_line(slide,sx(1),sy(0)-0.06,sx(1),sy(0)+0.06,MUTED,1.1)
draw_line(slide,sx(0)-0.06,sy(1),sx(0)+0.06,sy(1),MUTED,1.1)
tx(slide,"-A",sx(-1)-0.18,sy(0)+0.08,0.36,0.22,11,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"A",sx(1)-0.18,sy(0)+0.08,0.36,0.22,11,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"ωA",sx(0)-0.48,sy(1)-0.10,0.40,0.22,11,MUTED,True,PP_ALIGN.RIGHT,spa=0,scaled=False)

slide=new_slide("How it works","Theory of the S.H.M. ellipse",TEAL)
fml(slide,"v^2=omega^2(A^2-x^2)",XL,Y0,CW,LH,19)
fml(slide,"x^2/A^2 + v^2/(omega A)^2 = 1",XL,Y0+LH+NG,CW,LH,19)
body(slide,"At x=0, speed is maximum: vₘₐₓ=omega A.",XL,Y0+2*LH+2*NG,CW)
body(slide,"At x=+A or -A, speed is zero because the particle turns back.",XL,Y0+3*LH+3*NG,CW)
trap_card(slide,"Velocity can be positive or negative; speed is only the upper half.",XL,Y0+4*LH+4*NG,CW)

# EXPONENTIAL
slide=new_slide("New topic","Exponential graph: y = e^x",GREEN)
body(slide,"The variable x appears in the power.",XL,Y0,5.8)
fml(slide,"e is approximately 2.718",XL,Y0+LH+NG,5.8,LH,20)
body(slide,"The graph passes through (0,1) and rises rapidly.",XL,Y0+2*LH+2*NG,5.8)
sx,sy=graph_axes(slide,7.0,1.40,5.1,4.8,-3,2,-0.5,8)
plot_function(slide,lambda x:math.exp(x),sx,sy,-3,2,-0.5,8,GOLD)
axis_ticks(slide,sx,sy,[-1,1],[1])
mark_point(slide,sx,sy,0,1,"(0,1)",GREEN,0.10,-0.28)

slide=new_slide("How it works","Properties of y=e^x",TEAL)
fml(slide,"e^0=1, so y-intercept=(0,1)",XL,Y0,CW,LH,19)
body(slide,"e^x is positive for every real x, so the graph never crosses the x-axis.",XL,Y0+LH+NG,CW)
body(slide,"As x increases, e^x grows rapidly. As x becomes very negative, e^x approaches zero.",XL,Y0+2*LH+2*NG,CW)
fml(slide,"domain: all real x     range: y>0",XL,Y0+3*LH+3*NG,CW,LH,19)
result_card(slide,"Horizontal asymptote: y=0.",XL,Y0+4*LH+4*NG,CW,0.40,20)

slide=new_slide("Compare","y=e^x and y=e^(-x)",TEAL)
body(slide,"Gold: e^x     Teal: e^(-x). Both pass through (0,1).",XL,Y0,CW)
sx,sy=graph_axes(slide,1.2,1.75,10.8,4.7,-3,3,-0.5,8)
plot_function(slide,lambda x:math.exp(x),sx,sy,-3,3,-0.5,8,GOLD)
plot_function(slide,lambda x:math.exp(-x),sx,sy,-3,3,-0.5,8,TEAL)
axis_ticks(slide,sx,sy,[-1,1],[1])
mark_point(slide,sx,sy,0,1,"(0,1)",GREEN,0.10,-0.28)

slide=question_slide("Exponential values")
qtx(slide,"For y=e^x, find y at x=0, x=1 and x=-1.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"x=0: y=e^0=1",XL,Y0,CW,LH,20)
fml(slide,"x=1: y=e",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"x=-1: y=e^(-1)=1/e",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"Points: (0,1), (1,e), (-1,1/e).",XL,Y0+3*LH+3*NG,CW,0.40,21)

# DIFFERENTIATION
slide=new_slide("New topic","Differentiation: slope at a point",GREEN)
body(slide,"A curve has a different slope at different points.",XL,Y0,5.8)
body(slide,"The slope at one point is the slope of the tangent there.",XL,Y0+LH+NG,5.8)
sx,sy=graph_axes(slide,7.0,1.40,5.1,4.8,-2,3,-2,7)
plot_function(slide,lambda x:x*x,sx,sy,-2,3,-2,7,GOLD)
plot_function(slide,lambda x:2*x-1,sx,sy,-2,3,-2,7,TEAL)
tx(slide,"tangent at x=1",9.2,2.05,2.0,0.28,13,TEAL,True,spa=0)
axis_ticks(slide,sx,sy,[-1,1,2],[1,4])
mark_point(slide,sx,sy,1,1,"P(1,1)",GREEN,0.10,-0.28)

slide=new_slide("How it works","How a derivative gives slope",TEAL)
body(slide,"Take two nearby points on a curve. Their joining line is a secant.",XL,Y0,CW)
fml(slide,"average slope = change in y / change in x",XL,Y0+LH+NG,CW,LH,19)
body(slide,"Move the second point closer to the first. The secant becomes the tangent.",XL,Y0+2*LH+2*NG,CW)
ncert(slide,"dy/dx = lim(Δx→0)  Δy/Δx",XL,Y0+3*LH+3*NG,CW,1.05,22)
body(slide,"The derivative tells the instantaneous rate of change.",XL,Y0+4*LH+4*NG+0.95,CW)

slide=new_slide("Notation","Ways to write a derivative",TEAL)
fml(slide,"dy/dx     f'(x)     y'",XL,Y0,CW,LH,24)
body(slide,"All three mean the derivative of y with respect to x.",XL,Y0+LH+NG,CW)
fml(slide,"units of derivative = units of y / units of x",XL,Y0+2*LH+2*NG,CW,LH,19)
body(slide,"Example: slope of a position-time graph has unit metre per second (m s⁻¹) and represents velocity.",XL,Y0+3*LH+3*NG,CW)

slide=new_slide("Learn","Sign of the derivative",TEAL)
grid_table(slide,["curve behaviour","tangent slope","derivative"],[["rising","positive","dy/dx>0"],["falling","negative","dy/dx<0"],["horizontal tangent","zero","dy/dx=0"]],XL,Y0,[4.0,4.0,4.0],0.72)

slide=new_slide("Exact graphs","Positive, zero and negative slope",GREEN)
body(slide,"Gold curve: y=x^2. Left side falls, vertex has zero slope, right side rises.",XL,Y0,CW)
sx,sy=graph_axes(slide,1.2,1.75,10.8,4.7,-3,3,-1,7)
plot_function(slide,lambda x:x*x,sx,sy,-3,3,-1,7,GOLD)
tx(slide,"negative slope",1.8,3.0,1.8,0.25,13,RED,True,spa=0)
mark_point(slide,sx,sy,0,0,"(0,0)",GREEN,0.12,0.10)
tx(slide,"slope=0",6.45,5.38,1.2,0.25,13,GREEN,True,spa=0)
tx(slide,"positive slope",9.4,3.0,1.8,0.25,13,TEAL,True,spa=0)

slide=new_slide("Turning points","Maximum and minimum",TEAL)
body(slide,"At a smooth turning point, the tangent is horizontal.",XL,Y0,CW)
ncert(slide,"At a stationary point: dy/dx = 0",XL,Y0+LH+NG,CW,0.95,23)
body(slide,"Minimum: derivative changes from negative to positive.",XL,Y0+2*LH+2*NG+0.88,CW)
body(slide,"Maximum: derivative changes from positive to negative.",XL,Y0+3*LH+3*NG+0.88,CW)

slide=question_slide("Match slope at points A, B, C and D")
qtx(slide,"State the sign of the slope at A, B, C and D.",XL,Y0,5.6,0.62,19)
sx,sy=graph_axes(slide,7.0,1.40,5.1,4.8,-3,5,-2,5)
plot_function(slide,lambda x:-0.25*(x-1)*(x-1)+4,sx,sy,-3,5,-2,5,GOLD)
for px,label in [(-2,"A"),(-0.5,"B"),(1,"C"),(3.5,"D")]:
    py=-0.25*(px-1)*(px-1)+4
    mark_point(slide,sx,sy,px,py,label,GREEN)
slide=answer_slide()
fml(slide,"A: positive slope",XL,Y0,CW,LH,20)
fml(slide,"B: positive slope",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"C: zero slope because tangent is horizontal",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"D: negative slope",XL,Y0+3*LH+3*NG,CW,LH,20)
result_card(slide,"A positive, B positive, C zero, D negative.",XL,Y0+4*LH+4*NG,CW,0.40,20)

slide=new_slide("Basic rule","Power rule for differentiation",TEAL)
ncert(slide,"For constant n, where defined: if y=x^n, then dy/dx = n x^(n-1)",XL,Y0,CW,1.00,20)
fml(slide,"d(x^2)/dx=2x",XL,Y0+1.20,CW,LH,21)
fml(slide,"d(x^3)/dx=3x^2",XL,Y0+1.20+LH+NG,CW,LH,21)
fml(slide,"d(constant)/dx=0",XL,Y0+1.20+2*LH+2*NG,CW,LH,21)

slide=question_slide("Differentiate y=x^2+3x-4")
qtx(slide,"Find dy/dx and the slope at x=2.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"d(x^2)/dx=2x",XL,Y0,CW,LH,20)
fml(slide,"d(3x)/dx=3 and d(-4)/dx=0",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"dy/dx=2x+3",XL,Y0+2*LH+2*NG,CW,LH,21)
fml(slide,"At x=2: slope=2(2)+3=7",XL,Y0+3*LH+3*NG,CW,LH,20)
result_card(slide,"Slope at x=2 is 7.",XL,Y0+4*LH+4*NG,CW,0.40,22)

# QUADRATIC PRACTICE
slide=question_slide("Quadratic roots: x^2+7x+12=0")
qtx(slide,"Find both roots by factorisation.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"Find two numbers with product 12 and sum 7: 3 and 4",XL,Y0,CW,LH,18)
fml(slide,"x^2+7x+12=(x+3)(x+4)",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"(x+3)(x+4)=0",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"x=-3 or x=-4.",XL,Y0+3*LH+3*NG,CW,0.40,22)

slide=question_slide("Quadratic roots: 2x^2+3x-2=0")
qtx(slide,"Find both roots by factorisation.",XL,Y0,CW,0.45,18)
slide=answer_slide()
fml(slide,"Split 3x as 4x-x",XL,Y0,CW,LH,19)
fml(slide,"2x^2+4x-x-2=0",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"(2x-1)(x+2)=0",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"x=1/2 or x=-2.",XL,Y0+3*LH+3*NG,CW,0.40,22)

slide=new_slide("5-minute summary","Circle, ellipse and exponential",TEAL)
fml(slide,"circle: (x-h)^2+(y-k)^2=R^2",XL,Y0,CW,LH,19)
fml(slide,"origin-centred ellipse: x^2/a^2+y^2/b^2=1",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"e^x passes through (0,1), stays positive and has asymptote y=0",XL,Y0+2*LH+2*NG,CW,LH,18)

slide=new_slide("5-minute summary","Differentiation lock lines",TEAL)
fml(slide,"The derivative gives the tangent's slope and represents an instantaneous rate of change.",XL,Y0,CW,LH,17)
fml(slide,"rising: dy/dx>0     falling: dy/dx<0",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"stationary point: dy/dx=0",XL,Y0+2*LH+2*NG,CW,LH,20)
fml(slide,"d(x^n)/dx = n x^(n-1)",XL,Y0+3*LH+3*NG,CW,LH,21)

slide=new_slide("After class","Practice",AMBER)
body(slide,"Redraw each circle, ellipse and exponential graph from this lecture with labelled axes.",XL,Y0,CW)
body(slide,"Redo every question without looking at its answer slide.",XL,Y0+LH+NG,CW)
body(slide,"Learn the meaning and sign of a derivative before the next class.",XL,Y0+2*LH+2*NG,CW)

prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
