"""Units 03 — Dimensional applications, limitations, significant figures (PW 04)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Units_03_Dimensional_Applications_and_Significant_Figures.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³ᐟ²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units 03",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units and Measurement",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Dimensional Applications and Significant Figures",bx,by+1.62,bw,0.40,16,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"8 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"From PW Units and Measurement Lecture 04."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"From Units 02: derive a formula (pendulum)",
"Uses of dimensional analysis",
"sin / log / exp — find unknown constants",
"New fundamental quantities (F, a, t …)",
"Deducing relations by matching powers",
"Dimensional vs dimensionless constants",
"Limitations + significant figures",
]
for i,t in enumerate(topics):
    y=1.25+i*0.65
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(XL),Inches(y),Inches(0.42),Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.fill.background()
    tx(slide,f"{i+1:02d}",XL,y+0.05,0.42,0.32,14,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
    tx(slide,t,XL+0.60,y+0.05,10.5,0.35,20,TEXT,False,spa=0,scaled=False)
_footer(slide,2); after_content_slide()

slide=new_slide("Class note","Where we left Units 02",TEAL)
para(slide,"Last class we locked homogeneity: things that are added must match in dimensions. We also treated sin θ, e^x and log x as the number 1 for dimensions.",XL,Y0,CW,1.25,16)
para(slide,"Units 02 ended by opening the next idea — slide 51: use dimensions to find the powers in a formula. We finish that idea first, then go into exam applications and significant figures.",XL,Y0+1.40,CW,1.30,16)
ncert(slide,"Start today from Units 02 slide 51: derive the pendulum formula",XL,Y0+2.90,CW,0.95,15)
trap_card(slide,"Class notes will be saved after teaching. This deck is for the board.",XL,Y0+4.10,CW,0.45)

# FROM UNITS 02 SLIDE 51
slide=new_slide("From Units 02","Using dimensions to derive a formula",GREEN)
para(slide,"Sometimes we know which quantities a formula depends on, but we do not know the powers. Dimensional analysis can find those powers. We assume a product form, write dimensions on both sides, and match the powers of M, L and T one by one.",XL,Y0,CW,1.50,16)
para(slide,"Example idea: time period T of a simple pendulum depends on length l and acceleration due to gravity g. We write T = k l^a g^b, where k is a pure number. Then we find a and b. The method cannot give the value of k — experiment or deeper theory gives k = 2π.",XL,Y0+1.65,CW,1.45,16)
ncert(slide,"T ∝ l^a g^b   →   find a and b by comparing powers",XL,Y0+3.30,CW,0.95,15)

slide=new_slide("Theory","What we assume when we derive",TEAL)
para(slide,"We assume the unknown relation is a product of powers: quantity = constant × (first quantity)^a × (second)^b × …. This works when each variable appears once as a power. It fails if the true law has a sum of two different terms, or a trigonometric factor that we forgot to treat as dimensionless.",XL,Y0,CW,1.55,16)
para(slide,"Also, if a quantity depends on more variables than we can balance with M, L, T alone, dimensions may not give a unique answer. Then we need more physics. Keep that honesty for board answers under limitations.",XL,Y0+1.70,CW,1.35,16)
trap_card(slide,"k is a pure number here. Dimensions ignore it.",XL,Y0+3.30,CW,0.55)

slide=new_slide("How it works","Pendulum step by step",GREEN)
body(slide,"Write T = k l^a g^b. Now match dimensions of both sides.",XL,Y0,CW,0.42)
fml(slide,"Left side [T] = T^1. Right side = L^a × (L T^(-2))^b",XL,Y0+0.48,CW,LH,15)
fml(slide,"Right side becomes L^(a+b) T^(-2b)",XL,Y0+0.48+LH+NG,CW,LH,16)
fml(slide,"Match time power: 1 = −2b  ⇒  b = −1/2",XL,Y0+0.48+2*LH+2*NG,CW,LH,15)
fml(slide,"Match length power: 0 = a + b  ⇒  a = +1/2",XL,Y0+0.48+3*LH+3*NG,CW,LH,15)
result_card(slide,"So T = k √(l/g). Real formula uses k = 2π.",XL,Y0+0.48+4*LH+4*NG,CW,0.42)

slide=new_slide("Theory","Uses of dimensional analysis",TEAL)
para(slide,"Dimensional analysis has three standard school uses. First, check whether an equation is dimensionally correct. Second, convert a numerical value from one system of units to another using n1 u1 = n2 u2. Third, find how physical quantities are related — the powers in a formula — when we know which quantities matter.",XL,Y0,CW,1.55,16)
para(slide,"For a 3-mark CBSE answer, write the definition of homogeneity briefly, list these three uses, and add one limitation. That structure scores cleanly without long English.",XL,Y0+1.70,CW,1.25,16)
ncert(slide,"Check  ·  Convert  ·  Derive",XL,Y0+3.15,CW,0.95,20)

slide=new_slide("CBSE write","Uses of dimensional analysis (3 marks)",GREEN)
body(slide,"1. To check whether an equation is dimensionally correct.",XL,Y0,CW,0.50)
body(slide,"2. To convert a value from one system of units to another.",XL,Y0+0.55,CW,0.50)
body(slide,"3. To find the relation between physical quantities (derive formula).",XL,Y0+1.10,CW,0.55)
ncert(slide,"Check  ·  Convert  ·  Derive",XL,Y0+1.85,CW,1.00,20)
trap_card(slide,"Limitation: cannot find dimensionless constants like 2π or 1/2.",XL,Y0+3.10,CW,0.55)

slide=question_slide("NEET: time period from α and β")
qtx(slide,"Max acceleration of oscillating particle is α, max velocity is β. Time period is",XL,Y0,CW,1.10,15)
choices(slide,[
"2π α/β",
"2π α β",
"2π β/α",
"2π α^2 / β^2",
], XL, Y0+1.25, 10.5)
slide=answer_slide()
body(slide,"Easy way: check which option has dimension of time.",XL,Y0,CW,0.40)
fml(slide,"In SHM, max speed β = A ω, max acceleration α = A ω^2",XL,Y0+0.45,CW,LH,14)
fml(slide,"Divide: α/β = ω. Time period is T = 2π / ω",XL,Y0+0.45+LH+NG,CW,LH,15)
fml(slide,"So T = 2π × (β/α)",XL,Y0+0.45+2*LH+2*NG,CW,LH,18)
result_card(slide,"Final answer: 2π β/α   (option 3)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

slide=question_slide("Which does NOT have dimension of time?")
qtx(slide,"L, C, R are inductance, capacitance, resistance. Which has no dimension of time?",XL,Y0,CW,1.10,15)
choices(slide,["RC","L/R","√(LC)","L/C"], XL, Y0+1.25, 10.5)
slide=answer_slide()
body(slide,"Remember three time formulas from circuits.",XL,Y0,CW,0.40)
fml(slide,"RC has dimension of time. L/R has dimension of time.",XL,Y0+0.45,CW,LH,15)
fml(slide,"√(LC) also has dimension of time.",XL,Y0+0.45+LH+NG,CW,LH,16)
fml(slide,"L/C does not become time. That is the odd one out.",XL,Y0+0.45+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: L/C   (option 4)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

# THEORY: TRANSCENDENTAL
slide=new_slide("Theory","Inside sin, log and e^x — the hard rule",GREEN)
para(slide,"Any argument sitting inside sin, cos, tan, log or e^( ) must be dimensionless. That means its dimensions are M^0 L^0 T^0. If two pieces are added inside that argument, each piece must also be dimensionless — and they must match each other.",XL,Y0,CW,1.45,16)
para(slide,"Outside, the whole sin or log is just a pure number for dimensions. So if Force = A cos(Bx), then [A] = [Force], and [Bx] = 1.",XL,Y0+1.60,CW,1.25,16)
ncert(slide,"Inside = dimensionless    ·    Outside multiplier matches the left side",XL,Y0+3.05,CW,0.95,15)

slide=new_slide("Theory","How to attack such an equation",TEAL)
body(slide,"Step 1: Mark every + or − inside the sin / log / exp. Those pieces match.",XL,Y0,CW,0.48)
body(slide,"Step 2: Each piece alone must be dimensionless (equal to 1 for dimensions).",XL,Y0+0.55,CW,0.48)
body(slide,"Step 3: The factor in front of sin / log / exp matches the left-hand quantity.",XL,Y0+1.10,CW,0.48)
body(slide,"Step 4: Multiply or divide to get the combination the question asks for.",XL,Y0+1.65,CW,0.48)
ncert(slide,"Do not invent physics. Read only what the equation allows.",XL,Y0+2.40,CW,1.00,16)
trap_card(slide,"If λt sits inside sin, and λ is length, that equation is dimensionally wrong.",XL,Y0+3.65,CW,0.50)

slide=question_slide("Force with cos and sin")
qtx(slide,"F = A cos(Bx) + C sin(Dt). Find dimensions of (A D)/B.",XL,Y0,CW,1.05,16)
slide=answer_slide()
body(slide,"Idea: Bx and Dt are dimensionless. A and C carry the force.",XL,Y0,CW,0.42)
fml(slide,"[Bx] = 1 ⇒ [B] = 1/[x] = L^(-1).   [Dt] = 1 ⇒ [D] = T^(-1)",XL,Y0+0.48,CW,LH,14)
fml(slide,"[A] = [C] = [F] = M L T^(-2)",XL,Y0+0.48+LH+NG,CW,LH,15)
fml(slide,"[A D / B] = [F] × T^(-1) / L^(-1) = [F] × (L / T)",XL,Y0+0.48+2*LH+2*NG,CW,LH,14)
result_card(slide,"Final answer: same as Force × length / time",XL,Y0+0.48+3*LH+3*NG,CW,0.45)

slide=question_slide("Velocity with square-root density")
qtx(slide,"V = α / (β + √density). Find [α] and [β].",XL,Y0,CW,1.00,16)
slide=answer_slide()
body(slide,"Idea: β is added to √density, so they match. Then α / that piece = velocity.",XL,Y0,CW,0.45)
fml(slide,"[β] = [√density] = (M L^(-3))^(1/2) = M^(1/2) L^(-3/2)",XL,Y0+0.50,CW,LH,14)
fml(slide,"[α] = [V] × [β] = L T^(-1) × M^(1/2) L^(-3/2)",XL,Y0+0.50+LH+NG,CW,LH,14)
fml(slide,"[α] = M^(1/2) L^(-1/2) T^(-1)",XL,Y0+0.50+2*LH+2*NG,CW,LH,16)
result_card(slide,"Final: [β] = M^(1/2) L^(-3/2) ;  [α] = M^(1/2) L^(-1/2) T^(-1)",XL,Y0+0.50+3*LH+3*NG,CW,0.48)

slide=question_slide("Pressure with a log")
qtx(slide,"P = A log_e(B/t − C V). Find [A B] and [B C].",XL,Y0,CW,1.05,16)
slide=answer_slide()
body(slide,"Idea: log is a pure number, so [A] = [P]. Inside, each term is dimensionless.",XL,Y0,CW,0.45)
fml(slide,"[A] = [P] = M L^(-1) T^(-2).   B/t = 1 ⇒ [B] = T",XL,Y0+0.50,CW,LH,14)
fml(slide,"C V = 1 ⇒ [C] = 1/[V] = L^(-3)",XL,Y0+0.50+LH+NG,CW,LH,15)
fml(slide,"[A B] = M L^(-1) T^(-1) ;   [B C] = T L^(-3)",XL,Y0+0.50+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final: [AB] = M L^(-1) T^(-1) ;  [BC] = L^(-3) T",XL,Y0+0.50+3*LH+3*NG,CW,0.48)

slide=question_slide("NCERT: which displacement is wrong?")
qtx(slide,"Which relation for displacement y is dimensionally NOT correct?",XL,Y0,CW,0.90,16)
choices(slide,[
"y = a sin(2π t / T)",
"y = a sin(v t / λ)",
"y = √2 a sin(2π t / T) cos(2π x / λ)",
"y = (x/t) sin(λ t)",
], XL, Y0+1.15, 10.8)
slide=answer_slide()
body(slide,"Check the angle inside sin. It must be dimensionless.",XL,Y0,CW,0.40)
fml(slide,"Options 1–3: t/T, vt/λ, x/λ are all dimensionless. Fine.",XL,Y0+0.45,CW,LH,14)
fml(slide,"Option 4: λ t has dimensions L × T. Not an angle. Also x/t is velocity, not length.",XL,Y0+0.45+LH+NG,CW,LH,13)
result_card(slide,"Final answer: option 4 — y = (x/t) sin(λ t)",XL,Y0+0.45+2*LH+2*NG,CW,0.48)

# NEW FUNDAMENTALS
slide=new_slide("Theory","New fundamental quantities",GREEN)
para(slide,"Sometimes the question says: treat force, acceleration and time as the basic set. Then rewrite length or energy using only those letters. You assume length = F^α a^β t^γ, put dimensions of both sides, and match powers of M, L, T.",XL,Y0,CW,1.50,16)
para(slide,"A fast board shortcut: start from a formula you already trust. From s = (1/2) a t^2 you see length looks like a t^2. Force does not appear, so F^0 a^1 t^2.",XL,Y0+1.65,CW,1.35,16)
ncert(slide,"Assume powers → equate M, L, T → solve α, β, γ",XL,Y0+3.20,CW,0.95,16)

slide=question_slide("Length in terms of F, a, t")
qtx(slide,"If F, a and t are fundamental, dimension of length is",XL,Y0,CW,0.95,17)
choices(slide,["F^0 a^1 T^2","F a^2 T^2","F a^2 T^0","F^0 a T"], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Shortcut: length from kinematics looks like a t^2.",XL,Y0,CW,0.40)
fml(slide,"s ∝ a t^2  ⇒  [length] = F^0 a^1 T^2",XL,Y0+0.45,CW,LH,16)
fml(slide,"Full method: L = F^α a^β t^γ gives α=0, β=1, γ=2.",XL,Y0+0.45+LH+NG,CW,LH,14)
result_card(slide,"Final answer: F^0 a^1 T^2",XL,Y0+0.45+2*LH+2*NG,CW,0.48)

slide=question_slide("Energy in terms of F, a, t")
qtx(slide,"If force, acceleration and time are basic, dimension of energy is",XL,Y0,CW,1.00,16)
choices(slide,["F^2 A^(-1) T","F A T^2","F A T^(-2)","F A^(-1) T"], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Energy = force × distance, and distance ~ a t^2.",XL,Y0,CW,0.42)
fml(slide,"E = F × l = F × (a t^2) = F A T^2",XL,Y0+0.48,CW,LH,16)
result_card(slide,"Final answer: F A T^2",XL,Y0+0.48+LH+NG,CW,0.48)

slide=question_slide("Surface tension from E, V, T")
qtx(slide,"If energy E, velocity V and time T are fundamental, surface tension is",XL,Y0,CW,1.05,15)
choices(slide,["[E V^(-2) T^(-2)]","[E^(-2) V T^(-2)]","[E^(-2) V^(-2) T^2]","[E V^(-2) T^2]"], XL, Y0+1.20, 10.8)
slide=answer_slide()
body(slide,"Surface tension S = F/l. Rewrite using energy and length = V T.",XL,Y0,CW,0.45)
fml(slide,"S = F/l = (F l)/l^2 = E / l^2",XL,Y0+0.50,CW,LH,15)
fml(slide,"l = V T  ⇒  S = E / (V T)^2 = E V^(-2) T^(-2)",XL,Y0+0.50+LH+NG,CW,LH,14)
result_card(slide,"Final answer: [E V^(-2) T^(-2)]",XL,Y0+0.50+2*LH+2*NG,CW,0.48)

# DEDUCE RELATION
slide=new_slide("Theory","Deducing a relation by powers",TEAL)
para(slide,"Suppose critical velocity depends on viscosity η, density ρ and radius r. Write v_c = η^x ρ^y r^z. Put dimensions of every letter. Match powers of M, L and T. Solve the three equations for x, y, z.",XL,Y0,CW,1.50,16)
para(slide,"Board tip: often attack the mass power first. It is usually the simplest equation.",XL,Y0+1.65,CW,1.00,16)
ncert(slide,"Q = k A^a B^b C^c     →     equate [M], [L], [T]",XL,Y0+2.85,CW,0.95,16)

slide=question_slide("Critical velocity exponents")
qtx(slide,"If v_c = [η^x ρ^y r^z], find x, y, z.  (AIPMT style)",XL,Y0,CW,1.00,16)
choices(slide,["1, 1, 1","1, −1, −1","−1, −1, 1","−1, 1, −1"], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Need [η] first. From F = 6π η r v, [η] = M L^(-1) T^(-1).",XL,Y0,CW,0.45)
fml(slide,"L T^(-1) = (M L^(-1) T^(-1))^x (M L^(-3))^y (L)^z",XL,Y0+0.50,CW,LH,13)
fml(slide,"Time: −1 = −x ⇒ x = 1.  Mass: 0 = x+y ⇒ y = −1.",XL,Y0+0.50+LH+NG,CW,LH,14)
fml(slide,"Length: 1 = −x − 3y + z ⇒ z = −1.",XL,Y0+0.50+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: x=1, y=−1, z=−1",XL,Y0+0.50+3*LH+3*NG,CW,0.45)

slide=question_slide("Spring frequency powers")
qtx(slide,"f = c m^x K^y. Find x and y. (c dimensionless)",XL,Y0,CW,0.95,16)
choices(slide,[
"x = 1/2, y = 1/2",
"x = −1/2, y = −1/2",
"x = 1/2, y = −1/2",
"x = −1/2, y = 1/2",
], XL, Y0+1.10, 10.5)
slide=answer_slide()
body(slide,"[f] = T^(-1). From F = K x, [K] = M T^(-2).",XL,Y0,CW,0.42)
fml(slide,"T^(-1) = M^x (M T^(-2))^y  ⇒  −1 = −2y ⇒ y = 1/2",XL,Y0+0.48,CW,LH,14)
fml(slide,"Mass: 0 = x + y ⇒ x = −1/2",XL,Y0+0.48+LH+NG,CW,LH,16)
result_card(slide,"Final answer: x = −1/2 , y = 1/2",XL,Y0+0.48+2*LH+2*NG,CW,0.48)

slide=question_slide("NEET: length from h, G, c")
qtx(slide,"Which combination of h, G, c has dimensions of length? [NEET]",XL,Y0,CW,1.00,15)
choices(slide,[
"√(h G) / c^(3/2)",
"√(h G) / c^(5/2)",
"√(h c / G)",
"√(G c) / h^(3/2)",
], XL, Y0+1.15, 10.8)
slide=answer_slide()
body(slide,"Tip: first match mass. [h]=M L^2 T^(-1), [G]=M^(-1) L^3 T^(-2).",XL,Y0,CW,0.45)
fml(slide,"√(h G) kills mass. Remaining length and time need c^(3/2).",XL,Y0+0.50,CW,LH,14)
fml(slide,"Check: √(h G)/c^(3/2) → L^1. That is option 1.",XL,Y0+0.50+LH+NG,CW,LH,15)
result_card(slide,"Final answer: √(h G) / c^(3/2)",XL,Y0+0.50+2*LH+2*NG,CW,0.48)

# CONSTANTS CLASSIFICATION
slide=new_slide("Theory","Four kinds of quantities (board map)",GREEN)
para(slide,"Before significant figures, lock this language. A quantity may or may not have dimensions, and it may or may not change from problem to problem.",XL,Y0,CW,1.20,16)
mini_card(slide,"Dimensional variable","Has dimensions. Value can change. Example: force, velocity.",XL,Y0+1.40,5.80,1.55,TEAL)
mini_card(slide,"Dimensionless variable","No dimensions. Value can change. Example: strain, angle in a problem.",XL+6.20,Y0+1.40,5.80,1.55,ORANGE)
mini_card(slide,"Dimensional constant","Has dimensions. Fixed value. Example: G, h, c.",XL,Y0+3.15,5.80,1.55,GREEN)
mini_card(slide,"Dimensionless constant","No dimensions. Fixed number. Example: π, e, 2 in formulas.",XL+6.20,Y0+3.15,5.80,1.55,AMBER)

slide=question_slide("Which is a dimensional constant?")
qtx(slide,"Which of the following is a dimensional constant?",XL,Y0,CW,0.90,18)
choices(slide,["Relative density","Gravitational constant","Refractive index","Poisson’s ratio"], XL, Y0+1.10, 10.5)
slide=answer_slide()
body(slide,"Relative density, μ, Poisson’s ratio are ratios — dimensionless.",XL,Y0,CW,0.45)
fml(slide,"G has a fixed value and dimensions M^(-1) L^3 T^(-2).",XL,Y0+0.50,CW,LH,15)
result_card(slide,"Final answer: gravitational constant",XL,Y0+0.50+LH+NG,CW,0.48)

# LIMITATIONS
slide=new_slide("Theory","Limitations of dimensional analysis",TEAL)
para(slide,"Dimensions are powerful, but they are not magic. They cannot invent a pure number like 2π. They cannot invent sin or cos. They struggle when two different quantities share the same dimensions. And they cannot build a formula that is a sum of two different kinds of terms.",XL,Y0,CW,1.55,16)
body(slide,"1. Cannot find dimensionless constants (2π, 1/2, e).",XL,Y0+1.70,CW,0.42)
body(slide,"2. Cannot derive trig / exponential pieces of a formula.",XL,Y0+2.20,CW,0.42)
body(slide,"3. Fails if the quantity depends on more than three independent variables (with only M,L,T).",XL,Y0+2.70,CW,0.48)
body(slide,"4. Cannot derive formulas that need + or − of unlike structures (like s = ut + ½at^2).",XL,Y0+3.25,CW,0.50)
trap_card(slide,"Dimensionally possible does not mean physically true.",XL,Y0+3.95,CW,0.48)

slide=question_slide("Which equation can be derived dimensionally?")
qtx(slide,"Which of the following can be derived using dimensions alone?",XL,Y0,CW,0.95,16)
choices(slide,[
"s = v t − (1/2) a t^2",
"v^2 = u^2 − 2 a s",
"h = ω^2 r^2 / (2 g)",
"v = d / t",
], XL, Y0+1.15, 10.5)
slide=answer_slide()
body(slide,"Sums with two terms, or unknown numbers like 1/2 and 2, cannot be fixed by dimensions.",XL,Y0,CW,0.48)
fml(slide,"v = d/t is a single product/quotient. Dimensions can support it.",XL,Y0+0.55,CW,LH,14)
result_card(slide,"Best pick among these: v = d/t  (option 4)",XL,Y0+0.55+LH+NG,CW,0.48)

# SIGNIFICANT FIGURES
slide=new_slide("Idea","Significant figures — why they exist",GREEN)
para(slide,"A measurement is never infinitely precise. Significant figures tell us which digits in a reported number are trustworthy. They are the digits that carry real information from the measuring instrument.",XL,Y0,CW,1.40,16)
para(slide,"Example idea: if a scale can only read to 0.1 cm, writing 12.345 cm pretends we know more than we do. Significant-figure rules keep our writing honest.",XL,Y0+1.55,CW,1.30,16)
ncert(slide,"Significant digits = the digits that count in a measured value",XL,Y0+3.05,CW,0.95,15)

slide=new_slide("CBSE write","Rules to count significant digits",TEAL)
body(slide,"1. Every non-zero digit is significant.  (234 → 3)",XL,Y0,CW,0.45)
body(slide,"2. Zeros trapped between non-zero digits are significant.  (402 → 3)",XL,Y0+0.50,CW,0.48)
body(slide,"3. Trailing zeros with no decimal point are not significant.  (5280 → 3)",XL,Y0+1.00,CW,0.48)
body(slide,"4. Trailing zeros after a decimal point are significant.  (23.500 → 5)",XL,Y0+1.50,CW,0.48)
body(slide,"5. Leading zeros before the first non-zero digit are not significant.  (0.034 → 2)",XL,Y0+2.00,CW,0.50)
body(slide,"6. Exact counted numbers have infinite significant figures.  (12 eggs)",XL,Y0+2.50,CW,0.45)
body(slide,"7. In a × 10^n form, only the digits in a are significant.  (1.608×10^23 → 4)",XL,Y0+3.00,CW,0.50)
trap_card(slide,"The power of ten itself is not a significant digit.",XL,Y0+3.65,CW,0.48)

slide=new_slide("Copy","Counting table (board examples)",GREEN)
grid_table(slide,
["Value", "Significant digits", "Why"],
[
["23.500 m", "5", "zeros after decimal count"],
["402 cm", "3", "zero between non-zeros"],
["5280 kg", "3", "no decimal — trailing zero drops"],
["6.780", "4", "trailing zero after decimal"],
["0.034 g", "2", "leading zeros do not count"],
["3.070 g", "4", "middle + trailing after decimal"],
["1.608 × 10^23", "4", "only the mantissa"],
],
XL, Y0, [3.6, 3.2, 5.2], 0.48, 13, 14)

slide=question_slide("Count the significant digits")
qtx(slide,"How many significant digits are in 0.0340?",XL,Y0,CW,0.90,18)
choices(slide,["2","3","4","5"], XL, Y0+1.10, 10.5)
slide=answer_slide()
body(slide,"Leading zeros before 3 do not count. Trailing zero after decimal does.",XL,Y0,CW,0.45)
fml(slide,"Digits that count: 3, 4, 0  →  three significant digits.",XL,Y0+0.50,CW,LH,15)
result_card(slide,"Final answer: 3",XL,Y0+0.50+LH+NG,CW,0.48)

slide=question_slide("Scientific notation count")
qtx(slide,"Number of significant digits in 2.50 × 10^(-4) is",XL,Y0,CW,0.95,17)
choices(slide,["1","2","3","4"], XL, Y0+1.10, 10.5)
slide=answer_slide()
body(slide,"Look only at 2.50. The power −4 is not counted.",XL,Y0,CW,0.42)
fml(slide,"2, 5 and trailing 0 after decimal → three significant digits.",XL,Y0+0.48,CW,LH,15)
result_card(slide,"Final answer: 3",XL,Y0+0.48+LH+NG,CW,0.48)

# SUMMARY
slide=new_slide("Class summary","Units 03 lock lines",TEAL)
fml(slide,"Inside sin / log / exp → dimensionless. Outside factor matches left side.",XL,Y0,CW,LH,14)
fml(slide,"New fundamentals: assume powers, equate M L T, or use a known formula shortcut.",XL,Y0+LH+NG,CW,LH,13)
fml(slide,"Dimensional constant example: G. Dimensionless constant example: π.",XL,Y0+2*LH+2*NG,CW,LH,14)
fml(slide,"Limit: cannot get pure numbers or invent +/− multi-term formulas.",XL,Y0+3*LH+3*NG,CW,LH,14)
result_card(slide,"Significant figures: count only the trustworthy digits of a measurement.",XL,Y0+4*LH+4*NG,CW,0.50)

slide=new_slide("CBSE write","Ready theory lines",GREEN)
body(slide,"1. Arguments of trigonometric, logarithmic and exponential functions are dimensionless.",XL,Y0,CW,0.50)
body(slide,"2. Dimensional analysis can check a formula, convert units, and find powers in a product law.",XL,Y0+0.55,CW,0.50)
body(slide,"3. It cannot determine dimensionless constants or formulas with + / − of unlike terms.",XL,Y0+1.10,CW,0.50)
body(slide,"4. Significant figures are the meaningful digits in a measured number.",XL,Y0+1.65,CW,0.45)
body(slide,"5. Leading zeros are never significant; trailing zeros after a decimal are.",XL,Y0+2.20,CW,0.50)
trap_card(slide,"For 3 marks: definition + one use + one limitation of dimensional analysis.",XL,Y0+2.90,CW,0.50)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. If y = A e^(Bx + C/t), find [A], [B] and [C].",XL,Y0,CW,0.45)
body(slide,"2. If F, A, D are fundamental, write Young’s modulus in that set.",XL,Y0+0.55,CW,0.50)
body(slide,"3. Write five limitations of dimensional analysis.",XL,Y0+1.10,CW,0.45)
body(slide,"4. Count significant digits: 0.002040 and 2.040 × 10^3.",XL,Y0+1.65,CW,0.45)
body(slide,"5. Which is dimensionless: G, μ, strain, Planck’s constant?",XL,Y0+2.20,CW,0.45)
trap_card(slide,"Next: more on errors and measuring instruments (later PW lectures).",XL,Y0+2.90,CW,0.50)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
