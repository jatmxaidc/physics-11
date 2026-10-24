"""
Lecture 06 — PW Basic Maths and Calculus Class 05 (full class)
10-minute formula recap, then graph, coordinate, distance, slope and straight line.
Class 11 Physics  •  CBSE + NEET Foundation  •  2-hour session
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "class notes" / "Lecture_06_AP_GP_Straight_Line_Graphs.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.32
NG = 0.18
PT_SCALE = 0.80


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = (
        text.replace("PW 04 — ", "")
        .replace("PW Class 05 homework — ", "Homework — ")
        .replace("PW Class 05", "")
        .replace("PW intercept example: ", "")
        .replace("PW example: ", "")
    )
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


def question_slide(title):
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title="Simple solution"):
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.40, size=16):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, size, QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


# ── COVER ─────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx_3d(slide, "EXCELLENCE", XL, 0.55, CW, 1.05, 62, GOLD, RGBColor(210, 210, 220))
tx_3d(slide, "CLASS 11   ·   2026–27", XL, 1.62, CW, 0.55, 30, WHITE, RGBColor(90, 90, 100), font="Aptos")
_solid(slide, 4.15, 2.28, 5.00, 0.028, WHITE)
bx, by, bw = 2.05, 2.70, 9.20
_card_bg(slide, bx, by, bw, 2.35, WHITE, WHITE, 0)
_card_bg(slide, bx, by + 0.72, bw, 1.63, BLACK, RGBColor(220, 220, 220), 1.25)
_solid(slide, bx + bw / 2, by + 0.16, 0.018, 0.42, RGBColor(40, 40, 40))
tx(slide, "Physics", bx, by + 0.16, bw / 2 - 0.10, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Lecture - 06", bx + bw / 2, by + 0.16, bw / 2, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Basic Maths and Calculus", bx, by + 1.05, bw, 0.55, 32, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False, font="Georgia")
tx(slide, "A.P., G.P. and Straight-Line Graphs", bx, by + 1.62, bw, 0.40, 18, GOLD, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "By — Excellence Programme", XL, 5.30, CW, 0.42, 25, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "9 September 2026", XL, 5.80, CW, 0.32, 18, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, 1)
add_notes(slide, "Ten-minute recap, then sequences and straight-line graphs.")
after_content_slide()

slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx(slide, "Topics to be covered", XL, 0.38, CW, 0.55, 32, GOLD, True, spa=0, scaled=False)
_solid(slide, XL, 0.98, 3.40, 0.035, GOLD)
topics = [
    "Last class in 10 minutes — trigonometry, max–min, phase and superposition formulae",
    "Arithmetic progression — common difference, nth term and sum",
    "Geometric progression — common ratio, nth term and infinite sum",
    "Coordinates, distance and slope",
    "Equation of straight line y = mx + c; slope and intercept",
    "Draw lines and use them in velocity, ideal-gas and temperature graphs",
]
for i, item in enumerate(topics):
    yy = 1.18 + i * 0.78
    tx(slide, f"{i + 1:02d}", XL, yy, 0.70, 0.50, 23, GOLD, True, spa=0, scaled=False)
    tx(slide, item, XL + 0.90, yy, CW - 0.90, 0.70, 18, WHITE, False, spa=0, scaled=False)
_footer(slide, 2)
add_notes(slide, "Topics. Keep recap to ten minutes.")
after_content_slide()

# ── 10-MINUTE RECAP ───────────────────────────────────────────────────────────
slide = new_slide("10-minute recap", "Phase and superposition — formulae only", BLUE)
add_notes(slide, "No re-teaching. Say each line once.")
fml(slide, "phase = complete angle inside sin or cos", XL, Y0, CW, LH, 20)
fml(slide, "phase difference: first make both sine or both cosine, then subtract", XL, Y0 + LH + NG, CW, LH, 18)
ncert(slide, "a sin(ωt) + b cos(ωt) = R sin(ωt+φ)     R=√(a²+b²), tanφ=b/a", XL, Y0 + 2 * LH + 2 * NG, CW, 1.00, 17)
result_card(slide, "4 sin(ωt) + 3 cos(ωt) = 5 sin(ωt + 37°)", XL, Y0 + 3 * LH + 2 * NG + 0.90, CW)

slide = new_slide("10-minute recap", "Trigonometry and max–min — formulae only", BLUE)
add_notes(slide, "No AP or GP here because it was not taught. Say only previously completed formulae.")
fml(slide, "180° = π rad          arc = Rθ          θ must be radian", XL, Y0, CW, LH, 18)
fml(slide, "sin(A+B) = sinA cosB + cosA sinB", XL, Y0 + LH + NG, CW, LH, 18)
fml(slide, "sinθ ≈ θ, tanθ ≈ θ, cosθ ≈ 1     for θ≤5°, θ in radian", XL, Y0 + 2 * LH + NG, CW, LH, 17)
ncert(slide, "a sinθ + b cosθ has maximum √(a²+b²)", XL, Y0 + 3 * LH + 2 * NG, CW, 0.90, 20)

slide = new_slide("Learning path", "What we will learn today", AMBER)
add_notes(slide, "Introduce today's topics.")
body(slide, "1. Arithmetic progression: identify d, find a term and calculate a sum.", XL, Y0, CW)
body(slide, "2. Geometric progression: identify r, find a term and calculate an infinite sum.", XL, Y0 + LH + NG, CW)
body(slide, "3. Coordinates, distance and slope.", XL, Y0 + 2 * LH + 2 * NG, CW)
body(slide, "4. Equation and graph of a straight line.", XL, Y0 + 3 * LH + 3 * NG, CW)

# ── UNTaught PW LECTURE 04 FROM SLIDE 17 ──────────────────────────────────────
slide = new_slide("PW 04 — Copy", "A.P. — equal difference every time", TEAL)
add_notes(slide, "PW slide 17. Define arithmetic progression, nth term and sum.")
body(slide, "Arithmetic progression is a sequence in which the difference between consecutive terms is always constant.", XL, Y0, CW)
fml(slide, "a,  a+d,  a+2d,  a+3d,  ...", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "Tₙ = a + (n−1)d", XL, Y0 + 2 * LH + 2 * NG, 5.75, 0.95, 22)
ncert(slide, "Sₙ = n/2 [2a + (n−1)d]", XL + 6.0, Y0 + 2 * LH + 2 * NG, 6.05, 0.95, 19)
body(slide, "a = first term.  d = next term − previous term.  n = position.", XL, Y0 + 3 * LH + 2 * NG + 0.85, CW)

slide = question_slide("A.P. example: 2, 5, 8, 11, 14, ...")
add_notes(slide, "PW slide 18. a2 d3 T10 29 S10 155.")
qtx(slide, "Find a, d, the 10th term and the sum of the first 10 terms.", XL, Y0, CW, 0.42, 18)
slide = answer_slide()
fml(slide, "a=2          d=5−2=3", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "T₁₀ = 2 + 9(3) = 29", XL, Y0 + 0.52 + LH, CW, LH, 20)
fml(slide, "S₁₀ = 10/2 (2+29) = 5×31 = 155", XL, Y0 + 0.52 + 2 * LH, CW, LH, 20)
result_card(slide, "T₁₀=29          S₁₀=155", XL, Y0 + 0.52 + 3 * LH + NG, CW, 0.36, 22)

slide = question_slide("Question: 7, 11, 15, 19, ...")
add_notes(slide, "PW slide 19. T10 43 S10 250.")
qtx(slide, "Find the 10th term and sum of the first 10 terms.", XL, Y0, CW, 0.40, 18)
slide = answer_slide()
fml(slide, "a=7          d=4", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "T₁₀ = 7 + 9(4) = 43", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "S₁₀ = 10/2 (7+43) = 250", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "10th term=43          sum=250", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = question_slide("Which term of 27, 24, 21, ... is zero?")
add_notes(slide, "PW slide 20. a27 d-3 n10.")
qtx(slide, "Which term of the A.P. 27, 24, 21, ... is zero?", XL, Y0, CW, 0.40, 18)
slide = answer_slide()
fml(slide, "a=27          d=−3          Tₙ=0", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "0 = 27 + (n−1)(−3) = 30−3n", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "3n=30          n=10", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "The 10th term is zero.", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = question_slide("Question: 28, 22, x, y, 4")
add_notes(slide, "PW slide 21. d-6 x16 y10.")
qtx(slide, "28, 22, x, y, 4 is an A.P. Find x and y.", XL, Y0, CW, 0.40, 18)
slide = answer_slide()
fml(slide, "a=28          d=22−28=−6", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "x=T₃=28+2(−6)=16", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "y=T₄=28+3(−6)=10", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "x=16          y=10", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = question_slide("Sum of first n natural numbers")
add_notes(slide, "PW slide 22 MCQ. Option 2.")
qtx(slide, "Find 1+2+3+...+n.", XL, Y0, CW, 0.38, 18)
choices(slide, ["n²/2", "n(n+1)/2", "n(n−1)/2", "n²(n+1)/6"], XL, Y0 + 0.46)
slide = answer_slide()
fml(slide, "a=1, d=1, last term=n          Sₙ=n/2(1+n)", XL, Y0 + 2.20, CW, LH, 20)
result_card(slide, "Option 2     n(n+1)/2", XL, Y0 + 2.20 + LH + NG, CW, 0.36, 22)

slide = question_slide("Physics question: 25 tuning forks")
add_notes(slide, "PW slide 23. a160 Hz.")
qtx(slide, "25 frequencies form an increasing A.P. Difference=10 Hz; 25th frequency=400 Hz. Find the first.", XL, Y0, CW, 0.50, 16)
slide = answer_slide()
fml(slide, "T₂₅ = a + 24d", XL, Y0 + 0.60, CW, LH, 20)
fml(slide, "400 = a + 24(10)          a=400−240=160", XL, Y0 + 0.60 + LH, CW, LH, 20)
result_card(slide, "First tuning-fork frequency = 160 Hz", XL, Y0 + 0.60 + 2 * LH + NG, CW, 0.36, 22)

slide = new_slide("PW 04 — Copy", "G.P. — equal ratio every time", TEAL)
add_notes(slide, "PW slides 24-25. Define GP, nth term and infinite sum.")
body(slide, "Geometric progression is a sequence in which the ratio of consecutive terms is constant.", XL, Y0, CW)
fml(slide, "a,  ar,  ar²,  ar³,  ...", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "Tₙ = arⁿ⁻¹", XL, Y0 + 2 * LH + 2 * NG, 5.75, 0.95, 22)
ncert(slide, "S∞ = a/(1−r), only if |r|<1", XL + 6.0, Y0 + 2 * LH + 2 * NG, 6.05, 0.95, 18)
body(slide, "r = any term ÷ previous term.", XL, Y0 + 3 * LH + 2 * NG + 0.85, CW)

slide = question_slide("Which sequences are G.P.?")
add_notes(slide, "PW slide 26. Include all four rows.")
qtx(slide, "Check the ratio of consecutive terms. Which sequences are G.P.?", XL, Y0, CW, 0.42, 18)
body(slide, "A.  2, 4, 8, 12, 24, 48", XL, Y0 + 0.58, CW)
body(slide, "B.  2, 4, 8, 16, 32, 64", XL, Y0 + 1.08, CW)
body(slide, "C.  1, 1/2, 1/4, 1/8, ...", XL, Y0 + 1.58, CW)
body(slide, "D.  1, −1/2, 1/4, −1/8, ...", XL, Y0 + 2.08, CW)
slide = answer_slide()
grid_table(slide, ["sequence", "ratio", "answer"],
           [["2,4,8,12,24,48", "ratio changes", "not G.P."],
            ["2,4,8,16,32,64", "2", "G.P."],
            ["1,1/2,1/4,1/8,...", "1/2", "G.P., S∞=2"],
            ["1,−1/2,1/4,−1/8,...", "−1/2", "G.P., S∞=2/3"]],
           XL, Y0, [5.0, 3.0, 4.0], 0.62)
trap_card(slide, "A sequence is G.P. only if every consecutive ratio is the same.", XL, Y0 + 5 * 0.62 + 0.20, CW)

slide = question_slide("Infinite G.P.: 1 + 1/3 + 1/9 + ...")
add_notes(slide, "PW slide 27. Answer 3/2.")
qtx(slide, "Find 1 + 1/3 + 1/9 + 1/27 + ...", XL, Y0, CW, 0.40, 18)
slide = answer_slide()
fml(slide, "a=1          r=1/3          |r|<1", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "S∞ = a/(1−r) = 1/(1−1/3) = 1/(2/3)", XL, Y0 + 0.50 + LH, CW, LH, 20)
result_card(slide, "S∞ = 3/2", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 24)

slide = question_slide("Potential at origin — alternating charges")
add_notes(slide, "PW slides 28 and 32. Answer 2kq/3.")
qtx(slide, "+q, −q, +q, ... are at x=1,2,4,8,... . Find potential at origin.", XL, Y0, CW, 0.45, 17)
slide = answer_slide()
fml(slide, "V = kq [1 − 1/2 + 1/4 − 1/8 + ...]", XL, Y0 + 0.55, CW, LH, 20)
fml(slide, "a=1          r=−1/2          S∞=1/[1−(−1/2)]=2/3", XL, Y0 + 0.55 + LH, CW, LH, 18)
result_card(slide, "V = 2kq/3", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 24)

slide = question_slide("Infinite resistance network")
add_notes(slide, "PW slide 29. R+R/2+...=2R.")
qtx(slide, "Equivalent series blocks are R, R/2, R/4, R/8, ... . Find R_AB.", XL, Y0, CW, 0.42, 18)
slide = answer_slide()
fml(slide, "R_AB = R [1+1/2+1/4+1/8+...]", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "R_AB = R × 1/(1−1/2)", XL, Y0 + 0.52 + LH, CW, LH, 20)
result_card(slide, "R_AB = 2R", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.36, 24)

slide = question_slide("Force on q₀ at the origin")
add_notes(slide, "PW slide 30 homework now solved. Charges q at 1,2,4,... Force ratio is 1/4.")
qtx(slide, "Charges q lie at x=1,2,4,8,... . Find force on q₀ at origin.", XL, Y0, CW, 0.42, 18)
slide = answer_slide()
fml(slide, "F = kqq₀ [1/1² + 1/2² + 1/4² + ...]", XL, Y0 + 0.52, CW, LH, 18)
fml(slide, "F = kqq₀ [1 + 1/4 + 1/16 + ...] = kqq₀/(1−1/4)", XL, Y0 + 0.52 + LH, CW, LH, 18)
result_card(slide, "|F| = 4kqq₀/3     (toward −x for like charges)", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.36, 20)

slide = question_slide("Fifth resistance when each is half")
add_notes(slide, "PW slide 31 homework now solved. 10/16=5/8 ohm.")
qtx(slide, "Five resistors are in series. Each resistance is half the previous. First is 10 Ω. Find the fifth.", XL, Y0, CW, 0.48, 17)
slide = answer_slide()
fml(slide, "a=10 Ω          r=1/2          T₅=ar⁴", XL, Y0 + 0.58, CW, LH, 20)
fml(slide, "R₅ = 10(1/2)⁴ = 10/16", XL, Y0 + 0.58 + LH, CW, LH, 20)
result_card(slide, "R₅ = 5/8 Ω = 0.625 Ω", XL, Y0 + 0.58 + 2 * LH + NG, CW, 0.36, 22)

slide = new_slide("5-minute lock", "A.P. and G.P. are now complete", TEAL)
add_notes(slide, "Say all four lines before beginning PW Class 05.")
fml(slide, "A.P.:  Tₙ=a+(n−1)d          Sₙ=n/2[2a+(n−1)d]", XL, Y0, CW, LH, 18)
fml(slide, "G.P.:  Tₙ=arⁿ⁻¹          S∞=a/(1−r), |r|<1", XL, Y0 + LH + NG, CW, LH, 18)
result_card(slide, "A.P. → constant difference          G.P. → constant ratio", XL, Y0 + 2 * LH + 2 * NG, CW)
body(slide, "Use the difference for A.P. and the ratio for G.P.", XL, Y0 + 3 * LH + 3 * NG, CW)

# ── GRAPH AND COORDINATES ─────────────────────────────────────────────────────
slide = new_slide("Copy", "Graph — a picture of relation", TEAL)
add_notes(slide, "PW page 5. Define graph, axes and origin.")
body(slide, "A graph is a diagram in which a line or curve shows the relation between two physical quantities.", XL, Y0, 6.2)
fml(slide, "x-axis: horizontal          y-axis: vertical", XL, Y0 + LH + NG, 6.2, LH, 18)
fml(slide, "origin: O(0,0)", XL, Y0 + 2 * LH + NG, 6.2, LH, 20)
ox, oy, _, _ = axes(slide)
tx(slide, "O(0,0)", ox + 0.08, oy + 0.08, 0.8, 0.24, 13, GREEN, True, spa=0)
trap_card(slide, "Always write the physical quantity and its unit on each axis.", XL, Y0 + 3 * LH + 2 * NG, 6.2)

slide = new_slide("Copy", "Constant lines — do not swap them", TEAL)
add_notes(slide, "PW pages 5-6. y constant horizontal, x constant vertical.")
fml(slide, "y = 2          →          horizontal line", XL, Y0, 6.1, LH, 22)
body(slide, "x changes but y stays 2.", XL, Y0 + LH, 6.1)
fml(slide, "x = 3          →          vertical line", XL, Y0 + 2 * LH + NG, 6.1, LH, 22)
body(slide, "y changes but x stays 3.", XL, Y0 + 3 * LH + NG, 6.1)
ox, oy, _, _ = axes(slide)
line_on_axes(slide, ox, oy, -1.5, 1.0, 2.2, 1.0, TEAL)
line_on_axes(slide, ox, oy, 1.1, -1.8, 1.1, 1.8, GOLD)
tx(slide, "y = 2", ox + 1.55, oy - 1.25, 0.8, 0.24, 13, TEAL, True, spa=0)
tx(slide, "x = 3", ox + 1.20, oy + 1.45, 0.8, 0.24, 13, GOLD, True, spa=0)

slide = question_slide("Locate four points")
add_notes(slide, "PW page 7. A Q1, B Q2, C Q4, D Q3.")
qtx(slide, "Locate A(2,4), B(−2,3), C(4,−2), D(−1,−2).", XL, Y0, 6.1, 0.40, 18)
slide = answer_slide()
grid_table(slide, ["point", "x sign", "y sign", "quadrant"],
           [["A(2,4)", "+", "+", "I"], ["B(−2,3)", "−", "+", "II"],
            ["C(4,−2)", "+", "−", "IV"], ["D(−1,−2)", "−", "−", "III"]],
           XL, Y0 + 0.50, [2.2, 1.2, 1.2, 1.5], 0.55)
ox, oy, _, _ = axes(slide)
for label, dx, dy in [("A", 1.0, 1.4), ("B", -1.0, 1.1), ("C", 1.6, -0.8), ("D", -0.5, -0.8)]:
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ox + dx - 0.06), Inches(oy - dy - 0.06), Inches(0.12), Inches(0.12))
    dot.fill.solid(); dot.fill.fore_color.rgb = GREEN; dot.line.fill.background()
    tx(slide, label, ox + dx + 0.08, oy - dy - 0.12, 0.4, 0.22, 13, GREEN, True, spa=0)

slide = new_slide("Copy", "Distance formula — two dimensions", TEAL)
add_notes(slide, "PW page 8. Pythagoras.")
body(slide, "For P₁(x₁,y₁) and P₂(x₂,y₂), make a right triangle using horizontal and vertical changes.", XL, Y0, CW)
fml(slide, "Δx = x₂ − x₁          Δy = y₂ − y₁", XL, Y0 + LH + NG, CW, LH, 20)
ncert(slide, "P₁P₂ = √[(x₂−x₁)² + (y₂−y₁)²]", XL, Y0 + 2 * LH + 2 * NG, CW, 1.00, 23)
trap_card(slide, "Keep each difference inside a bracket before squaring.", XL, Y0 + 3 * LH + 2 * NG + 0.90, CW)

slide = question_slide("Distance in three dimensions")
add_notes(slide, "PW page 8. P(2,-1,0), Q(3,1,2), answer 3.")
qtx(slide, "Find distance between P(2,−1,0) and Q(3,1,2).", XL, Y0, CW, 0.40, 18)
slide = answer_slide()
ncert(slide, "PQ = √[(x₂−x₁)² + (y₂−y₁)² + (z₂−z₁)²]", XL, Y0 + 0.50, CW, 0.95, 20)
fml(slide, "PQ = √[(3−2)² + (1+1)² + (2−0)²]", XL, Y0 + 1.60, CW, LH, 20)
fml(slide, "PQ = √(1 + 4 + 4) = √9 = 3", XL, Y0 + 1.60 + LH, CW, LH, 20)
result_card(slide, "Distance = 3 units", XL, Y0 + 1.60 + 2 * LH + NG, CW, 0.36, 22)

# ── SLOPE ─────────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Slope — rise divided by run", TEAL)
add_notes(slide, "PW page 8. Angle is measured anticlockwise from positive x axis.")
body(slide, "Slope tells how quickly y changes when x changes.", XL, Y0, CW)
ncert(slide, "m = slope = tan θ = Δy/Δx = (y₂−y₁)/(x₂−x₁)", XL, Y0 + LH + NG, CW, 1.00, 21)
body(slide, "θ is the angle made by the line with the positive x-axis, measured anticlockwise.", XL, Y0 + 2 * LH + 2 * NG + 0.90, CW)

slide = new_slide("Copy", "A straight line has the same slope everywhere", TEAL)
add_notes(slide, "PW page 9. Any two points give same delta y/delta x.")
body(slide, "Choose any two points A and B on one straight line.", XL, Y0, 6.2)
fml(slide, "m_AB = (y_B−y_A)/(x_B−x_A)", XL, Y0 + LH + NG, 6.2, LH, 20)
body(slide, "Choose another pair C and D. The answer is unchanged.", XL, Y0 + 2 * LH + NG, 6.2)
fml(slide, "m_AB = m_CD", XL, Y0 + 3 * LH + 2 * NG, 6.2, LH, 24)
ox, oy, _, _ = axes(slide)
line_on_axes(slide, ox, oy, -1.5, -1.2, 2.0, 1.7, GREEN)
trap_card(slide, "This is special to a straight line. A curve generally has changing slope.", XL, Y0 + 4 * LH + 3 * NG, 6.2)

slide = new_slide("Copy", "Parallel straight lines have equal slope", TEAL)
add_notes(slide, "PW page 10. Same angle means same tan.")
body(slide, "Parallel lines make the same angle θ with the positive x-axis.", XL, Y0, 6.2)
fml(slide, "m₁ = tan θ          m₂ = tan θ", XL, Y0 + LH + NG, 6.2, LH, 20)
ncert(slide, "parallel lines  →  m₁ = m₂", XL, Y0 + 2 * LH + 2 * NG, 6.2, 0.90, 21)
ox, oy, _, _ = axes(slide)
line_on_axes(slide, ox, oy, -1.4, -0.8, 1.6, 1.5, GOLD)
line_on_axes(slide, ox, oy, -0.8, -1.5, 2.2, 0.8, TEAL)

slide = new_slide("Copy", "Slope values you must know", TEAL)
add_notes(slide, "PW pages 11-12. Use known tan values.")
grid_table(slide, ["angle θ", "slope tan θ", "line feel"],
           [["0°", "0", "horizontal"], ["30°", "1/√3", "gentle rise"],
            ["37°", "3/4", "rise"], ["53°", "4/3", "steeper rise"],
            ["90°", "not defined / infinite", "vertical"],
            ["120°", "−√3", "falls"], ["150°", "−1/√3", "gentle fall"]],
           XL, Y0, [3.0, 4.0, 5.0], 0.55)

slide = new_slide("Copy", "Sign of slope from the angle", TEAL)
add_notes(slide, "PW page 13. Acute positive, obtuse negative.")
ncert(slide, "0° < θ < 90°     →     m > 0", XL, Y0, 5.85, 0.85, 20)
ncert(slide, "90° < θ < 180°     →     m < 0", XL + 6.1, Y0, 5.95, 0.85, 20)
fml(slide, "θ = 0°  →  m = 0          θ = 90°  →  vertical slope is not defined", XL, Y0 + 1.15, CW, LH, 18)
body(slide, "A rising line has positive slope. A falling line has negative slope.", XL, Y0 + 1.15 + LH + NG, CW)

slide = question_slide("MR exam trap — slope is not the angle")
add_notes(slide, "PW page 14. Statement false. Slope equals tan theta.")
qtx(slide, "“The angle made by a straight line with +x-axis is called slope.” True or false?", XL, Y0, CW, 0.50, 17)
slide = answer_slide()
fml(slide, "angle = θ", XL, Y0 + 0.60, CW, LH, 22)
fml(slide, "slope = tan θ", XL, Y0 + 0.60 + LH, CW, LH, 22)
result_card(slide, "False — slope is tan θ, not θ.", XL, Y0 + 0.60 + 2 * LH + NG, CW, 0.36, 22)

slide = question_slide("Ratio of slopes — 60° line and 30° line")
add_notes(slide, "PW page 15 left. Ratio root3 over 1/root3 =3.")
qtx(slide, "Line A makes 60° and line B makes 30° with +x-axis. Find m_A/m_B.", XL, Y0, CW, 0.45, 17)
slide = answer_slide()
fml(slide, "m_A/m_B = tan60°/tan30°", XL, Y0 + 0.55, CW, LH, 22)
fml(slide, "m_A/m_B = √3 ÷ (1/√3) = 3", XL, Y0 + 0.55 + LH, CW, LH, 22)
result_card(slide, "m_A : m_B = 3 : 1", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 22)

slide = question_slide("Ratio of slopes — 30° line and 60° line")
add_notes(slide, "PW page 15 right. Reverse gives one-third.")
qtx(slide, "Now line A makes 30° and line B makes 60°. Find m_A/m_B.", XL, Y0, CW, 0.45, 17)
slide = answer_slide()
fml(slide, "m_A/m_B = tan30°/tan60°", XL, Y0 + 0.55, CW, LH, 22)
fml(slide, "m_A/m_B = (1/√3) ÷ √3 = 1/3", XL, Y0 + 0.55 + LH, CW, LH, 22)
result_card(slide, "m_A : m_B = 1 : 3", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 22)

# ── EQUATION y=mx+c ───────────────────────────────────────────────────────────
slide = new_slide("Copy", "Equation of a straight line", TEAL)
add_notes(slide, "PW page 16. Derive using point x,y and intercept c.")
body(slide, "Every non-vertical straight line can be written in slope-intercept form.", XL, Y0, CW)
ncert(slide, "y = mx + c", XL, Y0 + LH + NG, CW, 1.00, 28)
fml(slide, "m = slope          c = y-intercept = value of y when x=0", XL, Y0 + 2 * LH + 2 * NG + 0.90, CW, LH, 20)
trap_card(slide, "m and c are fixed for one straight line. x and y change from point to point.", XL, Y0 + 3 * LH + 3 * NG + 0.90, CW)

slide = new_slide("Copy", "Read slope and y-intercept directly", TEAL)
add_notes(slide, "PW page 17. Compare with y=mx+c.")
body(slide, "First arrange the equation so y is alone. Then compare term by term.", XL, Y0, CW)
fml(slide, "y = (4/3)x + 4          →          m = 4/3, c = 4", XL, Y0 + LH + NG, CW, LH, 20)
fml(slide, "y = −√3 x − 4          →          m = −√3, c = −4", XL, Y0 + 2 * LH + NG, CW, LH, 20)
trap_card(slide, "The constant c is where the line cuts the y-axis, not the x-axis.", XL, Y0 + 3 * LH + 2 * NG, CW)

slide = question_slide("PW example: c = 4 and θ = 53°")
add_notes(slide, "PW page 18. m=4/3. y at x=3 is 8.")
qtx(slide, "A line cuts the y-axis at 4 and makes 53° with +x-axis. Write its equation and find y at x=3.", XL, Y0, CW, 0.50, 16)
slide = answer_slide()
fml(slide, "m = tan53° = 4/3          c = 4", XL, Y0 + 0.60, CW, LH, 20)
fml(slide, "y = (4/3)x + 4", XL, Y0 + 0.60 + LH, CW, LH, 22)
fml(slide, "at x=3:  y = (4/3)(3) + 4 = 8", XL, Y0 + 0.60 + 2 * LH, CW, LH, 20)
result_card(slide, "Equation: y=(4/3)x+4          y(3)=8", XL, Y0 + 0.60 + 3 * LH + NG, CW)

slide = question_slide("PW example: c = −2 and θ = 30°")
add_notes(slide, "PW page 19. y=x/root3-2. x=2root3 gives zero.")
qtx(slide, "A line cuts y-axis at −2 and makes 30°. Write its equation. Find y at x=2√3.", XL, Y0, CW, 0.46, 17)
slide = answer_slide()
fml(slide, "m = tan30° = 1/√3          c = −2", XL, Y0 + 0.56, CW, LH, 20)
fml(slide, "y = x/√3 − 2", XL, Y0 + 0.56 + LH, CW, LH, 22)
fml(slide, "y(2√3) = 2√3/√3 − 2 = 0", XL, Y0 + 0.56 + 2 * LH, CW, LH, 20)
result_card(slide, "y = x/√3 − 2          y(2√3)=0", XL, Y0 + 0.56 + 3 * LH + NG, CW)

slide = new_slide("Board", "Signs of m and c — four quick pictures", GREEN)
add_notes(slide, "PW page 20. Rise/fall gives m. Crossing above/below origin gives c.")
grid_table(slide, ["line behaviour", "m", "y-axis crossing", "c"],
           [["rises, crosses above origin", "+", "above", "+"],
            ["falls, crosses above origin", "−", "above", "+"],
            ["rises, crosses below origin", "+", "below", "−"],
            ["falls, crosses below origin", "−", "below", "−"]],
           XL, Y0, [4.3, 1.6, 4.3, 1.6], 0.65)
trap_card(slide, "Do not use where the line cuts the x-axis to decide c.", XL, Y0 + 5 * 0.65 + 0.18, CW)

slide = question_slide("PW example: c = 5, θ = 150°, x = −4")
add_notes(slide, "PW page 21. m=-1/root3. y=4/root3+5.")
qtx(slide, "A line has y-intercept 5 and angle 150°. Write its equation and find y at x=−4.", XL, Y0, CW, 0.46, 17)
slide = answer_slide()
fml(slide, "m = tan150° = −1/√3          c = 5", XL, Y0 + 0.56, CW, LH, 20)
fml(slide, "y = −x/√3 + 5", XL, Y0 + 0.56 + LH, CW, LH, 22)
fml(slide, "y(−4) = 4/√3 + 5", XL, Y0 + 0.56 + 2 * LH, CW, LH, 20)
result_card(slide, "y = −x/√3 + 5          y(−4)=5+4/√3", XL, Y0 + 0.56 + 3 * LH + NG, CW)

slide = question_slide("PW example: c = −4 and θ = 120°")
add_notes(slide, "PW page 22. m=-root3.")
qtx(slide, "A line cuts the y-axis at −4 and makes angle 120°. Write its equation.", XL, Y0, CW, 0.42, 18)
slide = answer_slide()
fml(slide, "m = tan120° = −√3          c = −4", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "y = mx + c", XL, Y0 + 0.52 + LH, CW, LH, 20)
result_card(slide, "y = −√3 x − 4", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.40, 24)

slide = new_slide("Copy", "Horizontal and vertical line equations", TEAL)
add_notes(slide, "PW page 23. Horizontal y=4. Vertical x=3.")
fml(slide, "horizontal through (0,4):  m=0, c=4  →  y=4", XL, Y0, CW, LH, 20)
fml(slide, "vertical through (3,0):  x=3", XL, Y0 + LH + NG, CW, LH, 20)
trap_card(slide, "A vertical line cannot be written as y=mx+c because its slope is not defined.", XL, Y0 + 2 * LH + 2 * NG, CW)

slide = question_slide("PW intercept example: x-cut = −3, y-cut = 4")
add_notes(slide, "PW page 24. Rise 4 run3 so slope 4/3.")
qtx(slide, "A line crosses the x-axis at −3 and y-axis at 4. Write its equation.", XL, Y0, CW, 0.44, 18)
slide = answer_slide()
fml(slide, "c = 4          m = rise/run = 4/3", XL, Y0 + 0.54, CW, LH, 20)
result_card(slide, "y = (4/3)x + 4", XL, Y0 + 0.54 + LH + NG, CW, 0.40, 24)

slide = question_slide("PW intercept example: x-cut = 4, y-cut = 3")
add_notes(slide, "PW page 25. Falling line. m=-3/4.")
qtx(slide, "A line crosses the x-axis at 4 and y-axis at 3. Write its equation.", XL, Y0, CW, 0.44, 18)
slide = answer_slide()
fml(slide, "c = 3          m = (0−3)/(4−0) = −3/4", XL, Y0 + 0.54, CW, LH, 20)
result_card(slide, "y = −(3/4)x + 3", XL, Y0 + 0.54 + LH + NG, CW, 0.40, 24)

# ── DRAW AND COMPARE GRAPHS ───────────────────────────────────────────────────
slide = question_slide("Draw two straight-line graphs")
add_notes(slide, "PW page 26. Read m and c, then mark intercept and angle.")
qtx(slide, "Draw  y = √3x + 4  and  y = x. First identify m and c for each line.", XL, Y0, CW, 0.50, 17)
slide = answer_slide()
grid_table(slide, ["equation", "m", "c", "angle"],
           [["y=√3x+4", "√3", "4", "60°"], ["y=x", "1", "0", "45°"]],
           XL, Y0, [4.5, 2.0, 2.0, 3.5], 0.65)
body(slide, "Method: mark (0,c), use the angle or rise/run, then extend the straight line.", XL, Y0 + 3 * 0.65 + 0.20, CW)

slide = question_slide("Draw two more straight-line graphs")
add_notes(slide, "PW page 27. m=-1 c0; m=4/3 c=-5.")
qtx(slide, "Draw  y = −x  and  y = (4/3)x − 5. First identify m and c.", XL, Y0, CW, 0.50, 17)
slide = answer_slide()
grid_table(slide, ["equation", "m", "c", "line"],
           [["y=−x", "−1", "0", "falls through origin"],
            ["y=(4/3)x−5", "4/3", "−5", "rises from below origin"]],
           XL, Y0, [4.5, 2.0, 2.0, 3.5], 0.65)
trap_card(slide, "Negative c does not automatically mean negative slope. m and c tell different things.", XL, Y0 + 3 * 0.65 + 0.25, CW)

slide = question_slide("Find intercepts before drawing")
add_notes(slide, "PW page 28. y=2x-1 cuts x at 1/2. y=-2x+4 cuts x at 2.")
qtx(slide, "Draw  y=2x−1  and  y=−2x+4.", XL, Y0, CW, 0.36, 18)
slide = answer_slide()
grid_table(slide, ["equation", "y-intercept", "x-intercept", "slope"],
           [["y=2x−1", "−1", "1/2", "+2"],
            ["y=−2x+4", "4", "2", "−2"]],
           XL, Y0 + 0.44, [4.0, 2.7, 2.7, 2.6], 0.65)
body(slide, "For x-intercept, put y=0. For y-intercept, put x=0.", XL, Y0 + 3 * 0.65 + 0.25, CW)

slide = question_slide("Convert to y = mx + c first")
add_notes(slide, "PW page 29. Answers -4/3x-4 and -x+2.")
qtx(slide, "Write each in slope-intercept form.", XL, Y0, CW, 0.34, 18)
slide = answer_slide()
fml(slide, "(1)  x/3 + y/4 + 1 = 0", XL, Y0 + 0.42, CW, LH, 18)
fml(slide, "y/4 = −x/3 − 1          y = −(4/3)x − 4", XL, Y0 + 0.42 + LH, CW, LH, 18)
result_card(slide, "m=−4/3, c=−4", XL, Y0 + 0.42 + 2 * LH + NG, CW)
fml(slide, "(2)  (y−3)+(x−4)+5=0  →  x+y−2=0  →  y=−x+2", XL, Y0 + 0.42 + 3 * LH + 2 * NG, CW, LH, 18)
result_card(slide, "m=−1, c=2", XL, Y0 + 0.42 + 4 * LH + 3 * NG, CW)

slide = question_slide("Compare two lines without drawing")
add_notes(slide, "PW page 30. Both slopes positive. Line1 c negative. Line2 c positive.")
qtx(slide, "Two lines both rise. Line 1 crosses y-axis below origin; line 2 crosses above. State signs of m and c.", XL, Y0, CW, 0.48, 16)
slide = answer_slide()
grid_table(slide, ["line", "slope m", "intercept c"],
           [["1", "positive", "negative"], ["2", "positive", "positive"]],
           XL, Y0 + 0.58, [3.0, 4.5, 4.5], 0.70)
result_card(slide, "Both m are +ve. Only line 2 has +ve c.", XL, Y0 + 0.58 + 3 * 0.70 + 0.20, CW)

# ── EQUATION FROM GIVEN INFORMATION ───────────────────────────────────────────
slide = question_slide("Y-intercept 4 and point (2,6)")
add_notes(slide, "PW page 31. Put point into y=mx+4. m=1.")
qtx(slide, "Draw a line with y-intercept 4 passing through (2,6). Find its equation.", XL, Y0, CW, 0.44, 17)
slide = answer_slide()
fml(slide, "c=4          6 = m(2)+4          2m=2          m=1", XL, Y0 + 0.54, CW, LH, 20)
result_card(slide, "y = x + 4", XL, Y0 + 0.54 + LH + NG, CW, 0.40, 24)

slide = question_slide("Point (2,3) and slope 1")
add_notes(slide, "PW page 32 question. Solve now: 3=2+c, c=1.")
qtx(slide, "Draw a line passing through (2,3) with slope 1. Find its equation.", XL, Y0, CW, 0.44, 17)
slide = answer_slide()
fml(slide, "y=mx+c          3=1(2)+c          c=1", XL, Y0 + 0.54, CW, LH, 20)
result_card(slide, "y = x + 1", XL, Y0 + 0.54 + LH + NG, CW, 0.40, 24)

slide = question_slide("Line through (3,4) and (2,6)")
add_notes(slide, "PW page 33. m=-2 c10.")
qtx(slide, "Find equation of the straight line through (3,4) and (2,6).", XL, Y0, CW, 0.42, 18)
slide = answer_slide()
fml(slide, "m = (6−4)/(2−3) = 2/(−1) = −2", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "4 = −2(3)+c          c=10", XL, Y0 + 0.52 + LH, CW, LH, 20)
result_card(slide, "y = −2x + 10", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.40, 24)

# ── PHYSICS APPLICATIONS ──────────────────────────────────────────────────────
slide = question_slide("Physics graph: v = 2t − 4")
add_notes(slide, "PW page 34. v-t graph slope2 intercept -4, crosses t=2.")
qtx(slide, "Velocity is v=2t−4. Draw the velocity–time graph.", XL, Y0, 6.2, 0.42, 18)
slide = answer_slide()
fml(slide, "compare v=mt+c", XL, Y0 + 0.52, 6.2, LH, 20)
fml(slide, "slope = 2          v-intercept = −4", XL, Y0 + 0.52 + LH, 6.2, LH, 20)
fml(slide, "v=0  →  2t−4=0  →  t=2 s", XL, Y0 + 0.52 + 2 * LH, 6.2, LH, 20)
ox, oy, _, _ = axes(slide, xlab="t", ylab="v")
line_on_axes(slide, ox, oy, -0.2, -1.5, 2.0, 1.8, GREEN)
result_card(slide, "The line crosses the time-axis at 2 s.", XL, Y0 + 0.52 + 3 * LH + NG, 6.2)

slide = question_slide("Ideal gas: P versus T at constant volume")
add_notes(slide, "PW page 35. PV=nRT. P=nR/V T, through origin.")
qtx(slide, "If PV=nRT and volume is constant, draw P versus T.", XL, Y0, 6.2, 0.42, 18)
slide = answer_slide()
fml(slide, "P = (nR/V) T", XL, Y0 + 0.52, 6.2, LH, 22)
fml(slide, "m = nR/V > 0          c=0", XL, Y0 + 0.52 + LH, 6.2, LH, 20)
body(slide, "So it is a rising straight line through the origin.", XL, Y0 + 0.52 + 2 * LH, 6.2)
ox, oy, _, _ = axes(slide, xlab="T", ylab="P")
line_on_axes(slide, ox, oy, 0, 0, 1.8, 1.7, GREEN)

slide = question_slide("Celsius and kelvin — a straight line")
add_notes(slide, "PW page 36. C/100=(K-273)/100, so C=K-273.")
qtx(slide, "Given  C/100 = (K−273)/100, draw C versus K.", XL, Y0, 6.2, 0.44, 17)
slide = answer_slide()
fml(slide, "C = K − 273", XL, Y0 + 0.54, 6.2, LH, 24)
fml(slide, "slope = 1          C-intercept = −273", XL, Y0 + 0.54 + LH, 6.2, LH, 18)
fml(slide, "C=0 at K=273", XL, Y0 + 0.54 + 2 * LH, 6.2, LH, 20)
ox, oy, _, _ = axes(slide, xlab="K", ylab="C")
line_on_axes(slide, ox, oy, -0.2, -1.5, 2.0, 1.5, GOLD)

# ── TONIGHT: EXACT PW CLASS 05 QUESTIONS ──────────────────────────────────────
slide = question_slide("Intercept question")
add_notes(slide, "Allow students to choose before showing the next slide.")
qtx(slide, "Find the equation of the line whose x-intercept is +4 and y-intercept is +5.", XL, Y0, CW, 0.48, 17)
choices(slide, ["5x + 4y − 20 = 0", "4x + 5y − 20 = 0", "4y − 5x = −20", "4x + 5y + 20 = 0"], XL, Y0 + 0.58)
slide = answer_slide()
fml(slide, "Use the two points (4,0) and (0,5).", XL, Y0, CW, LH, 20)
fml(slide, "x/4 + y/5 = 1", XL, Y0 + LH + NG, CW, LH, 22)
fml(slide, "Multiply by 20:  5x + 4y = 20", XL, Y0 + 2 * LH + 2 * NG, CW, LH, 20)
result_card(slide, "Answer: 5x + 4y − 20 = 0     (Option 1)", XL, Y0 + 3 * LH + 3 * NG, CW, 0.38, 22)

slide = question_slide("Draw four graphs")
add_notes(slide, "Allow students to draw before showing the next two slides.")
qtx(slide, "Draw graph between y and x for each equation.", XL, Y0, CW, 0.40, 18)
qtx(slide, "(i)   y = √3x + 4", XL, Y0 + 0.55, CW, 0.34, 18)
qtx(slide, "(ii)  y = −x", XL, Y0 + 1.05, CW, 0.34, 18)
qtx(slide, "(iii) y = |x|", XL, Y0 + 1.55, CW, 0.34, 18)
qtx(slide, "(iv)  x/3 + y/4 + 1 = 0", XL, Y0 + 2.05, CW, 0.34, 18)
body(slide, "For each graph, mark at least two correct points.", XL, Y0 + 2.70, CW)

slide = answer_slide("Graphs (i) and (ii)")
fml(slide, "(i) y=√3x+4:  m=√3, c=4", XL, Y0, CW, LH, 20)
body(slide, "Mark (0,4) and (1,4+√3), then join them with a straight line.", XL, Y0 + LH + NG, CW)
fml(slide, "(ii) y=−x:  m=−1, c=0", XL, Y0 + 2 * LH + 2 * NG, CW, LH, 20)
body(slide, "Mark (0,0) and (1,−1), then join them with a straight line.", XL, Y0 + 3 * LH + 3 * NG, CW)

slide = answer_slide("Graphs (iii) and (iv)")
fml(slide, "(iii) y=|x|", XL, Y0, CW, LH, 20)
body(slide, "Mark (−1,1), (0,0) and (1,1). Join them to make a V shape.", XL, Y0 + LH + NG, CW)
fml(slide, "(iv) y=−(4/3)x−4", XL, Y0 + 2 * LH + 2 * NG, CW, LH, 20)
body(slide, "Mark (0,−4) and (−3,0), then join them with a straight line.", XL, Y0 + 3 * LH + 3 * NG, CW)

slide = new_slide("Tonight", "After class", AMBER)
add_notes(slide, "Daily test and two PW homework boards.")
body(slide, "Daily test before the next lecture.", XL, Y0, CW)
body(slide, "Copy every solved question once in your notes.", XL, Y0 + LH, CW)
body(slide, "Solve the intercept MCQ and four graph questions on rough pages.", XL, Y0 + 2 * LH + NG, CW)

# ── SUMMARY ───────────────────────────────────────────────────────────────────
slide = new_slide("5-minute summary", "Coordinates, distance and slope", TEAL)
add_notes(slide, "Students say all lines.")
fml(slide, "distance₂D = √[(x₂−x₁)² + (y₂−y₁)²]", XL, Y0, CW, LH, 18)
fml(slide, "distance₃D = √[(x₂−x₁)² + (y₂−y₁)² + (z₂−z₁)²]", XL, Y0 + LH, CW, LH, 18)
ncert(slide, "m = tanθ = Δy/Δx", XL, Y0 + 2 * LH + NG, CW, 0.90, 22)
result_card(slide, "rising → m positive       falling → m negative       horizontal → 0       vertical → undefined", XL, Y0 + 3 * LH + NG + 0.80, CW, 0.40, 16)

slide = new_slide("5-minute summary", "Straight-line lock-lines", TEAL)
add_notes(slide, "Students say y=mx+c, intercept methods.")
ncert(slide, "y = mx + c", XL, Y0, CW, 0.95, 28)
fml(slide, "m = slope          c = y-intercept", XL, Y0 + 1.15, CW, LH, 20)
fml(slide, "x-intercept: put y=0          y-intercept: put x=0", XL, Y0 + 1.15 + LH, CW, LH, 18)
fml(slide, "parallel lines have equal slope", XL, Y0 + 1.15 + 2 * LH + NG, CW, LH, 20)
trap_card(slide, "Slope is tanθ, not θ. A vertical line is x=constant.", XL, Y0 + 1.15 + 3 * LH + 2 * NG, CW)

prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
