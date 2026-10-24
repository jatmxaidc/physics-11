"""
Lecture 05 — Finish last class from slide 56, then phase (2-hour class)
Last class stopped at Lecture 04 slide 55 (y = 4 sin 3θ).
Hour 1: 10-min formula recap, then leftover max–min (PW lecture 4 / class-03).
Hour 2: phase, phasor, mixed boards.
Class 11 Physics  •  CBSE + NEET Foundation
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "Lecture_05_MaxMin_Phase_Phasor.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.32
NG = 0.18
PT_SCALE = 0.80


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(
    slide, text, x, y, w, h,
    size=15, color=TEXT, bold=False,
    align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
    spa=1, lsp=1.1, font="Aptos", scaled=True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    pt = size if not scaled else max(7, round(size * PT_SCALE, 1))
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


_CONTENT = 0


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working. Write here if a question needs more space.")
    return slide


def after_content_slide():
    global _CONTENT
    _CONTENT += 1
    if _CONTENT % 5 == 0:
        add_rough_work()


def new_slide(section, title, bar_color=BLUE):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide()
    return slide


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(
        slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
        size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15,
    )


def fml(slide, text, x, y, w, h=0.44, size=22):
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.40, size=16):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, size, QCOL, True, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    if w is None:
        w = CW
    row = 0.40
    for i, item in enumerate(items):
        cy = y + i * row
        circ = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30)
        )
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)
    return y + len(items) * row


def trap_card(slide, text, x, y, w, h):
    tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h, tsize=14):
    return tx(slide, text, x, y, w, 0.36, max(tsize, 18), GREEN, True, spa=0)


def warn_row(slide, wrong, correct, x, y, w):
    tx(slide, f"Wrong:  {wrong}", x, y, w, 0.30, 16, RED, False, spa=0, scaled=False)
    tx(slide, f"Right:  {correct}", x, y + LH, w, 0.30, 16, GREEN, False, spa=0, scaled=False)
    return y + LH + LH


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=11, csize=12, scaled=True):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(
            slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
            hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=scaled,
        )
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h)
            )
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(
                slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
                csize, WORK if j == 0 else TEXT, j == 0, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                spa=0, scaled=scaled,
            )


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def phasor_axes(slide, x, y, s, xlab="sin (ωt)", ylab="cos (ωt)"):
    cx, cy = x + s, y + s
    draw_line(slide, x, cy, x + 2 * s, cy, MUTED, 1.75)
    draw_line(slide, cx, y, cx, y + 2 * s, MUTED, 1.75)
    tx(slide, xlab, x + 2 * s - 1.15, cy + 0.06, 1.20, 0.22, 12, MUTED, True, spa=0)
    tx(slide, ylab, cx + 0.08, y - 0.02, 1.20, 0.22, 12, MUTED, True, spa=0)
    return cx, cy


# ── 01 COVER ──────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx_3d(slide, "EXCELLENCE", XL, 0.55, CW, 1.05, 62, GOLD, RGBColor(210, 210, 220))
tx_3d(slide, "CLASS 11   ·   2026–27", XL, 1.62, CW, 0.55, 30, WHITE, RGBColor(90, 90, 100), font="Aptos")
_solid(slide, 4.15, 2.28, 5.00, 0.028, WHITE)

bx, by, bw = 2.05, 2.70, 9.20
_card_bg(slide, bx, by, bw, 2.35, WHITE, WHITE, 0)
_card_bg(slide, bx, by + 0.72, bw, 1.63, BLACK, RGBColor(220, 220, 220), 1.25)
_solid(slide, bx + bw / 2, by + 0.16, 0.018, 0.42, RGBColor(40, 40, 40))
tx(slide, "Physics", bx, by + 0.16, bw / 2 - 0.10, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Lecture - 05", bx + bw / 2, by + 0.16, bw / 2, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Basic Maths and Calculus", bx, by + 1.05, bw, 0.55, 32, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False, font="Georgia")
tx(slide, "Last-week formulae, then PW Lecture 04", bx, by + 1.62, bw, 0.40, 18, GOLD, True, PP_ALIGN.CENTER, spa=0, scaled=False)

tx(slide, "By — Excellence Programme", XL, 5.30, CW, 0.42, 25, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "8 September 2026", XL, 5.80, CW, 0.32, 18, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, 1)
add_notes(slide, "Cover. Last class stopped at L04 slide 55. Recap formulae 10 min. Then leftover max-min and phase.")
after_content_slide()

# ── 02 TOPICS ─────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx(slide, "Topics to be covered", XL, 0.38, CW, 0.55, 32, GOLD, True, spa=0, scaled=False)
_solid(slide, XL, 0.98, 3.40, 0.035, GOLD)
topics = [
    "Last week in 10 minutes  —  lock-lines only, no new working",
    "Last class stopped at slide 55.  Today starts from slide 56  —  projectile range",
    "Finish maximum–minimum, then PW phase / phasor and superposition",
    "PW Lecture 04  —  angle tool, arithmetic progression and geometric progression",
    "Physics uses of G.P.  —  electric potential and repeating resistance blocks",
    "Graph basics  —  axes, constant lines and distance between two points",
]
for i, line in enumerate(topics):
    y = 1.18 + i * 0.78
    tx(slide, f"{i + 1:02d}", XL, y, 0.70, 0.50, 23, GOLD, True, spa=0, scaled=False)
    tx(slide, line, XL + 0.90, y, CW - 0.90, 0.70, 18, WHITE, False, spa=0, scaled=False)
_footer(slide, 2)
add_notes(slide, "Slide 2. Recap 10 min. Then leftover from L04 slide 56.")
after_content_slide()

# ── 10-MINUTE RECAP ───────────────────────────────────────────────────────────
slide = new_slide("10-minute recap", "Last week — lock-lines only", BLUE)
add_notes(slide, "Ten minutes. They say the lines. Do not re-derive. Last class stopped at 4 sin 3θ.")
body(slide, "Last class stopped at slide 55.  Today we do not re-teach that board.  Say these lines once.", XL, Y0, CW)
fml(slide, "180°  =  π rad          Arc  =  R θ          (θ in radian)", XL, Y0 + LH + NG, CW, LH, 18)
fml(slide, "sin (A + B)  =  sin A cos B  +  cos A sin B", XL, Y0 + 2 * LH + NG, CW, LH, 18)
fml(slide, "cos (A + B)  =  cos A cos B  −  sin A sin B", XL, Y0 + 3 * LH + NG, CW, LH, 18)
fml(slide, "sin θ  ≈  θ      tan θ  ≈  θ      cos θ  ≈  1      (θ in radian,  θ ≤ 5°)", XL, Y0 + 4 * LH + 2 * NG, CW, LH, 18)
trap_card(slide, "sin 2° is not 2. Convert first:  2 × (π/180)  =  π/90.", XL, Y0 + 5 * LH + 3 * NG, CW, 0.40)

slide = new_slide("10-minute recap", "ASTC and the max–min we already locked", BLUE)
add_notes(slide, "Three ASTC steps. ymax of A sin nθ is |A| at nθ=90. Last board was 4 sin 3θ.")
fml(slide, "1.  Sign from ASTC.     2.  Function changes at 90° and 270° only.     3.  sin ↔ cos.", XL, Y0, CW, LH, 16)
fml(slide, "(sin θ)max  =  +1  at  90°          (sin θ)min  =  −1  at  270°", XL, Y0 + LH + NG, CW, LH, 18)
ncert(slide, "y  =  A sin (nθ)          ymax  =  |A|          at  nθ  =  90°", XL, Y0 + 2 * LH + 2 * NG, CW, 0.90, 18)
result_card(slide, "Last board, slide 55:   y  =  4 sin (3θ)   →   ymax  =  4     at     θ  =  30°", XL, Y0 + 3 * LH + 2 * NG + 0.80, CW, 0.36, 16)
trap_card(slide, "Do not set θ = 90°. The sine is of 3θ. Divide 90° by 3.  We start from slide 56 now.", XL, Y0 + 4 * LH + 3 * NG + 0.80, CW, 0.40)

# ── LEFTOVER MAX-MIN FROM L04 SLIDE 56 ────────────────────────────────────────
slide = new_slide("Leftover", "Last class, slide 56 — range of a projectile", GREEN)
add_notes(slide, "Read. sin2θ max at 2θ=90, θ=45. Rmax=u²/g. This is where last class stopped.")
qtx(slide, "Last class, slide 56.  If  R  =  u² sin(2θ) / g,  with u and g constant, find θ at which R is maximum.", XL, Y0, CW, 0.50, 15)
fml(slide, "Rmax when  sin(2θ)  is maximum,  so  sin(2θ)  =  1          2θ  =  90°          θ  =  45°", XL, Y0 + 0.60, CW, LH, 18)
result_card(slide, "θ  =  45°          Rmax  =  u² / g", XL, Y0 + 0.60 + LH + NG, CW, 0.36, 20)
body(slide, "This is why a level-ground projectile goes farthest at 45°.", XL, Y0 + 0.60 + 2 * LH + 2 * NG, CW)

slide = new_slide("Board", "Add like sines first", GREEN)
add_notes(slide, "Read. 3sin+4sin=7sin. ymax=7 at 90.")
qtx(slide, "If  y  =  3 sin θ  +  4 sin θ,  find ymax.", XL, Y0, CW, 0.40, 18)
fml(slide, "y  =  7 sin θ          ymax  =  7 × 1  =  7          at  θ  =  90°", XL, Y0 + 0.50, CW, LH, 18)
result_card(slide, "ymax  =  7", XL, Y0 + 0.50 + LH + NG, CW, 0.36, 22)
trap_card(slide, "Do not open √(3²+4²) here. Both terms are sine of the same angle. Add the 3 and the 4.", XL, Y0 + 0.50 + 2 * LH + 2 * NG, CW, 0.45)

slide = new_slide("Board", "A constant sitting next to the sine", GREEN)
add_notes(slide, "Read. 3+2sin: 5 and 1. 3+4sin: 7 and −1. 3sin−4: −1 and −7.")
qtx(slide, "Find ymax and ymin.", XL, Y0, CW, 0.32, 16)
fml(slide, "y  =  3 + 2 sin θ          ymax  =  3 + 2(1)  =  5          ymin  =  3 + 2(−1)  =  1", XL, Y0 + 0.40, CW, LH, 16)
fml(slide, "y  =  3 + 4 sin θ          ymax  =  7          ymin  =  3 − 4  =  −1", XL, Y0 + 0.40 + LH, CW, LH, 16)
fml(slide, "y  =  3 sin θ − 4          ymax  =  3 − 4  =  −1          ymin  =  −3 − 4  =  −7", XL, Y0 + 0.40 + 2 * LH, CW, LH, 16)
result_card(slide, "3 + 2 sin θ  →  5  and  1          3 + 4 sin θ  →  7  and  −1          3 sin θ − 4  →  −1  and  −7", XL, Y0 + 0.40 + 3 * LH + NG, CW, 0.40, 16)

# ── a sin + b cos ─────────────────────────────────────────────────────────────
slide = new_slide("Copy", "When sine and cosine are added", TEAL)
add_notes(slide, "Copy. Amplitude is √(a²+b²). They cannot both be +1 at the same θ.")
body(slide, "sin θ and cos θ do not peak together. The combined amplitude is the hypotenuse of a and b.", XL, Y0, CW)
fml(slide, "y  =  a sin θ  +  b cos θ", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "ymax  =  + √(a² + b²)          ymin  =  − √(a² + b²)", XL, Y0 + 2 * LH + 2 * NG, CW, 1.00, 18)
trap_card(slide, "Do not write ymax = a + b. That would need sin = 1 and cos = 1 at the same instant, which never happens.", XL, Y0 + 3 * LH + 2 * NG + 0.90, CW, 0.45)

slide = new_slide("Board", "y  =  3 sin θ  +  4 cos θ", GREEN)
add_notes(slide, "Read. √25=5. min −5.")
qtx(slide, "If  y  =  3 sin θ  +  4 cos θ,  find ymax and ymin.", XL, Y0, CW, 0.40, 17)
fml(slide, "ymax  =  √(3² + 4²)  =  √(9 + 16)  =  √25  =  +5", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "ymin  =  −5", XL, Y0 + 0.50 + LH, CW, LH, 18)
result_card(slide, "ymax  =  +5          ymin  =  −5", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "Minimum magnitude of a fraction", GREEN)
add_notes(slide, "Read. Bigger denominator, smaller fraction. ymin mag = 4/√13.")
qtx(slide, "If  y  =  4 / (3 sin θ + 2 cos θ),  find the minimum magnitude of y.", XL, Y0, CW, 0.45, 16)
body(slide, "4/10 = 0.4.   4/100 = 0.04.   A larger denominator makes a smaller positive fraction.", XL, Y0 + 0.52, CW)
fml(slide, "ymin  =  4  /  (3 sin θ + 2 cos θ)max  =  4 / √(3² + 2²)  =  4 / √13", XL, Y0 + 0.52 + LH + NG, CW, LH, 18)
result_card(slide, "minimum magnitude  =  4 / √13", XL, Y0 + 0.52 + 2 * LH + 2 * NG, CW, 0.36, 20)

slide = new_slide("Board", "Five max–min — every line", GREEN)
add_notes(slide, "Read. 5, √13, 6 and −4, 3 and −1, 13.")
qtx(slide, "Find maximum and minimum of each.", XL, Y0, CW, 0.32, 16)
grid_table(
    slide,
    ["", "y", "ymax", "ymin"],
    [
        ["(a)", "4 sin θ + 3 cos θ", "√(16+9)  =  5", "−5"],
        ["(b)", "2 sin θ − 3 cos θ", "√(4+9)  =  √13", "−√13"],
        ["(c)", "1 + 5 sin θ", "1+5  =  6", "1−5  =  −4"],
        ["(d)", "1 − 2 sin θ", "1−2(−1)  =  3", "1−2(1)  =  −1"],
        ["(e)", "5 sin θ − 12 cos θ", "√(25+144)  =  13", "−13"],
    ],
    XL, Y0 + 0.40, [0.85, 4.20, 3.50, 3.50], 0.48, 14, 15, scaled=False,
)

# ── BREAK ─────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
_solid(slide, 0, 0, 13.333, 0.06, TEAL)
tx(slide, "BREAK", XL, 2.20, CW, 0.70, 48, GOLD, True, PP_ALIGN.CENTER, spa=4)
tx(slide, "Max–min leftover is done.  Next: phase and phasor", XL, 3.10, CW, 0.40, 18, BLUE, True, PP_ALIGN.CENTER, spa=0)
_footer(slide, len(prs.slides))
add_notes(slide, "Hour 1 leftover max-min done. Hour 2: phase, phasor, AC.")
after_content_slide()

# ── PHASE ─────────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Phase is the angle inside the ratio", TEAL)
add_notes(slide, "Copy. Phase is the whole argument. Used later in SHM, waves, AC, optics, EM, vectors.")
body(slide, "Phase is the angle sitting inside the trigonometric function. Not the 4 in front.", XL, Y0, CW)
fml(slide, "sin θ                         phase  =  θ", XL, Y0 + LH + NG, CW, LH, 20)
fml(slide, "sin (θ + α)                   phase  =  θ + α", XL, Y0 + 2 * LH + NG, CW, LH, 20)
fml(slide, "sin (ωt + α + 90°)            phase  =  ωt + α + 90°", XL, Y0 + 3 * LH + NG, CW, LH, 20)
ncert(slide, "Later:  AC,  SHM,  waves,  wave optics,  EM waves,  vectors.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.85, 16)

slide = new_slide("Copy", "Phase difference — same function first", TEAL)
add_notes(slide, "Copy. Same ratio. Δφ = φ2 − φ1. Plus is lead. Minus is lag.")
body(slide, "Angle difference between two of the same trigonometric function is called phase difference.", XL, Y0, CW)
fml(slide, "y1  =  sin θ                         φ1  =  θ", XL, Y0 + LH + NG, CW, LH, 18)
fml(slide, "y2  =  sin (θ + 60°)                 φ2  =  θ + 60°", XL, Y0 + 2 * LH + NG, CW, LH, 18)
fml(slide, "Δφ  =  φ2 − φ1  =  +60°     (y2 leads y1)          Δφ  =  φ1 − φ2  =  −60°     (lag)", XL, Y0 + 3 * LH + NG, CW, LH, 18)
trap_card(slide, "Do not subtract a cosine phase from a sine phase until both are written as the same ratio.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.45)

slide = new_slide("Board", "y1 leads y2 by 90°", GREEN)
add_notes(slide, "Read. +30 − (−60) = 90. Phasor: +30 and −60 on the sin axis.")
qtx(slide, "Find the phase difference.   y1  =  sin (ωt + 30°)     y2  =  sin (ωt − 60°)", XL, Y0, CW, 0.45, 16)
cx, cy = phasor_axes(slide, 8.55, 1.85, 1.55, "sin (ωt)", "")
draw_line(slide, cx, cy, cx + 1.20, cy - 0.70, GOLD, 2.25)
draw_line(slide, cx, cy, cx + 0.70, cy + 1.20, GOLD, 2.25)
tx(slide, "y1  +30°", cx + 1.10, cy - 0.95, 1.20, 0.22, 12, GOLD, True, spa=0)
tx(slide, "y2  −60°", cx + 0.55, cy + 1.22, 1.20, 0.22, 12, GOLD, True, spa=0)
tx(slide, "90°", cx + 0.55, cy - 0.18, 0.50, 0.22, 14, GREEN, True, spa=0)
fml(slide, "Δφ  =  φ1 − φ2  =  (ωt + 30°)  −  (ωt − 60°)  =  90°", XL, Y0 + 0.55, 7.6, LH, 17)
result_card(slide, "Δφ  =  90°     —  y1 is ahead of y2", XL, Y0 + 0.55 + LH + NG, 7.6, 0.36, 18)

slide = new_slide("Board", "Convert to the same ratio, or the answer is wrong", GREEN)
add_notes(slide, "Read. cos(ωt−90)=sin(ωt). Δφ=0. Naked subtract gives a fake 90°.")
qtx(slide, "Find Δφ.   y1  =  3 cos (ωt − 90°)     y2  =  4 sin (ωt)", XL, Y0, CW, 0.45, 16)
y = warn_row(slide, "Δφ  =  ωt − (ωt − 90°)  =  90°", "first write both as sine (or both as cosine)", XL, Y0 + 0.52, CW)
fml(slide, "cos (ωt − 90°)  =  sin ωt          so          y1  =  3 sin (ωt)", XL, y + NG, CW, LH, 18)
result_card(slide, "Δφ  =  0     —  they are in phase", XL, y + NG + LH + NG, CW, 0.36, 20)
trap_card(slide, "A minus cosine is a sine with a shift. Convert, then subtract.", XL, y + NG + 2 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Board", "cos (ωt + π/2)  points opposite to sine", GREEN)
add_notes(slide, "Read. cos(ωt+π/2)= −sin(ωt). Δφ=180°.")
qtx(slide, "Find Δφ.   y1  =  3 sin (ωt)     y2  =  4 cos (ωt + π/2)", XL, Y0, CW, 0.40, 16)
fml(slide, "cos (ωt + π/2)  =  − sin (ωt)          so          y2  =  − 4 sin (ωt)", XL, Y0 + 0.50, CW, LH, 18)
result_card(slide, "Δφ  =  180°     =     π rad     —  opposite on the phasor", XL, Y0 + 0.50 + LH + NG, CW, 0.36, 18)
cx, cy = phasor_axes(slide, 4.55, 3.55, 1.35, "sin (ωt)", "cos (ωt)")
draw_line(slide, cx, cy, cx + 1.30, cy, GOLD, 2.25)
draw_line(slide, cx, cy, cx - 1.30, cy, GREEN, 2.25)
tx(slide, "y1", cx + 1.15, cy - 0.28, 0.50, 0.22, 14, GOLD, True, spa=0)
tx(slide, "y2", cx - 1.55, cy - 0.28, 0.50, 0.22, 14, GREEN, True, spa=0)

slide = new_slide("Board", "Two more phase differences", GREEN)
add_notes(slide, "Read. 150° and 135°. Then 30° leftover as a one-liner.")
qtx(slide, "(1)  x1 = A cos(ωt − π/2)     y2 = A cos(ωt + π/3)", XL, Y0, CW, 0.36, 16)
fml(slide, "Same function already.   Δφ  =  (ωt + π/3) − (ωt − π/2)  =  5π/6  =  150°", XL, Y0 + 0.42, CW, LH, 16)
result_card(slide, "(1)   Δφ  =  150°", XL, Y0 + 0.42 + LH + NG, CW, 0.32, 18)
qtx(slide, "(2)  y1 = 4 sin(θ + π/4)     y2 = 3 sin(θ − π/2)", XL, Y0 + 0.42 + 2 * LH + 2 * NG, CW, 0.32, 16)
fml(slide, "Δφ  =  (θ + π/4) − (θ − π/2)  =  π/4 + π/2  =  135°", XL, Y0 + 0.42 + 3 * LH + 2 * NG, CW, LH, 16)
result_card(slide, "(2)   Δφ  =  135°          Also:  y1 = 4 sin(θ+π/3), y2 = 3 sin(θ+π/6)  →  30°", XL, Y0 + 0.42 + 4 * LH + 3 * NG, CW, 0.32, 16)

slide = new_slide("Board", "AC current and voltage", GREEN)
add_notes(slide, "Read. Convert V to sine. 90−(−30)=120=2π/3.")
qtx(slide, "I  =  I0 sin(ωt − 30°)     V  =  V0 cos(ωt).   Find the phase difference.", XL, Y0, CW, 0.45, 16)
fml(slide, "V  =  V0 cos(ωt)  =  V0 sin(ωt + 90°)", XL, Y0 + 0.55, CW, LH, 18)
fml(slide, "Δφ  =  90° − (−30°)  =  120°  =  2π/3", XL, Y0 + 0.55 + LH, CW, LH, 18)
result_card(slide, "Δφ  =  2π/3     —  voltage leads current by 120°", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "Negative cosine is a shifted sine", GREEN)
add_notes(slide, "Read. I=sin(ωt−π/2). Δφ=π/3. Option 2.")
qtx(slide, "I  =  − I0 cos(ωt)     V  =  V0 sin(ωt − π/6).   Find Δφ.", XL, Y0, CW, 0.40, 16)
choices(slide, ["2π/3", "π/3", "π/6", "5π/6"], XL, Y0 + 0.48)
fml(slide, "− cos(ωt)  =  sin(ωt − π/2)          Δφ  =  (−π/6) − (−π/2)  =  π/3", XL, Y0 + 2.20, CW, LH, 16)
result_card(slide, "Option 2     Δφ  =  π/3", XL, Y0 + 2.20 + LH + NG, CW, 0.36, 20)

# ── MIXED LIVE ────────────────────────────────────────────────────────────────
slide = new_slide("Board", "Person walks toward a 40 m building", GREEN)
add_notes(slide, "Read. d = 40(√3 − 1/√3)=80/√3. Option 2.")
qtx(slide, "A person sees a 40 m building at elevation 30°, then at 60°. Distance moved toward the building?", XL, Y0, CW, 0.50, 15)
choices(slide, ["40/√3", "80/√3", "60/√3", "120/√3"], XL, Y0 + 0.58, 5.4)
fml(slide, "x1  =  40 cot 30°  =  40√3          x2  =  40 cot 60°  =  40/√3", XL, Y0 + 2.30, CW, LH, 16)
fml(slide, "d  =  x1 − x2  =  40√3 − 40/√3  =  40(3 − 1)/√3  =  80/√3", XL, Y0 + 2.30 + LH, CW, LH, 16)
result_card(slide, "Option 2     d  =  80/√3", XL, Y0 + 2.30 + 2 * LH + NG, CW, 0.36, 18)

slide = new_slide("Board", "2 (sin 15° + sin 75°)²", GREEN)
add_notes(slide, "Read. sin75=cos15. 2(1+sin30)=3.")
qtx(slide, "What is the value of  2 (sin 15° + sin 75°)² ?", XL, Y0, CW, 0.40, 18)
fml(slide, "sin 75°  =  cos 15°          so          2 (sin 15° + cos 15°)²", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "        =  2 ( sin²15° + cos²15° + 2 sin 15° cos 15° )  =  2 ( 1 + sin 30° )", XL, Y0 + 0.50 + LH, CW, LH, 16)
fml(slide, "        =  2 ( 1 + 1/2 )  =  3", XL, Y0 + 0.50 + 2 * LH, CW, LH, 18)
result_card(slide, "2 (sin 15° + sin 75°)²  =  3", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "If sin θ = 1/2, first quadrant", GREEN)
add_notes(slide, "Read. θ=30. √3(√3/2 + 2/√3)=7/2. Option 3.")
qtx(slide, "If  sin θ  =  1/2,  the value of  √3 (cos θ + 2 tan θ)  is:", XL, Y0, CW, 0.40, 16)
choices(slide, ["7", "15/2", "7/2", "none"], XL, Y0 + 0.50)
fml(slide, "First quadrant:  θ  =  30°          cos 30°  =  √3/2          tan 30°  =  1/√3", XL, Y0 + 2.20, CW, LH, 16)
fml(slide, "√3 ( √3/2  +  2/√3 )  =  3/2  +  2  =  7/2", XL, Y0 + 2.20 + LH, CW, LH, 18)
result_card(slide, "Option 3     7/2", XL, Y0 + 2.20 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "If sin θ = √5 / 3", GREEN)
add_notes(slide, "Read. cos=2/3. tan=√5/2. sec=3/2. Answer √5/2 + 6.")
qtx(slide, "If  sin θ  =  √5 / 3,  find  tan θ  +  4 sec θ.   Take θ acute.", XL, Y0, CW, 0.45, 16)
fml(slide, "cos θ  =  √(1 − 5/9)  =  √(4/9)  =  2/3          tan θ  =  (√5/3) / (2/3)  =  √5 / 2", XL, Y0 + 0.55, CW, LH, 16)
fml(slide, "sec θ  =  3/2          tan θ + 4 sec θ  =  √5/2  +  4 × (3/2)  =  √5/2  +  6", XL, Y0 + 0.55 + LH, CW, LH, 16)
result_card(slide, "tan θ  +  4 sec θ  =  6 + √5/2", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "tan⁻¹(1)  —  principal value", GREEN)
add_notes(slide, "Read. Range of tan⁻¹ is (−π/2, π/2). Only π/4. Not 5π/4.")
qtx(slide, "The value of  tan⁻¹(1)  is", XL, Y0, CW, 0.36, 18)
choices(slide, ["π/4", "5π/4", "both (1) and (2)", "none of these"], XL, Y0 + 0.45)
y = warn_row(slide, "tan(5π/4) = 1, so both", "tan⁻¹ returns only the principal value, in (−π/2, π/2)", XL, Y0 + 2.15, CW)
result_card(slide, "Option 1     π/4", XL, y + NG, CW, 0.36, 20)
trap_card(slide, "5π/4 is an angle whose tangent is 1. It is not the inverse-tangent of 1.", XL, y + NG + 0.42, CW, 0.40)

# ── PW LECTURE 04: PHASE TABLE AND SUPERPOSITION ─────────────────────────────
slide = new_slide("PW Lecture 04", "Seven phase rows — answers from the PW board", GREEN)
add_notes(slide, "This was last-class homework. Read each answer. Convert cosine to sine before comparing.")
qtx(slide, "Complete the phase-difference table from last class.", XL, Y0, CW, 0.38, 16)
grid_table(
    slide,
    ["Equation 1", "Equation 2", "phase difference"],
    [
        ["I₀ sin(θ + π/3)", "I₀ sin(θ − π/6)", "90°"],
        ["I₀ sin(θ + π/3)", "I₀ cos(θ + π/3)", "90°"],
        ["4 sin(ωt)", "4 cos(ωt − π/3)", "30°"],
        ["3 sin(ωt + π/2)", "5 cos(ωt)", "0°"],
        ["4 cos(ωt − π/2)", "4 sin(ωt + π/2)", "90°"],
        ["3 sin(ωt + π/3)", "3 cos(ωt)", "30°"],
        ["4 sin(θ − 60°)", "4 cos(θ − 30°)", "120°"],
    ],
    XL, Y0 + 0.46, [4.15, 4.40, 3.50], 0.42, 13, 14, scaled=False,
)
trap_card(slide, "Amplitude does not decide phase. Compare only the angles after both functions are the same.", XL, Y0 + 0.46 + 8 * 0.42 + 0.10, CW, 0.34)

slide = new_slide("Copy", "Superposition — add two oscillations as phasors", TEAL)
add_notes(slide, "Copy. Same omega. Horizontal sine coefficient a, vertical cosine coefficient b.")
body(slide, "For the same angular frequency, add the sine part and cosine part like perpendicular sides.", XL, Y0, CW)
fml(slide, "x  =  a sin(ωt)  +  b cos(ωt)  =  R sin(ωt + φ)", XL, Y0 + LH + NG, CW, LH, 20)
ncert(slide, "R  =  √(a² + b²)          tan φ  =  b/a", XL, Y0 + 2 * LH + 2 * NG, CW, 0.95, 20)
body(slide, "R is the new amplitude. φ tells how far the resultant phasor is above the sine axis.", XL, Y0 + 3 * LH + 2 * NG + 0.85, CW)
trap_card(slide, "This shortcut needs the same ω in both terms. Do not combine sin(2ωt) with cos(ωt).", XL, Y0 + 4 * LH + 3 * NG + 0.85, CW, 0.40)

slide = new_slide("Board", "PW example:  4 sin(ωt) + 3 cos(ωt)", GREEN)
add_notes(slide, "Read. R=5. tan phi=3/4, phi about 37 degrees.")
qtx(slide, "Two oscillations are  x₁ = 4 sin(ωt)  and  x₂ = 3 cos(ωt). Find their superposition.", XL, Y0, CW, 0.42, 16)
fml(slide, "x  =  x₁ + x₂  =  4 sin(ωt) + 3 cos(ωt)", XL, Y0 + 0.52, CW, LH, 18)
fml(slide, "R  =  √(4² + 3²)  =  5          tan φ  =  3/4          φ  ≈  37°", XL, Y0 + 0.52 + LH, CW, LH, 18)
result_card(slide, "x  =  5 sin(ωt + 37°)", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.36, 22)

slide = new_slide("Board", "PW example: equal sine and cosine parts", GREEN)
add_notes(slide, "Read. Both 4. R=4 root2. tan phi=1, phi=45.")
qtx(slide, "Find the resultant:   y = 4 sin(ωt) + 4 cos(ωt).", XL, Y0, CW, 0.40, 18)
fml(slide, "R  =  √(4² + 4²)  =  √32  =  4√2", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "tan φ  =  4/4  =  1          so          φ  =  45°", XL, Y0 + 0.50 + LH, CW, LH, 20)
result_card(slide, "y  =  4√2 sin(ωt + 45°)", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 22)

# ── PW LECTURE 04: ANGLE TOOL ─────────────────────────────────────────────────
slide = new_slide("Copy", "Angle tool — incline and parallel lines", TEAL)
add_notes(slide, "PW pages 15-16. Normal is perpendicular to incline. Alternate angles are equal.")
body(slide, "A line perpendicular to an incline is called the normal. If the incline is θ above horizontal:", XL, Y0, CW)
fml(slide, "angle between vertical and normal  =  θ", XL, Y0 + LH + NG, CW, LH, 20)
body(slide, "For two parallel horizontal lines cut by one slant line, alternate interior angles are equal.", XL, Y0 + 2 * LH + 2 * NG, CW)
fml(slide, "alternate angles:  α = α          triangle angles:  A + B + C = 180°", XL, Y0 + 3 * LH + 2 * NG, CW, LH, 18)
trap_card(slide, "Do not use the angle with the incline when the force component is measured from the normal.", XL, Y0 + 4 * LH + 3 * NG, CW, 0.40)

slide = new_slide("Board", "PW angle check — find α, β and γ", GREEN)
add_notes(slide, "Parallel lines give alpha 30 and gamma 60. Triangle gives beta 90.")
qtx(slide, "Two strings make 30° and 60° with a horizontal ceiling. Find α, β and γ at the lower point.", XL, Y0, CW, 0.46, 16)
fml(slide, "alternate angles  →  α = 30°          γ = 60°", XL, Y0 + 0.56, CW, LH, 20)
fml(slide, "α + β + γ = 180°          β = 180° − 30° − 60° = 90°", XL, Y0 + 0.56 + LH, CW, LH, 18)
result_card(slide, "α = 30°          β = 90°          γ = 60°", XL, Y0 + 0.56 + 2 * LH + NG, CW, 0.36, 22)

# ── PW LECTURE 04: A.P. ───────────────────────────────────────────────────────
slide = new_slide("Copy", "A.P. — equal difference every time", TEAL)
add_notes(slide, "Copy. Arithmetic progression has constant difference.")
body(slide, "Arithmetic progression is a sequence in which the difference between consecutive terms is constant.", XL, Y0, CW)
fml(slide, "a,  a+d,  a+2d,  a+3d,  ...", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "nth term:  Tₙ = a + (n−1)d", XL, Y0 + 2 * LH + 2 * NG, 5.75, 0.95, 20)
ncert(slide, "sum:  Sₙ = n/2 [2a + (n−1)d]", XL + 6.0, Y0 + 2 * LH + 2 * NG, 6.05, 0.95, 18)
body(slide, "a = first term.  d = next term − previous term.  n = position of the term.", XL, Y0 + 3 * LH + 2 * NG + 0.85, CW)

slide = new_slide("Board", "PW A.P.:  2, 5, 8, 11, 14, ...", GREEN)
add_notes(slide, "Read. a=2, d=3, T10=29, S10=155.")
qtx(slide, "Find the first term, common difference, 10th term and sum of first 10 terms.", XL, Y0, CW, 0.42, 16)
fml(slide, "a = 2          d = 11 − 8 = 3", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "T₁₀ = a + 9d = 2 + 9(3) = 29", XL, Y0 + 0.52 + LH, CW, LH, 20)
fml(slide, "S₁₀ = 10/2 [2 + 29] = 5(31) = 155", XL, Y0 + 0.52 + 2 * LH, CW, LH, 20)
result_card(slide, "T₁₀ = 29          S₁₀ = 155", XL, Y0 + 0.52 + 3 * LH + NG, CW, 0.36, 22)

slide = new_slide("Board", "PW A.P.:  7, 11, 15, 19, ...", GREEN)
add_notes(slide, "Read. a=7 d=4 T10=43 S10=250.")
qtx(slide, "Find the 10th term and the sum of the first 10 terms.", XL, Y0, CW, 0.40, 18)
fml(slide, "a = 7          d = 19 − 15 = 4", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "T₁₀ = 7 + 9(4) = 43", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "S₁₀ = 10/2 (7 + 43) = 5(50) = 250", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "10th term = 43          sum = 250", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = new_slide("Board", "Which term of 27, 24, 21, ... is zero?", GREEN)
add_notes(slide, "Read. a=27 d=-3. Set Tn=0. n=10.")
qtx(slide, "Which term of the A.P.  27, 24, 21, ...  is zero?", XL, Y0, CW, 0.40, 18)
fml(slide, "a = 27          d = 24 − 27 = −3", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "0 = a + (n−1)d = 27 + (n−1)(−3)", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "0 = 30 − 3n          3n = 30          n = 10", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "The 10th term is zero.", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = new_slide("Board", "PW A.P.:  28, 22, x, y, 4", GREEN)
add_notes(slide, "Read. d=-6. x=16 y=10.")
qtx(slide, "28, 22, x, y, 4 is an A.P. Find x and y.", XL, Y0, CW, 0.40, 18)
fml(slide, "a = 28          d = 22 − 28 = −6", XL, Y0 + 0.50, CW, LH, 20)
fml(slide, "x = T₃ = 28 + 2(−6) = 16", XL, Y0 + 0.50 + LH, CW, LH, 20)
fml(slide, "y = T₄ = 28 + 3(−6) = 10", XL, Y0 + 0.50 + 2 * LH, CW, LH, 20)
result_card(slide, "x = 16          y = 10", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 22)

slide = new_slide("Copy", "Natural numbers are an A.P.", TEAL)
add_notes(slide, "Copy. 1 to n has a=1 d=1 last term n.")
body(slide, "1, 2, 3, ..., n is an A.P. with first term 1, last term n and n terms.", XL, Y0, CW)
ncert(slide, "1 + 2 + 3 + ... + n  =  n(n+1)/2", XL, Y0 + LH + NG, CW, 1.00, 24)
body(slide, "Quick check: n = 2 gives 1 + 2 = 3. The formula gives 2×3/2 = 3.", XL, Y0 + 2 * LH + 2 * NG + 0.90, CW)

slide = new_slide("Board", "PW physics use — tuning-fork frequencies", GREEN)
add_notes(slide, "Read. 25 terms, d=10, T25=400. a=160 Hz.")
qtx(slide, "25 tuning forks are in increasing A.P. order. Difference is 10 Hz. The 25th is 400 Hz. Find the first.", XL, Y0, CW, 0.50, 15)
fml(slide, "T₂₅ = a + 24d          400 = a + 24(10)", XL, Y0 + 0.60, CW, LH, 20)
fml(slide, "a = 400 − 240 = 160 Hz", XL, Y0 + 0.60 + LH, CW, LH, 20)
result_card(slide, "Frequency of the first tuning fork = 160 Hz", XL, Y0 + 0.60 + 2 * LH + NG, CW, 0.36, 20)

# ── PW LECTURE 04: G.P. ───────────────────────────────────────────────────────
slide = new_slide("Copy", "G.P. — equal ratio every time", TEAL)
add_notes(slide, "Copy. Geometric progression has constant ratio.")
body(slide, "Geometric progression is a sequence in which the ratio of consecutive terms is constant.", XL, Y0, CW)
fml(slide, "a,  ar,  ar²,  ar³,  ...", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "nth term:  Tₙ = arⁿ⁻¹", XL, Y0 + 2 * LH + 2 * NG, 5.75, 0.95, 21)
ncert(slide, "infinite sum:  S∞ = a/(1−r), only if |r| < 1", XL + 6.0, Y0 + 2 * LH + 2 * NG, 6.05, 0.95, 17)
body(slide, "r = any term ÷ previous term. If that ratio changes, the sequence is not a G.P.", XL, Y0 + 3 * LH + 2 * NG + 0.85, CW)

slide = new_slide("Board", "Is it a G.P.? Check the ratio", GREEN)
add_notes(slide, "Read. First not GP. Second r=2. Third r=1/2. Fourth r=-1/2.")
qtx(slide, "Decide which sequences are G.P.", XL, Y0, CW, 0.34, 17)
grid_table(
    slide,
    ["sequence", "ratio check", "answer"],
    [
        ["2, 4, 8, 12, 24, 48", "4/2=2, but 12/8≠2", "not G.P."],
        ["2, 4, 8, 16, 32, 64", "each ratio = 2", "G.P., r=2"],
        ["1, 1/2, 1/4, 1/8, ...", "each ratio = 1/2", "G.P., r=1/2"],
        ["1, −1/2, 1/4, −1/8, ...", "each ratio = −1/2", "G.P., r=−1/2"],
    ],
    XL, Y0 + 0.40, [4.50, 4.20, 3.35], 0.55, 14, 15, scaled=False,
)

slide = new_slide("Board", "Infinite G.P. sums — only when |r| < 1", GREEN)
add_notes(slide, "Read. Sum 1+1/2+...=2. Alternating sum=2/3. Odd reciprocal powers sum=3/2.")
qtx(slide, "Find each infinite sum.", XL, Y0, CW, 0.34, 17)
fml(slide, "(1)  1 + 1/2 + 1/4 + ...   =   1/(1−1/2)   =   2", XL, Y0 + 0.42, CW, LH, 18)
fml(slide, "(2)  1 − 1/2 + 1/4 − ...   =   1/[1−(−1/2)]   =   2/3", XL, Y0 + 0.42 + LH, CW, LH, 18)
fml(slide, "(3)  1 + 1/3 + 1/9 + ...   =   1/(1−1/3)   =   3/2", XL, Y0 + 0.42 + 2 * LH, CW, LH, 18)
trap_card(slide, "2 + 4 + 8 + ... has r = 2. It has no finite infinite sum because |r| is not below 1.", XL, Y0 + 0.42 + 3 * LH + NG, CW, 0.45)

slide = new_slide("Board", "PW physics use — potential from alternating charges", GREEN)
add_notes(slide, "Read. Distances 1,2,4,... signs alternate. V=kq(1-1/2+...) = 2kq/3.")
qtx(slide, "Charges +q, −q, +q, ... lie at x = 1, 2, 4, 8, ... . Find electric potential at the origin.", XL, Y0, CW, 0.48, 15)
fml(slide, "V = kq/1 − kq/2 + kq/4 − kq/8 + ...", XL, Y0 + 0.58, CW, LH, 18)
fml(slide, "V = kq [1 − 1/2 + 1/4 − ...]          a=1, r=−1/2", XL, Y0 + 0.58 + LH, CW, LH, 18)
fml(slide, "V = kq / [1 − (−1/2)] = 2kq/3", XL, Y0 + 0.58 + 2 * LH, CW, LH, 20)
result_card(slide, "Potential at origin = 2kq/3", XL, Y0 + 0.58 + 3 * LH + NG, CW, 0.36, 22)

slide = new_slide("Board", "PW physics use — repeating resistance blocks", GREEN)
add_notes(slide, "Read. Effective series blocks are R, R/2, R/4,... Sum=2R.")
qtx(slide, "Equivalent series resistance becomes  R + R/2 + R/4 + R/8 + ... . Find R_AB.", XL, Y0, CW, 0.42, 17)
fml(slide, "R_AB = R [1 + 1/2 + 1/4 + 1/8 + ...]", XL, Y0 + 0.52, CW, LH, 20)
fml(slide, "a = 1, r = 1/2          sum = 1/(1−1/2) = 2", XL, Y0 + 0.52 + LH, CW, LH, 20)
result_card(slide, "R_AB = 2R", XL, Y0 + 0.52 + 2 * LH + NG, CW, 0.36, 24)

# ── PW LECTURE 04: GRAPH BASICS ───────────────────────────────────────────────
slide = new_slide("Copy", "Graph — relation between two physical quantities", TEAL)
add_notes(slide, "Copy. Independent on x, dependent on y. Axes and origin.")
body(slide, "A graph is a diagram in which a line or curve shows the relation between two physical quantities.", XL, Y0, CW)
fml(slide, "horizontal line  =  x-axis          vertical line  =  y-axis          meeting point  =  origin (0,0)", XL, Y0 + LH + NG, CW, LH, 17)
body(slide, "Usually x is the quantity we choose or change. y is the quantity that responds.", XL, Y0 + 2 * LH + 2 * NG, CW)
ncert(slide, "Always write: quantity + unit on each axis.", XL, Y0 + 3 * LH + 3 * NG, CW, 0.85, 20)

slide = new_slide("Copy", "Constant graphs — horizontal and vertical", TEAL)
add_notes(slide, "Copy. y=2 horizontal. x=constant vertical.")
fml(slide, "y = 2          →          horizontal line through y = 2", XL, Y0, CW, LH, 20)
body(slide, "x can change, but y stays 2. So the graph moves left-right at the same height.", XL, Y0 + LH, CW)
fml(slide, "x = 3          →          vertical line through x = 3", XL, Y0 + 2 * LH + NG, CW, LH, 20)
body(slide, "y can change, but x stays 3. So the graph moves up-down at the same x-position.", XL, Y0 + 3 * LH + NG, CW)
trap_card(slide, "y = constant is horizontal. x = constant is vertical. Students commonly swap them.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.42)

slide = new_slide("Copy", "Distance between two points on a graph", TEAL)
add_notes(slide, "Copy. Draw horizontal and vertical differences. Pythagoras.")
body(slide, "For P₁(x₁,y₁) and P₂(x₂,y₂), the horizontal change is x₂−x₁ and vertical change is y₂−y₁.", XL, Y0, CW)
ncert(slide, "P₁P₂  =  √[(x₂−x₁)² + (y₂−y₁)²]", XL, Y0 + LH + NG, CW, 1.00, 22)
body(slide, "It is simply Pythagoras on the small right triangle made by the two changes.", XL, Y0 + 2 * LH + 2 * NG + 0.90, CW)

# ── TONIGHT ───────────────────────────────────────────────────────────────────
slide = new_slide("Tonight", "PW Lecture 04 homework — three questions", AMBER)
add_notes(slide, "Do not reveal answers. These are the three homework questions from PW pages 30-32.")
qtx(slide, "(1) Charges q lie at x = 1, 2, 4, 8, ... . Find the force on charge q₀ at the origin.", XL, Y0, CW, 0.55, 16)
qtx(slide, "(2) Five resistances are in series. Each is half the previous one. First is 10 Ω. Find the fifth resistance.", XL, Y0 + 0.72, CW, 0.55, 16)
qtx(slide, "(3) Charges +q, −q, +q, ... lie at x = 1, 2, 4, 8, ... . Find potential at the origin.", XL, Y0 + 1.44, CW, 0.55, 16)
body(slide, "Hint only: identify a, identify r, then choose Tₙ or S∞.", XL, Y0 + 2.25, CW)
trap_card(slide, "Do not solve these on this slide. They are tonight’s work.", XL, Y0 + 3.00, CW, 0.40)

slide = new_slide("Tonight", "After class", AMBER)
add_notes(slide, "Daily test before next lecture. AP and GP formulae ten times.")
body(slide, "Daily test before the next lecture.", XL, Y0, CW)
body(slide, "Write A.P. nth-term and sum formula five times.", XL, Y0 + LH, CW)
body(slide, "Write G.P. nth-term and infinite-sum formula five times.", XL, Y0 + 2 * LH, CW)
body(slide, "Tonight: the three PW Lecture 04 homework questions.", XL, Y0 + 3 * LH + NG, CW)

# ── 5-MINUTE SUMMARY ──────────────────────────────────────────────────────────
slide = new_slide("5-minute summary", "Max–min — the two boxes", TEAL)
add_notes(slide, "They say both boxes. Do not add a and b.")
fml(slide, "y  =  A sin (nθ)          ymax  =  |A|          at  nθ  =  90°", XL, Y0, CW, LH, 18)
ncert(slide, "y  =  a sin θ + b cos θ          ymax  =  √(a² + b²)", XL, Y0 + LH + NG, CW, 0.90, 18)
result_card(slide, "3 sin θ + 4 cos θ  →  ±5          4 / (3 sin θ + 2 cos θ)  →  4/√13", XL, Y0 + 2 * LH + NG + 0.80, CW, 0.36, 16)
trap_card(slide, "3 sin θ + 4 sin θ is 7 sin θ, not √(9+16). Add like ratios first.", XL, Y0 + 3 * LH + 2 * NG + 0.80, CW, 0.40)

slide = new_slide("5-minute summary", "Phase — the one lock-line", TEAL)
add_notes(slide, "Same function first. Then subtract. Fake 90° if you skip the convert.")
body(slide, "Phase is the angle inside the ratio. Phase difference needs the same trigonometric function.", XL, Y0, CW)
fml(slide, "Δφ  =  φ2 − φ1          after both are sine, or both are cosine", XL, Y0 + LH + NG, CW, LH, 18)
result_card(slide, "3 cos(ωt − 90°)  and  4 sin(ωt)   →   Δφ  =  0", XL, Y0 + 2 * LH + 2 * NG, CW, 0.36, 18)
trap_card(slide, "Never subtract a cosine-phase from a sine-phase. Convert, then subtract.", XL, Y0 + 3 * LH + 3 * NG, CW, 0.45)
body(slide, "Daily test before the next lecture. Three PW A.P./G.P. questions tonight.", XL, Y0 + 4 * LH + 4 * NG, CW)

slide = new_slide("5-minute summary", "A.P., G.P. and graph — lock-lines", TEAL)
add_notes(slide, "They say all five lines. Homework is three PW questions.")
fml(slide, "A.P.:   Tₙ = a + (n−1)d          Sₙ = n/2 [2a + (n−1)d]", XL, Y0, CW, LH, 17)
fml(slide, "G.P.:   Tₙ = arⁿ⁻¹          S∞ = a/(1−r), only if |r| < 1", XL, Y0 + LH, CW, LH, 17)
fml(slide, "1 + 2 + ... + n = n(n+1)/2", XL, Y0 + 2 * LH + NG, CW, LH, 18)
fml(slide, "distance = √[(x₂−x₁)² + (y₂−y₁)²]", XL, Y0 + 3 * LH + NG, CW, LH, 18)
result_card(slide, "y = constant → horizontal          x = constant → vertical", XL, Y0 + 4 * LH + 2 * NG, CW, 0.36, 18)

prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
