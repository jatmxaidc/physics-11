"""Lecture 10 — rate of change, quadratic review and differentiation rules."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_10_Differentiation_Rules_and_Applications.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)


# ── COVER ─────────────────────────────────────────────────────────────────────
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 10",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Differentiation Rules and Applications",bx,by+1.62,bw,0.40,18,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"15 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Begin differentiation from the first idea of tangent slope, then continue through the rules and applications."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["Differentiation: tangent slope at a point","Rate of change and the derivative","Quadratic equations and graph application","Basic differentiation formulae","Sum, difference, product and quotient rules","Mixed practice with simple solutions"]
for i,item in enumerate(topics):
    yy=1.18+i*0.78
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.70,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Differentiation was not started in the previous class, so teach it from the beginning."); after_content_slide()

# DIFFERENTIATION FROM THE BEGINNING
slide=new_slide("New topic","Differentiation: slope at a point",GREEN)
body(slide,"Differentiation finds how quickly one quantity changes with another.",XL,Y0,5.8,0.70)
body(slide,"On a graph, it gives the slope of the tangent at the chosen point.",XL,Y0+0.72,5.8,0.65)
ncert(slide,"slope = Δy/Δx     tangent slope = dy/dx",XL,Y0+1.62,5.8,0.95,21)
sx,sy=graph_axes(slide,7.25,1.50,4.65,4.75,-2.5,2.5,-4,4)
plot_function(slide,lambda x:x**3-3*x,sx,sy,-2.2,2.2,-4,4,GOLD)
mark_point(slide,sx,sy,-1,2,"maximum",GREEN,-0.95,-0.40)
mark_point(slide,sx,sy,1,-2,"minimum",GREEN,0.15,0.12)
tx(slide,"dy/dx = 0 at both turning points",7.55,6.35,4.30,0.30,14,WORK,True,PP_ALIGN.CENTER,spa=0,scaled=False)

slide=new_slide("How it works","Rate of change: from two points to one point",TEAL)
body(slide,"The average rate of change between two points is the secant's slope.",XL,Y0,5.8,0.82)
ncert(slide,"average rate = Δy/Δx = (y₂-y₁)/(x₂-x₁)",XL,Y0+0.88,5.8,1.00,20)
body(slide,"To find the rate at one point, bring the second point extremely close to the first.",XL,Y0+2.12,5.8,0.75)
sx,sy=graph_axes(slide,7.35,1.55,4.45,4.60,0,4,0,10)
plot_function(slide,lambda x:x*x,sx,sy,0,3.15,0,10,GOLD)
mark_point(slide,sx,sy,1,1,"P(1,1)",GREEN,-0.82,0.10)
mark_point(slide,sx,sy,3,9,"Q(3,9)",TEAL,-0.20,-0.35)
draw_line(slide,sx(0.65),sy(-0.4),sx(3.2),sy(9.8),WORK,1.8)

slide=new_slide("Complete theory","The derivative is an instantaneous rate",TEAL)
body(slide,"Let x change by a small amount Δx. The corresponding change in y is Δy.",XL,Y0,CW,0.55)
body(slide,"As Δx approaches zero, the secant approaches the tangent.",XL,Y0+0.72,CW,0.55)
ncert(slide,"dy/dx = lim(Δx→0)  Δy/Δx",XL,Y0+1.55,CW,1.10,24)
body(slide,"Therefore dy/dx gives the tangent's slope and the instantaneous rate of change.",XL,Y0+2.92,CW,0.60)
trap_card(slide,"Δx approaches zero; we do not divide by zero.",XL,Y0+3.78,CW,0.58)

slide=question_slide("Change in y for a linear relation")
qtx(slide,"If y = 2x+4, find the change in y and the average rate of change when x changes from 1 to 3.",XL,Y0,CW,1.10,21)
slide=answer_slide()
fml(slide,"At x₁=1: y₁=2(1)+4=6",XL,Y0,CW,LH,19)
fml(slide,"At x₂=3: y₂=2(3)+4=10",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"Δy = y₂-y₁ = 10-6 = 4",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"Δy/Δx = 4/(3-1) = 2",XL,Y0+3*LH+3*NG,CW,LH,19)
result_card(slide,"Change in y = 4; average rate of change = 2.",XL,Y0+4*LH+4*NG,CW,0.55)

slide=new_slide("Graph meaning","Increasing, decreasing and stationary",TEAL)
grid_table(slide,["Graph behaviour","Derivative sign","Meaning"],[
    ["Increasing","dy/dx > 0","y rises as x increases"],
    ["Stationary","dy/dx = 0","horizontal tangent"],
    ["Decreasing","dy/dx < 0","y falls as x increases"],
],XL,Y0,[3.5,3.2,5.35],0.64,16,16)
body(slide,"The magnitude |dy/dx| tells how steep the graph is.",XL,Y0+3.00,CW,0.50)

# QUADRATIC REVIEW AND APPLICATION
slide=new_slide("Quick review","Quadratic equation and its roots",BLUE)
body(slide,"A quadratic equation has the form ax²+bx+c=0, where a≠0.",XL,Y0,CW,0.50)
ncert(slide,"x = (-b ± sqrt(b²-4ac))/(2a)",XL,Y0+0.72,CW,1.05,23)
fml(slide,"sum of roots = -b/a",XL,Y0+2.08,5.75,LH,19)
fml(slide,"product of roots = c/a",6.95,Y0+2.08,5.75,LH,19)
grid_table(slide,["Discriminant D=b²-4ac","Nature of real roots"],[
    ["D > 0","two distinct real roots"],
    ["D = 0","one repeated real root"],
    ["D < 0","no real roots"],
],XL,Y0+2.90,[5.6,6.45],0.52,15,15)

slide=question_slide("Quadratic roots: x²-7x+12=0")
qtx(slide,"Find both roots by factorisation.",XL,Y0,CW,0.70,22)
slide=answer_slide()
fml(slide,"Find two numbers with product 12 and sum -7: -3 and -4",XL,Y0,CW,LH,18)
fml(slide,"x²-7x+12 = (x-3)(x-4)",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"(x-3)(x-4)=0",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"x = 3 or x = 4.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quadratic roots: x²+7x-18=0")
qtx(slide,"Find both roots by factorisation.",XL,Y0,CW,0.70,22)
slide=answer_slide()
fml(slide,"Split 7x as 9x-2x",XL,Y0,CW,LH,19)
fml(slide,"x²+9x-2x-18=0",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"(x+9)(x-2)=0",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"x = -9 or x = 2.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quadratic roots: x²+8x+15=0")
qtx(slide,"Find both roots by factorisation.",XL,Y0,CW,0.70,22)
slide=answer_slide()
fml(slide,"Find two numbers with product 15 and sum 8: 5 and 3",XL,Y0,CW,LH,18)
fml(slide,"x²+8x+15 = (x+5)(x+3)",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"(x+5)(x+3)=0",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"x = -5 or x = -3.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=new_slide("Exact graph","Parabola y=x²-3x-10",GREEN)
body(slide,"x-intercepts: solve x²-3x-10=(x-5)(x+2)=0.",XL,Y0,5.55,0.70)
fml(slide,"roots: x=-2 and x=5",XL,Y0+0.86,5.55,LH,18)
fml(slide,"y-intercept: (0,-10)",XL,Y0+1.47,5.55,LH,18)
fml(slide,"turning point: x=3/2",XL,Y0+2.08,5.55,LH,18)
body(slide,"At x=3/2, y=-49/4=-12.25.",XL,Y0+2.70,5.55,0.55)
sx,sy=graph_axes(slide,7.15,1.50,4.80,4.85,-3,6,-14,6)
plot_function(slide,lambda x:x*x-3*x-10,sx,sy,-3,6,-14,6,GOLD)
axis_ticks(slide,sx,sy,[-2,5],[-10])
mark_point(slide,sx,sy,1.5,-12.25,"vertex",GREEN,-0.25,0.12)

slide=question_slide("Physics application: downward motion")
qtx(slide,"A ball is thrown vertically downward at 10 m s⁻¹ from a height of 75 m. Taking g=10 m s⁻², find its time of flight.",XL,Y0,CW,1.25,20)
slide=answer_slide()
fml(slide,"Choose downward as positive: u=10, a=10, s=75",XL,Y0,CW,LH,18)
fml(slide,"Use s=ut+1/2 at²",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"75=10t+5t², so t²+2t-15=0",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"(t+5)(t-3)=0",XL,Y0+3*LH+3*NG,CW,LH,19)
result_card(slide,"Time cannot be negative, so t = 3 s.",XL,Y0+4*LH+4*NG,CW,0.55)

slide=new_slide("Worked example","Average rate becomes a derivative",TEAL)
body(slide,"For y=3x-2:",XL,Y0,CW,0.45)
fml(slide,"y(1)=1 and y(3)=7, so Δy=6",XL,Y0+0.62,CW,LH,18)
fml(slide,"Δy/Δx = 6/2 = 3",XL,Y0+1.23,CW,LH,18)
body(slide,"Because the graph is a straight line, the rate is 3 everywhere.",XL,Y0+1.90,CW,0.50)
ncert(slide,"dy/dx = d(3x-2)/dx = 3",XL,Y0+2.64,CW,0.95,22)

# DIFFERENTIATION RULES
slide=new_slide("Notation","Read the differential operator correctly",TEAL)
body(slide,"d/dx means: differentiate the expression that follows with respect to x.",XL,Y0,CW,0.58)
fml(slide,"d(y)/dx: rate of change of y with respect to x",XL,Y0+0.80,CW,LH,18)
fml(slide,"d(P)/dx: rate of change of P with respect to x",XL,Y0+1.42,CW,LH,18)
fml(slide,"d(x)/dt: rate of change of x with respect to time t",XL,Y0+2.04,CW,LH,18)
trap_card(slide,"Always identify both the changing quantity and the independent variable.",XL,Y0+2.92,CW,0.60)

slide=new_slide("Rule 1","Derivative of a constant",TEAL)
ncert(slide,"If y=C, then dy/dx = 0",XL,Y0,CW,1.05,24)
body(slide,"A constant does not change when x changes, so its rate of change is zero.",XL,Y0+1.34,CW,0.58)
fml(slide,"d(6)/dx=0",XL,Y0+2.18,3.60,LH,19)
fml(slide,"d(sin²θ+cos²θ)/dθ = d(1)/dθ = 0",4.65,Y0+2.18,7.95,LH,18)
fml(slide,"d(e⁰)/dx = d(1)/dx = 0",XL,Y0+2.82,CW,LH,18)

slide=new_slide("Rule 2","Power rule for algebraic functions",TEAL)
ncert(slide,"For constant n: d(x^n)/dx = n x^(n-1)",XL,Y0,CW,1.05,23)
body(slide,"Multiply by the old power, then reduce the power by one.",XL,Y0+1.34,CW,0.55)
fml(slide,"d(x²)/dx = 2x",XL,Y0+2.12,3.75,LH,20)
fml(slide,"d(x³)/dx = 3x²",4.85,Y0+2.12,3.75,LH,20)
fml(slide,"d(1/x²)/dx = d(x⁻²)/dx = -2x⁻³",8.95,Y0+2.12,3.75,LH,18)

slide=new_slide("Rule 3","Exponential and trigonometric derivatives",TEAL)
grid_table(slide,["Function","Derivative"],[
    ["e^x","e^x"],
    ["sin x","cos x"],
    ["cos x","-sin x"],
    ["tan x","sec²x"],
    ["logₑx","1/x"],
],2.15,Y0,[4.2,4.8],0.70,17,18)
body(slide,"Angles must be in radians when using these standard trigonometric derivatives.",XL,5.65,CW,0.55)

slide=new_slide("Rule 4","Constant multiple rule",TEAL)
ncert(slide,"d[k f(x)]/dx = k f'(x)",XL,Y0,CW,1.00,23)
body(slide,"A constant multiplier stays outside differentiation.",XL,Y0+1.28,CW,0.50)
fml(slide,"If y=3x⁵, then dy/dx=3(5x⁴)=15x⁴",XL,Y0+2.08,CW,LH,19)
fml(slide,"If y=7sin x, then dy/dx=7cos x",XL,Y0+2.72,CW,LH,19)

slide=new_slide("Rule 5","Sum and difference rules",TEAL)
ncert(slide,"d(A±B)/dx = dA/dx ± dB/dx",XL,Y0,CW,1.00,23)
body(slide,"Differentiate every term separately and keep its sign.",XL,Y0+1.30,CW,0.50)
fml(slide,"d(sin x+x²)/dx = cos x+2x",XL,Y0+2.08,CW,LH,19)
fml(slide,"d(e^x-cos x)/dx = e^x+sin x",XL,Y0+2.72,CW,LH,19)

slide=new_slide("Rule 6","Product and quotient rules",TEAL)
tx(slide,"PRODUCT",XL,Y0,5.75,0.35,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
ncert(slide,"d(AB)/dx = A dB/dx + B dA/dx",XL,Y0+0.45,5.75,1.08,18)
body(slide,"Differentiate one factor at a time.",XL,Y0+1.75,5.75,0.50)
tx(slide,"QUOTIENT",6.95,Y0,5.75,0.35,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
ncert(slide,"d(A/B)/dx = (B dA/dx - A dB/dx)/B²",6.95,Y0+0.45,5.75,1.08,17)
body(slide,"Keep the denominator squared.",6.95,Y0+1.75,5.75,0.50)
trap_card(slide,"For a quotient, remember: denominator × derivative of numerator minus numerator × derivative of denominator.",XL,Y0+2.62,CW,0.78)

slide=question_slide("Differentiate a simple polynomial")
qtx(slide,"For y=4x³-5x²+2x-7, find dy/dx.",XL,Y0,CW,0.72,22)
slide=answer_slide()
fml(slide,"d(4x³)/dx = 12x²",XL,Y0,CW,LH,19)
fml(slide,"d(-5x²)/dx = -10x",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"d(2x)/dx = 2 and d(-7)/dx = 0",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = 12x²-10x+2.",XL,Y0+3*LH+3*NG,CW,0.55)

# MIXED PRACTICE FROM THE SOURCE CLASS
slide=question_slide("Product-rule practice")
qtx(slide,"Differentiate y = tan x · logₑx.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Use d(AB)/dx = A dB/dx + B dA/dx",XL,Y0,CW,LH,18)
fml(slide,"d(tan x)/dx=sec²x and d(logₑx)/dx=1/x",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"dy/dx = tan x/x + logₑx·sec²x.",XL,Y0+2*LH+2*NG,CW,0.60)

slide=question_slide("Stationary-point practice")
qtx(slide,"If y=A/x³-B/x⁴, find x for which dy/dx=0. Assume A and B are non-zero constants.",XL,Y0,CW,1.00,21)
slide=answer_slide()
fml(slide,"Write y=Ax⁻³-Bx⁻⁴",XL,Y0,CW,LH,19)
fml(slide,"dy/dx=-3Ax⁻⁴+4Bx⁻⁵",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"Set zero and multiply by x⁵: -3Ax+4B=0",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"x = 4B/(3A).",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quotient-rule practice: exponential")
qtx(slide,"Differentiate y=e^x/x², where x≠0.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Take A=e^x and B=x²",XL,Y0,CW,LH,19)
fml(slide,"dy/dx = [x²e^x-e^x(2x)]/x⁴",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"Factor e^x x from the numerator",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"dy/dx = e^x(x-2)/x³.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Quotient-rule practice: algebraic")
qtx(slide,"Differentiate y=(x²-1)/(x²+1).",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Use the quotient rule with A=x²-1 and B=x²+1",XL,Y0,CW,LH,18)
fml(slide,"dy/dx = [(x²+1)(2x)-(x²-1)(2x)]/(x²+1)²",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"The numerator becomes 2x[(x²+1)-(x²-1)]=4x",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"dy/dx = 4x/(x²+1)².",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Mixed trigonometric-algebraic practice")
qtx(slide,"Differentiate y=cos x+2/x³, where x≠0.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"Write 2/x³ as 2x⁻³",XL,Y0,CW,LH,19)
fml(slide,"d(cos x)/dx=-sin x",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"d(2x⁻³)/dx=-6x⁻⁴",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = -sin x-6/x⁴.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Simplify before differentiating")
qtx(slide,"Differentiate y=(sqrt(x)+1/sqrt(x))² for x>0.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"Expand first: y=x+2+1/x",XL,Y0,CW,LH,19)
fml(slide,"Write 1/x as x⁻¹",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"dy/dx=1+0-x⁻²",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"dy/dx = 1-1/x².",XL,Y0+3*LH+3*NG,CW,0.55)

slide=new_slide("Class summary","Differentiation lock lines",TEAL)
fml(slide,"dy/dx is the instantaneous rate of change and tangent slope",XL,Y0,CW,LH,17)
fml(slide,"d(C)/dx=0     d(x^n)/dx=nx^(n-1)",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"differentiate sums term by term; constants stay outside",XL,Y0+2*LH+2*NG,CW,LH,17)
fml(slide,"product: A B'+B A'     quotient: (B A'-A B')/B²",XL,Y0+3*LH+3*NG,CW,LH,17)
body(slide,"At a smooth maximum or minimum, dy/dx=0.",XL,Y0+4*LH+4*NG,CW,0.48)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Redo every question without looking at its solution.",XL,Y0,CW,0.52)
body(slide,"2. Write the six basic derivative formulae from memory.",XL,Y0+0.72,CW,0.52)
body(slide,"3. Draw one curve and label positive, zero and negative slope regions.",XL,Y0+1.44,CW,0.52)
body(slide,"4. Check signs carefully in product and quotient questions.",XL,Y0+2.16,CW,0.52)
trap_card(slide,"Always simplify the function first when it makes differentiation shorter.",XL,Y0+3.08,CW,0.62)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)


# ── COVER ─────────────────────────────────────────────────────────────────────
