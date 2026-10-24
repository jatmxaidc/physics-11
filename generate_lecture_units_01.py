"""Units 01 — Physical quantities, dimensions and unit conversion (PW 01+02 merge)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Units_01_Physical_Quantities_and_Conversion.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(1/2)", "½"),
        ("^(-12)", "⁻¹²"),
        ("^(-7)", "⁻⁷"),
        ("^(-4)", "⁻⁴"),
        ("^(-2)", "⁻²"),
        ("^(-n)", "⁻ⁿ"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^1", "¹"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units 01",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units and Measurement",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Physical Quantities, Dimensions and Conversion",bx,by+1.62,bw,0.40,18,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"6 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Merged PW Units 01 + 02. Vectors and Basic Maths already done."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Physical quantity and why we measure",
"Fundamental, derived, supplementary",
"SI units and how to write them (CBSE)",
"Dimension — simple meaning",
"Must-know dimensions (v, a, F, W, P)",
"Unit conversion method",
"New-system of units",
"Board practice",
]
for i,t in enumerate(topics):
    y=1.25+i*0.58
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(XL),Inches(y),Inches(0.42),Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.fill.background()
    tx(slide,f"{i+1:02d}",XL,y+0.05,0.42,0.32,14,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
    tx(slide,t,XL+0.60,y+0.05,10.5,0.35,20,TEXT,False,spa=0,scaled=False)
_footer(slide,2); after_content_slide()

slide=new_slide("Class note","Where we are today",TEAL)
body(slide,"Basic Maths and Vectors are complete. First NCERT physics chapter starts now.",XL,Y0,CW,0.55)
body(slide,"This class merges PW Units 01 and Units 02 into one teaching hour.",XL,Y0+0.70,CW,0.55)
ncert(slide,"Physical quantity  ->  dimension  ->  unit conversion",XL,Y0+1.50,CW,1.10,18)
trap_card(slide,"Do not memorise later-chapter formulas today. Learn the method.",XL,Y0+2.90,CW,0.55)

slide=new_slide("Chapter map","Units and Measurement",GREEN)
body(slide,"1. Type of physical quantity and unit conversion   (today)",XL,Y0,CW,0.50)
body(slide,"2. Dimensional analysis",XL,Y0+0.60,CW,0.45)
body(slide,"3. Significant figures",XL,Y0+1.15,CW,0.45)
body(slide,"4. Errors",XL,Y0+1.70,CW,0.45)
body(slide,"5. Measuring instruments",XL,Y0+2.25,CW,0.45)
trap_card(slide,"If a later chapter formula appears, do not panic. The method is the same.",XL,Y0+3.05,CW,0.55)

# PHYSICAL QUANTITY
slide=new_slide("Idea","Physical quantity",GREEN)
ncert(slide,"Physical quantity = a quantity that can be measured",XL,Y0,CW,1.00,20)
body(slide,"Examples: mass, length, time, force, energy.",XL,Y0+1.25,CW,0.50)
body(slide,"A measurement is written as  n u   (number times unit).",XL,Y0+1.85,CW,0.50)
trap_card(slide,"Beauty, intelligence, honesty are not physical quantities.",XL,Y0+2.60,CW,0.55)

slide=new_slide("CBSE write","Need for measurement (2 marks)",TEAL)
body(slide,"Write this in simple English. Do not make it long.",XL,Y0,CW,0.40)
body(slide,"1. Physics is the study of nature using numbers.",XL,Y0+0.50,CW,0.45)
body(slide,"2. To compare two things, we must measure them.",XL,Y0+1.00,CW,0.45)
body(slide,"3. Measurement = how many times a standard unit fits into the quantity.",XL,Y0+1.50,CW,0.50)
ncert(slide,"Physical quantity = numerical value  x  unit",XL,Y0+2.20,CW,1.05,18)
trap_card(slide,"Always write the unit. A number without a unit is incomplete.",XL,Y0+3.50,CW,0.50)

slide=new_slide("CBSE write","How to write any 3-mark theory answer",GREEN)
body(slide,"Step 1: Definition in one line (easy words).",XL,Y0,CW,0.45)
body(slide,"Step 2: Formula or symbol, if any.",XL,Y0+0.55,CW,0.45)
body(slide,"Step 3: One example.",XL,Y0+1.10,CW,0.45)
body(slide,"Step 4: One extra point (use / difference / trap).",XL,Y0+1.65,CW,0.50)
ncert(slide,"Definition  +  formula  +  example  =  full marks",XL,Y0+2.35,CW,1.05,18)
trap_card(slide,"Do not copy hard English. Examiner wants clear points, not big words.",XL,Y0+3.65,CW,0.50)

slide=new_slide("Copy","Three types of physical quantity",TEAL)
mini_card(slide,"Fundamental","Independent. All other quantities are built from these.",XL,Y0,3.80,1.70,GREEN)
mini_card(slide,"Derived","Made from fundamentals: v = s/t, F = m a.",XL+4.10,Y0,3.80,1.70,TEAL)
mini_card(slide,"Supplementary","Angle (rad) and solid angle (sr). Have a unit, no dimension.",XL+8.20,Y0,3.80,1.70,ORANGE)
body(slide,"Fundamental means: you cannot derive one of them from the others.",XL,Y0+2.00,CW,0.50)
trap_card(slide,"Two quantities can share a unit and still be different (work and torque).",XL,Y0+2.70,CW,0.55)

slide=new_slide("CBSE write","Fundamental vs derived (3 marks)",GREEN)
grid_table(slide,
["Point", "Fundamental", "Derived"],
[
["Meaning", "Independent", "Made from fundamentals"],
["Count in SI", "Seven only", "Many"],
["Example", "mass, time", "speed, force"],
["Unit", "kg, s, m ...", "m/s, N, J ..."],
["Can we skip one?", "No", "Yes, if defined"],
],
XL, Y0, [2.6, 4.7, 4.7], 0.48, 13, 14)
trap_card(slide,"End the answer with: All derived quantities are built from the seven base ones.",XL,Y0+3.10,CW,0.55)

slide=new_slide("Copy","Seven SI base quantities",GREEN)
grid_table(slide,
["Quantity", "SI unit", "Symbol"],
[
["Mass", "kilogram", "kg"],
["Length", "metre", "m"],
["Time", "second", "s"],
["Temperature", "kelvin", "K"],
["Current", "ampere", "A"],
["Amount of substance", "mole", "mol"],
["Luminous intensity", "candela", "cd"],
],
XL, Y0, [4.4, 3.8, 3.8], 0.48, 14, 15)

slide=new_slide("CBSE write","Systems of units (2–3 marks)",TEAL)
body(slide,"CGS: centimetre, gram, second.",XL,Y0,CW,0.42)
body(slide,"FPS: foot, pound, second.   MKS: metre, kilogram, second.",XL,Y0+0.50,CW,0.48)
body(slide,"SI: international system. It starts from MKS and adds A, K, mol, cd.",XL,Y0+1.05,CW,0.50)
ncert(slide,"We use SI because it is one system for the whole world.",XL,Y0+1.75,CW,1.05,17)
body(slide,"Advantage: same unit everywhere, easy to convert, coherent (no extra numbers).",XL,Y0+2.95,CW,0.55)
trap_card(slide,"Write SI, not S.I. in running text. Say 'SI unit of force is newton (N)'.",XL,Y0+3.65,CW,0.50)

slide=new_slide("CBSE write","Rules for writing units (copy this)",GREEN)
body(slide,"1. Unit name is not capital: metre, newton, joule.  (Exception: Celsius.)",XL,Y0,CW,0.50)
body(slide,"2. Symbol may be capital if named after a person: N, J, W, K.",XL,Y0+0.55,CW,0.50)
body(slide,"3. Never write plural: 10 kg, not 10 kgs. Do not put a full stop after the symbol.",XL,Y0+1.10,CW,0.50)
body(slide,"4. Leave a gap: 5 m, 10 N. Write m/s or m s^(-1), not mps.",XL,Y0+1.65,CW,0.50)
body(slide,"5. Quantity, unit and symbol are three different things.",XL,Y0+2.20,CW,0.50)
ncert(slide,"Quantity: force     Unit: newton     Symbol: N",XL,Y0+2.85,CW,1.05,18)
trap_card(slide,"Wrong: 10 Newton / 10 n / 10 nts.  Right: 10 N.",XL,Y0+4.10,CW,0.48)

slide=new_slide("Copy","Derived quantities from definition",TEAL)
body(slide,"Velocity = displacement / time   ->   m/s",XL,Y0,CW,0.45)
body(slide,"Acceleration = change in velocity / time   ->   m/s^2",XL,Y0+0.50,CW,0.45)
body(slide,"Force = mass x acceleration   ->   kg m/s^2 = 1 N",XL,Y0+1.00,CW,0.45)
body(slide,"Momentum = mass x velocity   ->   kg m/s",XL,Y0+1.50,CW,0.45)
body(slide,"Energy = (1/2) m v^2   ->   kg m^2/s^2 = 1 J",XL,Y0+2.00,CW,0.45)
body(slide,"Power = energy / time   ->   J/s = W",XL,Y0+2.50,CW,0.45)
trap_card(slide,"Same unit does not prove same physical quantity.",XL,Y0+3.15,CW,0.50)

slide=question_slide("Which group can be fundamental?")
qtx(slide,"Which of the following groups may be treated as a set of fundamental quantities?",XL,Y0,CW,1.05,16)
choices(slide,[
"mass, momentum, velocity",
"force, acceleration, mass",
"force, time, velocity",
"none of these",
], XL, Y0+1.25, 10.8)
slide=answer_slide()
fml(slide,"Fundamental means independent: none can be written using only the others.",XL,Y0,CW,LH,15)
fml(slide,"(a) p = m v, so momentum is derived. (b) F = m a, so they are linked.",XL,Y0+LH+NG,CW,LH,15)
fml(slide,"(c) F, t, v cannot make each other. They can form a valid base set.",XL,Y0+2*LH+2*NG,CW,LH,15)
result_card(slide,"Answer: force, time, velocity.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=question_slide("Characteristic of a unit")
qtx(slide,"Which of the following is a characteristic of a unit?",XL,Y0,CW,0.90,18)
choices(slide,[
"universally accepted",
"invariable and well defined",
"suitable size and easily available",
"all of the above",
], XL, Y0+1.15, 10.8)
slide=answer_slide()
result_card(slide,"Answer: all of the above.",XL,Y0,CW,0.55)

slide=new_slide("CBSE write","A good unit — 3 points to remember",TEAL)
body(slide,"Write these three lines. That is the full 3-mark answer.",XL,Y0,CW,0.45)
body(slide,"1. It must be the same for everyone (universally accepted).",XL,Y0+0.55,CW,0.50)
body(slide,"2. It must not change with time or place (invariable, well defined).",XL,Y0+1.10,CW,0.50)
body(slide,"3. It should be of handy size and easy to copy (easily available).",XL,Y0+1.65,CW,0.50)
ncert(slide,"Accepted  +  does not change  +  easy to use",XL,Y0+2.35,CW,1.05,18)
trap_card(slide,"Hand-span and foot-step failed because they change from person to person.",XL,Y0+3.60,CW,0.50)

# DIMENSIONS
slide=new_slide("Idea","What is a dimension?",GREEN)
ncert(slide,"Dimension = power of a base quantity in the expression of a P.Q.",XL,Y0,CW,1.10,18)
body(slide,"Mass -> M^1     Length -> L^1     Time -> T^1     Current -> A^1",XL,Y0+1.35,CW,0.50)
body(slide,"Temperature -> K^1 (or theta^1)     Amount -> mol     Luminous intensity -> cd",XL,Y0+1.95,CW,0.50)
trap_card(slide,"The letter M is mass, not metre. Metre is the unit of length L.",XL,Y0+2.65,CW,0.55)

slide=new_slide("CBSE write","Dimensional formula — how to say it",TEAL)
body(slide,"Dimensional formula: the expression that shows powers of M, L, T.",XL,Y0,CW,0.50)
body(slide,"Example: [F] = M L T^(-2). This is the dimensional formula of force.",XL,Y0+0.55,CW,0.50)
body(slide,"Dimensional equation: quantity = its dimensional formula.",XL,Y0+1.10,CW,0.50)
body(slide,"Example: [force] = [M L T^(-2)].",XL,Y0+1.65,CW,0.45)
ncert(slide,"Write [ ] around the quantity. Then write M, L, T with powers.",XL,Y0+2.30,CW,1.05,17)
trap_card(slide,"CBSE likes the word 'dimensional formula'. Use it, not only 'dimension'.",XL,Y0+3.55,CW,0.50)

slide=new_slide("CBSE write","Unit vs dimension (2 marks)",GREEN)
grid_table(slide,
["", "Unit", "Dimension"],
[
["Meaning", "standard used to measure", "powers of base quantities"],
["Can change?", "yes (m or cm)", "no (length is always L)"],
["Example", "newton, joule", "M L T^(-2)"],
["Depends on system?", "yes (SI / CGS)", "no"],
],
XL, Y0, [3.2, 4.4, 4.4], 0.52, 13, 14)
trap_card(slide,"Force in SI is N, in CGS is dyne. Dimension is M L T^(-2) in both.",XL,Y0+3.20,CW,0.55)

slide=new_slide("CBSE write","Dimensionless quantities (2 marks)",TEAL)
body(slide,"A quantity with all powers zero is dimensionless: M^0 L^0 T^0.",XL,Y0,CW,0.50)
body(slide,"It may still have a unit. Angle has unit radian, but no dimension.",XL,Y0+0.55,CW,0.50)
body(slide,"Examples: angle, solid angle, strain, refractive index, relative density.",XL,Y0+1.10,CW,0.50)
ncert(slide,"No dimension  does not always mean  no unit",XL,Y0+1.85,CW,1.05,18)
trap_card(slide,"Strain = change in length / original length = L/L = 1. Pure number.",XL,Y0+3.15,CW,0.50)

slide=new_slide("Copy","Anything that is a length",TEAL)
body(slide,"Displacement, distance, height, width, diameter, amplitude,",XL,Y0,CW,0.50)
body(slide,"wavelength, radius of curvature, range of projectile, mean free path,",XL,Y0+0.55,CW,0.50)
body(slide,"focal length, circumference, fringe width, position, elongation.",XL,Y0+1.10,CW,0.50)
ncert(slide,"All of these have dimension L^1 M^0 T^0",XL,Y0+1.85,CW,1.05,20)
trap_card(slide,"If it is a measured length, write [L]. Do not overthink the name.",XL,Y0+3.15,CW,0.55)

slide=new_slide("Copy","Velocity family vs angular velocity",GREEN)
ncert(slide,"[velocity] = L T^(-1)",XL,Y0,CW,0.95,22)
body(slide,"Same dimension: speed, average / orbital / escape / terminal / drift velocity.",XL,Y0+1.20,CW,0.55)
ncert(slide,"[angular velocity omega] = T^(-1)   because omega = theta / t",XL,Y0+1.95,CW,1.10,18)
trap_card(slide,"omega is not a velocity. v = r omega mixes L into it.",XL,Y0+3.30,CW,0.55)

slide=new_slide("Copy","Mechanics dimensions you must own",TEAL)
grid_table(slide,
["Quantity", "From", "Dimension"],
[
["Acceleration", "v / t", "L T^(-2)"],
["Force", "m a", "M L T^(-2)"],
["Momentum / impulse", "m v", "M L T^(-1)"],
["Work / energy / torque", "F s", "M L^2 T^(-2)"],
["Power", "W / t", "M L^2 T^(-3)"],
["Pressure / stress", "F / A", "M L^(-1) T^(-2)"],
["Density", "m / V", "M L^(-3)"],
],
XL, Y0, [4.2, 2.6, 5.2], 0.48, 14, 15)
trap_card(slide,"Work and torque share dimension. They are still different quantities.",XL,Y0+4.45,CW,0.48)

slide=new_slide("Copy","Angular set and moment of inertia",GREEN)
body(slide,"Angle theta has no dimension: M^0 L^0 T^0. Unit is radian.",XL,Y0,CW,0.50)
body(slide,"Angular velocity omega = theta / t  ->  T^(-1)",XL,Y0+0.55,CW,0.45)
body(slide,"Angular acceleration = omega / t  ->  T^(-2)",XL,Y0+1.10,CW,0.45)
body(slide,"Moment of inertia I = m r^2  ->  M L^2",XL,Y0+1.65,CW,0.45)
ncert(slide,"Angular momentum L = r x p  ->  M L^2 T^(-1)",XL,Y0+2.30,CW,1.05,18)
trap_card(slide,"Torque = r x F has the same dimension as energy.",XL,Y0+3.55,CW,0.50)

slide=new_slide("Copy","Gradients and frequency",TEAL)
body(slide,"Gradient means change per unit length.",XL,Y0,CW,0.45)
body(slide,"Velocity gradient = v / L  ->  T^(-1)   (same as frequency)",XL,Y0+0.55,CW,0.50)
body(slide,"Temperature gradient = K / L  ->  K L^(-1)",XL,Y0+1.10,CW,0.50)
body(slide,"Frequency = 1/T  ->  T^(-1)     Charge = I t  ->  A T",XL,Y0+1.65,CW,0.50)
ncert(slide,"Area = L^2     Volume = L^3     Density = M L^(-3)",XL,Y0+2.35,CW,1.05,18)
trap_card(slide,"Same dimension (T^(-1)) does not mean same quantity.",XL,Y0+3.60,CW,0.50)

slide=new_slide("Copy","Every force is still a force",GREEN)
body(slide,"Friction, tension, spring force, weight, buoyancy, electrostatic force,",XL,Y0,CW,0.50)
body(slide,"magnetic force, nuclear force, gravitational force: all [F] = M L T^(-2).",XL,Y0+0.55,CW,0.55)
ncert(slide,"EMF is not a force. Volt = J/C = M L^2 T^(-3) A^(-1)",XL,Y0+1.35,CW,1.10,18)
body(slide,"Every energy (KE, PE, heat, work function) has [E] = M L^2 T^(-2).",XL,Y0+2.70,CW,0.55)
trap_card(slide,"Name can trick you. Dimension follows the definition.",XL,Y0+3.45,CW,0.50)

# UNIT CONVERSION
slide=new_slide("Idea","What a unit is",TEAL)
ncert(slide,"Unit = a reference value. Measurement = n u",XL,Y0,CW,1.05,20)
body(slide,"Example: 5 m means 5 times the reference 1 m. Also 500 cm.",XL,Y0+1.30,CW,0.50)
ncert(slide,"n u = constant, so  n1 u1 = n2 u2",XL,Y0+2.05,CW,1.05,22)
trap_card(slide,"If the unit becomes larger, the number becomes smaller.",XL,Y0+3.35,CW,0.50)

slide=new_slide("CBSE write","How to convert units (write these steps)",GREEN)
body(slide,"In a numerical, write these four lines. Examiner ticks each line.",XL,Y0,CW,0.48)
body(slide,"1. Write the given value with its unit.",XL,Y0+0.55,CW,0.45)
body(slide,"2. Replace each unit by the new unit (1 m = 100 cm, 1 kg = 1000 g).",XL,Y0+1.10,CW,0.50)
body(slide,"3. Keep the power on the unit: m^2 becomes (100 cm)^2.",XL,Y0+1.65,CW,0.50)
body(slide,"4. Multiply and write the final number with the new unit.",XL,Y0+2.20,CW,0.50)
ncert(slide,"n1 u1 = n2 u2     (the actual amount does not change)",XL,Y0+2.90,CW,1.05,18)
trap_card(slide,"Bigger unit -> smaller number. If your number grew with a bigger unit, check again.",XL,Y0+4.15,CW,0.48)

slide=new_slide("Copy","SI vs CGS you will convert",GREEN)
grid_table(slide,
["Quantity", "SI", "CGS"],
[
["Mass", "kg", "g"],
["Length", "m", "cm"],
["Time", "s", "s"],
["Force", "N = kg m/s^2", "dyne = g cm/s^2"],
["Energy", "J = kg m^2/s^2", "erg = g cm^2/s^2"],
["Power", "W = J/s", "erg/s"],
["Momentum", "kg m/s", "g cm/s"],
],
XL, Y0, [3.0, 4.5, 4.5], 0.48, 14, 14)
trap_card(slide,"1 N = 10^5 dyne.   1 J = 10^7 erg.",XL,Y0+4.45,CW,0.48)

slide=new_slide("How it works","Conversion method",TEAL)
body(slide,"Write the quantity with its SI unit, then replace each unit.",XL,Y0,CW,0.50)
body(slide,"1 m = 100 cm,  1 kg = 1000 g,  time unit is the same second.",XL,Y0+0.55,CW,0.50)
ncert(slide,"Area 5 m^2 = 5 (100 cm)^2 = 5 x 10^4 cm^2",XL,Y0+1.30,CW,1.10,18)
body(slide,"Powers travel with the unit: m^2 becomes (100 cm)^2, not 100 cm^2.",XL,Y0+2.65,CW,0.55)
trap_card(slide,"Do not convert 5 and then forget the power on the unit.",XL,Y0+3.40,CW,0.50)

slide=question_slide("Area 5 m^2 in CGS")
qtx(slide,"If area of an object is 5 m^2, find its value in CGS unit.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"5 m^2 = 5 (1 m)^2 = 5 (100 cm)^2",XL,Y0,CW,LH,18)
fml(slide,"= 5 x 10^4 cm^2",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"5 x 10^4 cm^2.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("25 m/s in CGS")
qtx(slide,"Convert 25 m/s into CGS unit.",XL,Y0,CW,0.85,20)
slide=answer_slide()
fml(slide,"25 m/s = 25 (100 cm / s) = 2500 cm/s",XL,Y0,CW,LH,18)
result_card(slide,"2500 cm/s.",XL,Y0+LH+NG,CW,0.48)

slide=question_slide("36 km/h into m/s")
qtx(slide,"Convert 36 km/h into m/s.",XL,Y0,CW,0.80,20)
slide=answer_slide()
fml(slide,"36 km/h = 36 x (1000 m) / (3600 s) = 10 m/s",XL,Y0,CW,LH,17)
fml(slide,"Shortcut: x km/h = x x (5/18) m/s",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"10 m/s.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("1 newton into dyne")
qtx(slide,"Convert 1 N into dyne.",XL,Y0,CW,0.80,20)
slide=answer_slide()
fml(slide,"1 N = 1 kg m / s^2 = (10^3 g)(10^2 cm) / s^2",XL,Y0,CW,LH,16)
fml(slide,"= 10^5 g cm / s^2 = 10^5 dyne",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"1 N = 10^5 dyne.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("10 J in CGS")
qtx(slide,"Convert 10 J into CGS unit.",XL,Y0,CW,0.85,20)
slide=answer_slide()
fml(slide,"10 J = 10 kg m^2 / s^2 = 10 (10^3 g)(10^2 cm)^2 / s^2",XL,Y0,CW,LH,15)
fml(slide,"= 10 x 10^7 erg = 10^8 erg",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"10^8 erg.  (1 J = 10^7 erg)",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("KE in SI and CGS")
qtx(slide,"Mass m = 10 kg, velocity 4 m/s. Find KE in SI and in CGS.",XL,Y0,CW,1.00,17)
slide=answer_slide()
fml(slide,"KE = (1/2) m v^2 = (1/2)(10)(16) = 80 J",XL,Y0,CW,LH,18)
fml(slide,"80 J = 80 x 10^7 erg = 8 x 10^8 erg",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"80 J  or  8 x 10^8 erg.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("If the length unit is doubled")
qtx(slide,"If the unit of length becomes double the SI metre, the area 5 m^2 in the new unit is",XL,Y0,CW,1.10,16)
slide=answer_slide()
fml(slide,"n u = constant. New length unit L' = 2 m, so 1 m = (1/2) L'",XL,Y0,CW,LH,15)
fml(slide,"5 m^2 = 5 (1/2 L')^2 = 5/4 = 1.25 new area units",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"1.25 (because area scales as L^2).",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("New unit of force")
qtx(slide,"New system: 1 unit of length = 2 m, mass = 10 kg, time = 4 s. Find the unit of force.",XL,Y0,CW,1.15,15)
slide=answer_slide()
fml(slide,"[F] = M L T^(-2), so 1 new force unit = 10 kg x 2 m / (4 s)^2",XL,Y0,CW,LH,15)
fml(slide,"= 20 / 16 = 1.25 N",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"1 new force unit = 1.25 N.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("10 J in a new system")
qtx(slide,"New units: length 10 m, mass 5 kg, time 2 s. Value of 10 J in this system?",XL,Y0,CW,1.10,16)
slide=answer_slide()
fml(slide,"1 new energy unit = M L^2 T^(-2) = 5 x (10)^2 / (2)^2 = 125 J",XL,Y0,CW,LH,15)
fml(slide,"n_new = 10 / 125 = 0.08",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"0.08 new energy units.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("Density 5 kg/m^3 in a new system")
qtx(slide,"New system: 1 length unit = x m, 1 mass unit = y kg. Value of density 5 kg/m^3?",XL,Y0,CW,1.15,15)
slide=answer_slide()
fml(slide,"1 m = 1/x of the new length,  1 kg = 1/y of the new mass",XL,Y0,CW,LH,15)
fml(slide,"5 kg/m^3 = 5 (1/y) / (1/x)^3 = (5 x^3 / y) new density units",XL,Y0+LH+NG,CW,LH,15)
result_card(slide,"(5 x^3 / y) new units.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Match the dimensions")
qtx(slide,"Match: A angular momentum, B torque, C stress, D pressure gradient.",XL,Y0,CW,1.05,16)
choices(slide,[
"A-I, B-IV, C-III, D-II",
"A-III, B-I, C-IV, D-II",
"A-II, B-III, C-IV, D-I",
"A-IV, B-II, C-I, D-III",
], XL, Y0+1.20, 10.8)
body(slide,"I [M L^2 T^(-2)]   II [M L^(-2) T^(-2)]   III [M L^2 T^(-1)]   IV [M L^(-1) T^(-2)]",XL,Y0+3.00,CW,0.55,14)
slide=answer_slide()
fml(slide,"L = r p -> M L^2 T^(-1) = III.   Torque = r F -> M L^2 T^(-2) = I.",XL,Y0,CW,LH,14)
fml(slide,"Stress = F/A -> M L^(-1) T^(-2) = IV.   Pressure gradient = P/L -> II.",XL,Y0+LH+NG,CW,LH,14)
result_card(slide,"Answer: A-III, B-I, C-IV, D-II.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Density in a new CGS-like system")
qtx(slide,"Density is 4 g/cm^3. New units: length 10 cm, mass 100 g. New density value is",XL,Y0,CW,1.15,15)
choices(slide,["0.04","0.4","40","400"], XL, Y0+1.30, 10.8)
slide=answer_slide()
fml(slide,"4 g/cm^3 = 4 (100 g / 100) / (10 cm / 10)^3 wait: 1 new mass = 100 g, 1 new length = 10 cm",XL,Y0,CW,LH,13)
fml(slide,"n' = 4 x (1/100) / (1/10)^3 = 4 x 0.01 x 1000 = 40",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"Answer: 40.",XL,Y0+2*LH+2*NG,CW,0.48)

# SUMMARY
slide=new_slide("Class summary","Units 01 lock lines",TEAL)
fml(slide,"Seven base quantities. Derived = built from these.",XL,Y0,CW,LH,16)
fml(slide,"Dimension = powers of M, L, T, ... in the formula.",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"[F] = M L T^(-2).  [W] = [torque] = M L^2 T^(-2).  [P] = M L^2 T^(-3).",XL,Y0+2*LH+2*NG,CW,LH,14)
fml(slide,"n u = constant. Larger unit -> smaller number. Powers travel with units.",XL,Y0+3*LH+3*NG,CW,LH,14)
result_card(slide,"1 N = 10^5 dyne.  1 J = 10^7 erg.  x km/h = x x (5/18) m/s.",XL,Y0+4*LH+4*NG,CW,0.55)

slide=new_slide("CBSE write","Theory lines to copy in the paper",GREEN)
body(slide,"1. Physical quantity = numerical value x unit.",XL,Y0,CW,0.45)
body(slide,"2. A unit must be accepted, unchanging, and of suitable size.",XL,Y0+0.50,CW,0.48)
body(slide,"3. SI has seven base units. All other units are derived.",XL,Y0+1.00,CW,0.45)
body(slide,"4. Dimension does not change with the system of units. Unit does.",XL,Y0+1.50,CW,0.50)
body(slide,"5. Angle is dimensionless but has a unit (radian).",XL,Y0+2.00,CW,0.45)
body(slide,"6. Work and torque have the same dimension, not the same meaning.",XL,Y0+2.50,CW,0.50)
trap_card(slide,"For 3 marks: definition + formula/example + one extra line.",XL,Y0+3.20,CW,0.50)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Write dimensions of velocity, acceleration, force, energy, power, pressure.",XL,Y0,CW,0.50)
body(slide,"2. Convert 2 N into dyne and 5 J into erg.",XL,Y0+0.55,CW,0.45)
body(slide,"3. Convert 72 km/h into m/s.",XL,Y0+1.10,CW,0.45)
body(slide,"4. New units: L = 2 m, M = 2 kg, T = 2 s. Find 1 new unit of energy in joule.",XL,Y0+1.65,CW,0.50)
body(slide,"5. Why can work and torque have the same dimension but remain different?",XL,Y0+2.20,CW,0.50)
body(slide,"6. Density 8 g/cm^3. New mass unit 200 g, new length unit 5 cm. Find new density.",XL,Y0+2.75,CW,0.50)
trap_card(slide,"Next class: uses of dimensional analysis (checking, converting, guessing formulas).",XL,Y0+3.40,CW,0.50)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
