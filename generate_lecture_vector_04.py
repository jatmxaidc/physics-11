"""Vector 04 — from Vector 02 slide 91, then fresh practice (no earlier repeats)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Vector_04_Revision_and_Practice.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(1/2)", "½"),
        ("^(-12)", "⁻¹²"),
        ("^(-7)", "⁻⁷"),
        ("^(-4)", "⁻⁴"),
        ("^(-2)", "⁻²"),
        ("^(-n)", "⁻ⁿ"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vector 04",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vectors",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Components, Dot Product and Practice",bx,by+1.62,bw,0.40,18,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"3 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Lean merge: V02 from s91 + unique V04 practice."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Adding by components",
"Change in velocity and rotation",
"Dot product",
"Unit vector practice",
"Angle and direction from components",
"Fresh JEE-style practice",
"Cross product (added)",
]
for i,item in enumerate(topics):
    yy=1.10+i*0.70
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.46,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.50,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Repeats from Vector 01 and Vector 02 before slide 91 removed."); after_content_slide()

slide=new_slide("Class note","What we are skipping today",TEAL)
body(slide,"Already done in Vector 01 and Vector 02 before slide 91:",XL,Y0,CW,0.50)
body(slide,"· scalar multiply, compass language, reverse/scale angles",XL,Y0+0.65,CW,0.50)
body(slide,"· rectangular components and magnitude",XL,Y0+1.20,CW,0.45)
body(slide,"· triangle law, resultant formula, range, equal-magnitude tricks",XL,Y0+1.75,CW,0.55)
body(slide,"· subtraction idea, and the repeated resultant drills",XL,Y0+2.35,CW,0.50)
trap_card(slide,"Today starts at adding by components, then moves into dot product and new practice. Cross product is added at the end.",XL,Y0+3.10,CW,0.60)

# PART A — Vector 02 from slide 91

slide=new_slide("Visual idea","Adding by components is adding columns",GREEN)
body(slide,"The x-parts combine with x-parts. The y-parts combine with y-parts. This is the cleanest method for many vectors.",XL,Y0,CW,0.55)
ox,oy=axis_diagram(slide,1.10,2.05,4.7,3.7)
vec(slide,ox,oy,ox+1.80,oy-1.15,"A",GOLD); draw_line(slide,ox+1.80,oy-1.15,ox+1.80,oy,WORK,1.2); draw_line(slide,ox,oy,ox+1.80,oy,WORK,1.2)
vec(slide,ox,oy,ox+1.25,oy+1.45,"B",GREEN); draw_line(slide,ox+1.25,oy+1.45,ox+1.25,oy,ORANGE,1.2); draw_line(slide,ox,oy,ox+1.25,oy,ORANGE,1.2)
ncert(slide,"A + B = (Aₓ + Bₓ) î + (Aᵧ + Bᵧ) ĵ",6.60,2.10,5.45,0.90,19)
mini_card(slide,"Column idea","x-column decides left-right. y-column decides up-down.",6.60,3.55,5.45,1.15,TEAL)
mini_card(slide,"3D same rule","A + B = (Aₓ+Bₓ)î + (Aᵧ+Bᵧ)ĵ + (A_z+B_z)k̂",6.60,5.00,5.45,1.15,GREEN)

slide=question_slide("Add by components")
qtx(slide,"A = 4 i-hat + 3 j-hat and B = 3 i-hat - 4 j-hat. Find A + B, A - B, and |A| + |B|.",XL,Y0,CW,1.00,17)
slide=answer_slide()
fml(slide,"A + B = 7 i-hat - 1 j-hat",XL,Y0,CW,LH,18)
fml(slide,"A - B = 1 i-hat + 7 j-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"|A| = 5 and |B| = 5, so |A| + |B| = 10",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"Sum of magnitudes is 10. Magnitude of the sum is sqrt(50).",XL,Y0+3*LH+3*NG,CW,0.55)
ox,oy=axis_diagram(slide,8.25,3.05,3.6,2.9)
vec(slide,ox,oy,ox+1.45,oy-1.05,"A",GOLD)
vec(slide,ox,oy,ox+1.10,oy+1.35,"B",GREEN)
vec(slide,ox,oy,ox+2.05,oy+0.30,"A+B",WORK)

slide=question_slide("Change in velocity")
qtx(slide,"Velocity changes from 10 m/s east to 10 m/s north. Find the change in velocity.",XL,Y0,CW,0.90,18)
slide=answer_slide()
fml(slide,"Change = final - initial = 10 j-hat - 10 i-hat",XL,Y0,CW,LH,18)
fml(slide,"Magnitude = sqrt(100 + 100) = 10 sqrt(2)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"10 sqrt(2) m/s, pointing north of west.",XL,Y0+2*LH+2*NG,CW,0.52)
ox,oy=axis_diagram(slide,8.4,3.15,3.4,2.8)
vec(slide,ox,oy,ox+1.55,oy,"vᵢ",GOLD)
vec(slide,ox,oy,ox,oy-1.55,"vᶠ",GREEN)
vec(slide,ox+1.55,oy,ox,oy-1.55,"Δv",WORK)

slide=question_slide("Rotate 2 i-hat")
qtx(slide,"A = 2 i-hat is rotated by 90 degrees, 180 degrees, 270 degrees and 360 degrees anticlockwise. Write each new vector.",XL,Y0,CW,0.95,17)
slide=answer_slide()
fml(slide,"90 degrees: 2 j-hat",XL,Y0,CW,LH,18)
fml(slide,"180 degrees: -2 i-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"270 degrees: -2 j-hat.   360 degrees: 2 i-hat again",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"Rotation keeps the magnitude and changes only the direction.",XL,Y0+3*LH+3*NG,CW,0.52)
ox,oy=axis_diagram(slide,8.4,3.05,3.4,2.9)
vec(slide,ox,oy,ox+1.25,oy,"0°",GOLD)
vec(slide,ox,oy,ox,oy-1.25,"90°",GREEN)
vec(slide,ox,oy,ox-1.25,oy,"180°",WORK)
vec(slide,ox,oy,ox,oy+1.25,"270°",ORANGE)

# DOT
slide=new_slide("New topic","Dot product",GREEN)
ncert(slide,"A dot B = A B cos theta = Aₓ Bₓ + Aᵧ Bᵧ + A_z B_z",XL,Y0,CW,1.00,20)
body(slide,"If the dot product is 0, the vectors are perpendicular.",XL,Y0+1.20,CW,0.50)
body(slide,"cos theta = (A dot B) / (A B). That is how the angle is found from components.",XL,Y0+1.80,CW,0.55)
body(slide,"Component of B along A = (A dot B) / A.",XL,Y0+2.45,CW,0.50)
body(slide,"i-hat dot i-hat = 1. i-hat dot j-hat = 0. The same rule holds for j-hat and k-hat.",XL,Y0+3.10,CW,0.55)
trap_card(slide,"A unit vector dotted with itself is 1.",XL,Y0+3.80,CW,0.48)
vec(slide,8.15,5.10,10.50,5.10,"A",GOLD)
vec(slide,8.15,5.10,9.65,3.70,"B",GREEN)
draw_line(slide,9.65,3.70,9.65,5.10,WORK,1.2)
arrow_line(slide,8.15,5.10,9.65,5.10,WORK,2.0)
tx(slide,"B cos θ",8.55,5.22,1.25,0.28,13,WORK,True,spa=0,scaled=False)
tx(slide,"projection of B along A",9.30,5.42,2.5,0.25,13,MUTED,True,spa=0,scaled=False)

slide=new_slide("Visual idea","Dot product tells alignment",GREEN)
body(slide,"Dot product is positive when vectors face mostly the same way, zero at 90°, and negative when they face opposite way.",XL,Y0,CW,0.55)
vec(slide,1.25,3.00,3.10,3.00,"A",GOLD); vec(slide,1.25,3.00,2.85,2.25,"B",GREEN); angle_mark(slide,1.25,3.00,0.60,0,25,"θ<90°")
tx(slide,"A · B > 0",1.35,3.55,1.8,0.30,17,GREEN,True,spa=0,scaled=False)
vec(slide,5.35,3.25,7.05,3.25,"A",GOLD); vec(slide,5.35,3.25,5.35,1.60,"B",GREEN); angle_mark(slide,5.35,3.25,0.60,0,90,"90°")
tx(slide,"A · B = 0",5.35,3.75,1.8,0.30,17,WORK,True,spa=0,scaled=False)
vec(slide,9.15,3.00,10.95,3.00,"A",GOLD); vec(slide,9.15,3.00,7.85,2.20,"B",GREEN); angle_mark(slide,9.15,3.00,0.65,140,180,"θ>90°")
tx(slide,"A · B < 0",9.05,3.55,1.8,0.30,17,ORANGE,True,spa=0,scaled=False)
mini_card(slide,"Why this matters","It quickly tests perpendicular vectors and helps find the angle between two vectors.",3.90,5.20,5.50,1.00,TEAL)

slide=question_slide("Are these two vectors perpendicular?")
qtx(slide,"A = 3 i-hat + 4 j-hat + 5 k-hat and B = 3 i-hat + 4 j-hat - 5 k-hat. What is the angle between them?",XL,Y0,CW,1.05,17)
slide=answer_slide()
fml(slide,"A dot B = 9 + 16 - 25 = 0",XL,Y0,CW,LH,18)
result_card(slide,"The angle is 90 degrees.",XL,Y0+LH+NG,CW,0.50)

slide=question_slide("Angle with 4 i-hat")
qtx(slide,"A = 2 i-hat + 6 j-hat + 3 k-hat and B = 4 i-hat. Find cos theta.",XL,Y0,CW,0.90,19)
slide=answer_slide()
fml(slide,"A dot B = 8",XL,Y0,CW,LH,18)
fml(slide,"|A| = sqrt(4 + 36 + 9) = 7 and |B| = 4",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"cos theta = 8 / 28 = 2/7",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"theta = cos inverse of 2/7.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=new_slide("Copy","Unit vectors from components",TEAL)
body(slide,"i-hat itself is already a unit vector.",XL,Y0,CW,0.45)
body(slide,"Unit vector of i-hat + j-hat is (i-hat + j-hat) / sqrt(2).",XL,Y0+0.60,CW,0.50)
body(slide,"Unit vector of i-hat + j-hat - k-hat is (i-hat + j-hat - k-hat) / sqrt(3).",XL,Y0+1.25,CW,0.55)
body(slide,"sin theta i-hat - cos theta j-hat already has magnitude 1, because sin^2 + cos^2 = 1.",XL,Y0+1.95,CW,0.60)
body(slide,"Unit vector of 3 i-hat + 4 j-hat is (3 i-hat + 4 j-hat) / 5.",XL,Y0+2.70,CW,0.55)

slide=new_slide("Visual idea","Unit vector keeps direction and removes size",TEAL)
body(slide,"A unit vector is only the direction of a vector. It has length 1 and no physical unit.",XL,Y0,CW,0.50)
ox,oy=axis_diagram(slide,1.20,2.20,4.8,3.5)
vec(slide,ox,oy,ox+2.40,oy-1.60,"A",GOLD,3.0)
vec(slide,ox,oy,ox+1.20,oy-0.80,"Â",GREEN,2.4)
draw_line(slide,ox+1.20,oy-0.80,ox+2.40,oy-1.60,MUTED,1.0)
ncert(slide,"Â = A / |A|",6.75,2.30,4.60,0.90,26)
mini_card(slide,"Example","For A = 3î + 4ĵ, |A| = 5, so Â = (3î + 4ĵ)/5.",6.75,3.80,4.60,1.25,GREEN)
trap_card(slide,"Do not attach N, m/s or kg to a unit vector.",6.75,5.35,4.60,0.55)

slide=question_slide("Perpendicular forces A and A/2")
qtx(slide,"Two forces of magnitudes A and A/2 are perpendicular. Find the magnitude of their resultant.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"cos 90 = 0, so R = sqrt(A^2 + A^2/4)",XL,Y0,CW,LH,18)
fml(slide,"R = sqrt(5 A^2 / 4) = (A/2) sqrt(5)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"(sqrt(5) / 2) times A.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Resultant is half the larger vector")
qtx(slide,"The resultant of two vectors is half the larger one, and it is perpendicular to the smaller one. Find the angle between the vectors.",XL,Y0,CW,1.10,17)
slide=answer_slide()
fml(slide,"R perpendicular to the smaller vector means (A + B) dot B = 0",XL,Y0,CW,LH,15)
fml(slide,"cos theta = -B/A",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"R = A/2 then gives B = (sqrt(3)/2) A",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"cos theta = -sqrt(3)/2, so theta = 150 degrees.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=question_slide("Acceleration of a 4 kg body")
qtx(slide,"F1 = 5 i-hat + 8 j-hat + 7 k-hat and F2 = 3 i-hat - 4 j-hat - 3 k-hat act on 4 kg. Find the acceleration.",XL,Y0,CW,1.05,16)
slide=answer_slide()
fml(slide,"Net force = 8 i-hat + 4 j-hat + 4 k-hat",XL,Y0,CW,LH,18)
fml(slide,"a = net force / 4",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"a = 2 i-hat + 1 j-hat + 1 k-hat.",XL,Y0+2*LH+2*NG,CW,0.52)

# PART B — unique Vector 04 only
slide=new_slide("Copy","Angle conditions to lock",TEAL)
grid_table(slide,["Condition","Angle"],[
    ["|A| = |B| = |A + B|","120°"],
    ["|A| = |B| = |A - B|","60°"],
    ["|A + B| = |A - B|","90°"],
    ["|A + B| = |A| + |B|","0°"],
    ["|A - B| = |A| + |B|","180°"],
    ["|A + B| = |A| - |B|","180°"],
    ["|A - B| = |A| - |B|","0°"],
],XL,Y0,[7.2,4.8], row_h=0.55, hsize=16, csize=16)

# COMPONENTS / UNIT VECTOR

slide=question_slide("A = 6 i-hat and B = 8 j-hat")
qtx(slide,"Find A + B, A - B, |A| + |B|, and |A + B|.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"A + B = 6 i-hat + 8 j-hat",XL,Y0,CW,LH,18)
fml(slide,"A - B = 6 i-hat - 8 j-hat",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"|A| + |B| = 6 + 8 = 14, and |A + B| = sqrt(36 + 64) = 10",XL,Y0+2*LH+2*NG,CW,LH,15)
result_card(slide,"Sum of lengths is 14. Length of the sum is 10.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Direction of A + B from A")
qtx(slide,"A = 6 i-hat and B = 8 j-hat. Find the direction of A + B from A.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"A + B = 6 i-hat + 8 j-hat",XL,Y0,CW,LH,18)
fml(slide,"tan alpha = 8 / 6 = 4 / 3",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"alpha = tan inverse of 4/3 from A.",XL,Y0+2*LH+2*NG,CW,0.52)
ox,oy=axis_diagram(slide,8.4,3.20,3.4,2.8)
vec(slide,ox,oy,ox+1.50,oy,"A",GOLD)
vec(slide,ox,oy,ox,oy-1.70,"B",GREEN)
vec(slide,ox,oy,ox+1.50,oy-1.70,"R",WORK)
tx(slide,"α",ox+0.45,oy-0.30,0.35,0.25,14,ORANGE,True,spa=0,scaled=False)

slide=new_slide("Copy","Unit vector list",TEAL)
body(slide,"î itself is already a unit vector.",XL,Y0,CW,0.45)
body(slide,"Unit vector of î + ĵ is (î + ĵ) / √2.",XL,Y0+0.60,CW,0.50)
body(slide,"Unit vector of î + ĵ - k̂ is (î + ĵ - k̂) / √3.",XL,Y0+1.25,CW,0.55)
body(slide,"sin θ î - cos θ ĵ already has magnitude 1.",XL,Y0+1.90,CW,0.50)
body(slide,"Unit vector of 3î + 4ĵ is (3î + 4ĵ) / 5.",XL,Y0+2.55,CW,0.50)
ncert(slide,"Â = A / |A|",XL,Y0+3.30,CW,0.80,22)

slide=question_slide("Find beta for a unit vector")
qtx(slide,"A = 0.4 i-hat + 0.3 j-hat + beta k-hat is a unit vector. Find beta.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"For a unit vector, |A| = 1",XL,Y0,CW,LH,18)
fml(slide,"sqrt(0.16 + 0.09 + beta^2) = 1",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"0.25 + beta^2 = 1, so beta^2 = 0.75, beta = ±√3 / 2",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"beta = ± (√3)/2.",XL,Y0+3*LH+3*NG,CW,0.50)

# DOT PRODUCT

slide=question_slide("Find alpha if A is perpendicular to B")
qtx(slide,"A = 2 i-hat + 3 j-hat + alpha k-hat and B = i-hat - 2 j-hat + 4 k-hat. If A is perpendicular to B, find alpha.",XL,Y0,CW,1.10,16)
slide=answer_slide()
fml(slide,"A · B = 0",XL,Y0,CW,LH,18)
fml(slide,"2 - 6 + 4 alpha = 0",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"4 alpha = 4, so alpha = 1",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"alpha = 1.",XL,Y0+3*LH+3*NG,CW,0.48)

slide=question_slide("Angle of a vector from the x-axis")
qtx(slide,"A = 5 i-hat + 5 sqrt(3) j-hat. Find the angle of A from the x-axis.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"tan alpha = Aᵧ / Aₓ = 5√3 / 5 = √3",XL,Y0,CW,LH,18)
fml(slide,"alpha = 60°",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"60° from the +x axis.",XL,Y0+2*LH+2*NG,CW,0.50)
ox,oy=axis_diagram(slide,8.4,3.20,3.4,2.8)
vec(slide,ox,oy,ox+1.20,oy-2.05,"A",GOLD)
tx(slide,"60°",ox+0.45,oy-0.35,0.55,0.25,14,ORANGE,True,spa=0,scaled=False)

# JEE PRACTICE

slide=question_slide("Find the smaller vector")
qtx(slide,"Resultant is perpendicular to the smaller vector and half the larger vector. Find the smaller magnitude in terms of the larger B.",XL,Y0,CW,1.15,16)
slide=answer_slide()
fml(slide,"Right triangle: A² + (B/2)² = B²",XL,Y0,CW,LH,18)
fml(slide,"A² = B² - B²/4 = 3B²/4",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"A = (√3 / 2) B.",XL,Y0+2*LH+2*NG,CW,0.50)
ox,oy=axis_diagram(slide,8.5,3.20,3.3,2.7)
vec(slide,ox,oy,ox+1.60,oy,"A",GOLD)
vec(slide,ox,oy,ox,oy-1.20,"R",WORK)
vec(slide,ox+1.60,oy,ox,oy-1.20,"B",GREEN)

slide=question_slide("Three equal vectors OP, OQ, OR")
qtx(slide,"Three vectors each of magnitude A act as OP along +x, OQ along +y, and OR 45° below +x. The resultant is A sqrt(x). Find x.",XL,Y0,CW,1.15,15)
slide=answer_slide()
fml(slide,"OP + OQ = A√2 at 45° above +x",XL,Y0,CW,LH,16)
fml(slide,"OR is 45° below +x, so that intermediate resultant and OR are at 90°",XL,Y0+LH+NG,CW,LH,14)
fml(slide,"R = sqrt((A√2)² + A²) = A√3",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"x = 3.",XL,Y0+3*LH+3*NG,CW,0.48)

slide=question_slide("Vector at 30° from the y-axis")
qtx(slide,"A vector in the x-y plane makes 30° with the y-axis. Its y-component is 2 sqrt(3). Find the x-component.",XL,Y0,CW,1.05,17)
slide=answer_slide()
fml(slide,"Aᵧ = A cos 30° = 2 sqrt(3)",XL,Y0,CW,LH,18)
fml(slide,"A (√3 / 2) = 2 √3, so A = 4",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"Aₓ = A sin 30° = 4 × 1/2 = 2",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"x-component = 2.",XL,Y0+3*LH+3*NG,CW,0.48)

# CROSS PRODUCT
slide=new_slide("Added topic","Cross product — the easy picture",GREEN)
body(slide,"Dot product gave a number. Cross product gives a new vector.",XL,Y0,CW,0.50)
ncert(slide,"A × B is perpendicular to both A and B",XL,Y0+0.65,CW,1.00,22)
body(slide,"Think: A and B lie on a table. A × B stands straight up from the table.",XL,Y0+1.90,CW,0.55)
trap_card(slide,"If A and B are parallel, A × B = 0. There is no unique upright direction.",XL,Y0+2.65,CW,0.60)

slide=new_slide("Visual idea","Magnitude first: A B sin θ",TEAL)
body(slide,"Only the turning part of B counts. That part is B sin θ.",XL,Y0,CW,0.50)
ncert(slide,"|A × B| = A B sin θ",XL,Y0+0.65,CW,0.95,24)
body(slide,"This is also the area of the parallelogram made by A and B.",XL,Y0+1.85,CW,0.50)
vec(slide,1.80,5.20,5.00,5.20,"A",GOLD)
vec(slide,1.80,5.20,3.40,3.40,"B",GREEN)
vec(slide,5.00,5.20,6.60,3.40,"",MUTED)
vec(slide,3.40,3.40,6.60,3.40,"",MUTED)
angle_mark(slide,1.80,5.20,0.70,0,50,"θ")
mini_card(slide,"Memory","sin for cross, cos for dot.",7.40,3.20,4.70,1.20,ORANGE)
mini_card(slide,"Max / zero","Max at 90°. Zero at 0° or 180°.",7.40,4.70,4.70,1.20,TEAL)

slide=new_slide("Visual idea","Direction: right-hand rule",GREEN)
body(slide,"Point fingers along A. Curl them toward B through the smaller angle. Thumb shows A × B.",XL,Y0,CW,0.60)
vec(slide,2.40,4.80,5.40,4.80,"A",GOLD)
vec(slide,2.40,4.80,3.80,2.90,"B",GREEN)
arrow_line(slide,2.40,4.80,2.40,2.40,WORK,3.2)
tx(slide,"A × B",1.55,2.55,1.40,0.35,18,WORK,True,spa=0,scaled=False)
angle_mark(slide,2.40,4.80,0.75,0,55,"θ")
ncert(slide,"A × B = − (B × A)",7.00,2.40,5.20,0.90,22)
mini_card(slide,"Order matters","Swap the two arrows and the result flips.",7.00,3.70,5.20,1.15,ORANGE)
trap_card(slide,"Curl through the smaller angle between A and B.",7.00,5.15,5.20,0.55)

slide=new_slide("Copy","Dot vs cross — one glance",TEAL)
grid_table(slide,
["", "Dot product", "Cross product"],
[
["Result", "a number (scalar)", "a vector"],
["Formula size", "A B cos θ", "A B sin θ"],
["Zero when", "perpendicular (90°)", "parallel (0° or 180°)"],
["Order", "A · B = B · A", "A × B = − B × A"],
],
XL, Y0, [2.2, 4.7, 5.1], 0.55, 15, 16)
trap_card(slide,"Same pair of arrows. Cos asks 'how parallel?'. Sin asks 'how perpendicular?'.",XL,Y0+3.55,CW,0.55)

slide=new_slide("Copy","Unit-vector cycle (easy)",GREEN)
ncert(slide,"î × ĵ = k̂ ,   ĵ × k̂ = î ,   k̂ × î = ĵ",XL,Y0,CW,1.00,22)
body(slide,"Go forward in the cycle i → j → k → i: answer is positive.",XL,Y0+1.25,CW,0.50)
body(slide,"Go backward: answer gets a minus sign.",XL,Y0+1.85,CW,0.45)
ncert(slide,"î × î = 0 ,  ĵ × ĵ = 0 ,  k̂ × k̂ = 0",XL,Y0+2.50,CW,1.00,20)
trap_card(slide,"Same letter crossed with itself is always zero.",XL,Y0+3.75,CW,0.55)

slide=new_slide("How it works","Component formula without fear",TEAL)
body(slide,"Write the 3 × 3 pattern once. Read the three rows for î, ĵ, k̂.",XL,Y0,CW,0.50)
ncert(slide,"A × B = (Ay Bz − Az By) î − (Ax Bz − Az Bx) ĵ + (Ax By − Ay Bx) k̂",XL,Y0+0.65,CW,1.15,16)
body(slide,"Shortcut for xy-plane: if A = Ax î + Ay ĵ and B = Bx î + By ĵ,",XL,Y0+2.05,CW,0.55)
ncert(slide,"A × B = (Ax By − Ay Bx) k̂",XL,Y0+2.70,CW,1.00,22)
trap_card(slide,"In the plane, the cross product points along ± k̂ only.",XL,Y0+3.95,CW,0.55)

slide=question_slide("Find |A × B|")
qtx(slide,"A = 3 and B = 4. Angle between them is 30°. Find |A × B|.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"|A × B| = A B sin θ = 3 × 4 × sin 30°",XL,Y0,CW,LH,17)
fml(slide,"sin 30° = 1/2, so |A × B| = 6",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"|A × B| = 6.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("î × ĵ and ĵ × î")
qtx(slide,"Find î × ĵ and ĵ × î.",XL,Y0,CW,0.85,20)
slide=answer_slide()
fml(slide,"î × ĵ = k̂",XL,Y0,CW,LH,20)
fml(slide,"ĵ × î = − (î × ĵ) = − k̂",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"k̂ and −k̂. Order flips the sign.",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Cross product in the xy-plane")
qtx(slide,"A = 2 î + 3 ĵ and B = 4 î − ĵ. Find A × B.",XL,Y0,CW,0.95,18)
slide=answer_slide()
fml(slide,"A × B = (Ax By − Ay Bx) k̂",XL,Y0,CW,LH,18)
fml(slide,"= (2)(-1) − (3)(4) = −2 − 12 = −14",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"A × B = −14 k̂.",XL,Y0+2*LH+2*NG,CW,0.48)

slide=question_slide("When is A × B zero?")
qtx(slide,"A = 2 î + 4 ĵ and B = î + 2 ĵ. Find A × B. What does the answer mean?",XL,Y0,CW,1.05,17)
slide=answer_slide()
fml(slide,"A × B = (2×2 − 4×1) k̂ = (4 − 4) k̂ = 0",XL,Y0,CW,LH,17)
fml(slide,"B = (1/2) A, so the arrows are parallel",XL,Y0+LH+NG,CW,LH,17)
result_card(slide,"A × B = 0. Parallel vectors give zero cross product.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Unit vector along A × B")
qtx(slide,"A = î + ĵ and B = î − ĵ. Find a unit vector along A × B.",XL,Y0,CW,1.00,17)
slide=answer_slide()
fml(slide,"A × B = (1×(-1) − 1×1) k̂ = −2 k̂",XL,Y0,CW,LH,17)
fml(slide,"|A × B| = 2, so unit vector = (−2 k̂)/2 = − k̂",XL,Y0+LH+NG,CW,LH,16)
result_card(slide,"Unit vector = − k̂.",XL,Y0+2*LH+2*NG,CW,0.48)


# CLASS SUMMARY
slide=new_slide("Class summary","Vector 04 lock lines",TEAL)
fml(slide,"Add by columns: (Ax + Bx) î + (Ay + By) ĵ",XL,Y0,CW,LH,15)
fml(slide,"Delta v = v_f - v_i. Same speed can still give large Delta v.",XL,Y0+LH+NG,CW,LH,14)
fml(slide,"A · B = A B cos θ. Unit vector = vector / magnitude.",XL,Y0+2*LH+2*NG,CW,LH,14)
fml(slide,"|A × B| = A B sin θ. Right-hand rule. A × B = − B × A.",XL,Y0+3*LH+3*NG,CW,LH,14)
result_card(slide,"In the xy-plane: A × B = (Ax By − Ay Bx) k̂.",XL,Y0+4*LH+4*NG,CW,0.50)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Redo A = 4 î + 3 ĵ and B = 3 î - 4 ĵ for sum and difference.",XL,Y0,CW,0.50)
body(slide,"2. East→north at 10 m/s: write Delta v and |Delta v|.",XL,Y0+0.55,CW,0.45)
body(slide,"3. Find beta if 0.4î + 0.3ĵ + beta k̂ is a unit vector.",XL,Y0+1.10,CW,0.50)
body(slide,"4. Redo the OP, OQ, OR resultant question.",XL,Y0+1.65,CW,0.45)
body(slide,"5. Find A × B for A = 3 î + ĵ and B = 2 î + 4 ĵ.",XL,Y0+2.20,CW,0.50)
body(slide,"6. Write î × ĵ, ĵ × k̂, and ĵ × î.",XL,Y0+2.75,CW,0.45)
trap_card(slide,"Dot uses cos. Cross uses sin. Do not mix them.",XL,Y0+3.35,CW,0.50)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
