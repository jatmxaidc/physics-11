"""Vector 01 — log recap, then feel of a vector."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Vector_01_Feel_of_Vectors.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(1/2)", "½"),
        ("^(-12)", "⁻¹²"),
        ("^(-7)", "⁻⁷"),
        ("^(-4)", "⁻⁴"),
        ("^(-2)", "⁻²"),
        ("^(-n)", "⁻ⁿ"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)



# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vector 01",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Vectors",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Feel of a Vector",bx,by+1.62,bw,0.40,20,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"29 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Log and binomial warmup, then the feel of a vector."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["Log formulae and loudness questions","Decaying current","Binomial approximation when 1 >> x","Physical quantity: scalar and vector","Writing a vector as magnitude times unit vector","Angle between two vectors","Types of vectors"]
for i,item in enumerate(topics):
    yy=1.10+i*0.78
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.66,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Same order as the Vector 01 class notes."); after_content_slide()

# LOG
slide=new_slide("Copy","Log lock lines",TEAL)
ncert(slide,"log (x^n) = n log x",XL,Y0,CW,0.78,22)
ncert(slide,"log (x to the 1/m) = (1/m) log x",XL,Y0+0.92,CW,0.78,20)
ncert(slide,"log(xy) = log x + log y      log(x/y) = log x - log y",XL,Y0+1.84,CW,0.82,18)
body(slide,"log_b a = log a / log b, taken in the same base.   log 1 = 0.   b^0 = 1.",XL,Y0+2.90,CW,0.55)
trap_card(slide,"Keep the base written. Change of base is a fraction of two logs, not a product.",XL,Y0+3.60,CW,0.55)

slide=question_slide("log base 8 of 16")
qtx(slide,"Find log base 8 of 16.",XL,Y0,CW,0.70,23)
slide=answer_slide()
fml(slide,"16 = 2^4 and 8 = 2^3",XL,Y0,CW,LH,20)
fml(slide,"log8(16) = log(2^4) / log(2^3) = 4/3",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"4/3",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("log base 27 of 3")
qtx(slide,"Find log base 27 of 3.",XL,Y0,CW,0.70,23)
slide=answer_slide()
fml(slide,"27 = 3^3, so log27(3) = log(3) / log(3^3) = 1/3",XL,Y0,CW,LH,18)
result_card(slide,"1/3",XL,Y0+LH+NG,CW,0.50)

slide=question_slide("log base 100 of 1000")
qtx(slide,"Find log base 100 of 1000.",XL,Y0,CW,0.70,23)
slide=answer_slide()
fml(slide,"1000 = 10^3 and 100 = 10^2",XL,Y0,CW,LH,20)
fml(slide,"log100(1000) = 3/2",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"3/2",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("log10 of 0.0001")
qtx(slide,"Find log10(0.0001).",XL,Y0,CW,0.70,23)
slide=answer_slide()
fml(slide,"0.0001 = 10^(-4)",XL,Y0,CW,LH,20)
result_card(slide,"log10(0.0001) = -4",XL,Y0+LH+NG,CW,0.50)

slide=question_slide("log base 0.01 of 10")
qtx(slide,"Find log base 0.01 of 10.",XL,Y0,CW,0.70,23)
slide=answer_slide()
fml(slide,"0.01 = 10^(-2), so the base is 10^(-2)",XL,Y0,CW,LH,18)
fml(slide,"log of 10 in base 10^(-2) = 1/(-2) = -1/2",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"-1/2",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Loudness 50 dB")
qtx(slide,"Loudness is 50 dB.  L = 10 log10(I/I0) with I0 = 10^(-12) W/m^2. Find I.",XL,Y0,CW,1.05,20)
slide=answer_slide()
fml(slide,"50 = 10 log10(I/I0), so 5 = log10(I/I0)",XL,Y0,CW,LH,18)
fml(slide,"I/I0 = 10^5",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"I = 10^5 times 10^(-12) = 10^(-7)",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"I = 10^(-7) W/m^2",XL,Y0+3*LH+3*NG,CW,0.50)

slide=question_slide("Current falling to 10 A")
qtx(slide,"I = 20 times e to the power (-lambda t), lambda = 2.303. Find t when I = 10 A.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"10 = 20 times e to the power (-lambda t), so 1/2 equals that exponential",XL,Y0,CW,LH,16)
fml(slide,"Take natural log: ln(1/2) = -lambda t",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"t = ln 2 / lambda.   ln 2 = 2.303 log10(2)",XL,Y0+2*LH+2*NG,CW,LH,17)
fml(slide,"t = 2.303 log10(2) / 2.303 = log10(2)",XL,Y0+3*LH+3*NG,CW,LH,17)
result_card(slide,"t = 0.30 s",XL,Y0+4*LH+4*NG,CW,0.48)

slide=question_slide("Loudness when intensity becomes I/4")
qtx(slide,"L = 16 dB at intensity I. Intensity becomes I/4. Find the new loudness.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"L = 10 log10(I/I0),   L' = 10 log10((I/4)/I0)",XL,Y0,CW,LH,17)
fml(slide,"L - L' = 10 log10(4) = 10 times 0.60 = 6",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"L' = 16 - 6 = 10",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"New loudness = 10 dB",XL,Y0+3*LH+3*NG,CW,0.50)

# BINOMIAL
slide=new_slide("New topic","Binomial when 1 is much larger than x",GREEN)
ncert(slide,"(1 + x)^n = 1 + n x     if |x| << 1",XL,Y0,CW,0.95,22)
ncert(slide,"(1 - x)^n = 1 - n x",XL,Y0+1.15,CW,0.85,22)
ncert(slide,"(1 + x)^(-n) = 1 - n x",XL,Y0+2.20,CW,0.85,22)
trap_card(slide,"Drop x^2 and higher powers only after the first term is 1 and |x| is much smaller than 1.",XL,Y0+3.25,CW,0.60)

slide=question_slide("(1.007)^5")
qtx(slide,"Find (1.007)^5 using the small-x approximation.",XL,Y0,CW,0.80,21)
slide=answer_slide()
fml(slide,"1.007 = 1 + 0.007, and 0.007 << 1",XL,Y0,CW,LH,19)
fml(slide,"(1 + 0.007)^5 = 1 + 5 times 0.007 = 1 + 0.035",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"1.035",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Square root of 0.99")
qtx(slide,"Find sqrt(0.99) using (1 - x)^n.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"0.99 = 1 - 0.01, and the power is 1/2",XL,Y0,CW,LH,19)
fml(slide,"(1 - 0.01)^(1/2) = 1 - (1/2)(0.01) = 1 - 0.005",XL,Y0+LH+NG,CW,LH,17)
result_card(slide,"0.995",XL,Y0+2*LH+2*NG,CW,0.50)

slide=question_slide("Gravity at a small height")
qtx(slide,"g_h = GM/(R+h)^2. If h << R, write a simple expression for g_h.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"R + h = R(1 + h/R), so (R+h)^2 = R^2 (1 + h/R)^2",XL,Y0,CW,LH,17)
fml(slide,"g_h = (GM/R^2) times (1 + h/R)^(-2)",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"(1 + h/R)^(-2) = 1 - 2(h/R)   because h/R << 1",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"g_h = g (1 - 2h/R)   where g = GM/R^2",XL,Y0+3*LH+3*NG,CW,0.52)
trap_card(slide,"Do not write the denominator as (R + 2h)^2. That is a different expansion.",XL,Y0+4*LH+4*NG,CW,0.48)

slide=question_slide("Trap: (2.01)^3 is not 2.03")
qtx(slide,"Why is (2.01)^3 = 2.03 a wrong use of (1+x)^n = 1 + n x?",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"First write 2.01 = 2(1 + 0.005), not (1 + 0.01)",XL,Y0,CW,LH,18)
fml(slide,"The formula needs the leading 1. Here the leading number is 2.",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"(2.01)^3 = 8 (1.005)^3 = 8(1 + 0.015) = 8.12",XL,Y0+2*LH+2*NG,CW,0.55)

# VECTORS
slide=new_slide("New topic","Physical quantity",GREEN)
body(slide,"A physical quantity can be measured. Love and hate cannot, so they are not physical quantities.",XL,Y0,CW,0.70)
body(slide,"A measurable quantity is either a scalar or a vector.",XL,Y0+0.85,CW,0.50)
ncert(slide,"Scalar: magnitude only.   Vector: magnitude, direction, and triangle law.",XL,Y0+1.60,CW,1.05,18)

slide=new_slide("Copy","Scalar and vector",TEAL)
grid_table(slide,["","Scalar","Vector"],[
    ["Needs","magnitude only","magnitude and direction"],
    ["Addition","ordinary + and -","triangle law"],
    ["Sign","plus or minus is a value","plus or minus is a direction"],
    ["Examples","mass, temperature, current, speed, energy","force, velocity, displacement, weight, electric field"],
],XL,Y0,[2.4,4.7,4.9], row_h=0.70, hsize=14, csize=14)
trap_card(slide,"Current has a conventional direction, but it does not follow the triangle law, so current is a scalar.",XL,Y0+3.7,CW,0.55)

slide=new_slide("How it works","Plus and minus on a vector",TEAL)
body(slide,"On a scalar, +10 and -10 differ in value. On a vector, the sign is the direction along a chosen line.",XL,Y0,CW,0.70)
ncert(slide,"F1 = +20 N i-hat     and     F2 = -10 N i-hat",XL,Y0+0.95,CW,0.95,20)
body(slide,"Same line, opposite senses. The unit vector i-hat carries the direction. The number is the magnitude with a sign.",XL,Y0+2.15,CW,0.70)

slide=new_slide("Copy","How we write a vector",TEAL)
ncert(slide,"A arrow = A times A-hat",XL,Y0,CW,0.95,24)
body(slide,"A is the magnitude (a positive number, with unit).",XL,Y0+1.20,CW,0.50)
body(slide,"A-hat is the unit vector. It has magnitude 1 and no unit. It only points the direction.",XL,Y0+1.80,CW,0.60)
body(slide,"Force is written F arrow, velocity v arrow, electric field E arrow. Mass, energy and speed stay without an arrow.",XL,Y0+2.55,CW,0.70)

slide=new_slide("Copy","Arrow on paper",TEAL)
body(slide,"Tail to head. The head shows the direction. The length of the arrow shows the magnitude.",XL,Y0,CW,0.70)
body(slide,"Two arrows of the same direction but different lengths are not equal.",XL,Y0+0.90,CW,0.55)
body(slide,"Two arrows of the same length but different directions are not equal.",XL,Y0+1.55,CW,0.55)
body(slide,"East is +i, west is -i, north is +j, south is -j, upward is +k.",XL,Y0+2.20,CW,0.55)
trap_card(slide,"A free vector may be slid parallel to itself. Magnitude and direction must stay the same. Turning it changes the vector.",XL,Y0+3.00,CW,0.70)

slide=question_slide("Which one is a vector?")
qtx(slide,"Which of the following is a vector?   emf, surface tension, weight, focal length.",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"emf is a potential difference, a scalar",XL,Y0,CW,LH,18)
fml(slide,"surface tension = force per length, a scalar",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"focal length is a length, a scalar",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"weight = mg downward, so magnitude and direction",XL,Y0+3*LH+3*NG,CW,LH,18)
result_card(slide,"Weight is the vector.",XL,Y0+4*LH+4*NG,CW,0.48)

slide=question_slide("Which symbol is written correctly?")
qtx(slide,"Work with an arrow, force without an arrow, or energy without an arrow?",XL,Y0,CW,0.85,20)
slide=answer_slide()
fml(slide,"Work and energy are scalars, so no arrow",XL,Y0,CW,LH,18)
fml(slide,"Force is a vector, so it needs an arrow or bold type",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"Energy (E) is the correct writing among the three.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Direction does not automatically make a vector")
qtx(slide,"A: no direction means scalar.   B: a direction means vector. Which is correct?",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"A scalar really has no direction, so A is true",XL,Y0,CW,LH,18)
fml(slide,"Current has a direction but does not obey the triangle law",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"So having a direction is not enough. B is false",XL,Y0+2*LH+2*NG,CW,LH,18)
result_card(slide,"A true, B false.",XL,Y0+3*LH+3*NG,CW,0.50)

slide=new_slide("New topic","Angle between two vectors",GREEN)
ncert(slide,"Join tail with tail, or head with head. Take the smaller angle.",XL,Y0,CW,1.00,20)
body(slide,"The bigger gap and the smaller gap add to 180 degrees only when the arrows are on crossing lines. The angle between the vectors is the smaller one, and it is never more than 180 degrees.",XL,Y0+1.25,CW,0.90)
trap_card(slide,"Do not measure the angle between a tail and a head when the arrows are placed carelessly. First bring tails together.",XL,Y0+2.40,CW,0.70)

slide=question_slide("Two arrows at 60 degrees, one reversed")
qtx(slide,"Arrow A points right. Arrow B points down-right and the marked corner is 60 degrees. After reversing B onto the same tails, what is the angle between them?",XL,Y0,CW,1.15,18)
slide=answer_slide()
fml(slide,"Place both tails at the same point",XL,Y0,CW,LH,19)
fml(slide,"The supplement of 60 degrees is 180 - 60 = 120",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"Angle between the vectors = 120 degrees.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=new_slide("Copy","Negative of a vector",TEAL)
ncert(slide,"-A has the same magnitude and the opposite direction.",XL,Y0,CW,0.95,20)
body(slide,"If A = 10 i-hat, then -A = -10 i-hat.",XL,Y0+1.20,CW,0.50)
body(slide,"On the paper, flip the head of the arrow. Do not change the length.",XL,Y0+1.85,CW,0.55)

slide=question_slide("Unit vector of a 40 N force")
qtx(slide,"A force of 40 N points north. What is the unit vector of that force?",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"F arrow = 40 N times (north unit vector)",XL,Y0,CW,LH,18)
fml(slide,"Unit vector = F arrow / magnitude",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"The newton cancels. Direction does not carry a unit.",XL,Y0+2*LH+2*NG,CW,LH,17)
result_card(slide,"Unit vector = 1 north, not 1 N north.",XL,Y0+3*LH+3*NG,CW,0.52)

slide=new_slide("Copy","Types of vector — equality and lines",TEAL)
body(slide,"Equal: same magnitude, same direction, same kind of quantity.",XL,Y0,CW,0.50)
body(slide,"Parallel: same direction. Magnitudes may differ. Every equal pair is parallel. A parallel pair need not be equal.",XL,Y0+0.70,CW,0.70)
body(slide,"Negative: same magnitude, exactly opposite direction.",XL,Y0+1.55,CW,0.50)
body(slide,"Anti-parallel: opposite direction, magnitudes may differ. Every negative pair is anti-parallel. An anti-parallel pair need not be a negative pair.",XL,Y0+2.20,CW,0.75)

slide=new_slide("Copy","Types of vector — size and plane",TEAL)
body(slide,"Unit vector: magnitude 1, no unit, direction only.",XL,Y0,CW,0.50)
body(slide,"Zero vector: magnitude 0. Example: velocity of a car at rest.",XL,Y0+0.65,CW,0.50)
body(slide,"Collinear: all lie on the same straight line.",XL,Y0+1.30,CW,0.45)
body(slide,"Coplanar: all lie in the same plane.",XL,Y0+1.90,CW,0.45)
body(slide,"Orthogonal: angle between them is 90 degrees. Magnitudes may differ.",XL,Y0+2.50,CW,0.50)
body(slide,"Concurrent: arrows act at the same point, even if magnitudes and directions differ.",XL,Y0+3.15,CW,0.55)
body(slide,"Axial: direction is along the axis of rotation. Examples: torque, angular velocity, angular acceleration, angular momentum.",XL,Y0+3.80,CW,0.70)

slide=new_slide("Class summary","Vector 01 lock lines",TEAL)
fml(slide,"log change of base is a ratio of logs in one common base",XL,Y0,CW,LH,16)
fml(slide,"(1+x)^n = 1 + n x only when the first term is 1 and |x| is tiny",XL,Y0+LH+NG,CW,LH,15)
fml(slide,"Vector = magnitude and direction, and it obeys the triangle law",XL,Y0+2*LH+2*NG,CW,LH,16)
fml(slide,"A arrow = A times A-hat, and the angle is the smaller tail-to-tail angle",XL,Y0+3*LH+3*NG,CW,LH,15)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Redo the five log values without looking.",XL,Y0,CW,0.50)
body(slide,"2. Redo both loudness questions and the decaying current.",XL,Y0+0.65,CW,0.50)
body(slide,"3. Write five scalars and five vectors from memory.",XL,Y0+1.30,CW,0.50)
body(slide,"4. Sketch equal, parallel, negative and orthogonal pairs.",XL,Y0+1.95,CW,0.50)
trap_card(slide,"Do not put an arrow on work or energy. Do not call current a vector.",XL,Y0+2.70,CW,0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
