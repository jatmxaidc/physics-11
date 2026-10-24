"""Lecture 18 - area between curves, average value and revision."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_18_Area_Between_Curves_and_Average_Value.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "â¿â»Â¹"),
        ("^(n+1)", "â¿âºÂ¹"),
        ("^(3/2)", "Â³áŸÂ²"),
        ("^(1/2)", "Â¹áŸÂ²"),
        ("^(-2)", "â»Â²"),
        ("^(-1)", "â»Â¹"),
        ("^(-x)", "â»Ë£"),
        ("^n", "â¿"),
        ("^x", "Ë£"),
        ("^6", "â¶"),
        ("^5", "âµ"),
        ("^4", "â´"),
        ("^3", "Â³"),
        ("^2", "Â²"),
        ("^0", "â°"),
        ("sqrt", "âˆš"),
        ("omega", "Ï‰"),
        ("lambda", "Î»"),
        (" proportional to ", " âˆ "),
        (">=", "â‰¥"),
        ("<=", "â‰¤"),
        ("->", "â†’"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)



# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 18",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Area between Curves, Average Value and Revision",bx,by+1.62,bw,0.40,15,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"24 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Close the calculus block with area between curves and average value."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["5-minute substitution recap","Area between two curves","Which curve is above?","Average value of a function","Physics average velocity / mean value","Mixed indefinite + definite revision","Lock lines for the calculus block"]
for i,item in enumerate(topics):
    yy=1.14+i*0.76
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.66,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Keep formulas short; sketch every region."); after_content_slide()

# RECAP
slide=new_slide("5-minute recap","Substitution from last class",BLUE)
fml(slide,"Choose u, write du, rewrite the whole integral",XL,Y0,CW,LH,18)
fml(slide,"Definite: change both limits with u",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"∫_a^b f = F(b) − F(a)     (no +C)",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"Check by differentiating",XL,Y0+3*LH+3*NG,CW,LH,19)

slide=question_slide("Quick check: substitution")
qtx(slide,"Find ∫ 2x cos(x²) dx.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"u = x²  →  du = 2x dx",XL,Y0,CW,LH,20)
result_card(slide,"sin(x²) + C.",XL,Y0+LH+NG,CW,0.55)

# AREA BETWEEN CURVES
slide=new_slide("New topic","Area between two curves",GREEN)
body(slide,"When one curve lies above another from x = a to x = b, the area between them is the difference of integrals.",XL,Y0,CW,0.70)
ncert(slide,"Area = ∫_a^b [y_upper − y_lower] dx",XL,Y0+0.95,CW,1.05,22)
trap_card(slide,"Always upper minus lower. Do not subtract in the wrong order.",XL,Y0+2.30,CW,0.55)

slide=new_slide("How it works","Three board steps",TEAL)
fml(slide,"1. Sketch both curves and mark a and b",XL,Y0,CW,LH,19)
fml(slide,"2. Decide which is upper on [a, b]",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"3. Integrate (upper − lower) from a to b",XL,Y0+2*LH+2*NG,CW,LH,19)
body(slide,"If they cross, split the integral at the intersection.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=new_slide("Exact graph","Region between y = x and y = x²",TEAL)
body(slide,"On [0, 1], the line y = x is above the parabola y = x².",XL,Y0,5.4,0.60)
ncert(slide,"Area = ∫_0^1 (x − x²) dx",XL,Y0+0.85,5.4,0.90,18)
sx,sy=graph_axes(slide,7.05,1.45,5.20,4.55,-0.2,1.3,-0.2,1.3,"x","y")
axis_ticks(slide,sx,sy,[0,0.5,1],[0,0.5,1])
plot_function(slide,lambda x:x,sx,sy,0,1,-0.2,1.3,GOLD,40,2.5)
plot_function(slide,lambda x:x*x,sx,sy,0,1,-0.2,1.3,TEAL,80,2.5)
mark_point(slide,sx,sy,0,0,"0",GREEN,0.08,0.10)
mark_point(slide,sx,sy,1,1,"(1,1)",GREEN,0.08,-0.28)

slide=question_slide("Area between line and parabola")
qtx(slide,"Find the area between y = x and y = x² from x = 0 to x = 1.",XL,Y0,CW,0.90,21)
slide=answer_slide()
fml(slide,"Upper = x, lower = x²",XL,Y0,CW,LH,20)
fml(slide,"Area = ∫_0^1 (x − x²) dx = [x²/2 − x³/3]_0^1",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"= 1/2 − 1/3 = 1/6",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"Area = 1/6 square units.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Area between line and axis")
qtx(slide,"Find the area between y = 2x and the x-axis from x = 1 to x = 3.",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"This is still ∫_1^3 2x dx  (upper = 2x, lower = 0)",XL,Y0,CW,LH,18)
fml(slide,"= [x²]_1^3 = 9 − 1 = 8",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Area = 8 square units.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Area between two lines")
qtx(slide,"Find the area between y = 4 and y = x from x = 1 to x = 3.",XL,Y0,CW,0.90,21)
slide=answer_slide()
fml(slide,"On [1, 3], y = 4 is above y = x",XL,Y0,CW,LH,19)
fml(slide,"Area = ∫_1^3 (4 − x) dx = [4x − x²/2]_1^3",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"= (12 − 9/2) − (4 − 1/2) = 15/2 − 7/2 = 4",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"Area = 4 square units.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Find intersection first")
qtx(slide,"Find the area between y = x² and y = 4x − x².",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"Intersection: x² = 4x − x²  →  2x² − 4x = 0  →  2x(x − 2) = 0",XL,Y0,CW,LH,16)
fml(slide,"x = 0 or x = 2.  On [0, 2], upper = 4x − x²",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"Area = ∫_0^2 [(4x − x²) − x²] dx = ∫_0^2 (4x − 2x²) dx",XL,Y0+2*LH+2*NG,CW,LH,15)
fml(slide,"= [2x² − (2/3)x³]_0^2 = 8 − 16/3 = 8/3",XL,Y0+3*LH+3*NG,CW,LH,17)
result_card(slide,"Area = 8/3 square units.",XL,Y0+4*LH+4*NG,CW,0.50)

# AVERAGE VALUE
slide=new_slide("New topic","Average value of a function",GREEN)
body(slide,"The average height of y = f(x) on [a, b] is the area under the curve divided by the width.",XL,Y0,CW,0.65)
ncert(slide,"f_avg = (1/(b − a)) ∫_a^b f(x) dx",XL,Y0+0.90,CW,1.05,22)
body(slide,"This is the mean value of f on that interval.",XL,Y0+2.20,CW,0.55)
trap_card(slide,"Do not forget to divide by (b − a).",XL,Y0+3.00,CW,0.50)

slide=new_slide("How it works","Why this formula?",TEAL)
body(slide,"Area under f equals the area of a rectangle of width (b − a) and height f_avg.",XL,Y0,CW,0.65)
ncert(slide,"∫_a^b f(x) dx = f_avg · (b − a)",XL,Y0+0.90,CW,1.00,20)
body(slide,"Solve for f_avg to get the formula.",XL,Y0+2.20,CW,0.50)

slide=question_slide("Average of a line")
qtx(slide,"Find the average value of f(x) = 2x on [0, 4].",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"f_avg = (1/4) ∫_0^4 2x dx = (1/4)[x²]_0^4",XL,Y0,CW,LH,18)
fml(slide,"= (1/4)(16) = 4",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Average value = 4.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Average of a parabola")
qtx(slide,"Find the average value of f(x) = x² on [0, 3].",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"f_avg = (1/3) ∫_0^3 x² dx = (1/3)[x³/3]_0^3",XL,Y0,CW,LH,18)
fml(slide,"= (1/3)(9) = 3",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Average value = 3.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Average of cosine")
qtx(slide,"Find the average value of f(x) = cos x on [0, π/2].",XL,Y0,CW,0.90,20)
slide=answer_slide()
fml(slide,"Width = π/2 − 0 = π/2",XL,Y0,CW,LH,19)
fml(slide,"f_avg = (2/π) ∫_0^{π/2} cos x dx = (2/π)[sin x]_0^{π/2}",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"= (2/π)(1) = 2/π",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"Average value = 2/π.",XL,Y0+3*LH+3*NG,CW,0.55)

# PHYSICS AVERAGE
slide=new_slide("New topic","Physics: average velocity",GREEN)
body(slide,"If v(t) is given, the average velocity over [t₁, t₂] is the mean value of v.",XL,Y0,CW,0.65)
ncert(slide,"v_avg = (1/(t₂ − t₁)) ∫_{t₁}^{t₂} v(t) dt = Δx / Δt",XL,Y0+0.90,CW,1.05,18)
body(slide,"Same formula as average value — now with time as the variable.",XL,Y0+2.25,CW,0.55)

slide=question_slide("Average velocity from v(t)")
qtx(slide,"If v = 2t (m s⁻¹) from t = 0 to t = 4 s, find average velocity.",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"v_avg = (1/4) ∫_0^4 2t dt = (1/4)[t²]_0^4 = 4",XL,Y0,CW,LH,17)
fml(slide,"Check: Δx = 16 m, Δt = 4 s → v_avg = 4 m s⁻¹",XL,Y0+LH+NG,CW,LH,17)
result_card(slide,"v_avg = 4 m s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Average of constant acceleration speed-up")
qtx(slide,"If v = 3 + 2t from t = 0 to t = 3, find v_avg.",XL,Y0,CW,0.90,21)
slide=answer_slide()
fml(slide,"v_avg = (1/3) ∫_0^3 (3 + 2t) dt = (1/3)[3t + t²]_0^3",XL,Y0,CW,LH,16)
fml(slide,"= (1/3)(9 + 9) = 6",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"v_avg = 6 m s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Mean of a linear a")
qtx(slide,"If a = 4t from t = 1 to t = 3, find average acceleration.",XL,Y0,CW,0.90,21)
slide=answer_slide()
fml(slide,"a_avg = (1/2) ∫_1^3 4t dt = (1/2)[2t²]_1^3",XL,Y0,CW,LH,17)
fml(slide,"= (1/2)(18 − 2) = 8",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"a_avg = 8 m s⁻².",XL,Y0+2*LH+2*NG,CW,0.55)

# MIXED REVISION
slide=new_slide("Revision","Indefinite integral lock lines",TEAL)
fml(slide,"∫ x^n dx = x^(n+1)/(n+1) + C     (n ≠ −1)",XL,Y0,CW,LH,17)
fml(slide,"∫ sin x dx = −cos x + C     ·     ∫ cos x dx = sin x + C",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"∫ e^x dx = e^x + C     ·     ∫ dx/x = ln|x| + C",XL,Y0+2*LH+2*NG,CW,LH,16)
fml(slide,"Substitution undoes the chain rule",XL,Y0+3*LH+3*NG,CW,LH,18)

slide=question_slide("Revision: indefinite")
qtx(slide,"Find ∫ (4x³ − 1/x) dx  for x > 0.",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"∫ 4x³ dx = x⁴",XL,Y0,CW,LH,20)
fml(slide,"∫ −1/x dx = −ln x",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"x⁴ − ln x + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Revision: definite")
qtx(slide,"Evaluate ∫_0^{π/4} cos(2x) · 2 dx.",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"u = 2x  →  du = 2 dx;   limits 0 → π/2",XL,Y0,CW,LH,18)
fml(slide,"∫_0^{π/2} cos u du = 1",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"1.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Revision: area between")
qtx(slide,"Find the area between y = √x and y = x from x = 0 to x = 1.",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"On [0, 1], √x ≥ x",XL,Y0,CW,LH,20)
fml(slide,"Area = ∫_0^1 (√x − x) dx = [ (2/3)x^(3/2) − x²/2 ]_0^1",XL,Y0+LH+NG,CW,LH,15)
fml(slide,"= 2/3 − 1/2 = 1/6",XL,Y0+2*LH+2*NG,CW,LH,20)
result_card(slide,"Area = 1/6 square units.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Revision: average value")
qtx(slide,"Find the average value of f(x) = 3x² on [1, 2].",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"f_avg = (1/1) ∫_1^2 3x² dx = [x³]_1^2 = 8 − 1",XL,Y0,CW,LH,17)
result_card(slide,"Average value = 7.",XL,Y0+LH+NG,CW,0.55)

slide=question_slide("Homework: between curves")
qtx(slide,"Find the area between y = 6 − x and y = x from x = 1 to x = 2.",XL,Y0,CW,0.95,20)
slide=answer_slide()
fml(slide,"Upper = 6 − x on this interval",XL,Y0,CW,LH,19)
fml(slide,"Area = ∫_1^2 (6 − x − x) dx = ∫_1^2 (6 − 2x) dx",XL,Y0+LH+NG,CW,LH,16)
fml(slide,"= [6x − x²]_1^2 = (12 − 4) − (6 − 1) = 3",XL,Y0+2*LH+2*NG,CW,LH,16)
result_card(slide,"Area = 3 square units.",XL,Y0+3*LH+3*NG,CW,0.55)

slide=question_slide("Homework: average velocity")
qtx(slide,"If v = t² from t = 0 to t = 3, find v_avg.",XL,Y0,CW,0.85,21)
slide=answer_slide()
fml(slide,"v_avg = (1/3) ∫_0^3 t² dt = (1/3)[t³/3]_0^3 = 3",XL,Y0,CW,LH,17)
result_card(slide,"v_avg = 3 m s⁻¹.",XL,Y0+LH+NG,CW,0.55)

# SUMMARY
slide=new_slide("Class summary","Calculus block lock lines",TEAL)
fml(slide,"Area between curves: ∫ (upper − lower) dx",XL,Y0,CW,LH,18)
fml(slide,"Average value: f_avg = (1/(b−a)) ∫_a^b f dx",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"v_avg = (1/Δt) ∫ v dt = Δx/Δt",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"Differentiation ↔ Integration are reverse operations",XL,Y0+3*LH+3*NG,CW,LH,17)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Sketch every area-between-curves question before integrating.",XL,Y0,CW,0.55)
body(slide,"2. Memorise the average-value formula with (b − a) in the denominator.",XL,Y0+0.72,CW,0.55)
body(slide,"3. Redo one substitution + one definite + one average question.",XL,Y0+1.44,CW,0.55)
body(slide,"4. Write a one-page summary of the whole calculus block.",XL,Y0+2.16,CW,0.55)
trap_card(slide,"Next chapter after this block: Vectors.",XL,Y0+3.00,CW,0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
