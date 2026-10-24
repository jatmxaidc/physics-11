"""Lecture 14 — applications of differentiation and maxima-minima."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_14_Applications_of_Differentiation_and_Maxima_Minima.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)


# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 14",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Applications of Differentiation and Maxima-Minima",bx,by+1.62,bw,0.40,16,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"18 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Use only familiar Class 11 applications."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["5-minute differentiation recap","Position, velocity and acceleration","Exact velocity–time graphs","Related rates: circle, square and sphere","Stationary points","Second derivative test","Maxima and minima practice"]
for i,item in enumerate(topics):
    yy=1.14+i*0.76
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.66,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Preview the lesson."); after_content_slide()

# RECAP AND MOTION
slide=new_slide("5-minute recap","Rates of change",BLUE)
fml(slide,"position x(t) → velocity v=dx/dt",XL,Y0,CW,LH,20)
fml(slide,"velocity v(t) → acceleration a=dv/dt=d²x/dt²",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"stationary point: dy/dx=0",XL,Y0+2*LH+2*NG,CW,LH,20)
fml(slide,"related rates: differentiate both sides with respect to time",XL,Y0+3*LH+3*NG,CW,LH,18)

slide=new_slide("New topic","From position to a velocity graph",GREEN)
body(slide,"Differentiate the position function to obtain velocity.",XL,Y0,CW,0.55)
body(slide,"Then plot v against t using its exact equation.",XL,Y0+0.70,CW,0.52)
ncert(slide,"x(t) → v(t)=dx/dt → v–t graph",XL,Y0+1.50,CW,1.00,23)
body(slide,"Where v=0, the object is instantaneously at rest.",XL,Y0+2.78,CW,0.55)

slide=question_slide("Motion: increasing velocity line")
qtx(slide,"If x=t²-4t+8, find velocity and draw the exact v–t graph.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"v=dx/dt=2t-4",XL,Y0,5.25,LH,20)
fml(slide,"v(0)=-4 and v=0 at t=2",XL,Y0+LH+NG,5.25,LH,19)
sx,sy=graph_axes(slide,7.10,1.42,4.75,4.55,0,6,-5,9,"t","v")
axis_ticks(slide,sx,sy,[0,2,4,6],[-4,0,4,8])
plot_function(slide,lambda t:2*t-4,sx,sy,0,6,-5,9,GOLD,80,2.6)
mark_point(slide,sx,sy,2,0,"rest",GREEN,0.10,-0.28)

slide=question_slide("Motion: decreasing velocity line")
qtx(slide,"If x=-t²+4t+8, find velocity and draw the exact v–t graph.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"v=dx/dt=-2t+4",XL,Y0,5.25,LH,20)
fml(slide,"v(0)=4 and v=0 at t=2",XL,Y0+LH+NG,5.25,LH,19)
sx,sy=graph_axes(slide,7.10,1.42,4.75,4.55,0,6,-9,5,"t","v")
axis_ticks(slide,sx,sy,[0,2,4,6],[-8,-4,0,4])
plot_function(slide,lambda t:-2*t+4,sx,sy,0,6,-9,5,GOLD,80,2.6)
mark_point(slide,sx,sy,2,0,"rest",GREEN,0.10,-0.28)

slide=question_slide("Acceleration from velocity")
qtx(slide,"If v=t³/3-3t²+8t, find acceleration and draw the exact a–t graph.",XL,Y0,CW,0.95,22)
slide=answer_slide()
fml(slide,"a=dv/dt=t²-6t+8",XL,Y0,5.25,LH,20)
fml(slide,"a=(t-2)(t-4)",XL,Y0+LH+NG,5.25,LH,20)
sx,sy=graph_axes(slide,7.10,1.42,4.75,4.55,0,6,-2,9,"t","a")
axis_ticks(slide,sx,sy,[0,2,3,4,6],[-1,0,4,8])
plot_function(slide,lambda t:t*t-6*t+8,sx,sy,0,6,-2,9,GOLD,140,2.6)
mark_point(slide,sx,sy,2,0,"2",GREEN,0.06,-0.26); mark_point(slide,sx,sy,4,0,"4",GREEN,0.06,-0.26)

slide=question_slide("Changing x with time")
qtx(slide,"If y=x² and x changes with time, express dy/dt in terms of x and dx/dt.",XL,Y0,CW,0.95,22)
slide=answer_slide()
fml(slide,"dy/dx=2x",XL,Y0,CW,LH,21)
fml(slide,"dy/dt=(dy/dx)(dx/dt)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dt=2x(dx/dt).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Changing angle with time")
qtx(slide,"If y=sin θ and θ changes with time, express dy/dt.",XL,Y0,CW,0.85,23)
slide=answer_slide()
fml(slide,"dy/dθ=cos θ",XL,Y0,CW,LH,21)
fml(slide,"dy/dt=(dy/dθ)(dθ/dt)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dt=cos θ(dθ/dt).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Nested time dependence")
qtx(slide,"If y=sin(t²+2t), find dy/dt.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"Outer derivative: cos(t²+2t)",XL,Y0,CW,LH,20)
fml(slide,"Inner derivative: 2t+2",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dy/dt=(2t+2)cos(t²+2t).",XL,Y0+2*LH+2*NG,CW,0.55)

# RELATED RATES
slide=new_slide("New topic","Related rates",GREEN)
body(slide,"When dimensions change with time, geometrical quantities also change.",XL,Y0,CW,0.60)
body(slide,"Write the geometry formula first, then differentiate with respect to time.",XL,Y0+0.80,CW,0.62)
ncert(slide,"formula → differentiate with respect to t → substitute values with units",XL,Y0+1.70,CW,1.00,21)
trap_card(slide,"Substitute numerical values only after differentiating.",XL,Y0+2.96,CW,0.60)

slide=new_slide("How it works","Circle, square and sphere formulae",TEAL)
fml(slide,"circle: A=πr² → dA/dt=2πr(dr/dt)",XL,Y0,CW,LH,19)
fml(slide,"square: P=4l → dP/dt=4(dl/dt)",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"sphere: V=(4/3)πr³ → dV/dt=4πr²(dr/dt)",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"sphere: S=4πr² → dS/dt=8πr(dr/dt)",XL,Y0+3*LH+3*NG,CW,LH,19)

slide=question_slide("Sphere volume rate")
qtx(slide,"For V=(4/3)πr³, express dV/dt in terms of r and dr/dt.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"Differentiate with respect to t",XL,Y0,CW,LH,20)
result_card(slide,"dV/dt=4πr²(dr/dt).",XL,Y0+LH+NG,CW,0.58)

slide=question_slide("Sphere surface-area rate")
qtx(slide,"For S=4πr², express dS/dt in terms of r and dr/dt.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"d(r²)/dt=2r(dr/dt)",XL,Y0,CW,LH,20)
result_card(slide,"dS/dt=8πr(dr/dt).",XL,Y0+LH+NG,CW,0.58)

slide=question_slide("Expanding circle")
qtx(slide,"A circle's radius increases at 1/π m s⁻¹. Find dA/dt when r=4 m.",XL,Y0,CW,0.95,22)
slide=answer_slide()
fml(slide,"dA/dt=2πr(dr/dt)",XL,Y0,CW,LH,20)
fml(slide,"=2π(4)(1/π)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dA/dt=8 m² s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Growing square perimeter")
qtx(slide,"The side of a square increases at 0.2 cm s⁻¹. Find the rate of increase of its perimeter.",XL,Y0,CW,1.00,21)
slide=answer_slide()
fml(slide,"P=4l",XL,Y0,CW,LH,21)
fml(slide,"dP/dt=4(dl/dt)=4(0.2)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dP/dt=0.8 cm s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Area change with radius")
qtx(slide,"For a circle A=πr², find dA/dr.",XL,Y0,CW,0.75,24)
slide=answer_slide(); result_card(slide,"dA/dr=2πr.",XL,Y0,CW,0.58)

slide=question_slide("Inflating balloon")
qtx(slide,"A spherical balloon's volume increases at 10 cm³ s⁻¹. Find dr/dt when r=5 cm.",XL,Y0,CW,0.95,22)
slide=answer_slide()
fml(slide,"dV/dt=4πr²(dr/dt)",XL,Y0,CW,LH,20)
fml(slide,"10=4π(25)(dr/dt)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dr/dt=1/(10π) cm s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.58)

slide=question_slide("Growing square area")
qtx(slide,"A square's side increases at 2 cm s⁻¹. Find dA/dt when the side is 10 cm.",XL,Y0,CW,0.95,22)
slide=answer_slide()
fml(slide,"A=l² → dA/dt=2l(dl/dt)",XL,Y0,CW,LH,20)
fml(slide,"=2(10)(2)",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"dA/dt=40 cm² s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

# SHORT MOTION PRACTICE
slide=question_slide("When is the particle at rest?")
qtx(slide,"If v(t)=t²-3t, when is the particle at rest?",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"At rest: v=0",XL,Y0,CW,LH,20)
fml(slide,"t(t-3)=0",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"t=0 s or 3 s.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Acceleration from displacement")
qtx(slide,"If s(t)=3t²-4t+2, find acceleration at t=1 s.",XL,Y0,CW,0.88,22)
slide=answer_slide()
fml(slide,"v=ds/dt=6t-4",XL,Y0,CW,LH,20)
fml(slide,"a=dv/dt=6",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Acceleration=6 m s⁻².",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Velocity from displacement")
qtx(slide,"If x(t)=4t³+3t²-6t+1, find velocity at t=2 s.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"v=dx/dt=12t²+6t-6",XL,Y0,CW,LH,20)
fml(slide,"v(2)=48+12-6",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Velocity=54 m s⁻¹.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Exponential position")
qtx(slide,"If x(t)=e²ᵗ+5, find velocity.",XL,Y0,CW,0.80,23)
slide=answer_slide(); result_card(slide,"v=dx/dt=2e²ᵗ.",XL,Y0,CW,0.58)

# MAXIMA AND MINIMA
slide=new_slide("New topic","Maxima and minima",GREEN)
body(slide,"At a smooth local maximum or minimum, the tangent is horizontal.",XL,Y0,CW,0.60)
ncert(slide,"Step 1: solve dy/dx=0 to find stationary points",XL,Y0+0.84,CW,0.95,21)
body(slide,"A stationary point may be a maximum, minimum or neither. Test it using the second derivative.",XL,Y0+2.04,CW,0.72)

slide=new_slide("How it works","Second derivative test",TEAL)
ncert(slide,"d²y/dx²>0 at a stationary point → local minimum",XL,Y0,CW,1.00,21)
ncert(slide,"d²y/dx²<0 at a stationary point → local maximum",XL,Y0+1.26,CW,1.00,21)
trap_card(slide,"If d²y/dx²=0, this test is inconclusive; another test is needed.",XL,Y0+2.56,CW,0.66)

slide=new_slide("Exact graph","Maximum and minimum shapes",TEAL)
sx1,sy1=graph_axes(slide,0.95,1.55,5.25,4.35,-3,3,-1,6)
axis_ticks(slide,sx1,sy1,[-2,0,2],[0,2,4])
plot_function(slide,lambda x:5-x*x,sx1,sy1,-2.4,2.4,-1,6,GOLD,150,2.6)
mark_point(slide,sx1,sy1,0,5,"maximum",RED,0.10,-0.30)
sx2,sy2=graph_axes(slide,7.05,1.55,5.25,4.35,-3,3,-1,6)
axis_ticks(slide,sx2,sy2,[-2,0,2],[0,2,4])
plot_function(slide,lambda x:x*x+1,sx2,sy2,-2.2,2.2,-1,6,TEAL,150,2.6)
mark_point(slide,sx2,sy2,0,1,"minimum",GREEN,0.10,-0.30)

slide=question_slide("Minimum of a quadratic")
qtx(slide,"Find the maxima and minima of y=x²+5.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"dy/dx=2x=0 → x=0",XL,Y0,CW,LH,20)
fml(slide,"d²y/dx²=2>0 → minimum",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Minimum value y=5 at x=0; no maximum.",XL,Y0+2*LH+2*NG,CW,0.58)

slide=question_slide("Maximum of a quadratic")
qtx(slide,"Find the maxima and minima of y=4x-x².",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"dy/dx=4-2x=0 → x=2",XL,Y0,CW,LH,20)
fml(slide,"d²y/dx²=-2<0 → maximum",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Maximum value y=4 at x=2; no minimum.",XL,Y0+2*LH+2*NG,CW,0.58)

slide=question_slide("Two stationary points")
qtx(slide,"For y=x³-3x²+4, find its local maximum and minimum values.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"dy/dx=3x(x-2)=0 → x=0,2",XL,Y0,CW,LH,19)
fml(slide,"d²y/dx²=6x-6",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"x=0 gives maximum; x=2 gives minimum",XL,Y0+2*LH+2*NG,CW,LH,19)
result_card(slide,"Local maximum=4; local minimum=0.",XL,Y0+3*LH+3*NG,CW,0.58)

slide=question_slide("Maximum and minimum velocity")
qtx(slide,"If v=t³-6t²+12, find its local maximum and minimum values.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"dv/dt=3t(t-4)=0 → t=0,4",XL,Y0,CW,LH,19)
fml(slide,"d²v/dt²=6t-12",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Local maximum v=12 at t=0; local minimum v=-20 at t=4.",XL,Y0+2*LH+2*NG,CW,0.62)

slide=question_slide("Minimum by completing the test")
qtx(slide,"Find the maxima and minima of y=x²-4x+5.",XL,Y0,CW,0.80,23)
slide=answer_slide()
fml(slide,"dy/dx=2x-4=0 → x=2",XL,Y0,CW,LH,20)
fml(slide,"d²y/dx²=2>0",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Minimum value y=1 at x=2; no maximum.",XL,Y0+2*LH+2*NG,CW,0.58)

slide=question_slide("Cubic maxima and minima")
qtx(slide,"For y=x³-12x²+36x, find the local maximum and minimum values.",XL,Y0,CW,0.90,22)
slide=answer_slide()
fml(slide,"dy/dx=3(x-2)(x-6)=0 → x=2,6",XL,Y0,CW,LH,19)
fml(slide,"d²y/dx²=6x-24",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Local maximum=32 at x=2; local minimum=0 at x=6.",XL,Y0+2*LH+2*NG,CW,0.62)

slide=new_slide("Class summary","Applications of differentiation",TEAL)
fml(slide,"v=dx/dt,     a=dv/dt=d²x/dt²",XL,Y0,CW,LH,20)
fml(slide,"related rates: differentiate the geometry formula with respect to t",XL,Y0+LH+NG,CW,LH,17)
fml(slide,"stationary point: dy/dx=0",XL,Y0+2*LH+2*NG,CW,LH,20)
fml(slide,"d²y/dx²>0 minimum; d²y/dx²<0 maximum",XL,Y0+3*LH+3*NG,CW,LH,19)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Redraw all three motion graphs using their exact equations.",XL,Y0,CW,0.58)
body(slide,"2. Memorise the related-rate formulae for a circle, square and sphere.",XL,Y0+0.78,CW,0.62)
body(slide,"3. Redo all maxima–minima questions without seeing the answers.",XL,Y0+1.58,CW,0.58)
body(slide,"4. Include units in every rate answer.",XL,Y0+2.34,CW,0.52)
trap_card(slide,"Never call dy/dx=0 alone proof of a maximum or minimum.",XL,Y0+3.18,CW,0.62)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)


