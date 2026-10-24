"""Units 02 — Principle of homogeneity and dimensional applications (PW 03)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Units_02_Principle_of_Homogeneity.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³ᐟ²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units 02",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Units and Measurement",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Principle of Homogeneity and Applications",bx,by+1.62,bw,0.40,18,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"7 October 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"From PW Units and Measurement Lecture 03."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=[
"Warm-up theory: why n1 u1 = n2 u2",
"Unitless vs dimensionless (full theory)",
"Principle of homogeneity (paragraphs)",
"Finding dimensions of constants",
"Dimensionless functions (θ, sin, e^x, log)",
"Derive a formula + limitations",
"Board practice + CBSE write lines",
]
for i,t in enumerate(topics):
    y=1.25+i*0.58
    circ=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(XL),Inches(y),Inches(0.42),Inches(0.42))
    circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.fill.background()
    tx(slide,f"{i+1:02d}",XL,y+0.05,0.42,0.32,14,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
    tx(slide,t,XL+0.60,y+0.05,10.5,0.35,20,TEXT,False,spa=0,scaled=False)
_footer(slide,2); after_content_slide()

slide=new_slide("Class note","Where we left Units 01",TEAL)
para(slide,"In Units 01 we learned that a physical quantity is something we can measure. Every measurement is written as a number times a unit. The size of the quantity itself does not change when we change the unit — only the number changes.",XL,Y0,CW,1.35,16)
para(slide,"We also wrote dimensions using M, L and T. Today we use those dimensions as a tool: to check equations, to find unknown constants, and to derive the form of a formula.",XL,Y0+1.50,CW,1.20,16)
ncert(slide,"Today: check equations  ·  find unknown dimensions  ·  derive formulas",XL,Y0+2.90,CW,1.00,15)
trap_card(slide,"Class notes will be saved after teaching. This deck is for the board.",XL,Y0+4.15,CW,0.48)

slide=new_slide("Theory","What dimensional analysis is",GREEN)
para(slide,"Dimensional analysis means studying a physical equation by looking only at the dimensions of each quantity. We do not put numbers first. We ask: do the powers of mass, length and time match on both sides?",XL,Y0,CW,1.40,16)
para(slide,"If the dimensions do not match, the equation cannot be correct. If they do match, the equation is dimensionally correct. That does not always mean it is the true physical law — a wrong constant or a missing factor can still hide inside.",XL,Y0+1.55,CW,1.45,16)
trap_card(slide,"Dimensions are a filter. They catch many mistakes, but not every mistake.",XL,Y0+3.25,CW,0.55)

# WARM-UP THEORY
slide=new_slide("Theory","Why conversion works",TEAL)
para(slide,"The amount of a physical quantity is fixed. When we measure the same amount in a different unit, the number and the unit change together so that their product stays the same. That is why we write n1 u1 = n2 u2.",XL,Y0,CW,1.40,16)
para(slide,"If the new unit is larger, the number becomes smaller. If the new unit is smaller, the number becomes larger. Density, force, torque — all follow this same idea once you know their dimensions.",XL,Y0+1.55,CW,1.35,16)
ncert(slide,"n1 u1 = n2 u2",XL,Y0+3.15,CW,0.95,22)

slide=new_slide("Theory","Two conversion habits (continued)",GREEN)
para(slide,"Habit A — SI value given, new system asked. Keep the SI number. Replace each SI unit (kg, m, s, …) by how many new-system units fit into it. Then simplify. You are rewriting the same amount in new clothing.",XL,Y0,CW,1.35,16)
para(slide,"Habit B — find how large one new unit is in SI. Build one new unit from the new mass, length and time using the dimensions of the quantity. Multiply those SI sizes. The result is the SI value of one new unit.",XL,Y0+1.50,CW,1.40,16)
bullet(slide,"Both habits use: the real amount does not change.",XL,Y0+3.10,CW,0.45,16)
bullet(slide,"Always write dimensions first for derived quantities (density, torque, …).",XL,Y0+3.55,CW,0.45,16)

slide=new_slide("Warm-up","Two conversion habits (board)",TEAL)
mini_card(slide,"SI given → new asked","Write SI value. Replace each SI unit by the new unit.",XL,Y0,5.80,1.50,TEAL)
mini_card(slide,"New unit value in SI","Write 1 new unit using M, L, T. Convert that to SI.",XL+6.20,Y0,5.80,1.50,ORANGE)
body(slide,"Both use the same idea: the actual amount does not change.",XL,Y0+1.80,CW,0.50)
ncert(slide,"n1 u1 = n2 u2",XL,Y0+2.50,CW,1.00,22)

slide=question_slide("Density in a new system")
qtx(slide,"Density is 4 g/cm^3. New units: length = 10 cm, mass = 100 g. New density value?",XL,Y0,CW,1.10,16)
choices(slide,["0.04","0.4","40","400"], XL, Y0+1.30, 10.5)
slide=answer_slide()
body(slide,"Idea: rewrite 4 g/cm^3 using the new mass and new length.",XL,Y0,CW,0.40)
fml(slide,"New mass is bigger: 1 new mass = 100 g  ⇒  1 g = 1/100 new mass",XL,Y0+0.45,CW,LH,14)
fml(slide,"New length is bigger: 1 new length = 10 cm  ⇒  1 cm = 1/10 new length",XL,Y0+0.45+LH+NG,CW,LH,14)
fml(slide,"Density = mass / length^3 = 4 × (1/100) ÷ (1/10)^3",XL,Y0+0.45+2*LH+2*NG,CW,LH,14)
fml(slide,"= 4 × (1/100) × 1000 = 40",XL,Y0+0.45+3*LH+3*NG,CW,LH,16)
result_card(slide,"Final answer: 40",XL,Y0+0.45+4*LH+4*NG,CW,0.42)

slide=question_slide("New unit of torque")
qtx(slide,"New system: length = 10 m, time = 2 s, mass = 5 kg. Find 1 new unit of torque in SI.",XL,Y0,CW,1.15,15)
choices(slide,["125 N m","0.125 N m","8 N m","8 × 10^(-3) N m"], XL, Y0+1.30, 10.5)
slide=answer_slide()
body(slide,"Idea: one new torque unit = (new mass)(new length)^2 / (new time)^2.",XL,Y0,CW,0.42)
fml(slide,"Torque has the same dimensions as energy: M L^2 T^(-2)",XL,Y0+0.48,CW,LH,15)
fml(slide,"Put the new values: 5 kg × (10 m)^2 / (2 s)^2",XL,Y0+0.48+LH+NG,CW,LH,15)
fml(slide,"= 5 × 100 / 4 = 125",XL,Y0+0.48+2*LH+2*NG,CW,LH,18)
result_card(slide,"Final answer: 125 N m",XL,Y0+0.48+3*LH+3*NG,CW,0.45)

# UNITLESS VS DIMENSIONLESS
slide=new_slide("Theory","Unitless and dimensionless — two floors",TEAL)
para(slide,"Students often mix two words: unitless and dimensionless. They sound alike, but they live on different floors of the same building. Floor 1 is the unit — the name we write after the number. Floor 2 is the dimension — the powers of M, L and T hidden inside the quantity.",XL,Y0,CW,1.55,16)
para(slide,"Unitless means we do not write a unit after the number. Dimensionless means every power is zero: M^0 L^0 T^0. These two ideas can travel together, but they are not forced to stay together.",XL,Y0+1.70,CW,1.40,16)
trap_card(slide,"Learn both words separately. Do not treat them as synonyms.",XL,Y0+3.35,CW,0.55)

slide=new_slide("Theory","Unitless and dimensionless (continued)",GREEN)
para(slide,"If a quantity is unitless, it must also be dimensionless. There is nowhere for a hidden M, L or T to hide if there is no unit at all. Strain and refractive index are common examples: they are pure numbers.",XL,Y0,CW,1.40,16)
para(slide,"The reverse is not always true. A quantity can be dimensionless and still have a unit. Plane angle is the best school example. Its dimension is none, but its SI unit is radian. Degree is another unit of the same dimensionless quantity.",XL,Y0+1.55,CW,1.45,16)
ncert(slide,"Angle has a unit (radian) but no dimension.",XL,Y0+3.20,CW,0.95,17)

slide=new_slide("Theory","What follows from what",TEAL)
para(slide,"Has dimension ⇒ must have a unit. If mass, length or time appear in the formula, we need a unit to measure that quantity. You cannot measure force without newton (or dyne, or some other force unit).",XL,Y0,CW,1.35,16)
para(slide,"Has a unit ⇒ may or may not have a dimension. Angle has a unit. Force has a unit. Only force has a dimension in the M-L-T sense. So writing a unit does not automatically create a dimension.",XL,Y0+1.50,CW,1.40,16)
bullet(slide,"Unitless ⇒ dimensionless (always).",XL,Y0+3.10,CW,0.42,16)
bullet(slide,"Dimension ⇒ unit (always).",XL,Y0+3.55,CW,0.42,16)
bullet(slide,"Unit ⇒ dimension? Not always. Angle is the counter-example.",XL,Y0+4.00,CW,0.42,16)

slide=new_slide("CBSE write","True / false lines (copy)",GREEN)
body(slide,"1. A unitless quantity may have dimension.  → False",XL,Y0,CW,0.48)
body(slide,"2. A dimensionless quantity may be unitless.  → True  (μ, strain)",XL,Y0+0.55,CW,0.50)
body(slide,"3. A quantity that has a unit must have a dimension.  → False  (angle)",XL,Y0+1.10,CW,0.55)
body(slide,"4. A quantity that has a dimension must have a unit.  → True",XL,Y0+1.65,CW,0.50)
ncert(slide,"Has dimension ⇒ must have unit.   Unitless ⇒ must be dimensionless.",XL,Y0+2.35,CW,1.10,16)
trap_card(slide,"Best CBSE example: plane angle. Unit = radian. Dimension = none.",XL,Y0+3.70,CW,0.50)

slide=new_slide("Copy","Decision table (easy)",TEAL)
grid_table(slide,
["If you know…", "Then…"],
[
["no unit (unitless)", "must be dimensionless"],
["has a unit", "may or may not have dimension"],
["no dimension", "may or may not have a unit"],
["has dimension", "must have a unit"],
],
XL, Y0, [5.5, 6.5], 0.55, 15, 16)
trap_card(slide,"Floor 1 = unit. Floor 2 = dimension. Unit can exist without dimension.",XL,Y0+2.90,CW,0.55)

# PRINCIPLE OF HOMOGENEITY
slide=new_slide("Theory","Principle of homogeneity — meaning",GREEN)
para(slide,"The principle of homogeneity says that in any correct physical equation, every term that is added, subtracted or compared must have the same dimensions. Left side and right side must also match. Nature does not add metres to kilograms.",XL,Y0,CW,1.50,16)
para(slide,"Think of dimensions as the \"type\" of a quantity. You may multiply or divide different types and get a new type. You may not add different types. That is the whole rule in one sentence.",XL,Y0+1.65,CW,1.30,16)
ncert(slide,"In any correct equation, every term must have the same dimensions.",XL,Y0+3.15,CW,0.95,16)

slide=new_slide("Theory","Principle of homogeneity (continued)",TEAL)
para(slide,"Why does this rule exist? Because an equation is a statement about equality of physical amounts. Two sides can be equal only if they are the same kind of thing. Force equals mass times acceleration — both sides are M L T^(-2). Force cannot equal pressure, because pressure is force per area.",XL,Y0,CW,1.55,16)
para(slide,"Multiplication and division are allowed for unlike dimensions because they create a new quantity. Speed is length divided by time. Density is mass divided by volume. The new quantity simply carries the combined dimensions.",XL,Y0+1.70,CW,1.40,16)
trap_card(slide,"4 kg + 8 m is meaningless. 5 m / 2 s = 2.5 m/s is fine.",XL,Y0+3.35,CW,0.55)

slide=new_slide("Theory","What homogeneity allows and forbids",GREEN)
para(slide,"Allowed: add or subtract quantities with the same dimensions. Compare them with >, < or =. Multiply or divide any dimensions. Raise a quantity to a power (the dimensions take that power).",XL,Y0,CW,1.25,16)
para(slide,"Forbidden: add force to pressure, add energy to torque as if they were different physical ideas that somehow mix, or write an equation where one side is time and the other is length. Same dimensions do not prove same quantity — work and torque share M L^2 T^(-2) but they are not the same.",XL,Y0+1.40,CW,1.55,16)
bullet(slide,"Same dimensions ⇒ may add / subtract / compare.",XL,Y0+3.15,CW,0.42,16)
bullet(slide,"Same dimensions ≠ same physical quantity.",XL,Y0+3.60,CW,0.42,16)

slide=new_slide("CBSE write","Principle of homogeneity (3 marks)",TEAL)
para(slide,"Definition: The dimensions of each term on both sides of a physical equation must be the same. This is called the principle of homogeneity of dimensions.",XL,Y0,CW,1.15,16)
body(slide,"Use 1: Check if a formula is dimensionally correct.",XL,Y0+1.30,CW,0.42)
body(slide,"Use 2: Find dimensions of unknown constants in an equation.",XL,Y0+1.75,CW,0.45)
body(slide,"Use 3: Derive the form of a formula (pendulum, later chapters).",XL,Y0+2.25,CW,0.45)
ncert(slide,"Same dimensions on every + or − term. That is the whole idea.",XL,Y0+2.90,CW,1.00,15)
trap_card(slide,"Dimensionally correct does not always mean physically correct.",XL,Y0+4.15,CW,0.48)

slide=new_slide("Theory","Board shortcut for checking an equation",GREEN)
para(slide,"When you only want to check dimensions, you may mentally replace +, −, > and < by =. You are not changing the real maths. You are only saying: every piece that sits in those places must carry the same dimensions.",XL,Y0,CW,1.40,16)
para(slide,"Example: v = u + a t. For dimensions, write [v] = [u] = [a t]. Check each piece. Velocity, initial velocity, and acceleration times time must all be L T^(-1).",XL,Y0+1.55,CW,1.35,16)
ncert(slide,"For dimension check only: treat +, −, >, < like =",XL,Y0+3.15,CW,0.95,16)

slide=new_slide("Copy","What you can and cannot do",TEAL)
mini_card(slide,"Allowed","Add / subtract same dimensions. Multiply / divide any.",XL,Y0,5.80,1.60,GREEN)
mini_card(slide,"Not allowed","Add force to pressure. Compare 4 kg with 2 m.",XL+6.20,Y0,5.80,1.60,ORANGE)
body(slide,"4 kg + 8 kg = 12 kg.   3 kg / 5 m = 0.6 kg/m (linear density).",XL,Y0+1.95,CW,0.55)
trap_card(slide,"Different names can share dimensions. Same dimensions do not mean same quantity.",XL,Y0+2.75,CW,0.55)

slide=question_slide("Assertion–Reason")
qtx(slide,"Assertion: Force is not added with pressure.",XL,Y0,CW,0.90,17)
qtx(slide,"Reason: Quantities with different dimensions are not added or subtracted.",XL,Y0+0.70,CW,1.00,16)
choices(slide,[
"Both true, Reason explains Assertion",
"Both true, Reason does not explain",
"Assertion true, Reason false",
"Both false",
], XL, Y0+1.90, 10.5)
slide=answer_slide()
body(slide,"Check Assertion first, then Reason, then the link.",XL,Y0,CW,0.40)
fml(slide,"Force dimension = M L T^(-2). Pressure = F/A = M L^(-1) T^(-2).",XL,Y0+0.45,CW,LH,14)
fml(slide,"These are different, so force and pressure cannot be added. Assertion true.",XL,Y0+0.45+LH+NG,CW,LH,14)
fml(slide,"Reason says the same rule. So Reason is true and it explains Assertion.",XL,Y0+0.45+2*LH+2*NG,CW,LH,14)
result_card(slide,"Final answer: Option 1 — both true, Reason explains Assertion",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

# FINDING CONSTANTS
slide=new_slide("Theory","Finding dimensions of unknown constants",GREEN)
para(slide,"Often an equation contains letters a, b, α, β whose dimensions are not given. Homogeneity still works. Wherever two things are added or subtracted, their dimensions match. Wherever something equals the left side, it matches the left side.",XL,Y0,CW,1.50,16)
para(slide,"So the equation itself becomes the teacher. You do not invent a new formula. You read the structure: who is added to whom, who is multiplied by whom, and who sits alone on one side.",XL,Y0+1.65,CW,1.30,16)
ncert(slide,"If A + B = C, then [A] = [B] = [C]",XL,Y0+3.15,CW,0.95,18)

slide=new_slide("Theory","Finding constants (method continued)",TEAL)
para(slide,"Step by step method. First mark every + and −. Those pieces must have equal dimensions. Next equate each such piece with the left-hand side if the whole sum equals the left side. Finally solve for the unknown using known dimensions of force, pressure, velocity, and so on.",XL,Y0,CW,1.50,16)
para(slide,"If a quantity is divided or multiplied, rearrange. Example: if F = x / (b t^2), then b = x / (F t^2). Write dimensions of the right side and simplify. Pure numbers like 2 or π do not change dimensions.",XL,Y0+1.65,CW,1.40,16)
trap_card(slide,"Do not invent formulas. Use only what the equation itself tells you.",XL,Y0+3.30,CW,0.55)

slide=new_slide("How it works","Find dimensions — board steps",GREEN)
body(slide,"Step 1: Mark every + or −. Those pieces must match.",XL,Y0,CW,0.45)
body(slide,"Step 2: Equate each piece with the left side (or with its neighbour).",XL,Y0+0.55,CW,0.50)
body(slide,"Step 3: Solve for the unknown using known dimensions.",XL,Y0+1.10,CW,0.50)
ncert(slide,"If A + B = C, then [A] = [B] = [C]",XL,Y0+1.80,CW,1.05,20)
trap_card(slide,"Do not invent formulas. Use only what the equation itself tells you.",XL,Y0+3.10,CW,0.55)

slide=question_slide("Electric field constants")
qtx(slide,"E = (A / x^2) i-hat + (B / y^3) j-hat. SI units of A and B are",XL,Y0,CW,1.05,16)
choices(slide,[
"N m^3 C^(-1), N m^2 C^(-1)",
"N m^2 C^(-1), N m^3 C^(-1)",
"N m^3 C, N m^2 C",
"N m^2 C, N m^3 C",
], XL, Y0+1.25, 10.5)
slide=answer_slide()
body(slide,"Idea: each term on the right must have the same unit as E.",XL,Y0,CW,0.40)
fml(slide,"Electric field E = force / charge, so unit of E = N/C",XL,Y0+0.45,CW,LH,15)
fml(slide,"A / x^2 = E  ⇒  A = E × x^2 = (N/C) × m^2",XL,Y0+0.45+LH+NG,CW,LH,15)
fml(slide,"B / y^3 = E  ⇒  B = E × y^3 = (N/C) × m^3",XL,Y0+0.45+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: A → N m^2 C^(-1),   B → N m^3 C^(-1)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

slide=question_slide("Pressure-style equation")
qtx(slide,"(P + b / V^2) = a. P is pressure, V is volume. Find [a] and [b].",XL,Y0,CW,1.05,16)
slide=answer_slide()
body(slide,"Idea: things added must match. Right side must match left side.",XL,Y0,CW,0.42)
fml(slide,"a sits alone on the right, so [a] = [P] = M L^(-1) T^(-2)",XL,Y0+0.48,CW,LH,15)
fml(slide,"P and b/V^2 are added, so [b/V^2] = [P]",XL,Y0+0.48+LH+NG,CW,LH,15)
fml(slide,"[b] = [P] × [V]^2 = M L^(-1) T^(-2) × L^6 = M L^5 T^(-2)",XL,Y0+0.48+2*LH+2*NG,CW,LH,14)
result_card(slide,"Final answer: [a] = M L^(-1) T^(-2),  [b] = M L^5 T^(-2)",XL,Y0+0.48+3*LH+3*NG,CW,0.45)

slide=question_slide("v = αt + βx − γ x t")
qtx(slide,"If velocity v = α t + β x − γ x t, find dimensions of α, β and γ.",XL,Y0,CW,1.05,16)
slide=answer_slide()
body(slide,"Idea: treat each term like it equals velocity. [v] = L T^(-1).",XL,Y0,CW,0.42)
fml(slide,"α × time = velocity  ⇒  [α] = velocity / time = L T^(-2)",XL,Y0+0.48,CW,LH,14)
fml(slide,"β × length = velocity  ⇒  [β] = velocity / length = T^(-1)",XL,Y0+0.48+LH+NG,CW,LH,14)
fml(slide,"γ × length × time = velocity  ⇒  [γ] = T^(-2)",XL,Y0+0.48+2*LH+2*NG,CW,LH,14)
result_card(slide,"Final answer: [α] = L T^(-2),  [β] = T^(-1),  [γ] = T^(-2)",XL,Y0+0.48+3*LH+3*NG,CW,0.45)

slide=question_slide("Force formula F = (a − x) / (b t^2)")
qtx(slide,"Equation of force: F = (a − x)/(b t^2). Find dimensions of a and b.",XL,Y0,CW,1.05,16)
slide=answer_slide()
body(slide,"Idea: first use the subtraction, then rearrange for b.",XL,Y0,CW,0.40)
fml(slide,"a − x means [a] = [x]. Distance x is length, so [a] = L",XL,Y0+0.45,CW,LH,15)
fml(slide,"From F = x / (b t^2), we get b = x / (F t^2)",XL,Y0+0.45+LH+NG,CW,LH,15)
fml(slide,"[b] = L / (M L T^(-2) × T^2) = 1/M = M^(-1)",XL,Y0+0.45+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: [a] = L,  [b] = M^(-1)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

# DIMENSIONLESS FUNCTIONS
slide=new_slide("Theory","Dimensionless functions — why they matter",TEAL)
para(slide,"Trigonometric functions, exponential and logarithm appear often in waves, SHM and decay. For dimensions we treat them in a special way. Their values are pure numbers, so they do not add M, L or T of their own.",XL,Y0,CW,1.40,16)
para(slide,"There is a strict rule: the thing sitting inside sin, cos, tan, e^( ), or log( ) must itself be dimensionless. You cannot take sin of a length or log of a force. The argument must reduce to M^0 L^0 T^0.",XL,Y0+1.55,CW,1.45,16)
trap_card(slide,"Inside sin, cos, tan, exp, log — that argument must be dimensionless.",XL,Y0+3.25,CW,0.55)

slide=new_slide("Theory","Dimensionless functions (continued)",GREEN)
para(slide,"Angle θ is dimensionless because θ = arc length / radius = L/L. So [θ] = 1. Then sin θ, cos θ and tan θ are ratios of sides of a triangle — again length over length — so they are dimensionless too.",XL,Y0,CW,1.40,16)
para(slide,"When you only check dimensions, write θ, sin θ, cos θ, e^x and log x as the number 1. They do not change the dimensions of whatever they multiply. If P = α sin β, then [α] = [P] and [β] = 1.",XL,Y0+1.55,CW,1.40,16)
ncert(slide,"For dimensions: write θ, sin θ, e^x, log x as 1",XL,Y0+3.20,CW,0.95,17)

slide=new_slide("CBSE write","Why angle is dimensionless (2 marks)",GREEN)
para(slide,"Plane angle is defined as the ratio of arc length to radius. Both arc and radius are lengths, so their ratio has no dimension: θ = L/L = 1. Therefore [θ] = M^0 L^0 T^0.",XL,Y0,CW,1.25,16)
body(slide,"Still, angle has a unit: radian (or degree).",XL,Y0+1.40,CW,0.45)
ncert(slide,"Dimensionless  does not mean  no unit",XL,Y0+2.05,CW,1.00,18)
trap_card(slide,"Same idea: strain = ΔL / L. Dimensionless. Usually unitless too.",XL,Y0+3.30,CW,0.55)

slide=question_slide("P = α sin β")
qtx(slide,"If momentum P = α sin β, find dimensions of α and β.",XL,Y0,CW,0.95,17)
slide=answer_slide()
body(slide,"Idea: sin(anything) is just a pure number for dimensions.",XL,Y0,CW,0.42)
fml(slide,"Inside sin must be dimensionless. So [β] = M^0 L^0 T^0 = 1",XL,Y0+0.48,CW,LH,15)
fml(slide,"sin β = 1 (dimensionally). So P = α × 1  ⇒  [α] = [P]",XL,Y0+0.48+LH+NG,CW,LH,15)
fml(slide,"Momentum P = m v, so [α] = M L T^(-1)",XL,Y0+0.48+2*LH+2*NG,CW,LH,16)
result_card(slide,"Final answer: [α] = M L T^(-1),  [β] = dimensionless",XL,Y0+0.48+3*LH+3*NG,CW,0.45)

slide=question_slide("Wave equation: find [c]")
qtx(slide,"Y = A sin[(2π/λ)(c t − x)]. Find dimension of c.",XL,Y0,CW,1.00,17)
slide=answer_slide()
body(slide,"Idea: look only at (c t − x). Those two must match.",XL,Y0,CW,0.42)
fml(slide,"c t and x are subtracted, so [c t] = [x]",XL,Y0+0.48,CW,LH,16)
fml(slide,"[c] × time = length  ⇒  [c] = length / time = L T^(-1)",XL,Y0+0.48+LH+NG,CW,LH,15)
result_card(slide,"Final answer: [c] = L T^(-1)  (like speed)",XL,Y0+0.48+2*LH+2*NG,CW,0.48)

# APPLICATION: DERIVE FORMULA
slide=new_slide("Theory","Using dimensions to derive a formula",GREEN)
para(slide,"Sometimes we know which quantities a formula depends on, but we do not know the powers. Dimensional analysis can find those powers. We assume a product form, write dimensions on both sides, and match the powers of M, L and T one by one.",XL,Y0,CW,1.50,16)
para(slide,"Example idea: time period T of a simple pendulum depends on length l and acceleration due to gravity g. We write T = k l^a g^b, where k is a pure number. Then we find a and b. The method cannot give the value of k — experiment or deeper theory gives k = 2π.",XL,Y0+1.65,CW,1.45,16)
ncert(slide,"T ∝ l^a g^b   →   find a and b by comparing powers",XL,Y0+3.30,CW,0.95,15)

slide=new_slide("Theory","Derive formula — what we assume",TEAL)
para(slide,"We assume the unknown relation is a product of powers: quantity = constant × (first quantity)^a × (second)^b × …. This works when each variable appears once as a power. It fails if the true law has a sum of two different terms, or a trigonometric factor that we forgot to treat as dimensionless.",XL,Y0,CW,1.55,16)
para(slide,"Also, if a quantity depends on more variables than we can balance with M, L, T alone, dimensions may not give a unique answer. Then we need more physics. Keep that honesty for board answers under \"limitations\".",XL,Y0+1.70,CW,1.35,16)
trap_card(slide,"k is a pure number here. Dimensions ignore it.",XL,Y0+3.30,CW,0.55)

slide=new_slide("How it works","Pendulum step by step",GREEN)
body(slide,"Write T = k l^a g^b. Now match dimensions of both sides.",XL,Y0,CW,0.42)
fml(slide,"Left side [T] = T^1. Right side = L^a × (L T^(-2))^b",XL,Y0+0.48,CW,LH,15)
fml(slide,"Right side becomes L^(a+b) T^(-2b)",XL,Y0+0.48+LH+NG,CW,LH,16)
fml(slide,"Match time power: 1 = −2b  ⇒  b = −1/2",XL,Y0+0.48+2*LH+2*NG,CW,LH,15)
fml(slide,"Match length power: 0 = a + b  ⇒  a = +1/2",XL,Y0+0.48+3*LH+3*NG,CW,LH,15)
result_card(slide,"So T = k √(l/g). Real formula uses k = 2π.",XL,Y0+0.48+4*LH+4*NG,CW,0.42)

slide=new_slide("Theory","Uses of dimensional analysis",TEAL)
para(slide,"Dimensional analysis has three standard school uses. First, check whether an equation is dimensionally correct. Second, convert a numerical value from one system of units to another using n1 u1 = n2 u2. Third, find how physical quantities are related — the powers in a formula — when we know which quantities matter.",XL,Y0,CW,1.55,16)
para(slide,"For a 3-mark CBSE answer, write the definition of homogeneity briefly, list these three uses, and add one limitation. That structure scores cleanly without long English.",XL,Y0+1.70,CW,1.25,16)
ncert(slide,"Check  ·  Convert  ·  Derive",XL,Y0+3.15,CW,0.95,20)

slide=new_slide("Theory","Limitations of dimensional analysis",GREEN)
para(slide,"Dimensional analysis cannot find the value of a dimensionless constant such as 2π, 1/2 or e. Those numbers have no M, L or T, so dimensions never see them. The real formula must come from experiment or from a fuller derivation.",XL,Y0,CW,1.40,16)
para(slide,"It also cannot decide between two formulas that already have the same dimensions. And if a quantity depends on more independent variables than the number of base dimensions we use, we may not get unique powers. Write these limits honestly in board answers.",XL,Y0+1.55,CW,1.45,16)
bullet(slide,"Cannot give pure numbers like 2π.",XL,Y0+3.20,CW,0.42,16)
bullet(slide,"Cannot prove a formula is physically true — only dimensionally possible.",XL,Y0+3.65,CW,0.42,16)
bullet(slide,"Fails when too many variables or when the law is not a simple product.",XL,Y0+4.10,CW,0.42,16)

slide=new_slide("CBSE write","Uses of dimensional analysis (3 marks)",TEAL)
body(slide,"1. To check whether an equation is dimensionally correct.",XL,Y0,CW,0.50)
body(slide,"2. To convert a value from one system of units to another.",XL,Y0+0.55,CW,0.50)
body(slide,"3. To find the relation between physical quantities (derive formula).",XL,Y0+1.10,CW,0.55)
ncert(slide,"Check  ·  Convert  ·  Derive",XL,Y0+1.85,CW,1.00,20)
trap_card(slide,"Limitation: cannot find dimensionless constants like 2π or 1/2.",XL,Y0+3.10,CW,0.55)

slide=question_slide("NEET: time period from α and β")
qtx(slide,"Max acceleration of oscillating particle is α, max velocity is β. Time period is",XL,Y0,CW,1.10,15)
choices(slide,[
"2π α/β",
"2π α β",
"2π β/α",
"2π α^2 / β^2",
], XL, Y0+1.25, 10.5)
slide=answer_slide()
body(slide,"Easy way: check which option has dimension of time.",XL,Y0,CW,0.40)
fml(slide,"In SHM, max speed β = A ω, max acceleration α = A ω^2",XL,Y0+0.45,CW,LH,14)
fml(slide,"Divide: α/β = ω. Time period is T = 2π / ω",XL,Y0+0.45+LH+NG,CW,LH,15)
fml(slide,"So T = 2π × (β/α)",XL,Y0+0.45+2*LH+2*NG,CW,LH,18)
result_card(slide,"Final answer: 2π β/α   (option 3)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

slide=question_slide("Which does NOT have dimension of time?")
qtx(slide,"L, C, R are inductance, capacitance, resistance. Which has no dimension of time?",XL,Y0,CW,1.10,15)
choices(slide,["RC","L/R","√(LC)","L/C"], XL, Y0+1.25, 10.5)
slide=answer_slide()
body(slide,"Remember three time formulas from circuits.",XL,Y0,CW,0.40)
fml(slide,"RC has dimension of time. L/R has dimension of time.",XL,Y0+0.45,CW,LH,15)
fml(slide,"√(LC) also has dimension of time.",XL,Y0+0.45+LH+NG,CW,LH,16)
fml(slide,"L/C does not become time. That is the odd one out.",XL,Y0+0.45+2*LH+2*NG,CW,LH,15)
result_card(slide,"Final answer: L/C   (option 4)",XL,Y0+0.45+3*LH+3*NG,CW,0.45)

# SUMMARY
slide=new_slide("Class summary","Units 02 — theory in one place",TEAL)
para(slide,"Homogeneity: every term that is added, subtracted or compared must have the same dimensions. Unitless always means dimensionless; having a dimension always means having a unit; having a unit does not always mean having a dimension (angle).",XL,Y0,CW,1.40,16)
para(slide,"For dimensions, treat θ, sin θ, e^x and log x as 1. Uses of dimensional analysis: check a formula, convert units, derive a relation. Limitation: cannot find pure numbers like 2π, and cannot prove a formula is physically true.",XL,Y0+1.55,CW,1.45,16)
result_card(slide,"Lock these four ideas before board practice.",XL,Y0+3.25,CW,0.48)

slide=new_slide("Class summary","Units 02 lock lines",GREEN)
fml(slide,"Homogeneity: every + / − / compare term has the same dimensions.",XL,Y0,CW,LH,15)
fml(slide,"Unitless ⇒ dimensionless. Has dimension ⇒ must have unit.",XL,Y0+LH+NG,CW,LH,15)
fml(slide,"θ, sin θ, e^x, log x → treat as 1 for dimensions.",XL,Y0+2*LH+2*NG,CW,LH,15)
fml(slide,"Uses: check formula, convert units, derive relation.",XL,Y0+3*LH+3*NG,CW,LH,15)
result_card(slide,"Limitation: cannot get pure numbers like 2π.",XL,Y0+4*LH+4*NG,CW,0.50)

slide=new_slide("CBSE write","Ready theory lines",TEAL)
body(slide,"1. Principle of homogeneity: dimensions of all terms in an equation are equal.",XL,Y0,CW,0.50)
body(slide,"2. Only quantities of the same dimensions can be added or subtracted.",XL,Y0+0.55,CW,0.50)
body(slide,"3. Angle is dimensionless but has unit radian.",XL,Y0+1.10,CW,0.45)
body(slide,"4. Arguments of sin, cos, tan, log and exp must be dimensionless.",XL,Y0+1.65,CW,0.50)
body(slide,"5. Dimensional analysis cannot give the value of a dimensionless constant.",XL,Y0+2.20,CW,0.55)
trap_card(slide,"For 3 marks: definition + one use + one limitation.",XL,Y0+2.95,CW,0.50)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. If F = a / t + b x^2, find [a] and [b].",XL,Y0,CW,0.45)
body(slide,"2. If y = A cos(ωt − kx), show [ω]/[k] = L T^(-1).",XL,Y0+0.55,CW,0.50)
body(slide,"3. Write three uses and two limitations of dimensional analysis.",XL,Y0+1.10,CW,0.50)
body(slide,"4. Derive T ∝ √(l/g) for a simple pendulum.",XL,Y0+1.65,CW,0.45)
body(slide,"5. True/False: A quantity with a unit must have a dimension.",XL,Y0+2.20,CW,0.50)
body(slide,"6. If E = α / β, and E is energy, β is time, find [α].",XL,Y0+2.75,CW,0.45)
trap_card(slide,"Next: significant figures and errors (PW later lectures).",XL,Y0+3.40,CW,0.50)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
