"""Motion 03 — new problems only: cube, ant, carrom, rolling, x(t), speed (PW 03)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_Motion_03_Advanced_and_Speed.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90

_SUP_MAP = str.maketrans({
    "0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
    "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
    "+": "⁺", "-": "⁻", "=": "⁼", "(": "⁽", ")": "⁾",
    "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ",
    "f": "ᶠ", "g": "ᵍ", "h": "ʰ", "i": "ⁱ", "j": "ʲ",
    "k": "ᵏ", "l": "ˡ", "m": "ᵐ", "n": "ⁿ", "o": "ᵒ",
    "p": "ᵖ", "r": "ʳ", "s": "ˢ", "t": "ᵗ", "u": "ᵘ",
    "v": "ᵛ", "w": "ʷ", "x": "ˣ", "y": "ʸ", "z": "ᶻ",
    "A": "ᴬ", "B": "ᴮ", "D": "ᴰ", "E": "ᴱ", "G": "ᴳ",
    "H": "ᴴ", "I": "ᴵ", "J": "ᴶ", "K": "ᴷ", "L": "ᴸ",
    "M": "ᴹ", "N": "ᴺ", "O": "ᴼ", "P": "ᴾ", "R": "ᴿ",
    "T": "ᵀ", "U": "ᵁ", "V": "ⱽ", "W": "ᵂ",
    "/": "ᐟ",
})


def _to_sup(s: str) -> str:
    return s.translate(_SUP_MAP)


_GREEK_SUP = {"α": "ᵅ", "β": "ᵝ", "γ": "ᵞ", "δ": "ᵟ", "ε": "ᵋ", "θ": "ᶿ"}


_SUB_CHARS = {
    "0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄",
    "5": "₅", "6": "₆", "7": "₇", "8": "₈", "9": "₉",
    "a": "ₐ", "e": "ₑ", "h": "ₕ", "i": "ᵢ", "j": "ⱼ",
    "k": "ₖ", "l": "ₗ", "m": "ₘ", "n": "ₙ", "o": "ₒ",
    "p": "ₚ", "r": "ᵣ", "s": "ₛ", "t": "ₜ", "u": "ᵤ",
    "v": "ᵥ", "x": "ₓ", "y": "ᵧ",
    # no reliable Unicode subscript f — use PPT native subscript run instead
}


def _format_powers(text: str) -> str:
    """Turn caret powers into Unicode superscripts: 10^(-3) → 10⁻³, L^2 → L²."""
    text = text.replace("^(-1/2)", "⁻½").replace("^(-3/2)", "⁻³⁄²")
    text = text.replace("^(-5/2)", "⁻⁵⁄²").replace("^(-7/2)", "⁻⁷⁄²")
    text = text.replace("^(1/2)", "½").replace("^(3/2)", "³⁄²")
    text = text.replace("^(5/2)", "⁵⁄²").replace("^(7/2)", "⁷⁄²")
    text = re.sub(r"\^\(([^)]+)\)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^(-?\d+)", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([A-Za-z])", lambda m: _to_sup(m.group(1)), text)
    text = re.sub(r"\^([αβγδεθ])", lambda m: _GREEK_SUP[m.group(1)], text)
    return text


def _format_subs_unicode(text: str) -> str:
    """Convert _i, _{fi}, _x into Unicode subscripts; unknown letters → PPT subscript markers."""
    def brace(m):
        out = []
        for ch in m.group(1):
            if ch in _SUB_CHARS:
                out.append(_SUB_CHARS[ch])
            elif ch.isalnum():
                out.append(f"\0{ch}\0")
            else:
                out.append(ch)
        return "".join(out)

    text = re.sub(r"_\{([^}]+)\}", brace, text)
    text = re.sub(
        r"_([A-Za-z0-9])",
        lambda m: _SUB_CHARS[m.group(1)] if m.group(1) in _SUB_CHARS else f"\0{m.group(1)}\0",
        text,
    )
    return text


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    text = _format_powers(text)
    for old, new in (
        ("<<", "≪"),
        (">>", "≫"),
        ("sqrt(2500)", "√2500"),
        ("sqrt(200)", "√200"),
        ("sqrt(125)", "√125"),
        ("sqrt(100 + 100)", "√(100 + 100)"),
        ("sqrt(50)", "√50"),
        ("sqrt(48)", "√48"),
        ("sqrt(37)", "√37"),
        ("sqrt(25)", "√25"),
        ("sqrt(13)", "√13"),
        ("sqrt(8)", "√8"),
        ("sqrt(5)", "√5"),
        ("sqrt(3)", "√3"),
        ("sqrt(2)", "√2"),
        ("sqrt", "√"),
        ("cos 90", "cos 90°"),
        ("cos 60", "cos 60°"),
        ("cos 53", "cos 53°"),
        ("cos 45", "cos 45°"),
        ("cos 37", "cos 37°"),
        ("sin(60)", "sin 60°"),
        ("cos(60)", "cos 60°"),
        ("sin 60", "sin 60°"),
        ("sin 53", "sin 53°"),
        ("sin 45", "sin 45°"),
        ("sin 37", "sin 37°"),
        ("i-hat", "î"),
        ("j-hat", "ĵ"),
        ("k-hat", "k̂"),
        ("A-hat", "Â"),
        (" dot ", " · "),
        ("· product", "dot product"),
        (" times ", " × "),
        ("*", "×"),
        (" degrees", "°"),
        (" degree", "°"),
        ("°°", "°"),
        ("theta", "θ"),
        ("alpha", "α"),
        ("beta", "β"),
        ("cos inverse", "cos⁻¹"),
        ("tan inverse", "tan⁻¹"),
        ("F1", "F₁"),
        ("F2", "F₂"),
        ("Fx", "Fₓ"),
        ("Fy", "Fᵧ"),
        ("Fz", "Fz"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    # Subscripts: prefer Unicode; letters without a Unicode form use a PPT subscript run.
    text = _format_subs_unicode(text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    pt_size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))

    def _add_run(chunk: str, as_sub: bool = False):
        if not chunk:
            return
        run = p.add_run()
        run.text = chunk
        run.font.name = font
        run.font.size = pt_size
        run.font.bold = bold
        run.font.color.rgb = color
        if as_sub:
            # python-pptx Font.subscript is unreliable — set OOXML baseline directly
            rPr = run._r.get_or_add_rPr()
            for old in rPr.findall(qn("a:baseline")):
                rPr.remove(old)
            bl = OxmlElement("a:baseline")
            bl.set("val", "-25000")
            rPr.append(bl)

    # \0X\0 markers = letters that need PPT native subscript (e.g. f)
    parts = re.split(r"\0([A-Za-z0-9])\0", text)
    if len(parts) == 1:
        _add_run(text)
    else:
        for i, part in enumerate(parts):
            _add_run(part, as_sub=(i % 2 == 1))
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def para(slide, text, x, y, w, h=0.95, size=16):
    """Long theory paragraph — wraps; use taller h and spill to next slide when full."""
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, lsp=1.28, scaled=False)


def bullet(slide, text, x, y, w, h=0.42, size=16):
    return tx(slide, "•  " + text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    step = 0.52
    for i, item in enumerate(items):
        cy = y + i * step
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.06), Inches(0.32), Inches(0.32))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.08, 0.32, 0.28, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.46, cy + 0.04, w - 0.46, 0.42, 17, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def arrow_line(slide, x1, y1, x2, y2, color=GOLD, pt=2.4):
    c = draw_line(slide, x1, y1, x2, y2, color, pt)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    c.line._ln.append(head)
    return c


def vec(slide, x1, y1, x2, y2, label, color=GOLD, pt=2.8):
    arrow_line(slide, x1, y1, x2, y2, color, pt)
    tx(slide, label, (x1 + x2) / 2 + 0.08, (y1 + y2) / 2 - 0.28, 1.2, 0.30,
       16, color, True, spa=0, scaled=False)


def axis_diagram(slide, x=8.0, y=1.55, w=4.0, h=3.6, coord=True):
    ox, oy = x + w * 0.40, y + h * 0.58
    arrow_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    arrow_line(slide, ox, y + h, ox, y, MUTED, 1.5)
    if coord:
        tx(slide, "+x", x + w - 0.05, oy + 0.06, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
        tx(slide, "+y", ox + 0.08, y, 0.45, 0.24, 12, MUTED, True, spa=0, scaled=False)
    return ox, oy


def small_dot(slide, x, y, color=WHITE):
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.045), Inches(y - 0.045), Inches(0.09), Inches(0.09))
    d.fill.solid()
    d.fill.fore_color.rgb = color
    d.line.fill.background()
    return d


def angle_mark(slide, cx, cy, r, a1, a2, label=None, color=ORANGE):
    pts = []
    steps = 18
    for i in range(steps + 1):
        a = math.radians(a1 + (a2 - a1) * i / steps)
        pts.append((cx + r * math.cos(a), cy - r * math.sin(a)))
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        draw_line(slide, x1, y1, x2, y2, color, 1.4)
    if label:
        mid = math.radians((a1 + a2) / 2)
        tx(slide, label, cx + (r + 0.08) * math.cos(mid), cy - (r + 0.08) * math.sin(mid),
           0.75, 0.25, 13, color, True, spa=0, scaled=False)


def mini_card(slide, title, detail, x, y, w=3.6, h=1.15, color=TEAL):
    _card_bg(slide, x, y, w, h, RGBColor(12, 14, 18), color, 1.0, True)
    tx(slide, title, x + 0.15, y + 0.12, w - 0.3, 0.30, 15, color, True, spa=0, scaled=False)
    tx(slide, detail, x + 0.15, y + 0.48, w - 0.3, h - 0.58, 12, TEXT, False, spa=0, scaled=False)


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)









# ===== Figure helpers (Motion 01) =====

def stick(slide, x, y, color=WHITE, h=0.55, label=None, lab_color=None):
    """Simple standing person: head + body."""
    r = 0.09
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - h), Inches(2 * r), Inches(2 * r))
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    draw_line(slide, x, y - h + 2 * r, x, y, color, 2.0)
    draw_line(slide, x, y - h * 0.45, x - 0.16, y - h * 0.15, color, 1.6)
    draw_line(slide, x, y - h * 0.45, x + 0.16, y - h * 0.15, color, 1.6)
    draw_line(slide, x, y, x - 0.12, y + 0.22, color, 1.6)
    draw_line(slide, x, y, x + 0.12, y + 0.22, color, 1.6)
    if label:
        tx(slide, label, x - 0.35, y + 0.26, 0.70, 0.28, 13, lab_color or color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)


def ground_line(slide, x1, x2, y, color=MUTED, label="ground"):
    draw_line(slide, x1, y, x2, y, color, 2.0)
    for i in range(8):
        gx = x1 + 0.25 + i * ((x2 - x1 - 0.4) / 7)
        draw_line(slide, gx, y, gx - 0.12, y + 0.14, color, 1.2)
    tx(slide, label, (x1 + x2) / 2 - 0.5, y + 0.16, 1.0, 0.26, 12, color, True,
       PP_ALIGN.CENTER, spa=0, scaled=False)


def rounded_box(slide, x, y, w, h, fill, border, pt=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(pt)
    return s


def number_line(slide, x, y, w, xmin, xmax, ticks=None):
    """Horizontal number line; returns function pos -> screen x."""
    draw_line(slide, x, y, x + w, y, MUTED, 2.0)
    arrow_line(slide, x + w - 0.15, y, x + w + 0.05, y, MUTED, 1.6)

    def sx(v):
        return x + (v - xmin) * w / (xmax - xmin)

    ticks = ticks if ticks is not None else list(range(xmin, xmax + 1))
    for v in ticks:
        px = sx(v)
        draw_line(slide, px, y - 0.08, px, y + 0.08, MUTED, 1.2)
        tx(slide, str(v), px - 0.18, y + 0.10, 0.36, 0.22, 11, MUTED, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    return sx


def mark_point(slide, px, y, label, color=GOLD, above=True):
    small_dot(slide, px, y, color)
    dy = -0.32 if above else 0.18
    tx(slide, label, px - 0.35, y + dy, 0.70, 0.28, 13, color, True,
       PP_ALIGN.CENTER, spa=0, scaled=False)


def path_polyline(slide, points, color=GOLD, pt=2.6):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        arrow_line(slide, x1, y1, x2, y2, color, pt)


# ===== COVER =====
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx_3d(slide, "EXCELLENCE", XL, 0.52, CW, 1.10, 62, GOLD, RGBColor(210, 210, 220))
tx_3d(slide, "CLASS 11   -   2026-27", XL, 1.64, CW, 0.58, 30, WHITE, RGBColor(90, 90, 100), font="Aptos")
_solid(slide, 4.15, 2.28, 5.00, 0.028, WHITE)
bx, by, bw = 2.05, 2.70, 9.20
_card_bg(slide, bx, by, bw, 2.35, WHITE, WHITE, 0)
_card_bg(slide, bx, by + 0.72, bw, 1.63, BLACK, RGBColor(220, 220, 220), 1.25)
_solid(slide, bx + bw / 2, by + 0.16, 0.018, 0.42, RGBColor(40, 40, 40))
tx(slide, "Physics", bx, by + 0.14, bw / 2 - 0.10, 0.50, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Motion 03", bx + bw / 2, by + 0.14, bw / 2, 0.50, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Motion in a Straight Line", bx, by + 1.02, bw, 0.55, 28, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False, font="Georgia")
tx(slide, "New problems · Cube & paths · Speed and average velocity", bx, by + 1.62, bw, 0.40, 14, GOLD, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "By - Excellence Programme", XL, 5.28, CW, 0.50, 25, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "22 October 2026", XL, 5.80, CW, 0.32, 18, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, 1)
add_notes(slide, "PW Motion 03. Only ideas not already taught in Motion 01 and 02.")
after_content_slide()

slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx(slide, "What is new today", XL, 0.35, CW, 0.64, 32, GOLD, True, spa=0, scaled=False)
_solid(slide, XL, 1.00, 3.40, 0.035, GOLD)
topics = [
    "Cube: fly through air vs crawl on faces",
    "Carrom path by the mirror method",
    "Drunkard steps and a rolling wheel",
    "x(t): catch the turning point inside the time",
    "Speed, average speed, average velocity",
    "Two average-speed recipes and their uses",
]
for i, t in enumerate(topics):
    y = 1.20 + i * 0.72
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(XL), Inches(y), Inches(0.42), Inches(0.42))
    circ.fill.solid()
    circ.fill.fore_color.rgb = TEAL
    circ.line.fill.background()
    tx(slide, f"{i + 1:02d}", XL, y + 0.05, 0.42, 0.32, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
    tx(slide, t, XL + 0.60, y + 0.05, 10.5, 0.35, 19, TEXT, False, spa=0, scaled=False)
_footer(slide, 2)
after_content_slide()

slide = new_slide("Class note", "What we will not redo", TEAL)
para(slide, "Motion 01 and 02 already fixed frames, distance versus displacement, polygons, circular arcs, and how to read an x–t graph. Those pages stay closed today.", XL, Y0, CW, 1.25, 16)
para(slide, "Today starts where those lectures stop: shortest path on a surface, a reflected carrom shot, a point on a rolling wheel, motion given by x(t), and the first speed–velocity formulas.", XL, Y0 + 1.45, CW, 1.30, 16)

# ===== CUBE =====
slide = new_slide("Figure", "Butterfly through a cube", ORANGE)
para(slide, "A butterfly can fly. It starts at one corner of a cubical room of side ℓ and reaches the opposite corner. The flight path is unknown, so distance cannot be found. Displacement is the body diagonal.", XL, Y0, 7.0, 1.45, 15)
# simple cube
a = 2.2
ox, oy = 9.3, 4.8
# front face
front = [(ox, oy), (ox + a, oy), (ox + a, oy - a), (ox, oy - a)]
for p, q in zip(front, front[1:] + [front[0]]):
    draw_line(slide, p[0], p[1], q[0], q[1], WHITE, 1.6)
# back offset
dx, dy = 0.7, -0.55
back = [(p[0] + dx, p[1] + dy) for p in front]
for p, q in zip(back, back[1:] + [back[0]]):
    draw_line(slide, p[0], p[1], q[0], q[1], MUTED, 1.3)
for p, b in zip(front, back):
    draw_line(slide, p[0], p[1], b[0], b[1], MUTED, 1.2)
arrow_line(slide, front[0][0], front[0][1], back[2][0], back[2][1], GREEN, 2.6)
tx(slide, "(0,0,0)", front[0][0] - 1.1, front[0][1] - 0.05, 1.1, 0.25, 11, GREEN, True, spa=0, scaled=False)
tx(slide, "(ℓ,ℓ,ℓ)", back[2][0] + 0.05, back[2][1] - 0.35, 1.2, 0.25, 11, GOLD, True, spa=0, scaled=False)
fml(slide, "|Δr| = √(ℓ^2 + ℓ^2 + ℓ^2) = ℓ√3", XL, Y0 + 1.70, 7.0, LH, 16)
result_card(slide, "Displacement = ℓ√3    ·    distance unknown", XL, Y0 + 2.55, 7.0, 0.50, 14)

slide = new_slide("Figure", "Ant on the faces — shortest crawl", GREEN)
para(slide, "An ant cannot fly. It must stay on the walls. Displacement is still the body diagonal ℓ√3. The shortest distance is found by opening two faces into one flat rectangle and drawing a straight line.", XL, Y0, CW, 1.25, 15)
# net:  l wide, 2l tall
x0, y0, s = 2.2, 5.6, 1.15
draw_line(slide, x0, y0, x0 + s, y0, WHITE, 1.8)
draw_line(slide, x0 + s, y0, x0 + s, y0 - 2 * s, WHITE, 1.8)
draw_line(slide, x0 + s, y0 - 2 * s, x0, y0 - 2 * s, WHITE, 1.8)
draw_line(slide, x0, y0 - 2 * s, x0, y0, WHITE, 1.8)
draw_line(slide, x0, y0 - s, x0 + s, y0 - s, MUTED, 1.2)
arrow_line(slide, x0, y0, x0 + s, y0 - 2 * s, GOLD, 2.6)
tx(slide, "start", x0 - 0.7, y0 - 0.1, 0.7, 0.25, 12, GREEN, True, spa=0, scaled=False)
tx(slide, "end", x0 + s + 0.08, y0 - 2 * s - 0.05, 0.6, 0.25, 12, GOLD, True, spa=0, scaled=False)
tx(slide, "ℓ", x0 + 0.35, y0 + 0.08, 0.4, 0.22, 12, MUTED, True, spa=0, scaled=False)
tx(slide, "2ℓ", x0 - 0.45, y0 - s, 0.45, 0.22, 12, MUTED, True, spa=0, scaled=False)
fml(slide, "min distance = √((2ℓ)^2 + ℓ^2) = ℓ√5", 6.6, 2.6, 5.8, LH, 16)
fml(slide, "displacement stays ℓ√3", 6.6, 2.6 + LH + NG, 5.8, LH, 16)
trap_card(slide, "3ℓ along three edges is a possible crawl, not the shortest.", 6.6, 4.3, 5.8, 0.60)

slide = new_slide("Figure", "Carrom — mirror the board", TEAL)
para(slide, "A 4 ft × 4 ft board. The queen starts at the centre, hits the front edge, and falls into the hole behind the striking line. Reflect the board across that edge. The straight line on the double board is the real path length.", XL, Y0, CW, 1.20, 15)
# square 4 and reflected
x0, y0, u = 2.0, 5.5, 0.55
# lower square (real) and upper (mirror)
draw_line(slide, x0, y0, x0 + 4 * u, y0, ORANGE, 2.0)
draw_line(slide, x0, y0, x0, y0 - 4 * u, WHITE, 1.6)
draw_line(slide, x0 + 4 * u, y0, x0 + 4 * u, y0 - 4 * u, WHITE, 1.6)
draw_line(slide, x0, y0 - 4 * u, x0 + 4 * u, y0 - 4 * u, WHITE, 1.6)
draw_line(slide, x0, y0, x0, y0 + 4 * u, MUTED, 1.3)
draw_line(slide, x0 + 4 * u, y0, x0 + 4 * u, y0 + 4 * u, MUTED, 1.3)
draw_line(slide, x0, y0 + 4 * u, x0 + 4 * u, y0 + 4 * u, MUTED, 1.3)
# center of lower
cx, cy = x0 + 2 * u, y0 - 2 * u
# hole image at top of mirror, 2 ft from left? PW used 6 by 2: vertical 6 ft (2+4), horizontal 2 ft
small_dot(slide, cx, cy, GOLD)
tx(slide, "O", cx + 0.1, cy - 0.05, 0.3, 0.22, 12, GOLD, True, spa=0, scaled=False)
# image of hole: from O, 2 ft sideways and 6 ft up in the unfolded picture
# draw OQ
arrow_line(slide, cx, cy, x0 + 2 * u, y0 + 2 * u, GREEN, 2.4)
tx(slide, "image", x0 + 2 * u + 0.1, y0 + 2 * u - 0.1, 0.8, 0.22, 11, GREEN, True, spa=0, scaled=False)
fml(slide, "path = √(6^2 + 2^2) = √40 = 2√10 ft", 7.4, 2.5, 5.2, LH, 15)
fml(slide, "displacement = centre to corner = 2√2 ft", 7.4, 2.5 + LH + NG, 5.2, LH, 14)
result_card(slide, "distance 2√10 ft   ·   |displacement| 2√2 ft", 7.4, 4.3, 5.2, 0.55, 13)

# ===== DRUNKARD + WHEEL =====
slide = new_slide("Example", "Drunkard and the pit", ORANGE)
para(slide, "He takes 5 steps forward and 3 steps back, again and again. Each step is 1 m and takes 1 s. A pit is 11 m ahead.", XL, Y0, CW, 1.00, 16)
fml(slide, "One cycle: +5 then −3  →  net +2 m in 8 s", XL, Y0 + 1.20, CW, LH, 16)
fml(slide, "Three cycles: 6 m in 24 s. He is still short of the pit.", XL, Y0 + 1.20 + LH + NG, CW, LH, 15)
fml(slide, "Next 5 forward steps: 6 + 5 = 11 m. He falls in. No return.", XL, Y0 + 1.20 + 2 * (LH + NG), CW, LH, 15)
result_card(slide, "Time = 24 + 5 = 29 s", XL, Y0 + 3.20, CW, 0.50, 18)
trap_card(slide, "Do not add a backward 3 after he has already reached 11 m.", XL, Y0 + 3.90, CW, 0.50)

slide = question_slide("MCQ — time to the pit")
qtx(slide, "The drunkard falls into the pit after", XL, Y0, CW, 0.55, 17)
choices(slide, ["21 s", "29 s", "31 s", "37 s"], XL, Y0 + 0.80)

slide = answer_slide()
fml(slide, "8 + 8 + 8 + 5 = 29 s", XL, Y0, CW, LH, 18)
result_card(slide, "Answer: 29 s", XL, Y0 + 0.90, CW, 0.45, 18)

slide = new_slide("Figure", "Point on a pure-rolling wheel", BLUE)
para(slide, "A wheel of radius R rolls without slipping. Point P starts at the bottom, touching the road. We want the displacement of P, not the path length of P.", XL, Y0, CW, 1.05, 15)
# two wheels
def wheel(cx, cy, r, color):
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(40, 18, 18)
    oval.line.color.rgb = color
    oval.line.width = Pt(1.8)

wheel(2.2, 4.6, 0.7, ORANGE)
wheel(6.2, 4.6, 0.7, ORANGE)
draw_line(slide, 1.2, 5.35, 8.2, 5.35, MUTED, 1.4)
small_dot(slide, 2.2, 5.25, GREEN)
small_dot(slide, 6.2, 5.25, GOLD)
arrow_line(slide, 2.4, 5.7, 6.0, 5.7, GREEN, 2.2)
tx(slide, "full turn: |Δr| = 2πR", 1.5, 5.85, 5.5, 0.30, 14, GREEN, True, spa=0, scaled=False)
tx(slide, "Half turn", 8.6, 2.3, 4.0, 0.35, 16, GOLD, True, spa=0, scaled=False)
fml(slide, "|Δr| = √((πR)^2 + (2R)^2)", 8.4, 2.8, 4.4, LH, 14)
fml(slide, "= R √(π^2 + 4)", 8.4, 2.8 + LH + NG, 4.4, LH, 16)
trap_card(slide, "Full turn brings P back to the road, but the wheel has moved forward.", 8.4, 4.5, 4.4, 0.85)

# ===== x(t) =====
slide = new_slide("Method", "Position given by x(t)", TEAL)
para(slide, "When the question gives x as a formula in t, do not assume the object keeps one direction. Find the instant when velocity is zero. If that instant lies inside the asked time, the object turned there.", XL, Y0, CW, 1.25, 15)
fml(slide, "1.  v = dx/dt. Set v = 0 and solve for t.", XL, Y0 + 1.45, CW, LH, 16)
fml(slide, "2.  Keep a root only if it lies between the start and end times.", XL, Y0 + 1.45 + LH + NG, CW, LH, 15)
fml(slide, "3.  Displacement = x(end) − x(start).", XL, Y0 + 1.45 + 2 * (LH + NG), CW, LH, 16)
fml(slide, "4.  Distance = sum of |Δx| on each piece split by the turn.", XL, Y0 + 1.45 + 3 * (LH + NG), CW, LH, 15)

slide = new_slide("Example", "x = t^2 − 4t + 2, from 0 to 3 s", ORANGE)
fml(slide, "v = 2t − 4 = 0  →  t = 2 s  (inside 0 to 3)", XL, Y0, CW, LH, 16)
fml(slide, "x(0) = 2 ,   x(2) = −2 ,   x(3) = −1", XL, Y0 + LH + NG, CW, LH, 16)
fml(slide, "displacement = −1 − 2 = −3 m", XL, Y0 + 2 * (LH + NG), CW, LH, 16)
fml(slide, "distance = |−2 − 2| + |−1 − (−2)| = 4 + 1 = 5 m", XL, Y0 + 3 * (LH + NG), CW, LH, 15)
# number line
sx = number_line(slide, 1.5, 5.3, 10.0, -3, 3, [-2, -1, 0, 2])
mark_point(slide, sx(2), 5.3, "t=0", GREEN)
mark_point(slide, sx(-2), 5.3, "t=2", ORANGE, above=False)
mark_point(slide, sx(-1), 5.3, "t=3", GOLD)

slide = new_slide("Example", "x = −t^2 + 6t + 4, from 0 to 4 s", BLUE)
fml(slide, "v = −2t + 6 = 0  →  t = 3 s  (inside 0 to 4)", XL, Y0, CW, LH, 16)
fml(slide, "x(0) = 4 ,   x(3) = 13 ,   x(4) = 12", XL, Y0 + LH + NG, CW, LH, 16)
fml(slide, "displacement = 12 − 4 = 8 m", XL, Y0 + 2 * (LH + NG), CW, LH, 16)
fml(slide, "distance = |13 − 4| + |12 − 13| = 9 + 1 = 10 m", XL, Y0 + 3 * (LH + NG), CW, LH, 15)
result_card(slide, "Turn inside the interval → distance is larger than |displacement|", XL, Y0 + 3.55, CW, 0.55, 14)

slide = new_slide("Example", "Arc plus two closing sides", GREEN)
para(slide, "Radius R = 21 cm. The object travels a 240° arc, then two straight pieces of length R each, and the loop closes.", XL, Y0, CW, 1.00, 15)
fml(slide, "arc = (240/360) × 2πR = (2/3) × 2π × 21", XL, Y0 + 1.20, CW, LH, 16)
fml(slide, "Use π = 22/7:  arc = 88 cm", XL, Y0 + 1.20 + LH + NG, CW, LH, 16)
fml(slide, "two straight pieces = 21 + 21 = 42 cm", XL, Y0 + 1.20 + 2 * (LH + NG), CW, LH, 16)
result_card(slide, "distance = 88 + 42 = 130 cm", XL, Y0 + 3.15, CW, 0.50, 17)

slide = new_slide("Example", "Same ends, two different routes", TEAL)
para(slide, "Both travellers start and finish at opposite ends of a diameter. A walks the semicircle. B walks three sides of the inscribed regular hexagon. Side of the hexagon equals R.", XL, Y0, CW, 1.15, 15)
fml(slide, "|displacement| is the same for both: the diameter 2R. Ratio = 1.", XL, Y0 + 1.35, CW, LH, 15)
fml(slide, "distance of A = πR ,   distance of B = 3R", XL, Y0 + 1.35 + LH + NG, CW, LH, 16)
result_card(slide, "distance ratio A : B = πR : 3R = π : 3", XL, Y0 + 2.80, CW, 0.55, 16)

# ===== SPEED =====
slide = new_slide("Idea", "Speed and velocity — first definitions", ORANGE)
para(slide, "Speed uses distance. Velocity uses displacement. That is the only split you need before the formulas.", XL, Y0, CW, 1.00, 16)
ncert(slide, "average speed = total distance / total time     (scalar)", XL, Y0 + 1.20, CW, 0.95, 15)
ncert(slide, "average velocity = total displacement / total time     (vector)", XL, Y0 + 2.30, CW, 0.95, 15)
fml(slide, "instantaneous speed = |dx/dt|     (size of the slope of x–t)", XL, Y0 + 3.50, CW, LH, 15)

slide = new_slide("Rule", "Two ways to mix speeds", BLUE)
para(slide, "If each piece is given by a speed and a time, weight by time. If each piece is given by a speed and a distance, weight by distance. Do not average the two speeds by adding and dividing by 2 unless the times are equal.", XL, Y0, CW, 1.30, 15)
fml(slide, "speeds for times:   (s_1 t_1 + s_2 t_2) / (t_1 + t_2)", XL, Y0 + 1.50, CW, LH, 16)
fml(slide, "speeds for distances:   (x_1 + x_2) / (x_1/s_1 + x_2/s_2)", XL, Y0 + 1.50 + LH + NG, CW, LH, 15)
trap_card(slide, "Equal times → ordinary mean. Equal distances → the distance formula, not (s_1+s_2)/2.", XL, Y0 + 3.10, CW, 0.70)

slide = new_slide("Example", "Four straight trips — one table", GREEN)
grid_table(
    slide,
    ["Trip", "avg speed", "|avg velocity|"],
    [
        ["0 → 20 m in 4 s", "5 m/s", "5 m/s"],
        ["0 → 20 → 5 m in 10 s", "3.5 m/s", "0.5 m/s"],
        ["0 → −20 → 10 m in 5 s", "10 m/s", "2 m/s"],
        ["out and back, 40 m path, 10 s", "4 m/s", "0"],
    ],
    XL, Y0, [5.2, 3.2, 3.4], row_h=0.62, hsize=13, csize=14,
)
trap_card(slide, "Path length goes in the speed. Only start-to-end goes in the velocity.", XL, Y0 + 3.50, CW, 0.55)

slide = new_slide("Example", "Semicircle in 10 s", TEAL)
para(slide, "Radius 4 m. The object runs the semicircle from one end of the diameter to the other in 10 s.", XL, Y0, CW, 0.90, 16)
fml(slide, "distance = πR = 4π m ,   |displacement| = 8 m", XL, Y0 + 1.15, CW, LH, 16)
fml(slide, "average speed = 4π / 10 = 0.4 π m/s", XL, Y0 + 1.15 + LH + NG, CW, LH, 16)
fml(slide, "|average velocity| = 8 / 10 = 0.8 m/s", XL, Y0 + 1.15 + 2 * (LH + NG), CW, LH, 16)

slide = new_slide("Rule", "Constant speed v on a circle", BLUE)
para(slide, "One lap takes time 2πR/v. Read each fraction of the lap from that.", XL, Y0, CW, 0.80, 15)
grid_table(
    slide,
    ["Part", "distance", "|disp|", "|v_avg|"],
    [
        ["1/4 lap", "πR/2", "R√2", " (R√2) / t "],
        ["1/2 lap", "πR", "2R", "2v/π"],
        ["full lap", "2πR", "0", "0"],
    ],
    XL, Y0 + 1.00, [2.6, 2.8, 2.6, 3.8], row_h=0.58, hsize=14, csize=15,
)
trap_card(slide, "Average speed for a full lap is still v. Average velocity is zero.", XL, Y0 + 3.55, CW, 0.55)

slide = new_slide("Example", "Half the distance at v, half at 2v", ORANGE)
fml(slide, "time = (d/2)/v + (d/2)/(2v) = 3d / (4v)", XL, Y0, CW, LH, 16)
fml(slide, "average speed = d / (3d/(4v)) = 4v/3", XL, Y0 + LH + NG, CW, LH, 17)
para(slide, "Same direction the whole way, so |average velocity| equals this average speed.", XL, Y0 + 1.70, CW, 0.90, 15)
result_card(slide, "Answer: 4v/3", XL, Y0 + 2.80, CW, 0.50, 18)

slide = new_slide("Example", "Out at 30 km/h, back at 60 km/h", GREEN)
para(slide, "Same road both ways. Distances are equal, so use the distance formula. Let the one-way distance be d.", XL, Y0, CW, 1.00, 16)
fml(slide, "average speed = 2d / (d/30 + d/60) = 2 / (1/30 + 1/60)", XL, Y0 + 1.20, CW, LH, 15)
fml(slide, "= 2 / (3/60) = 40 km/h", XL, Y0 + 1.20 + LH + NG, CW, LH, 17)
result_card(slide, "Round trip of two speeds: 2 s_1 s_2 / (s_1 + s_2)", XL, Y0 + 2.70, CW, 0.70, 15)
trap_card(slide, "Displacement for the round trip is zero, so average velocity is zero.", XL, Y0 + 3.60, CW, 0.55)

slide = question_slide("MCQ — three equal parts")
qtx(slide, "First third at 10 m/s, second at 20 m/s, last at 30 m/s, all same direction. Average velocity is", XL, Y0, CW, 0.80, 16)
choices(slide, ["17.8 m/s", "16.4 m/s", "18.3 m/s", "20.2 m/s"], XL, Y0 + 1.05)

slide = answer_slide()
fml(slide, "t = (x/10) + (x/20) + (x/30) = 11x/60", XL, Y0, CW, LH, 16)
fml(slide, "total distance = 3x ,   average = 3x / (11x/60) = 180/11", XL, Y0 + LH + NG, CW, LH, 15)
result_card(slide, "180/11 ≈ 16.4 m/s", XL, Y0 + 1.70, CW, 0.50, 18)

slide = new_slide("Example", "Unequal times, then a mixed half", TEAL)
fml(slide, "v_1 for time t/3 and v_2 for time 2t/3", XL, Y0, CW, LH, 16)
fml(slide, "average = (v_1 + 2 v_2) / 3", XL, Y0 + LH + NG, CW, LH, 17)
para(slide, "A different mix: first half of the distance at 5 m/s, second half at 15 m/s.", XL, Y0 + 1.80, CW, 0.85, 15)
fml(slide, "average = 2×5×15 / (5+15) = 7.5 m/s", XL, Y0 + 2.80, CW, LH, 16)

slide = new_slide("Example", "One full wave on x–t", BLUE)
para(slide, "Position swings between +10 m and −10 m and returns to the start after 8 s. That is one full oscillation.", XL, Y0, CW, 1.00, 16)
sx, sy = graph_axes(slide, 1.3, 2.5, 6.2, 3.2, 0, 8, -12, 12, "t", "x")
# rough sine via points
pts = []
for i in range(17):
    t = 8 * i / 16
    y = 10 * math.sin(2 * math.pi * t / 8)
    pts.append((sx(t), sy(y)))
curve(slide, pts, GOLD, 2.2)
fml(slide, "distance ≈ 4 × 10 = 40 m in 8 s", 8.0, 2.6, 4.6, LH, 14)
fml(slide, "average speed = 40/8 = 5 m/s", 8.0, 2.6 + LH + NG, 4.6, LH, 15)
fml(slide, "displacement = 0  →  average velocity = 0", 8.0, 2.6 + 2 * (LH + NG), 4.6, LH, 14)

slide = new_slide("Lock lines", "Only today’s lines", GREEN)
bullet(slide, "Fly the cube: |displacement| = ℓ√3. Crawl it: shortest distance = ℓ√5.", XL, Y0, CW, 0.55, 15)
bullet(slide, "A bounce becomes a straight line after you reflect the board.", XL, Y0 + 0.60, CW, 0.50, 15)
bullet(slide, "Drunkard: count full cycles, then the last forward run into the pit.", XL, Y0 + 1.15, CW, 0.50, 15)
bullet(slide, "Rolling point, half turn: R √(π^2 + 4).", XL, Y0 + 1.70, CW, 0.48, 15)
bullet(slide, "In x(t), split the path only at a turning time inside the interval.", XL, Y0 + 2.25, CW, 0.50, 15)
bullet(slide, "Equal times: ordinary mean. Equal distances: the distance formula.", XL, Y0 + 2.80, CW, 0.50, 15)
ncert(slide, "Next lecture continues velocity and acceleration. Today’s formulas stay.", XL, Y0 + 3.55, CW, 0.90, 14)

slide = prs.slides.add_slide(BLANK)
_bg(slide)
_bar(slide, "End", GOLD)
tx_3d(slide, "Thank you", XL, 2.4, CW, 1.0, 48, GOLD, RGBColor(120, 120, 130))
tx(slide, "Excellence  ·  Motion 03  ·  New paths · Speed · Average velocity", XL, 3.7, CW, 0.45, 15, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, len(prs.slides))
after_content_slide()

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(OUTPUT)
print("slides", len(prs.slides))
