"""Rebuild Meera Regar Jaisalmer circuit in Teena format using ONLY real photos from source PPTX."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
MEDIA = BASE / "_jaisal_media"
OUTPUT = BASE / "geography Jaisalmer tourism circuit - Teena format.pptx"

# Teena aspect
W, H = 10.0, 5.625

MUSTARD = RGBColor(214, 170, 58)
WHITE = RGBColor(255, 255, 255)
BROWN = RGBColor(166, 110, 58)
TEAL = RGBColor(20, 90, 85)
BLUE = RGBColor(40, 95, 180)
DARK = RGBColor(40, 40, 45)
SOFT = RGBColor(248, 248, 248)

# Real photos extracted from source PPTX (not full scanned pages)
P = {
    "logo": MEDIA / "s01_830286.jpg",
    "map": MEDIA / "s06_1429375.jpg",
    "intro_art": MEDIA / "s09_235716.png",
    "fort1": MEDIA / "s11_1015250.jpg",
    "fort2": MEDIA / "s11_818834.jpg",
    "fort3": MEDIA / "s12_1575645.jpg",
    "haveli1": MEDIA / "s14_788312.jpg",
    "haveli2": MEDIA / "s15_719606.jpg",
    "lake": MEDIA / "s17_651573.jpg",
    "dunes": MEDIA / "s19_808246.jpg",
    "safari": MEDIA / "s20_251059.jpg",
    "tanot": MEDIA / "s21_629897.jpg",
    "culture": MEDIA / "s23_715794.jpg",
    "food": MEDIA / "s25_699921.jpg",
    # extras from media folder if slide exports missing
    "extra1": MEDIA / "image406.jpg",
    "extra2": MEDIA / "image367.jpg",
    "extra3": MEDIA / "image605.jpg",
    "extra4": MEDIA / "image636.jpg",
    "extra5": MEDIA / "image532.jpg",
    "extra6": MEDIA / "image411.jpg",
}

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def solid(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def tx(slide, text, x, y, w, h, size=16, color=BROWN, bold=False, align=PP_ALIGN.LEFT,
       font="Nirmala UI", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def corner_stripes(slide, side="content"):
    if side == "cover":
        for i, off in enumerate((0.0, 0.18, 0.36)):
            s = slide.shapes.add_shape(
                MSO_SHAPE.PARALLELOGRAM, Inches(-0.4 + off), Inches(0.15 + i * 0.12),
                Inches(2.2), Inches(0.16),
            )
            s.fill.solid(); s.fill.fore_color.rgb = DARK; s.line.fill.background(); s.rotation = 28.0
        for i, off in enumerate((0.0, 0.18, 0.36)):
            s = slide.shapes.add_shape(
                MSO_SHAPE.PARALLELOGRAM, Inches(7.7 + off), Inches(4.6 + i * 0.10),
                Inches(2.4), Inches(0.16),
            )
            s.fill.solid(); s.fill.fore_color.rgb = MUSTARD; s.line.fill.background(); s.rotation = 28.0
    else:
        for i, off in enumerate((0.0, 0.16, 0.32)):
            s = slide.shapes.add_shape(
                MSO_SHAPE.PARALLELOGRAM, Inches(8.2 + off), Inches(-0.15 + i * 0.08),
                Inches(2.1), Inches(0.18),
            )
            s.fill.solid(); s.fill.fore_color.rgb = TEAL; s.line.fill.background(); s.rotation = 35.0
        for i, off in enumerate((0.0, 0.16, 0.32)):
            s = slide.shapes.add_shape(
                MSO_SHAPE.PARALLELOGRAM, Inches(-0.5 + off), Inches(4.55 + i * 0.08),
                Inches(2.1), Inches(0.18),
            )
            s.fill.solid(); s.fill.fore_color.rgb = BLUE; s.line.fill.background(); s.rotation = 35.0


def pic(slide, path: Path, x, y, w, h):
    if path and path.exists():
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return None


def content_slide(title: str, bullets: list[str], image: Path | None = None):
    slide = prs.slides.add_slide(BLANK)
    solid(slide, 0, 0, W, H, WHITE)
    corner_stripes(slide, "content")
    tx(slide, title, 0.4, 0.18, 9.2, 0.50, 20, BROWN, True, PP_ALIGN.CENTER)
    y = 0.85
    for b in bullets:
        tx(slide, "●  " + b, 0.40, y, 5.2, 0.70, 14, BROWN, False)
        y += 0.70
    if image and image.exists():
        solid(slide, 5.80, 0.85, 3.80, 4.00, SOFT)
        pic(slide, image, 5.90, 0.95, 3.60, 3.80)
    return slide


def text_block_slide(title: str, paragraphs: list[str], image: Path | None = None):
    slide = prs.slides.add_slide(BLANK)
    solid(slide, 0, 0, W, H, WHITE)
    corner_stripes(slide, "content")
    tx(slide, title, 0.4, 0.18, 9.2, 0.48, 20, BROWN, True, PP_ALIGN.CENTER)
    y = 0.80
    left_w = 5.2 if image else 9.1
    for p in paragraphs:
        tx(slide, p, 0.40, y, left_w, 1.10, 13, BROWN, False)
        y += 1.15
    if image and image.exists():
        solid(slide, 5.80, 0.85, 3.80, 4.00, SOFT)
        pic(slide, image, 5.90, 0.95, 3.60, 3.80)
    return slide


# COVER
slide = prs.slides.add_slide(BLANK)
solid(slide, 0, 0, W, H, MUSTARD)
white_tri = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(4.6), Inches(2.55), Inches(6.0), Inches(3.5))
white_tri.fill.solid(); white_tri.fill.fore_color.rgb = WHITE; white_tri.line.fill.background()
corner_stripes(slide, "cover")
# small logo only (not full page)
if P["logo"].exists():
    pic(slide, P["logo"], 7.55, 0.35, 1.90, 1.90)
tx(slide, "Acharya Shri mahapragyaa institute of excellence,", 0.55, 0.70, 6.8, 0.40, 15, WHITE, True)
tx(slide, "Asind ( Bhilwara)", 0.55, 1.10, 6.0, 0.35, 15, WHITE, True)
tx(slide, "Name :- Meera Regar", 0.55, 1.75, 6.5, 0.35, 16, WHITE, True)
tx(slide, "Class :- B.A. part - 3 Sem - 6", 0.55, 2.15, 6.5, 0.35, 16, WHITE, True)
tx(slide, "Roll.no. :- 24425671", 0.55, 2.55, 6.5, 0.35, 16, WHITE, True)
tx(slide, "Topic :- Jaisalmer Tourism Circuit", 0.55, 2.95, 6.5, 0.35, 16, WHITE, True)
tx(slide, "Guidance :- Sanwar Lal Prajapat", 0.55, 3.35, 6.5, 0.35, 15, WHITE, True)

# INDEX
slide = prs.slides.add_slide(BLANK)
solid(slide, 0, 0, W, H, WHITE)
corner_stripes(slide, "content")
tx(slide, "~ जैसलमेर पर्यटन सर्किट ~", 0.4, 0.20, 9.2, 0.45, 22, BROWN, True, PP_ALIGN.CENTER)
left = [
    "1. प्रस्तावना",
    "2. जैसलमेर का संक्षिप्त परिचय",
    "3. भौगोलिक एवं ऐतिहासिक परिचय",
    "4. पर्यटन सर्किट का परिचय",
    "5. स्वर्णिम किला (सोनार किला)",
    "6. पटवों की हवेली",
    "7. नाथमल की हवेली",
    "8. सलीम सिंह की हवेली",
]
right = [
    "9. गडीसर झील",
    "10. बड़ा बाग",
    "11. कुलधरा गाँव",
    "12. डेजर्ट नेशनल पार्क",
    "13. सम के रेत के टीले",
    "14. तनोट माता / लोंगेवाला",
    "15. संस्कृति, मेले व निष्कर्ष",
]
for i, line in enumerate(left):
    tx(slide, line, 0.50, 0.80 + i * 0.42, 4.5, 0.38, 14, BROWN, False)
for i, line in enumerate(right):
    tx(slide, line, 5.20, 0.80 + i * 0.42, 4.4, 0.38, 14, BROWN, False)

# CONTENT — photos only from source media
text_block_slide(
    "💫 प्रस्तावना 💫",
    [
        "जैसलमेर राजस्थान की स्वर्ण नगरी है। यह अपनी स्थापत्य कला, इतिहास, लोक संस्कृति और थार मरुस्थल के लिए प्रसिद्ध है।",
        "सोनार किला, पटवों की हवेली, गडीसर झील, बड़ा बाग, कुलधरा और सम के रेत के टीले इस सर्किट के मुख्य आकर्षण हैं।",
        "यह पर्यटन सर्किट राजस्थान की मरु-विरासत को समझने और स्थानीय रोज़गार बढ़ाने में सहायक है।",
    ],
    P["intro_art"],
)

text_block_slide(
    "✨ जैसलमेर का संक्षिप्त परिचय ✨",
    [
        "जैसलमेर को स्वर्ण नगरी कहा जाता है क्योंकि यहाँ की पीले बलुआ पत्थर की इमारतें सूर्य प्रकाश में सुनहरी दिखती हैं।",
        "प्रमुख आकर्षण:- सोनार किला, पटवों व नाथमल की हवेलियाँ, गडीसर झील, थार मरुस्थल और सम के धोरे।",
        "यहाँ लोक संगीत, कालबेलिया नृत्य और रेगिस्तानी संस्कृति पर्यटकों को आकर्षित करती है।",
    ],
    P["map"],
)

content_slide(
    "⚡ भौगोलिक एवं ऐतिहासिक परिचय ⚡",
    [
        "भौगोलिक :- पश्चिमी राजस्थान का सीमावर्ती जिला, थार मरुस्थल में स्थित।",
        "जलवायु :- ग्रीष्म में अत्यधिक गर्मी, शीत में ठंडी रातें, कम वर्षा।",
        "ऐतिहासिक :- 12वीं शताब्दी में रावल जैसल द्वारा स्थापना।",
        "प्रसिद्ध :- सोनार किला, हवेलियाँ, बड़ा बाग और कुलधरा।",
    ],
    P["fort1"],
)

content_slide(
    "🌟 पर्यटन सर्किट का परिचय 🌟",
    [
        "सर्किट में किला, हवेलियाँ, झील, गाँव, मरुस्थल और सीमा क्षेत्र जुड़े हैं।",
        "मुख्य स्थल :- सोनार किला, पटवों की हवेली, गडीसर झील, सम धोरे।",
        "गतिविधियाँ :- ऊँट सफारी, कैंपिंग, लोक नृत्य-संगीत, राजस्थानी खान-पान।",
        "यह सर्किट ऐतिहासिक, सांस्कृतिक और प्राकृतिक पर्यटन का संगम है।",
    ],
    P["fort2"],
)

content_slide(
    "🏰 स्वर्णिम किला (सोनार किला) 🏰",
    [
        "त्रिकूट पहाड़ी पर बना विश्व प्रसिद्ध जैसलमेर किला।",
        "पीले बलुआ पत्थर से निर्मित होने के कारण सोनार किला कहलाता है।",
        "किले के अंदर मंदिर, हवेलियाँ, बाज़ार और निवासी क्षेत्र हैं।",
        "सूर्यास्त के समय किला सुनहरे रंग में चमकता है।",
    ],
    P["fort3"],
)

content_slide(
    "🏛️ पटवों की हवेली 🏛️",
    [
        "जैसलमेर की सबसे भव्य और प्रसिद्ध हवेली समूह।",
        "पाँच हवेलियों का समूह, उत्कृष्ट नक्काशी और झरोखे।",
        "व्यापारी पटवा परिवार द्वारा बनवाई गई।",
        "संग्रहालय और वास्तुकला अध्ययन के लिए महत्वपूर्ण।",
    ],
    P["haveli1"],
)

content_slide(
    "🏡 नाथमल की हवेली 🏡",
    [
        "दो भाई शिल्पकारों द्वारा बनाई गई अनोखी हवेली।",
        "दोनों भाग लगभग समान, पर सूक्ष्म अंतर भी दिखाई देते हैं।",
        "पत्थर की नक्काशी और हाथी की मूर्तियाँ विशेष आकर्षण हैं।",
        "जैसलमेर की स्थापत्य कला का उत्कृष्ट उदाहरण।",
    ],
    P["extra1"],
)

content_slide(
    "🕌 सलीम सिंह की हवेली 🕌",
    [
        "विशेष मयूर शैली की छतों वाली हवेली।",
        "ऊपरी मंज़िल की नक्काशी और झरोखे अत्यंत आकर्षक।",
        "नगर के मध्य में स्थित प्रमुख दर्शनीय स्थल।",
        "जैसलमेर की शाही-व्यापारी वास्तुकला का प्रतीक।",
    ],
    P["haveli2"],
)

content_slide(
    "💧 गडीसर झील 💧",
    [
        "प्राचीन कृत्रिम झील और जल संरक्षण का उदाहरण।",
        "किनारे पर सुंदर छतरियाँ, मंदिर और घाट।",
        "नौका विहार और सूर्यास्त के दृश्य प्रसिद्ध हैं।",
        "शहर के निकट स्थित शांत पर्यटन स्थल।",
    ],
    P["lake"],
)

content_slide(
    "🌳 बड़ा बाग 🌳",
    [
        "भाटी शासकों की छतरियों और स्मारकों का समूह।",
        "सुनहरे बलुआ पत्थर की छतरियाँ मरुस्थलीय पृष्ठभूमि में सुंदर लगती हैं।",
        "सूर्योदय और सूर्यास्त फोटोग्राफी के लिए लोकप्रिय।",
        "जैसलमेर की शाही स्मृति और इतिहास से जुड़ा स्थल।",
    ],
    P["extra2"],
)

content_slide(
    "🏚️ कुलधरा गाँव 🏚️",
    [
        "रहस्यमय परित्यक्त गाँव, ऐतिहासिक कथाओं से जुड़ा।",
        "पुराने घरों के खंडहर और सुनसान गलियाँ आकर्षित करती हैं।",
        "पर्यटक यहाँ इतिहास और लोककथाओं को जानने आते हैं।",
        "जैसलमेर सर्किट का महत्वपूर्ण ऑफबीट स्थल।",
    ],
    P["extra6"],
)

content_slide(
    "🦌 डेजर्ट नेशनल पार्क 🦌",
    [
        "थार मरुस्थल का प्रमुख संरक्षित क्षेत्र।",
        "ग्रेट इंडियन बस्टर्ड सहित मरुस्थलीय वन्यजीव।",
        "जीप सफारी और रेगिस्तानी परिदृश्य का अनुभव।",
        "प्राकृतिक पर्यटन का महत्वपूर्ण केंद्र।",
    ],
    P["extra5"],
)

content_slide(
    "🐪 सम के रेत के टीले 🐪",
    [
        "जैसलमेर से लगभग 40–45 किमी दूर प्रसिद्ध रेत के धोरे।",
        "ऊँट सवारी, जीप सफारी और डेजर्ट कैंपिंग।",
        "लोक संगीत-नृत्य और सूर्यास्त का अद्भुत दृश्य।",
        "सर्किट का सबसे लोकप्रिय मरुस्थलीय अनुभव।",
    ],
    P["dunes"],
)

content_slide(
    "🛕 तनोट माता एवं लोंगेवाला 🛕",
    [
        "तनोट माता मंदिर सीमा क्षेत्र की आस्था का केंद्र।",
        "लोंगेवाला युद्ध स्मारक देशभक्ति और इतिहास से जुड़ा।",
        "युद्ध संग्रहालय और टैंक्स पर्यटकों को आकर्षित करते हैं।",
        "देशभक्ति पर्यटन का महत्वपूर्ण हिस्सा।",
    ],
    P["tanot"],
)

content_slide(
    "🎵 संस्कृति, खान-पान एवं मेले 🎵",
    [
        "लोक संगीत, कालबेलिया नृत्य और हस्तशिल्प प्रसिद्ध।",
        "दाल-बाटी, कैर-सांगरी और मिष्ठान स्थानीय स्वाद हैं।",
        "मरु महोत्सव जैसलमेर का प्रमुख उत्सव है।",
        "पर्यटकों के लिए सफारी, कैंपिंग और सांस्कृतिक कार्यक्रम।",
    ],
    P["culture"] if P["culture"].exists() else P["food"],
)

# conclusion
slide = prs.slides.add_slide(BLANK)
solid(slide, 0, 0, W, H, RGBColor(55, 60, 70))
solid(slide, 0.55, 0.45, 8.9, 4.70, WHITE)
corner_stripes(slide, "content")
tx(slide, "∆ निष्कर्ष ∆", 0.7, 1.20, 8.6, 0.55, 24, BROWN, True, PP_ALIGN.CENTER)
tx(
    slide,
    "जैसलमेर पर्यटन सर्किट राजस्थान की ऐतिहासिक, सांस्कृतिक और प्राकृतिक धरोहर को एक संगठित रूप में प्रस्तुत करता है। स्वर्णिम किला, हवेलियाँ, झील, मरुस्थल, लोक कला और सीमा क्षेत्र देश-विदेश के पर्यटकों को आकर्षित करते हैं तथा स्थानीय अर्थव्यवस्था को भी बल देते हैं।",
    1.1, 2.00, 7.8, 2.2, 15, BROWN, False, PP_ALIGN.CENTER,
)

# aabhar
slide = prs.slides.add_slide(BLANK)
solid(slide, 0, 0, W, H, WHITE)
corner_stripes(slide, "content")
tx(slide, "🙏 आभार 🙏", 0.5, 0.35, 9.0, 0.50, 22, BROWN, True, PP_ALIGN.CENTER)
tx(
    slide,
    "मैं आचार्य श्री महाप्रज्ञ इंस्टीट्यूट ऑफ एक्सीलेंस, आसींद (भीलवाड़ा) के भूगोल विभाग तथा मार्गदर्शक श्री सांवर लाल प्रजापत का हृदय से आभार व्यक्त करती हूँ। माता-पिता, साथियों और सभी सूचना स्रोतों का भी धन्यवाद।",
    0.8, 1.30, 8.4, 1.8, 15, BROWN, False, PP_ALIGN.LEFT,
)
tx(slide, "धन्यवाद!", 0.8, 3.30, 8.4, 0.45, 18, BROWN, True, PP_ALIGN.LEFT)
tx(slide, "प्रस्तुतकर्ता :- Meera Regar", 0.8, 4.00, 5.0, 0.40, 14, BROWN, False)
tx(slide, "Roll no. :- 24425671", 0.8, 4.40, 5.0, 0.40, 14, BROWN, False)

missing = [k for k, v in P.items() if not v.exists()]
if missing:
    print("Missing media keys:", missing)

prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
