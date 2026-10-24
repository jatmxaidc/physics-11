"""Lecture 15 — introduction to integration (reverse of differentiation)."""
from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "lecture slides" / "Lecture_15_Introduction_to_Integration.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.42
NG = 0.18
PT_SCALE = 0.90


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(slide, text, x, y, w, h, size=15, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spa=1, lsp=1.1,
       font="Aptos", scaled=True):
    for old, new in (
        ("^(n-1)", "ⁿ⁻¹"),
        ("^(n+1)", "ⁿ⁺¹"),
        ("^(3/2)", "³ᐟ²"),
        ("^(1/2)", "¹ᐟ²"),
        ("^(-2)", "⁻²"),
        ("^(-1)", "⁻¹"),
        ("^(-x)", "⁻ˣ"),
        ("^n", "ⁿ"),
        ("^x", "ˣ"),
        ("^6", "⁶"),
        ("^5", "⁵"),
        ("^4", "⁴"),
        ("^3", "³"),
        ("^2", "²"),
        ("^0", "⁰"),
        ("sqrt", "√"),
        ("omega", "ω"),
        ("lambda", "λ"),
        (" proportional to ", " ∝ "),
        (">=", "≥"),
        ("<=", "≤"),
        ("->", "→"),
    ):
        text = text.replace(old, new)
    text = re.sub(r"\s*=\s*", " = ", text)
    text = re.sub(r"\s*([<>])\s*", r" \1 ", text)
    text = re.sub(r",\s*", ", ", text)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size if not scaled else max(7, round(size * PT_SCALE, 1)))
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


_CONTENT = 0
_ROUGH_DUE = False
_ANSWER_STEPS = {}


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working.")


def after_content_slide(defer_rough=False):
    global _CONTENT, _ROUGH_DUE
    _CONTENT += 1
    if defer_rough and _CONTENT % 5 == 0:
        _ROUGH_DUE = True
    elif not defer_rough and (_ROUGH_DUE or _CONTENT % 5 == 0):
        add_rough_work()
        _ROUGH_DUE = False


def new_slide(section, title, bar_color=BLUE, defer_rough=False):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide(defer_rough)
    return slide


_LAST_QUESTION_TITLE = ""


def question_slide(title):
    global _LAST_QUESTION_TITLE
    _LAST_QUESTION_TITLE = title
    return new_slide("Question", title, ORANGE, defer_rough=True)


def answer_slide(title=None):
    title = title or f"Solution: {_LAST_QUESTION_TITLE}"
    slide = new_slide("Answer", title, GREEN)
    _ANSWER_STEPS[id(slide)] = 0
    return slide


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
       size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15)


def fml(slide, text, x, y, w, h=0.44, size=22):
    if id(slide) in _ANSWER_STEPS:
        _ANSWER_STEPS[id(slide)] += 1
        text = f"Step {_ANSWER_STEPS[id(slide)]}:  {text}"
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.44, size=17):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, max(size, 19), QCOL, True, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h=0.36, tsize=18):
    return tx(slide, text, x, y, w, h, max(tsize, 18), GREEN, True, spa=0)


def trap_card(slide, text, x, y, w, h=0.40):
    return tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    w = w or CW
    for i, item in enumerate(items):
        cy = y + i * 0.40
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=14, csize=15):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
           hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h))
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
               csize, WORK if j == 0 else TEXT, j == 0,
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=False)


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def axes(slide, x=7.5, y=1.55, w=4.5, h=4.5, xlab="x", ylab="y"):
    ox, oy = x + w * 0.40, y + h * 0.58
    draw_line(slide, x, oy, x + w, oy, MUTED, 1.5)
    draw_line(slide, ox, y, ox, y + h, MUTED, 1.5)
    tx(slide, xlab, x + w - 0.15, oy + 0.06, 0.3, 0.22, 13, MUTED, True, spa=0)
    tx(slide, ylab, ox + 0.08, y, 0.3, 0.22, 13, MUTED, True, spa=0)
    return ox, oy, w, h


def line_on_axes(slide, ox, oy, dx1, dy1, dx2, dy2, color=GOLD, pt=2.4):
    return draw_line(slide, ox + dx1, oy - dy1, ox + dx2, oy - dy2, color, pt)


def curve(slide, points, color=GOLD, pt=2.4):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        draw_line(slide, x1, y1, x2, y2, color, pt)


def graph_axes(slide, x, y, w, h, xmin, xmax, ymin, ymax, xlab="x", ylab="y"):
    def sx(value):
        return x + (value - xmin) * w / (xmax - xmin)

    def sy(value):
        return y + h - (value - ymin) * h / (ymax - ymin)

    if ymin <= 0 <= ymax:
        draw_line(slide, x, sy(0), x + w, sy(0), MUTED, 1.4)
    if xmin <= 0 <= xmax:
        draw_line(slide, sx(0), y, sx(0), y + h, MUTED, 1.4)
    tx(slide, xlab, x + w - 0.18, sy(0) + 0.05 if ymin <= 0 <= ymax else y + h - 0.22,
       0.35, 0.24, 13, MUTED, True, spa=0, scaled=False)
    tx(slide, ylab, sx(0) + 0.07 if xmin <= 0 <= xmax else x + 0.05, y,
       0.45, 0.24, 13, MUTED, True, spa=0, scaled=False)
    return sx, sy


def axis_ticks(slide, sx, sy, x_values, y_values=None, color=MUTED):
    """Mark readable values on one or both axes."""
    y_values = x_values if y_values is None else y_values
    for v in x_values:
        if v == 0:
            continue
        draw_line(slide, sx(v), sy(0) - 0.06, sx(v), sy(0) + 0.06, color, 1.1)
        tx(slide, str(v), sx(v) - 0.18, sy(0) + 0.08, 0.36, 0.22, 12, color, True,
           PP_ALIGN.CENTER, spa=0, scaled=False)
    for v in y_values:
        if v == 0:
            continue
        draw_line(slide, sx(0) - 0.06, sy(v), sx(0) + 0.06, sy(v), color, 1.1)
        tx(slide, str(v), sx(0) - 0.46, sy(v) - 0.11, 0.38, 0.22, 12, color, True,
           PP_ALIGN.RIGHT, spa=0, scaled=False)
    tx(slide, "0", sx(0) - 0.32, sy(0) + 0.08, 0.26, 0.22, 12, color, True,
       PP_ALIGN.RIGHT, spa=0, scaled=False)


def plot_circle(slide, sx, sy, h, k, r, color=GOLD, pt=2.4):
    """True geometric circle via oval (equal axis units required for round look)."""
    left, top = sx(h - r), sy(k + r)
    width = sx(h + r) - sx(h - r)
    height = sy(k - r) - sy(k + r)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def plot_ellipse(slide, sx, sy, h, k, a, b, color=GOLD, pt=2.4):
    left, top = sx(h - a), sy(k + b)
    width = sx(h + a) - sx(h - a)
    height = sy(k - b) - sy(k + b)
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    oval.fill.background()
    oval.line.color.rgb = color
    oval.line.width = Pt(pt)
    return oval


def mark_point(slide, sx, sy, x, y, label, color=GREEN, dx=0.10, dy=-0.26):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx(x) - 0.055), Inches(sy(y) - 0.055),
                                 Inches(0.11), Inches(0.11))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    tx(slide, label, sx(x) + dx, sy(y) + dy, 0.90, 0.24, 12, color, True, spa=0, scaled=False)


def plot_function(slide, fn, sx, sy, xmin, xmax, ymin, ymax, color=GOLD, samples=160, pt=2.2):
    previous = None
    for i in range(samples + 1):
        xv = xmin + (xmax - xmin) * i / samples
        try:
            yv = fn(xv)
        except (ValueError, ZeroDivisionError, OverflowError):
            previous = None
            continue
        if not math.isfinite(yv) or yv < ymin or yv > ymax:
            previous = None
            continue
        current = (sx(xv), sy(yv))
        if previous is not None:
            draw_line(slide, previous[0], previous[1], current[0], current[1], color, pt)
        previous = current


def plot_parametric(slide, fx, fy, sx, sy, t0, t1, color=GOLD, samples=180, pt=2.2):
    points = []
    for i in range(samples + 1):
        t = t0 + (t1 - t0) * i / samples
        points.append((sx(fx(t)), sy(fy(t))))
    curve(slide, points, color, pt)


# COVER
slide=prs.slides.add_slide(BLANK); _bg(slide)
tx_3d(slide,"EXCELLENCE",XL,0.52,CW,1.10,62,GOLD,RGBColor(210,210,220))
tx_3d(slide,"CLASS 11   -   2026-27",XL,1.64,CW,0.58,30,WHITE,RGBColor(90,90,100),font="Aptos")
_solid(slide,4.15,2.28,5.00,0.028,WHITE); bx,by,bw=2.05,2.70,9.20
_card_bg(slide,bx,by,bw,2.35,WHITE,WHITE,0); _card_bg(slide,bx,by+0.72,bw,1.63,BLACK,RGBColor(220,220,220),1.25)
_solid(slide,bx+bw/2,by+0.16,0.018,0.42,RGBColor(40,40,40))
tx(slide,"Physics",bx,by+0.14,bw/2-0.10,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Lecture - 15",bx+bw/2,by+0.14,bw/2,0.50,25,BLACK,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"Basic Maths and Calculus",bx,by+1.02,bw,0.62,32,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False,font="Georgia")
tx(slide,"Introduction to Integration",bx,by+1.62,bw,0.40,18,GOLD,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"By - Excellence Programme",XL,5.28,CW,0.50,25,WHITE,True,PP_ALIGN.CENTER,spa=0,scaled=False)
tx(slide,"23 September 2026",XL,5.80,CW,0.32,18,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)
_footer(slide,1); add_notes(slide,"Start integration as the reverse of differentiation."); after_content_slide()

slide=prs.slides.add_slide(BLANK); _bg(slide)
tx(slide,"Topics to be covered",XL,0.35,CW,0.64,32,GOLD,True,spa=0,scaled=False); _solid(slide,XL,1.00,3.40,0.035,GOLD)
topics=["5-minute differentiation recap","Integration as reverse differentiation","Indefinite integral and +C","Power-rule reverse","Standard integrals table","Easy indefinite-integral practice","Physics link: a → v → x by integrating"]
for i,item in enumerate(topics):
    yy=1.14+i*0.76
    tx(slide,f"{i+1:02d}",XL,yy,0.70,0.50,23,GOLD,True,spa=0,scaled=False)
    tx(slide,item,XL+0.90,yy,CW-0.90,0.66,18,WHITE,False,spa=0,scaled=False)
_footer(slide,2); add_notes(slide,"Keep the first idea very clear: integration undoes differentiation."); after_content_slide()

# RECAP
slide=new_slide("5-minute recap","What differentiation gave us",BLUE)
fml(slide,"d(x^n)/dx = n x^(n-1)",XL,Y0,CW,LH,20)
fml(slide,"d(sin x)/dx = cos x     ·     d(cos x)/dx = −sin x",XL,Y0+LH+NG,CW,LH,18)
fml(slide,"d(e^x)/dx = e^x     ·     d(constant)/dx = 0",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"v = dx/dt     ·     a = dv/dt",XL,Y0+3*LH+3*NG,CW,LH,20)

slide=question_slide("Quick check: differentiate")
qtx(slide,"Differentiate y = 3x⁴ − 2x + 5.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"d(3x⁴)/dx = 12x³",XL,Y0,CW,LH,20)
fml(slide,"d(−2x)/dx = −2     and     d(5)/dx = 0",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"dy/dx = 12x³ − 2.",XL,Y0+2*LH+2*NG,CW,0.55)

# NEW TOPIC
slide=new_slide("New topic","Integration: reverse of differentiation",GREEN)
body(slide,"Differentiation finds a rate from a quantity.",XL,Y0,CW,0.50)
body(slide,"Integration finds the original quantity from a rate.",XL,Y0+0.65,CW,0.50)
ncert(slide,"If  dy/dx = f(x)  then  y = ∫ f(x) dx",XL,Y0+1.40,CW,1.05,22)
body(slide,"We read ∫ f(x) dx as “the integral of f with respect to x”.",XL,Y0+2.70,CW,0.55)

slide=new_slide("How it works","The constant of integration",TEAL)
body(slide,"Many functions share the same derivative, because a constant disappears when we differentiate.",XL,Y0,CW,0.70)
ncert(slide,"∫ f(x) dx = F(x) + C",XL,Y0+0.95,CW,1.00,24)
body(slide,"F(x) is any antiderivative of f.  C is an arbitrary constant.",XL,Y0+2.20,CW,0.55)
trap_card(slide,"Never drop +C in an indefinite integral.",XL,Y0+3.00,CW,0.55)

slide=new_slide("Exact graph","Same slope, different heights",TEAL)
body(slide,"y = x²,  y = x² + 3 and  y = x² − 2 all have slope 2x.",XL,Y0,5.4,0.70)
ncert(slide,"d(x² + C)/dx = 2x  for every C",XL,Y0+0.95,5.4,0.95,18)
sx,sy=graph_axes(slide,7.05,1.45,5.20,4.55,-2.5,2.5,-3,8,"x","y")
axis_ticks(slide,sx,sy,[-2,0,2],[-2,0,2,4,6])
plot_function(slide,lambda x:x*x,sx,sy,-2.2,2.2,-3,8,GOLD,120,2.4)
plot_function(slide,lambda x:x*x+3,sx,sy,-2.2,2.2,-3,8,TEAL,120,2.0)
plot_function(slide,lambda x:x*x-2,sx,sy,-2.2,2.2,-3,8,ORANGE,120,2.0)
tx(slide,"same shape, shifted up/down",7.20,6.20,4.90,0.28,13,MUTED,True,PP_ALIGN.CENTER,spa=0,scaled=False)

# POWER RULE
slide=new_slide("New topic","Power-rule reverse",GREEN)
body(slide,"Differentiation lowers the power by 1 and multiplies by the old power.",XL,Y0,CW,0.55)
body(slide,"Integration raises the power by 1 and divides by the new power.",XL,Y0+0.70,CW,0.55)
ncert(slide,"∫ x^n dx = x^(n+1)/(n+1) + C     (n ≠ −1)",XL,Y0+1.50,CW,1.10,22)
trap_card(slide,"This formula fails when n = −1.  That special case is ln|x| + C (later).",XL,Y0+2.90,CW,0.60)

slide=question_slide("Integrate a power")
qtx(slide,"Find ∫ x³ dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"Raise power: x³ → x⁴",XL,Y0,CW,LH,20)
fml(slide,"Divide by new power: x⁴ / 4",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"∫ x³ dx = x⁴/4 + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Integrate with a coefficient")
qtx(slide,"Find ∫ 5x² dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"Constants come outside: 5 ∫ x² dx",XL,Y0,CW,LH,20)
fml(slide,"∫ x² dx = x³/3",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"∫ 5x² dx = 5x³/3 + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Polynomial term by term")
qtx(slide,"Find ∫ (4x³ − 6x + 2) dx.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"Integrate each term separately",XL,Y0,CW,LH,19)
fml(slide,"∫ 4x³ dx = 4 · x⁴/4 = x⁴",XL,Y0+LH+NG,CW,LH,19)
fml(slide,"∫ −6x dx = −6 · x²/2 = −3x²",XL,Y0+2*LH+2*NG,CW,LH,19)
fml(slide,"∫ 2 dx = 2x",XL,Y0+3*LH+3*NG,CW,LH,19)
result_card(slide,"x⁴ − 3x² + 2x + C.",XL,Y0+4*LH+4*NG,CW,0.55)

# STANDARD TABLE
slide=new_slide("Copy","Standard indefinite integrals",TEAL)
grid_table(slide,["Function f(x)","∫ f(x) dx"],[
    ["x^n (n ≠ −1)","x^(n+1)/(n+1) + C"],
    ["1  (constant)","x + C"],
    ["sin x","−cos x + C"],
    ["cos x","sin x + C"],
    ["e^x","e^x + C"],
    ["1/x  (x ≠ 0)","ln|x| + C"],
],XL,Y0, [5.80,6.25], row_h=0.52, hsize=15, csize=16)

slide=new_slide("How it works","Check by differentiating",TEAL)
body(slide,"The safest check: differentiate your answer. You must recover the integrand.",XL,Y0,CW,0.60)
ncert(slide,"d/dx [ ∫ f(x) dx ] = f(x)",XL,Y0+0.90,CW,1.00,22)
body(slide,"If the derivative does not match f(x), the integral is wrong.",XL,Y0+2.20,CW,0.55)
trap_card(slide,"Always check +C disappears when you differentiate — that is correct.",XL,Y0+3.00,CW,0.55)

slide=question_slide("Check an antiderivative")
qtx(slide,"Is F(x) = x³/3 + 7 an antiderivative of f(x) = x²?",XL,Y0,CW,0.85,22)
slide=answer_slide()
fml(slide,"Differentiate F: dF/dx = x²",XL,Y0,CW,LH,20)
fml(slide,"This equals f(x) = x²",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"Yes.  (Any constant in place of 7 also works.)",XL,Y0+2*LH+2*NG,CW,0.58)

# MORE PRACTICE
slide=question_slide("Integrate sine")
qtx(slide,"Find ∫ sin x dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"From the table: antiderivative of sin x is −cos x",XL,Y0,CW,LH,19)
result_card(slide,"∫ sin x dx = −cos x + C.",XL,Y0+LH+NG,CW,0.55)

slide=question_slide("Integrate cosine")
qtx(slide,"Find ∫ cos x dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
result_card(slide,"∫ cos x dx = sin x + C.",XL,Y0,CW,0.55)

slide=question_slide("Integrate exponential")
qtx(slide,"Find ∫ e^x dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"e^x is its own derivative and its own antiderivative",XL,Y0,CW,LH,19)
result_card(slide,"∫ e^x dx = e^x + C.",XL,Y0+LH+NG,CW,0.55)

slide=question_slide("Mixed easy integral")
qtx(slide,"Find ∫ (2 cos x − 3 sin x) dx.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"∫ 2 cos x dx = 2 sin x",XL,Y0,CW,LH,20)
fml(slide,"∫ −3 sin x dx = −3(−cos x) = 3 cos x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"2 sin x + 3 cos x + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Simple reciprocal")
qtx(slide,"Find ∫ (1/x) dx  for x > 0.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"From the table: ∫ dx/x = ln|x| + C",XL,Y0,CW,LH,19)
fml(slide,"For x > 0, |x| = x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"ln x + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Integrate √x")
qtx(slide,"Find ∫ √x dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"Write √x = x^(1/2)",XL,Y0,CW,LH,20)
fml(slide,"∫ x^(1/2) dx = x^(3/2) / (3/2) = (2/3) x^(3/2)",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"(2/3) x^(3/2) + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Integrate 1/x²")
qtx(slide,"Find ∫ (1/x²) dx.",XL,Y0,CW,0.70,24)
slide=answer_slide()
fml(slide,"1/x² = x^(−2)",XL,Y0,CW,LH,20)
fml(slide,"∫ x^(−2) dx = x^(−1)/(−1) = −1/x",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"−1/x + C.",XL,Y0+2*LH+2*NG,CW,0.55)

# PHYSICS LINK
slide=new_slide("New topic","Physics: integrating motion",GREEN)
body(slide,"Differentiation went downhill:  x → v → a.",XL,Y0,CW,0.50)
body(slide,"Integration goes uphill:  a → v → x.",XL,Y0+0.65,CW,0.50)
ncert(slide,"v = ∫ a dt + C₁      ·      x = ∫ v dt + C₂",XL,Y0+1.40,CW,1.05,20)
body(slide,"Initial conditions fix the constants C₁ and C₂.",XL,Y0+2.70,CW,0.55)

slide=question_slide("From constant acceleration to velocity")
qtx(slide,"If a = 4 m s⁻² (constant) and v = 0 at t = 0, find v(t).",XL,Y0,CW,0.95,21)
slide=answer_slide()
fml(slide,"v = ∫ a dt = ∫ 4 dt = 4t + C",XL,Y0,CW,LH,19)
fml(slide,"At t = 0, v = 0 → 0 = 0 + C → C = 0",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"v = 4t  (m s⁻¹).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("From velocity to position")
qtx(slide,"If v = 4t and x = 0 at t = 0, find x(t).",XL,Y0,CW,0.85,22)
slide=answer_slide()
fml(slide,"x = ∫ v dt = ∫ 4t dt = 2t² + C",XL,Y0,CW,LH,19)
fml(slide,"At t = 0, x = 0 → C = 0",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"x = 2t²  (m).",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Constant velocity displacement")
qtx(slide,"If v = 6 m s⁻¹ (constant) and x = 2 m at t = 0, find x(t).",XL,Y0,CW,0.95,21)
slide=answer_slide()
fml(slide,"x = ∫ 6 dt = 6t + C",XL,Y0,CW,LH,20)
fml(slide,"t = 0, x = 2 → C = 2",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"x = 6t + 2  (m).",XL,Y0+2*LH+2*NG,CW,0.55)

# AREA IDEA (light)
slide=new_slide("Idea for later","Area under a curve",TEAL)
body(slide,"A definite integral will later mean the signed area between the curve and the x-axis.",XL,Y0,CW,0.70)
ncert(slide,"Area idea:  ∫_a^b f(x) dx   (limits a to b)",XL,Y0+0.95,CW,1.00,20)
body(slide,"Tonight we stay with indefinite integrals only (no limits yet).",XL,Y0+2.20,CW,0.55)
trap_card(slide,"Do not mix indefinite (+C) with definite (number, no +C).",XL,Y0+3.00,CW,0.55)

# FINAL PRACTICE
slide=question_slide("Homework style: power")
qtx(slide,"Find ∫ (3x⁵ − x) dx.",XL,Y0,CW,0.75,23)
slide=answer_slide()
fml(slide,"∫ 3x⁵ dx = 3 · x⁶/6 = x⁶/2",XL,Y0,CW,LH,19)
fml(slide,"∫ −x dx = −x²/2",XL,Y0+LH+NG,CW,LH,19)
result_card(slide,"(1/2)x⁶ − (1/2)x² + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Homework style: trig mix")
qtx(slide,"Find ∫ (sin x + cos x) dx.",XL,Y0,CW,0.80,22)
slide=answer_slide()
fml(slide,"∫ sin x dx = −cos x",XL,Y0,CW,LH,20)
fml(slide,"∫ cos x dx = sin x",XL,Y0+LH+NG,CW,LH,20)
result_card(slide,"−cos x + sin x + C.",XL,Y0+2*LH+2*NG,CW,0.55)

slide=question_slide("Homework style: check")
qtx(slide,"Differentiate y = −cos x + sin x + 4 and confirm it matches sin x + cos x.",XL,Y0,CW,1.00,20)
slide=answer_slide()
fml(slide,"d(−cos x)/dx = sin x",XL,Y0,CW,LH,20)
fml(slide,"d(sin x)/dx = cos x     ·     d(4)/dx = 0",XL,Y0+LH+NG,CW,LH,18)
result_card(slide,"dy/dx = sin x + cos x.  Check successful.",XL,Y0+2*LH+2*NG,CW,0.58)

# SUMMARY
slide=new_slide("Class summary","Integration lock lines",TEAL)
fml(slide,"Integration reverses differentiation",XL,Y0,CW,LH,19)
fml(slide,"∫ f(x) dx = F(x) + C",XL,Y0+LH+NG,CW,LH,20)
fml(slide,"∫ x^n dx = x^(n+1)/(n+1) + C   (n ≠ −1)",XL,Y0+2*LH+2*NG,CW,LH,18)
fml(slide,"Check every answer by differentiating",XL,Y0+3*LH+3*NG,CW,LH,19)

slide=new_slide("After class","Practice",AMBER)
body(slide,"1. Memorise the standard integrals table from tonight.",XL,Y0,CW,0.55)
body(slide,"2. Redo every ∫ question without looking at the solution.",XL,Y0+0.72,CW,0.55)
body(slide,"3. For each answer, differentiate and recover the integrand.",XL,Y0+1.44,CW,0.55)
body(slide,"4. Write one motion example: given a, find v and x with +C fixed by data.",XL,Y0+2.16,CW,0.60)
trap_card(slide,"Missing +C is the most common indefinite-integral mistake.",XL,Y0+3.05,CW,0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
