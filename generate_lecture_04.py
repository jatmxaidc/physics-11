"""
Lecture 04 — Leftover L03, then ASTC, max–min, phase (2-hour class)
Hour 1: L03 leftover — 0–360 table, radian/arc, homework 21/23/25/33/35/37/49,
prefixes, compound, then pending slides 71–82 (small angle).
Hour 2: PW class-03 — ASTC method, max–min, phase / phasor.
Class 11 Physics  •  CBSE + NEET Foundation
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "Lecture_04_ASTC_MaxMin_Phase_Phasor.pptx"

NAVY = RGBColor(8, 30, 63)
BLUE = RGBColor(64, 156, 255)
TEAL = RGBColor(40, 180, 190)
GREEN = RGBColor(90, 214, 128)
AMBER = RGBColor(255, 204, 72)
ORANGE = RGBColor(255, 158, 64)
RED = RGBColor(255, 96, 96)
TEXT = RGBColor(255, 255, 255)
MUTED = RGBColor(176, 182, 192)
STROKE = RGBColor(72, 78, 90)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(22, 24, 28)
BLACK = RGBColor(0, 0, 0)
GOLD = RGBColor(255, 214, 70)
NCERT_FILL = RGBColor(255, 248, 225)
NCERT_LINE = RGBColor(218, 168, 64)
QCOL = ORANGE
WORK = RGBColor(90, 200, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

XL, CW, Y0, YF = 0.65, 12.05, 1.18, 7.12
LH = 0.32
NG = 0.18
PT_SCALE = 0.80


def _bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _solid(slide, x, y, w, h, fill, no_line=True):
    s = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if no_line:
        s.line.fill.background()
    else:
        s.line.color.rgb = fill
    return s


def _card_bg(slide, x, y, w, h, fill, border, line_pt=1.0, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(line_pt)
    return s


def tx(
    slide, text, x, y, w, h,
    size=15, color=TEXT, bold=False,
    align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
    spa=1, lsp=1.1, font="Aptos", scaled=True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(spa)
    p.line_spacing = lsp
    r = p.add_run()
    r.text = text
    r.font.name = font
    pt = size if not scaled else max(7, round(size * PT_SCALE, 1))
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _bar(slide, label, color=BLUE):
    _solid(slide, 0, 0, 13.333, 0.29, color)
    if label:
        tx(slide, label.upper(), XL, 0.05, 9, 0.20, 8, WHITE, True, spa=0)


def _title(slide, text):
    tx(slide, text, XL, 0.34, CW, 0.72, 26, GOLD, True, spa=0)


def _sep(slide, y=1.10):
    _solid(slide, XL, y, CW, 0.018, STROKE)


def _footer(slide, n):
    tx(slide, f"{n}", 12.25, YF, 0.8, 0.25, 11, MUTED, True, PP_ALIGN.RIGHT, spa=0)


_CONTENT = 0


def add_rough_work():
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, "Rough work", MUTED)
    tx(slide, "Rough work  /  extra work", XL, 0.34, CW, 0.50, 22, GOLD, True, spa=0)
    _sep(slide)
    _footer(slide, len(prs.slides))
    add_notes(slide, "Empty board for extra working. Write here if a question needs more space.")
    return slide


def after_content_slide():
    global _CONTENT
    _CONTENT += 1
    if _CONTENT % 5 == 0:
        add_rough_work()


def new_slide(section, title, bar_color=BLUE):
    slide = prs.slides.add_slide(BLANK)
    _bg(slide)
    _bar(slide, section, bar_color)
    _title(slide, title)
    _sep(slide)
    _footer(slide, len(prs.slides))
    after_content_slide()
    return slide


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def ncert(slide, formula, x, y, w, h, size=20):
    s = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = NCERT_FILL
    s.line.color.rgb = NCERT_LINE
    s.line.width = Pt(1.5)
    tx(
        slide, formula, x + 0.18, y + 0.16, w - 0.36, h - 0.32,
        size, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, lsp=1.15,
    )


def fml(slide, text, x, y, w, h=0.44, size=22):
    return tx(slide, text, x, y, w, h, size, WORK, True, spa=0)


def body(slide, text, x, y, w, h=0.40, size=16):
    return tx(slide, text, x, y, w, h, size, TEXT, False, spa=0, scaled=False)


def qtx(slide, text, x, y, w, h=0.44, size=18):
    return tx(slide, text, x, y, w, h, size, QCOL, True, spa=0, scaled=False)


def choices(slide, items, x, y, w=None):
    if w is None:
        w = CW
    row = 0.40
    for i, item in enumerate(items):
        cy = y + i * row
        circ = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.04), Inches(0.30), Inches(0.30)
        )
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tx(slide, str(i + 1), x, cy + 0.06, 0.30, 0.26, 14, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
        tx(slide, item, x + 0.42, cy + 0.06, w - 0.42, 0.28, 18, TEXT, False, spa=0, scaled=False)
    return y + len(items) * row


def trap_card(slide, text, x, y, w, h):
    tx(slide, text, x, y, w, h, 16, ORANGE, False, spa=0, scaled=False)


def result_card(slide, text, x, y, w, h, tsize=14):
    return tx(slide, text, x, y, w, 0.36, max(tsize, 18), GREEN, True, spa=0)


def warn_row(slide, wrong, correct, x, y, w):
    tx(slide, f"Wrong:  {wrong}", x, y, w, 0.30, 16, RED, False, spa=0, scaled=False)
    tx(slide, f"Right:  {correct}", x, y + LH, w, 0.30, 16, GREEN, False, spa=0, scaled=False)
    return y + LH + LH


def grid_table(slide, headers, rows, x, y, col_w, row_h=0.42, hsize=11, csize=12, scaled=True):
    col_x = [x]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    for j, ht in enumerate(headers):
        _solid(slide, col_x[j], y, col_w[j], row_h, RGBColor(28, 36, 52), False)
        tx(
            slide, ht, col_x[j] + 0.04, y, col_w[j] - 0.08, row_h,
            hsize, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spa=0, scaled=scaled,
        )
    for i, row in enumerate(rows):
        ry = y + row_h + i * row_h
        fill = RGBColor(16, 16, 20) if i % 2 == 0 else SOFT
        for j, val in enumerate(row):
            s = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(col_x[j]), Inches(ry), Inches(col_w[j]), Inches(row_h)
            )
            s.fill.solid()
            s.fill.fore_color.rgb = fill
            s.line.color.rgb = STROKE
            s.line.width = Pt(0.5)
            tx(
                slide, val, col_x[j] + 0.04, ry, col_w[j] - 0.08, row_h,
                csize, WORK if j == 0 else TEXT, j == 0, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
                spa=0, scaled=scaled,
            )


def tx_3d(slide, text, x, y, w, h, size, face, edge, font="Arial Black"):
    tx(slide, text, x + 0.045, y + 0.045, w, h, size, edge, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)
    tx(slide, text, x, y, w, h, size, face, True, PP_ALIGN.CENTER, spa=0, scaled=False, font=font)


def draw_line(slide, x1, y1, x2, y2, color=WHITE, pt=1.75):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def phasor_axes(slide, x, y, s, xlab="sin (ωt)", ylab="cos (ωt)"):
    cx, cy = x + s, y + s
    draw_line(slide, x, cy, x + 2 * s, cy, MUTED, 1.75)
    draw_line(slide, cx, y, cx, y + 2 * s, MUTED, 1.75)
    tx(slide, xlab, x + 2 * s - 1.15, cy + 0.06, 1.20, 0.22, 12, MUTED, True, spa=0)
    tx(slide, ylab, cx + 0.08, y - 0.02, 1.20, 0.22, 12, MUTED, True, spa=0)
    return cx, cy


# ── 01 COVER ──────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx_3d(slide, "EXCELLENCE", XL, 0.55, CW, 1.05, 62, GOLD, RGBColor(210, 210, 220))
tx_3d(slide, "CLASS 11   ·   2026–27", XL, 1.62, CW, 0.55, 30, WHITE, RGBColor(90, 90, 100), font="Aptos")
_solid(slide, 4.15, 2.28, 5.00, 0.028, WHITE)

bx, by, bw = 2.05, 2.70, 9.20
_card_bg(slide, bx, by, bw, 2.35, WHITE, WHITE, 0)
_card_bg(slide, bx, by + 0.72, bw, 1.63, BLACK, RGBColor(220, 220, 220), 1.25)
_solid(slide, bx + bw / 2, by + 0.16, 0.018, 0.42, RGBColor(40, 40, 40))
tx(slide, "Physics", bx, by + 0.16, bw / 2 - 0.10, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Lecture - 04", bx + bw / 2, by + 0.16, bw / 2, 0.44, 25, BLACK, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "Basic Maths and Calculus", bx, by + 1.05, bw, 0.55, 32, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False, font="Georgia")
tx(slide, "Leftover last class, then ASTC, max–min, phase", bx, by + 1.62, bw, 0.40, 20, GOLD, True, PP_ALIGN.CENTER, spa=0, scaled=False)

tx(slide, "By — Excellence Programme", XL, 5.30, CW, 0.42, 25, WHITE, True, PP_ALIGN.CENTER, spa=0, scaled=False)
tx(slide, "4 September 2026", XL, 5.80, CW, 0.32, 18, MUTED, True, PP_ALIGN.CENTER, spa=0, scaled=False)
_footer(slide, 1)
add_notes(slide, "Cover. Hour 1: leftover L03 tables, HW answers, prefixes, compound, small angle. Hour 2: ASTC method, max-min, phasor.")
after_content_slide()

# ── 02 TOPICS ─────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
tx(slide, "Topics to be covered", XL, 0.38, CW, 0.55, 32, GOLD, True, spa=0, scaled=False)
_solid(slide, XL, 0.98, 3.40, 0.035, GOLD)
topics = [
    "Full table  0° to 360°  for all six ratios, then arc and radian ↔ degree",
    "Last-class homework  —  slides 21, 23, 25, 33, 35, 37, 49  —  answers in simple lines",
    "Length prefixes with full names, then compound  sin / cos / tan",
    "Pending last class, slides 71–82  —  small angle, then allied homework",
    "ASTC method  —  sign first, then change the function only at 90° and 270°",
    "Maximum and minimum, then phase / phasor",
]
for i, line in enumerate(topics):
    y = 1.18 + i * 0.78
    tx(slide, f"{i + 1:02d}", XL, y, 0.70, 0.50, 23, GOLD, True, spa=0, scaled=False)
    tx(slide, line, XL + 0.90, y, CW - 0.90, 0.70, 18, WHITE, False, spa=0, scaled=False)
_footer(slide, 2)
add_notes(slide, "Slide 2. Hour 1 leftover L03. Hour 2 new ASTC method, max-min, phasor.")
after_content_slide()

# ── 0–360 TABLE ───────────────────────────────────────────────────────────────
TW = [1.70] + [2.07] * 5
RH = 0.62
slide = new_slide("Table", "All six ratios from 0° to 90°", TEAL)
add_notes(slide, "Copy. Reciprocal of 0 is not defined. tan 90 and sec 90 undefined.")
grid_table(
    slide,
    ["θ", "0°", "30°", "45°", "60°", "90°"],
    [
        ["sin", "0", "1/2", "1/√2", "√3/2", "1"],
        ["cos", "1", "√3/2", "1/√2", "1/2", "0"],
        ["tan", "0", "1/√3", "1", "√3", "undef."],
        ["cot", "undef.", "√3", "1", "1/√3", "0"],
        ["sec", "1", "2/√3", "√2", "2", "undef."],
        ["cosec", "undef.", "2", "√2", "2/√3", "1"],
    ],
    XL, Y0, TW, row_h=RH, hsize=16, csize=16, scaled=False,
)
trap_card(slide, "tan 90° and sec 90° are not defined. cosec 0° and cot 0° are not defined.", XL, Y0 + 7 * RH + 0.10, CW, 0.32)

slide = new_slide("Table", "All six ratios from 90° to 180°", TEAL)
add_notes(slide, "Copy. Second quadrant: only sin and cosec positive.")
grid_table(
    slide,
    ["θ", "90°", "120°", "135°", "150°", "180°"],
    [
        ["sin", "1", "√3/2", "1/√2", "1/2", "0"],
        ["cos", "0", "−1/2", "−1/√2", "−√3/2", "−1"],
        ["tan", "undef.", "−√3", "−1", "−1/√3", "0"],
        ["cot", "0", "−1/√3", "−1", "−√3", "undef."],
        ["sec", "undef.", "−2", "−√2", "−2/√3", "−1"],
        ["cosec", "1", "2/√3", "√2", "2", "undef."],
    ],
    XL, Y0, TW, row_h=RH, hsize=16, csize=16, scaled=False,
)
trap_card(slide, "Second quadrant: only sin and cosec are positive. cot 180° and cosec 180° are not defined.", XL, Y0 + 7 * RH + 0.10, CW, 0.32)

slide = new_slide("Table", "All six ratios from 180° to 270°", TEAL)
add_notes(slide, "Copy. Third quadrant: only tan and cot positive.")
grid_table(
    slide,
    ["θ", "180°", "210°", "225°", "240°", "270°"],
    [
        ["sin", "0", "−1/2", "−1/√2", "−√3/2", "−1"],
        ["cos", "−1", "−√3/2", "−1/√2", "−1/2", "0"],
        ["tan", "0", "1/√3", "1", "√3", "undef."],
        ["cot", "undef.", "√3", "1", "1/√3", "0"],
        ["sec", "−1", "−2/√3", "−√2", "−2", "undef."],
        ["cosec", "undef.", "−2", "−√2", "−2/√3", "−1"],
    ],
    XL, Y0, TW, row_h=RH, hsize=16, csize=16, scaled=False,
)
trap_card(slide, "Third quadrant: only tan and cot are positive. tan 270° and sec 270° are not defined.", XL, Y0 + 7 * RH + 0.10, CW, 0.32)

slide = new_slide("Table", "All six ratios from 270° to 360°", TEAL)
add_notes(slide, "Copy. Fourth quadrant: only cos and sec positive. 360° is the same as 0°.")
grid_table(
    slide,
    ["θ", "270°", "300°", "315°", "330°", "360°"],
    [
        ["sin", "−1", "−√3/2", "−1/√2", "−1/2", "0"],
        ["cos", "0", "1/2", "1/√2", "√3/2", "1"],
        ["tan", "undef.", "−√3", "−1", "−1/√3", "0"],
        ["cot", "0", "−1/√3", "−1", "−√3", "undef."],
        ["sec", "undef.", "2", "√2", "2/√3", "1"],
        ["cosec", "−1", "−2/√3", "−√2", "−2", "undef."],
    ],
    XL, Y0, TW, row_h=RH, hsize=16, csize=16, scaled=False,
)
trap_card(slide, "Fourth quadrant: only cos and sec are positive. 360° repeats 0°. Never write cosec 360°.", XL, Y0 + 7 * RH + 0.10, CW, 0.32)

# ── RADIAN AND ARC ────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Arc length — θ must be in radian", TEAL)
add_notes(slide, "Read. Arc = Rθ. Never put 90° into this formula.")
fml(slide, "Arc length  =  R θ", XL, Y0, 7.0, LH, 26)
ncert(slide, "θ in radian", XL + 7.15, Y0, 4.90, 0.90, 20)
body(slide, "Practical unit of angle: degree.     SI unit of angle: radian.", XL, Y0 + LH + NG + 0.40, CW)
body(slide, "+θ anti-clockwise.     −θ clockwise.", XL, Y0 + 2 * LH + NG + 0.40, CW)
trap_card(slide, "If the object turns 90° on a circle of radius R, first change 90° to π/2. Distance = πR/2. Never R × 90°.", XL, Y0 + 3 * LH + 2 * NG + 0.40, CW, 0.50)

slide = new_slide("Copy", "The one line:  180° = π rad", TEAL)
add_notes(slide, "They say it. Wherever you see π rad, put 180°.")
body(slide, "Lock this. Every conversion starts here.", XL, Y0, CW)
fml(slide, "180°  =  π rad", XL, Y0 + LH + NG, 7.0, LH, 28)
ncert(slide, "Wherever you see π rad, put 180°.", XL + 7.15, Y0 + LH + NG, 4.90, 1.10, 18)
fml(slide, "degree → radian:     multiply by  π/180", XL, Y0 + 2 * LH + NG + 0.80, CW, LH, 20)
fml(slide, "radian → degree:     multiply by  180/π", XL, Y0 + 3 * LH + NG + 0.80, CW, LH, 20)
trap_card(slide, "Never write π = 180°. Write π rad = 180°.  π with no degree mark is already radian.", XL, Y0 + 4 * LH + 2 * NG + 0.80, CW, 0.40)

# ── L03 HW 21, 23, 25 ─────────────────────────────────────────────────────────
slide = new_slide("Board", "Homework slide 21 — six conversions", GREEN)
add_notes(slide, "Read slowly. π rad = 180. Half of that is 90. Quarter is 45.")
qtx(slide, "Last class, slide 21. Convert each. Easy rule:  π rad  =  180°.", XL, Y0, CW, 0.40, 16)
body(slide, "If you see π, put 180°. If you see a degree, multiply by π/180.", XL, Y0 + 0.45, CW)
items = [
    "π/2 rad  =  180°/2  =  90°",
    "π/4 rad  =  180°/4  =  45°",
    "π/6 rad  =  180°/6  =  30°",
    "45°  =  45 × π/180  =  π/4 rad",
    "90°  =  90 × π/180  =  π/2 rad",
    "60°  =  60 × π/180  =  π/3 rad",
]
for i, item in enumerate(items):
    r, c = divmod(i, 2)
    result_card(slide, item, XL + c * 6.1, Y0 + 0.90 + r * LH, 5.85, LH, 16)

slide = new_slide("Board", "Homework slide 23 — radian into degree", GREEN)
add_notes(slide, "Read. Number times 180, divide by the bottom.")
qtx(slide, "Last class, slide 23.  Example:  5π/6  means  5 × 180° / 6.", XL, Y0, CW, 0.36, 16)
rad_a = [
    "5π/6  =  150°", "2π/3  =  120°", "4π/3  =  240°",
    "3π/5  =  108°", "5π  =  900°", "2π/5  =  72°",
    "7π/12  =  105°", "5π/2  =  450°", "3π  =  540°",
]
for i, item in enumerate(rad_a):
    r, c = divmod(i, 3)
    result_card(slide, item, XL + c * 4.0, Y0 + 0.48 + r * LH, 3.85, LH, 16)
body(slide, "How: cancel.  5π/6  →  5 × 180 / 6  →  5 × 30  =  150.  Same for every line.", XL, Y0 + 0.48 + 3 * LH + NG, CW)

slide = new_slide("Board", "Homework slide 25 — degree into radian", GREEN)
add_notes(slide, "Read. Multiply by π/180. π alone is already radian.")
qtx(slide, "Last class, slide 25.  Multiply the degree by π/180.  Cancel 180 if you can.", XL, Y0, CW, 0.40, 16)
deg_a = [
    "120°  =  2π/3 rad", "30°  =  π/6 rad", "150°  =  5π/6 rad",
    "270°  =  3π/2 rad", "135°  =  3π/4 rad", "π  is already radian",
]
for i, item in enumerate(deg_a):
    r, c = divmod(i, 3)
    result_card(slide, item, XL + c * 4.0, Y0 + 0.50 + r * LH, 3.85, LH, 16)
warn_row(slide, "π°  treated as π rad", "The board wrote π, not π°.  π is already radian.  π° would be π²/180 rad.", XL, Y0 + 0.50 + 2 * LH + NG, CW)

# ── PREFIXES ──────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Tiny lengths you will meet in physics", TEAL)
add_notes(slide, "They say each full name. fm nucleus. pm atom spacing. nm light. μm cell. Å atom.")
body(slide, "A prefix is a short name stuck in front of metre. It tells you how many metres.", XL, Y0, CW)
grid_table(
    slide,
    ["short", "full name", "in metres", "where you meet it"],
    [
        ["fm", "femtometre", "10⁻¹⁵ m", "nucleus of an atom"],
        ["pm", "picometre", "10⁻¹² m", "spacing inside an atom"],
        ["Å", "ångström", "10⁻¹⁰ m", "atom size, X-ray wavelength"],
        ["nm", "nanometre", "10⁻⁹ m", "visible light, DNA, thin films"],
        ["μm", "micrometre", "10⁻⁶ m", "cell, dust, width of a hair"],
    ],
    XL, Y0 + 0.45, [1.50, 2.70, 2.40, 5.45], 0.52, 15, 16, scaled=False,
)
trap_card(slide, "Å is not SI, but physics still uses it for atoms. Nuclear size is fm, never mm.", XL, Y0 + 0.45 + 6 * 0.52 + 0.12, CW, 0.36)

slide = new_slide("Copy", "Everyday and large lengths", TEAL)
add_notes(slide, "They say milli, centi, metre, kilo, mega. Earth radius is about 6.4 Mm.")
grid_table(
    slide,
    ["short", "full name", "in metres", "where you meet it"],
    [
        ["mm", "millimetre", "10⁻³ m", "small lab length, thickness"],
        ["cm", "centimetre", "10⁻² m", "ruler, height in class"],
        ["m", "metre", "1 m", "SI unit of length"],
        ["km", "kilometre", "10³ m", "road, tower, walking distance"],
        ["Mm", "megametre", "10⁶ m", "Earth radius ≈ 6.4 Mm"],
    ],
    XL, Y0, [1.50, 2.70, 2.40, 5.45], 0.52, 15, 16, scaled=False,
)
ncert(slide, "1 m  =  100 cm          1 km  =  1000 m          1 Å  =  10⁻¹⁰ m          1 fm  =  10⁻¹⁵ m", XL, Y0 + 6 * 0.52 + 0.18, CW, 0.90, 16)

slide = new_slide("Copy", "How to read a prefix in one line", TEAL)
add_notes(slide, "milli is thousandth. kilo is thousand. micro is millionth. nano is billionth. femto is a thousand-trillionth.")
fml(slide, "milli  =  one thousandth          kilo  =  one thousand", XL, Y0, CW, LH, 18)
fml(slide, "micro  =  one millionth          nano  =  one billionth", XL, Y0 + LH, CW, LH, 18)
fml(slide, "pico  =  one trillionth          femto  =  one thousand-trillionth", XL, Y0 + 2 * LH, CW, LH, 18)
body(slide, "1 mm is 1 m cut into 1000 equal pieces.  1 km is 1000 metres joined.", XL, Y0 + 3 * LH + NG, CW)
body(slide, "1 μm is 1 m cut into 1 000 000 pieces.  A red-blood-cell is a few μm.", XL, Y0 + 4 * LH + NG, CW)
trap_card(slide, "Never mix them: a nucleus is a few fm. An atom is about 1 Å. Light is a few hundred nm.", XL, Y0 + 5 * LH + 2 * NG, CW, 0.40)

# ── COMPOUND ──────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Compound sine — two angles, one sine", TEAL)
add_notes(slide, "Copy. Plus for sum. Minus for difference. Say it in words.")
body(slide, "A and B are two angles. Compound means we open the bracket into products.", XL, Y0, CW)
fml(slide, "sin (A + B)  =  sin A cos B  +  cos A sin B", XL, Y0 + LH + NG, CW, LH, 22)
body(slide, "In words: sine of a sum  =  sine-cosine  plus  cosine-sine.", XL, Y0 + 2 * LH + NG, CW)
fml(slide, "sin (A − B)  =  sin A cos B  −  cos A sin B", XL, Y0 + 3 * LH + 2 * NG, CW, LH, 22)
body(slide, "Same line, only the middle sign flips. Difference takes a minus.", XL, Y0 + 4 * LH + 2 * NG, CW)
trap_card(slide, "Do not write sin(A+B) = sin A + sin B. That is the classic trap.", XL, Y0 + 5 * LH + 3 * NG, CW, 0.40)

slide = new_slide("Copy", "Compound cosine — watch the minus", TEAL)
add_notes(slide, "Copy. Cosine of a sum has a minus. Cosine of a difference has a plus.")
fml(slide, "cos (A + B)  =  cos A cos B  −  sin A sin B", XL, Y0, CW, LH, 22)
body(slide, "In words: cosine of a sum  =  cosine-cosine  minus  sine-sine.", XL, Y0 + LH, CW)
fml(slide, "cos (A − B)  =  cos A cos B  +  sin A sin B", XL, Y0 + 2 * LH + NG, CW, LH, 22)
body(slide, "Difference of cosine takes a plus. Sum of cosine takes a minus. Opposite of sine.", XL, Y0 + 3 * LH + NG, CW)
trap_card(slide, "Swap those two signs and every 15°, 75°, 105° value will come out wrong.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Copy", "Compound tangent — one fraction", TEAL)
add_notes(slide, "Copy. Minus in the denominator for the sum.")
fml(slide, "tan (A + B)  =  (tan A + tan B) / (1 − tan A tan B)", XL, Y0, CW, LH, 20)
body(slide, "Top: add the tans.  Bottom: 1 minus the product.  Sum has a minus downstairs.", XL, Y0 + LH, CW)
fml(slide, "tan (A − B)  =  (tan A − tan B) / (1 + tan A tan B)", XL, Y0 + 2 * LH + NG, CW, LH, 20)
body(slide, "Difference: minus on the top, plus downstairs.", XL, Y0 + 3 * LH + NG, CW)
ncert(slide, "If tan A tan B = 1, then tan(A + B) is not defined.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.85, 18)

slide = new_slide("Copy", "When A and B are the same — double angle", TEAL)
add_notes(slide, "Copy. Put B = A in the compound lines. sin2θ = 2sin cos. cos2θ = 1 − 2sin².")
body(slide, "Put B = A in the compound lines. Two of the same angle become 2A.", XL, Y0, CW)
fml(slide, "sin 2A  =  2 sin A cos A", XL, Y0 + LH + NG, CW, LH, 22)
fml(slide, "cos 2A  =  cos² A  −  sin² A  =  1  −  2 sin² A  =  2 cos² A  −  1", XL, Y0 + 2 * LH + NG, CW, LH, 18)
fml(slide, "tan 2A  =  2 tan A / (1 − tan² A)", XL, Y0 + 3 * LH + 2 * NG, CW, LH, 20)
trap_card(slide, "sin 2θ is not 2 sin θ. The 2 sits with the angle, not in front of the sine.", XL, Y0 + 4 * LH + 3 * NG, CW, 0.40)

# ── L03 HW 33, 35, 37 ─────────────────────────────────────────────────────────
slide = new_slide("Board", "Homework slide 33 — shadow of a pole", GREEN)
add_notes(slide, "Read. tan = opposite/adjacent = h/l. So l = h/tanθ. Option 2.")
qtx(slide, "Last class, slide 33. Pole height h. Shadow l. Sun at angle θ above the ground. Find l.", XL, Y0, CW, 0.50, 15)
choices(slide, ["l  =  h tan θ", "l  =  h / tan θ", "l  =  h sin θ", "l  =  h / sin θ"], XL, Y0 + 0.55)
body(slide, "Picture: height is the wall, shadow is the floor.  tan θ  =  wall / floor  =  h / l.", XL, Y0 + 2.25, CW)
fml(slide, "So  l  =  h / tan θ     (same as  h cot θ)", XL, Y0 + 2.25 + LH, CW, LH, 18)
result_card(slide, "Option 2     l  =  h / tan θ", XL, Y0 + 2.25 + 2 * LH + NG, CW, 0.36, 18)
trap_card(slide, "Option 1 swaps height and shadow. tan is not shadow/height.", XL, Y0 + 2.25 + 3 * LH + 2 * NG, CW, 0.36)

slide = new_slide("Board", "Homework slide 35 — wire falls", GREEN)
add_notes(slide, "Read. Wire is hypotenuse. After fall, extra bit = wire − base. None of four options.")
qtx(slide, "Last class, slide 35. Wire from the top at angle θ breaks and falls. Find x from the base.", XL, Y0, CW, 0.45, 15)
choices(slide, ["x  =  h cot θ", "x  =  h tan θ", "x  =  h sin θ", "x  =  h sec θ"], XL, Y0 + 0.50)
body(slide, "Wire is the long side  =  h / sin θ.  Original base  =  h / tan θ.  After the fall, x is wire minus base.", XL, Y0 + 2.20, CW)
fml(slide, "x  =  h / sin θ  −  h / tan θ  =  h (cosec θ  −  cot θ)", XL, Y0 + 2.20 + LH, CW, LH, 16)
result_card(slide, "x  =  h (cosec θ  −  cot θ)      —  none of the four options", XL, Y0 + 2.20 + 2 * LH + NG, CW, 0.36, 16)
trap_card(slide, "Option 1 is only the original base, before the wire falls.", XL, Y0 + 2.20 + 3 * LH + 2 * NG, CW, 0.36)

slide = new_slide("Board", "Homework slide 37 — shadow on a slope", GREEN)
add_notes(slide, "Read. a=l/2, b=l√3/2. tan60=(10+a)/b. l=10 m. Option 2.")
qtx(slide, "Last class, slide 37. Pole 10 m. Slope 30°. Sun 60°. Shadow length l on the slope?", XL, Y0, CW, 0.45, 15)
choices(slide, ["5 m", "10 m", "10√3 m", "10 / (√3 − 1) m"], XL, Y0 + 0.50)
body(slide, "Rise of the slope over length l is l/2.  Horizontal run is l√3/2.  tan 60° looks at (10 + rise) over run.", XL, Y0 + 2.20, CW)
fml(slide, "√3  =  (10 + l/2) / (l √3 / 2)          3l/2  =  10 + l/2          l  =  10", XL, Y0 + 2.20 + LH, CW, LH, 16)
result_card(slide, "Option 2     l  =  10 m", XL, Y0 + 2.20 + 2 * LH + NG, CW, 0.36, 18)

# ── L03 HW 49 ─────────────────────────────────────────────────────────────────
slide = new_slide("Board", "Homework slide 49 — four values, easy route", GREEN)
add_notes(slide, "Read. Brother angle. 120=180-60. 150=180-30. 240=180+60.")
qtx(slide, "Last class, slide 49. Find  sin 120°,  cos 150°,  sin 240°,  cos 120°.", XL, Y0, CW, 0.40, 16)
body(slide, "Easy way: write each as 180° with a small angle. ASTC then gives the sign.", XL, Y0 + 0.45, CW)
fml(slide, "(1)  sin 120°  =  sin (180° − 60°)  =  + sin 60°  =  √3/2", XL, Y0 + 0.90, CW, LH, 17)
fml(slide, "(2)  cos 150°  =  cos (180° − 30°)  =  − cos 30°  =  −√3/2", XL, Y0 + 0.90 + LH, CW, LH, 17)
fml(slide, "(3)  sin 240°  =  sin (180° + 60°)  =  − sin 60°  =  −√3/2", XL, Y0 + 0.90 + 2 * LH, CW, LH, 17)
fml(slide, "(4)  cos 120°  =  cos (180° − 60°)  =  − cos 60°  =  −1/2", XL, Y0 + 0.90 + 3 * LH, CW, LH, 17)
result_card(slide, "√3/2          −√3/2          −√3/2          −1/2", XL, Y0 + 0.90 + 4 * LH + NG, CW, 0.36, 18)

# ── PENDING 71–82 SMALL ANGLE ─────────────────────────────────────────────────
slide = new_slide("Leftover", "Last class stopped at slide 71 — small angle now", AMBER)
add_notes(slide, "Do not skip. Class notes stop before small angle. Teach 71–82 here.")
body(slide, "Last class ended before small angle. Slides 71 to 82 of that deck are today’s leftover.", XL, Y0, CW)
fml(slide, "small-angle rule  →  sin 2°, tan 3°, cos 3°  →  tan 2° trap  →  allied homework", XL, Y0 + LH + NG, CW, LH, 18)
body(slide, "Then we start the new work: ASTC method, max–min, phase.", XL, Y0 + 2 * LH + 2 * NG, CW)

slide = new_slide("Small angle", "When θ is small, the arc looks like the opposite side", TEAL)
add_notes(slide, "θ ≤ 5°. sinθ≈θ, tanθ≈θ in radian. cosθ≈1.")
body(slide, "On a circle of radius R, a tiny angle makes opposite ≈ arc = Rθ, and adjacent ≈ R.", XL, Y0, CW)
fml(slide, "sin θ  ≈  θ          tan θ  ≈  θ          cos θ  ≈  1", XL, Y0 + LH + NG, 7.0, LH, 20)
ncert(slide, "θ in radian.   Use only if θ ≤ 5°.", XL + 7.15, Y0 + LH + NG, 4.90, 1.00, 16)
body(slide, "Easy words: change the degree to radian first, then drop the word radian. That number is the answer.", XL, Y0 + 2 * LH + 2 * NG, CW)
trap_card(slide, "sin 2° is not 2. 2° in radian is 2π/180 = π/90. That is the value.", XL, Y0 + 3 * LH + 3 * NG, CW, 0.45)

slide = new_slide("Board", "Pending slide 72 — sin 2°, tan 3°, cos 3°", GREEN)
add_notes(slide, "Read. π/90, π/60, 1. Convert first.")
qtx(slide, "Last class, slide 72. Using small-angle, find  sin 2°,  tan 3°,  cos 3°.", XL, Y0, CW, 0.40, 16)
body(slide, "Step 1: the angle is small, so sine and tan become the radian number. Cosine becomes 1.", XL, Y0 + 0.48, CW)
fml(slide, "(i)    sin 2°  ≈  2 × (π/180)  =  π/90", XL, Y0 + 0.95, CW, LH, 18)
fml(slide, "(ii)   tan 3°  ≈  3 × (π/180)  =  π/60", XL, Y0 + 0.95 + LH, CW, LH, 18)
fml(slide, "(iii)  cos 3°  ≈  1", XL, Y0 + 0.95 + 2 * LH, CW, LH, 18)
result_card(slide, "sin 2°  ≈  π/90          tan 3°  ≈  π/60          cos 3°  ≈  1", XL, Y0 + 0.95 + 3 * LH + NG, CW, 0.36, 18)

slide = new_slide("Board", "Pending slide 74 — which tan 2° is correct?", GREEN)
add_notes(slide, "Read. Option 2. Not 2°, not 2 rad.")
qtx(slide, "Last class, slide 74. Which is correct for  tan 2°  under small-angle?", XL, Y0, CW, 0.40, 16)
choices(slide, ["2°", "π / 90", "π / 30", "2 rad"], XL, Y0 + 0.48)
y = warn_row(slide, "tan 2°  =  2°", "tan 2°  ≈  2 × (π/180)  =  π/90", XL, Y0 + 2.18, CW)
result_card(slide, "Option 2     tan 2°  ≈  π/90", XL, y + NG, CW, 0.36, 20)
trap_card(slide, "Option 1 leaves the degree. Option 4 pretends 2° is already 2 radian. Both are wrong.", XL, y + NG + 0.42, CW, 0.40)

# ── LAST-CLASS HOMEWORK (slides 76–77) ────────────────────────────────────────
slide = new_slide("Board", "Pending slides 76–77 — two allied lines", GREEN)
add_notes(slide, "Read. (13) sin θ. (14) − cos θ. This was tonight of last class.")
qtx(slide, "Last class, slides 76–77. Simplify.   (13)  sin (180° − θ)     (14)  cos (180° + θ)", XL, Y0, CW, 0.45, 16)
fml(slide, "sin (180° − θ)  =  0 · cos θ  −  (−1) sin θ  =  sin θ", XL, Y0 + 0.55, CW, LH, 16)
result_card(slide, "(13)   sin (180° − θ)  =  sin θ", XL, Y0 + 0.55 + LH + NG, CW, 0.32, 18)
fml(slide, "cos (180° + θ)  =  (−1) cos θ  −  0  =  − cos θ", XL, Y0 + 0.55 + 2 * LH + 2 * NG, CW, LH, 16)
result_card(slide, "(14)   cos (180° + θ)  =  − cos θ", XL, Y0 + 0.55 + 3 * LH + 3 * NG, CW, 0.32, 18)

slide = new_slide("Board", "Pending slide 77 — match the column", GREEN)
add_notes(slide, "Read. A PT, B QT, C QT, D PS. Option 1.")
qtx(slide, "Last class, slide 77. Match.  Modulus means size, ignore the minus.", XL, Y0, CW, 0.36, 16)
grid_table(
    slide,
    ["Column I", "easy working", "match"],
    [
        ["(A)  sin 390°", "same as sin 30°  =  +1/2", "P and T"],
        ["(B)  sin (−30°)", "minus, so  −1/2", "Q and T"],
        ["(C)  cos 120°", "minus, so  −1/2", "Q and T"],
        ["(D)  tan (−120°)", "comes out  +√3  which is bigger than 1", "P and S"],
    ],
    XL, Y0 + 0.42, [3.4, 5.3, 3.35], 0.42, 14, 14, scaled=False,
)
choices(slide, [
    "A → PT,   B → QT,   C → QT,   D → PS",
    "A → PT,   B → QS,   C → QT,   D → PS",
    "A → QT,   B → QS,   C → PT,   D → PS",
    "A → QS,   B → PT,   C → QT,   D → PS",
], XL, Y0 + 2.55, 7.4)
result_card(slide, "Option 1     A → PT,   B → QT,   C → QT,   D → PS", XL, Y0 + 4.20, CW, 0.36, 18)
trap_card(slide, "|√3| > 1 so D is P and S, not T.  |1/2| < 1 so A, B, C take T.", XL, Y0 + 4.62, CW, 0.32)

slide = new_slide("Tonight leftover", "Last class, slide 78 — after that class", AMBER)
add_notes(slide, "Those two allied lines and the match-column are now on the board. Compound ten times still useful.")
body(slide, "That class asked you to write the six compound-angle lines ten times.", XL, Y0, CW)
body(slide, "The two allied lines and the match-column are now answered on the board. Copy them once more.", XL, Y0 + LH, CW)
result_card(slide, "sin (180° − θ)  =  sin θ          cos (180° + θ)  =  − cos θ", XL, Y0 + 2 * LH + NG, CW, 0.36, 16)
result_card(slide, "Match:  A → PT,   B → QT,   C → QT,   D → PS", XL, Y0 + 3 * LH + 2 * NG, CW, 0.36, 16)

slide = new_slide("5-minute leftover", "Radian and arc — now locked", TEAL)
add_notes(slide, "They say 180=π. Arc=Rθ in radian. This was last-class slides 79–80.")
fml(slide, "180°  =  π rad", XL, Y0, 7.0, LH, 28)
ncert(slide, "Wherever you see π rad, put 180°.", XL + 7.15, Y0, 4.90, 1.10, 18)
fml(slide, "degree → radian:     multiply by  π/180", XL, Y0 + LH + NG + 0.70, CW, LH, 20)
fml(slide, "radian → degree:     multiply by  180/π", XL, Y0 + 2 * LH + NG + 0.70, CW, LH, 20)
fml(slide, "Arc length  =  R θ     with θ in radian.     Never  R × 90°.", XL, Y0 + 3 * LH + NG + 0.70, CW, LH, 20)
result_card(slide, "90° on a circle of radius R  →  distance  πR/2", XL, Y0 + 4 * LH + 2 * NG + 0.70, CW, 0.36, 18)

slide = new_slide("5-minute leftover", "Compound and prefixes — now locked", TEAL)
add_notes(slide, "They say the three compound pairs. Prefix full names once.")
fml(slide, "sin (A + B)  =  sin A cos B  +  cos A sin B", XL, Y0, CW, LH, 18)
fml(slide, "cos (A + B)  =  cos A cos B  −  sin A sin B", XL, Y0 + LH, CW, LH, 18)
fml(slide, "tan (A + B)  =  (tan A + tan B) / (1 − tan A tan B)", XL, Y0 + 2 * LH, CW, LH, 18)
result_card(slide, "fm femtometre  10⁻¹⁵ m     nm nanometre  10⁻⁹ m     Å ångström  10⁻¹⁰ m", XL, Y0 + 3 * LH + NG, CW, 0.36, 16)
trap_card(slide, "Do not write sin(A+B) = sin A + sin B. Nuclear size is fm, not mm.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.40)

slide = new_slide("5-minute leftover", "Small angle — the one lock-line", TEAL)
add_notes(slide, "Last class slides 81–82. sinθ≈θ, tanθ≈θ in radian. cosθ≈1. tan2° is π/90.")
ncert(slide, "sin θ  ≈  θ      tan θ  ≈  θ      cos θ  ≈  1      (θ in radian,  θ ≤ 5°)", XL, Y0, CW, 1.05, 18)
fml(slide, "sin 2°  ≈  π/90          tan 3°  ≈  π/60          cos 3°  ≈  1", XL, Y0 + 1.25, CW, LH, 18)
trap_card(slide, "Never leave the degree number in the approximation. Convert, then drop the word radian.", XL, Y0 + 1.25 + LH + NG, CW, 0.45)
body(slide, "Leftover from last class is done. After the break: ASTC method, max–min, phase.", XL, Y0 + 1.25 + 2 * LH + 2 * NG, CW)

# ── BREAK ─────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
_solid(slide, 0, 0, 13.333, 0.06, TEAL)
tx(slide, "BREAK", XL, 2.20, CW, 0.70, 48, GOLD, True, PP_ALIGN.CENTER, spa=4)
tx(slide, "Leftover from last class is done.  Next: ASTC method, max–min, phase", XL, 3.10, CW, 0.40, 18, BLUE, True, PP_ALIGN.CENTER, spa=0)
_footer(slide, len(prs.slides))
add_notes(slide, "Hour 1 leftover L03 done. Hour 2: ASTC method, max-min, phasor.")
after_content_slide()

# ── ASTC METHOD ───────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Three steps for any allied angle", TEAL)
add_notes(slide, "Copy. Sign from ASTC. Function changes only at 90 and 270. Reciprocals swap in pairs.")
body(slide, "θ is acute. The three steps write the value without opening the compound line every time.", XL, Y0, CW)
y = Y0 + LH + NG
fml(slide, "1.   Decide the quadrant. Write the sign of that ratio from ASTC.", XL, y, CW, LH, 18)
fml(slide, "2.   At  90° ± θ  and  270° ± θ  the function changes.   At  180° ± θ  and  360° ± θ  it stays.", XL, y + LH, CW, LH, 18)
fml(slide, "3.   Change pairs:     sin  ↔  cos          tan  ↔  cot          sec  ↔  cosec", XL, y + 2 * LH, CW, LH, 18)
ncert(slide, "A   S   T   C      →      All    sin    tan    cos", XL, y + 3 * LH + NG, CW, 0.90, 18)
trap_card(slide, "Do not look at the sign of the result after the change. Write the ASTC sign first, then the brother.", XL, y + 3 * LH + NG + 1.05, CW, 0.40)

slide = new_slide("Copy", "Where each allied angle sits", TEAL)
add_notes(slide, "They point. 90+θ and 180−θ are second. 180+θ and 270−θ are third. Negative is fourth.")
quads = [
    (GREEN,  "I     0° to 90°",     "All +ve          90° − θ          360° + θ"),
    (BLUE,   "II    90° to 180°",   "sin, cosec +ve   90° + θ          180° − θ"),
    (AMBER,  "III   180° to 270°",  "tan, cot +ve     180° + θ          270° − θ"),
    (ORANGE, "IV    270° to 360°",  "cos, sec +ve     270° + θ          360° − θ          −θ"),
]
for i, (col, head, sub) in enumerate(quads):
    y = Y0 + i * 1.05
    _card_bg(slide, XL, y, CW, 0.95, RGBColor(12, 16, 26), col, 1.5)
    tx(slide, head, XL + 0.25, y + 0.28, 4.4, 0.40, 18, col, True, spa=0, scaled=False)
    tx(slide, sub, XL + 4.70, y + 0.28, 7.0, 0.40, 18, TEXT, False, spa=0, scaled=False)

slide = new_slide("Copy", "Allied lines — 90° and 180°", TEAL)
add_notes(slide, "Copy. They say each line. 90 changes the function. 180 keeps it.")
fml(slide, "sin (90° + θ)  =  + cos θ          sin (90° − θ)  =  + cos θ", XL, Y0, CW, LH, 18)
fml(slide, "cos (90° + θ)  =  − sin θ          cos (90° − θ)  =  + sin θ", XL, Y0 + LH, CW, LH, 18)
fml(slide, "tan (90° − θ)  =  + cot θ", XL, Y0 + 2 * LH, CW, LH, 18)
fml(slide, "sin (180° + θ)  =  − sin θ          sin (180° − θ)  =  + sin θ", XL, Y0 + 3 * LH + NG, CW, LH, 18)
fml(slide, "cos (180° + θ)  =  − cos θ          tan (180° − θ)  =  − tan θ", XL, Y0 + 4 * LH + NG, CW, LH, 18)
trap_card(slide, "cos (90° + θ) is minus sine. That minus is the ASTC sign of cosine in the second quadrant.", XL, Y0 + 5 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Copy", "Allied lines — 270°, 360°, 450°, 540°", TEAL)
add_notes(slide, "Copy. 450 is 360+90. 540 is 360+180. Reduce first, then the same three steps.")
fml(slide, "sin (270° + θ)  =  − cos θ          sin (270° − θ)  =  − cos θ", XL, Y0, CW, LH, 18)
fml(slide, "cos (270° + θ)  =  + sin θ          tan (360° + θ)  =  + tan θ", XL, Y0 + LH, CW, LH, 18)
fml(slide, "cos (360° − θ)  =  + cos θ          cos (−θ)  =  + cos θ", XL, Y0 + 2 * LH, CW, LH, 18)
fml(slide, "sin (450° + θ)  =  sin (90° + θ)  =  + cos θ", XL, Y0 + 3 * LH + NG, CW, LH, 18)
fml(slide, "tan (540° + θ)  =  tan (180° + θ)  =  + tan θ", XL, Y0 + 4 * LH + NG, CW, LH, 18)
ncert(slide, "Drop full turns of 360° first.  450° = 90°.   540° = 180°.   720° = 0°.", XL, Y0 + 5 * LH + 2 * NG, CW, 0.80, 16)

# ── ALLIED PRACTICE ───────────────────────────────────────────────────────────
slide = new_slide("Board", "sin 300°  and  cos 240°", GREEN)
add_notes(slide, "Read. Both −√3/2 and −1/2. Two routes, same value.")
qtx(slide, "Find  (1)  sin 300°     and     (2)  cos 240°.   Two routes each.", XL, Y0, CW, 0.40, 16)
fml(slide, "sin 300°  =  sin (270° + 30°)  =  − cos 30°  =  −√3/2", XL, Y0 + 0.48, CW, LH, 17)
fml(slide, "         =  sin (360° − 60°)  =  − sin 60°  =  −√3/2", XL, Y0 + 0.48 + LH, CW, LH, 17)
result_card(slide, "(1)   sin 300°  =  −√3/2", XL, Y0 + 0.48 + 2 * LH + NG, CW, 0.32, 18)
fml(slide, "cos 240°  =  cos (180° + 60°)  =  − cos 60°  =  −1/2", XL, Y0 + 0.48 + 3 * LH + 2 * NG, CW, LH, 17)
fml(slide, "         =  cos (270° − 30°)  =  − sin 30°  =  −1/2", XL, Y0 + 0.48 + 4 * LH + 2 * NG, CW, LH, 17)
result_card(slide, "(2)   cos 240°  =  −1/2", XL, Y0 + 0.48 + 5 * LH + 3 * NG, CW, 0.32, 18)

slide = new_slide("Board", "Past 360° — drop full turns first", GREEN)
add_notes(slide, "Read. 405=45. 480=120. 780=60.")
qtx(slide, "Find  sin 405°,   cos 480°   and   cos 780°.", XL, Y0, CW, 0.36, 17)
fml(slide, "sin 405°  =  sin (360° + 45°)  =  sin 45°  =  1/√2", XL, Y0 + 0.45, CW, LH, 18)
fml(slide, "cos 480°  =  cos (450° + 30°)  =  cos (90° + 30°)  =  − sin 30°  =  −1/2", XL, Y0 + 0.45 + LH, CW, LH, 18)
fml(slide, "cos 780°  =  cos (720° + 60°)  =  cos 60°  =  1/2", XL, Y0 + 0.45 + 2 * LH, CW, LH, 18)
result_card(slide, "sin 405°  =  1/√2          cos 480°  =  −1/2          cos 780°  =  1/2", XL, Y0 + 0.45 + 3 * LH + NG, CW, 0.36, 18)
trap_card(slide, "720° is two full turns. Cross it out. Do not open a compound line on 780°.", XL, Y0 + 0.45 + 4 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Board", "Ten values from ASTC — read down", GREEN)
add_notes(slide, "Read the green column. 510=150. cot 210 = cot 30.")
qtx(slide, "Find each value. Drop 360° first. Then the three steps.", XL, Y0, CW, 0.36, 16)
grid_table(
    slide,
    ["", "angle", "reduction", "value"],
    [
        ["(a)", "sin 30°", "table", "1/2"],
        ["(b)", "sin 45°", "table", "1/√2"],
        ["(c)", "sin 510°", "sin (360° + 150°)  =  sin 150°", "1/2"],
        ["(d)", "sin 240°", "sin (180° + 60°)", "−√3/2"],
        ["(e)", "sin 330°", "sin (360° − 30°)", "−1/2"],
        ["(f)", "cos 300°", "cos (360° − 60°)", "1/2"],
        ["(g)", "tan 135°", "tan (180° − 45°)", "−1"],
        ["(h)", "cot 210°", "cot (180° + 30°)  =  cot 30°", "√3"],
        ["(i)", "cos 510°", "cos 150°  =  − cos 30°", "−√3/2"],
        ["(j)", "sin 135°", "sin (180° − 45°)", "1/√2"],
    ],
    XL, Y0 + 0.40, [0.85, 2.15, 5.70, 3.35], 0.38, 13, 14, scaled=False,
)

# ── HALF ANGLE LEFTOVER ───────────────────────────────────────────────────────
slide = new_slide("Board", "sin 74°  from  2 × 37°", GREEN)
add_notes(slide, "Read. 2 sin37 cos37 = 24/25.")
qtx(slide, "Find  sin 74°.   Last class:  sin 37° = 3/5,  cos 37° = 4/5.", XL, Y0, CW, 0.40, 16)
fml(slide, "sin 74°  =  sin (2 × 37°)  =  2 sin 37° cos 37°", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "         =  2 × (3/5) × (4/5)  =  24/25", XL, Y0 + 0.50 + LH, CW, LH, 18)
result_card(slide, "sin 74°  =  24/25", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 20)
ncert(slide, "sin 2θ  =  2 sin θ cos θ", XL, Y0 + 0.50 + 3 * LH + 2 * NG, CW, 0.80, 18)

slide = new_slide("Board", "sin 22.5°  from the half-angle line", GREEN)
add_notes(slide, "Read. √[(1−cos45)/2] = √[(√2−1)/(2√2)].")
qtx(slide, "Find  sin 22.5°.   Last class:  sin θ  =  √[(1 − cos 2θ) / 2].", XL, Y0, CW, 0.40, 16)
fml(slide, "sin 22.5°  =  √[ (1 − cos 45°) / 2 ]  =  √[ (1 − 1/√2) / 2 ]", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "           =  √[ (√2 − 1) / (2√2) ]", XL, Y0 + 0.50 + LH, CW, LH, 18)
result_card(slide, "sin 22.5°  =  √[(√2 − 1) / (2√2)]", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 20)
body(slide, "Same line gives  cos 15°  =  (√6 + √2)/4     and     cos 105°  =  (√2 − √6)/4.", XL, Y0 + 0.50 + 3 * LH + 2 * NG, CW)

# ── BREAK ─────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(BLANK)
_bg(slide)
_solid(slide, 0, 0, 13.333, 0.06, TEAL)
tx(slide, "BREAK", XL, 2.20, CW, 0.70, 48, GOLD, True, PP_ALIGN.CENTER, spa=4)
tx(slide, "After the break:  maximum and minimum,  then phase and phasor", XL, 3.10, CW, 0.40, 18, BLUE, True, PP_ALIGN.CENTER, spa=0)
_footer(slide, len(prs.slides))
add_notes(slide, "Hour 1 done. Allied method locked. Hour 2 is max-min then phasor.")
after_content_slide()

# ── MAX MIN ───────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Maximum and minimum of sine and cosine", TEAL)
add_notes(slide, "Copy. sin max +1 at 90. cos max +1 at 0. sin min −1 at 270. cos min −1 at 180.")
fml(slide, "(sin θ)max  =  +1     at     θ  =  90°", XL, Y0, 7.0, LH, 20)
fml(slide, "(cos θ)max  =  +1     at     θ  =  0°", XL, Y0 + LH, 7.0, LH, 20)
ncert(slide, "−1  ≤  sin θ  ≤  1", XL + 7.15, Y0, 4.90, 1.05, 18)
fml(slide, "(sin θ)min  =  −1     at     θ  =  270°", XL, Y0 + 2 * LH + NG, 7.0, LH, 20)
fml(slide, "(cos θ)min  =  −1     at     θ  =  180°", XL, Y0 + 3 * LH + NG, 7.0, LH, 20)
ncert(slide, "−1  ≤  cos θ  ≤  1", XL + 7.15, Y0 + 2 * LH + NG, 4.90, 1.05, 18)
trap_card(slide, "A constant in front of sine is not the angle.  y = 6 sin θ  has ymax = 6, not 6 × 90°.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Board", "y  =  6 sin θ", GREEN)
add_notes(slide, "Read. ymax +6 at 90. ymin −6 at 270.")
qtx(slide, "If  y  =  6 sin θ,  find the maximum and minimum value of y.", XL, Y0, CW, 0.40, 17)
fml(slide, "ymax  =  6 × (sin θ)max  =  6 × 1  =  +6          at  θ  =  90°", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "ymin  =  6 × (sin θ)min  =  6 × (−1)  =  −6          at  θ  =  270°", XL, Y0 + 0.50 + LH, CW, LH, 18)
result_card(slide, "ymax  =  +6          ymin  =  −6", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "The 3θ is the angle, not the 4", GREEN)
add_notes(slide, "Read. ymax=4. Angle 3θ=90 so θ=30.")
qtx(slide, "If  y  =  4 sin (3θ),  find ymax and the θ at which y is maximum.", XL, Y0, CW, 0.45, 16)
fml(slide, "(sin 3θ)max  =  +1          so          ymax  =  4 × 1  =  +4", XL, Y0 + 0.55, CW, LH, 18)
fml(slide, "The angle that is 90° is  3θ,  not θ.          3θ  =  90°          θ  =  30°", XL, Y0 + 0.55 + LH, CW, LH, 18)
result_card(slide, "ymax  =  +4     at     θ  =  30°", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 20)
trap_card(slide, "Do not set θ = 90°. The sine is of 3θ. Divide 90° by 3.", XL, Y0 + 0.55 + 3 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Board", "Range of a projectile", GREEN)
add_notes(slide, "Read. sin2θ max at 2θ=90, θ=45. Rmax=u²/g.")
qtx(slide, "If  R  =  u² sin(2θ) / g,  with u and g constant, find θ at which R is maximum.", XL, Y0, CW, 0.50, 16)
fml(slide, "Rmax when  sin(2θ)  is maximum,  so  sin(2θ)  =  1          2θ  =  90°          θ  =  45°", XL, Y0 + 0.60, CW, LH, 18)
result_card(slide, "θ  =  45°          Rmax  =  u² / g", XL, Y0 + 0.60 + LH + NG, CW, 0.36, 20)
body(slide, "This is why a level-ground projectile goes farthest at 45°.", XL, Y0 + 0.60 + 2 * LH + 2 * NG, CW)

slide = new_slide("Board", "Add like sines first", GREEN)
add_notes(slide, "Read. 3sin+4sin=7sin. ymax=7 at 90.")
qtx(slide, "If  y  =  3 sin θ  +  4 sin θ,  find ymax.", XL, Y0, CW, 0.40, 18)
fml(slide, "y  =  7 sin θ          ymax  =  7 × 1  =  7          at  θ  =  90°", XL, Y0 + 0.50, CW, LH, 18)
result_card(slide, "ymax  =  7", XL, Y0 + 0.50 + LH + NG, CW, 0.36, 22)
trap_card(slide, "Do not open √(3²+4²) here. Both terms are sine of the same angle. Add the 3 and the 4.", XL, Y0 + 0.50 + 2 * LH + 2 * NG, CW, 0.45)

slide = new_slide("Board", "A constant sitting next to the sine", GREEN)
add_notes(slide, "Read. 3+2sin: 5 and +1. 3+4sin: 7 and −1. 3sin−4: −1 and −7.")
qtx(slide, "Find ymax and ymin.", XL, Y0, CW, 0.32, 16)
fml(slide, "y  =  3 + 2 sin θ          ymax  =  3 + 2(1)  =  5          ymin  =  3 + 2(−1)  =  1", XL, Y0 + 0.40, CW, LH, 16)
fml(slide, "y  =  3 + 4 sin θ          ymax  =  7          ymin  =  3 − 4  =  −1", XL, Y0 + 0.40 + LH, CW, LH, 16)
fml(slide, "y  =  3 sin θ − 4          ymax  =  3 − 4  =  −1          ymin  =  −3 − 4  =  −7", XL, Y0 + 0.40 + 2 * LH, CW, LH, 16)
result_card(slide, "3 + 2 sin θ  →  5  and  1          3 + 4 sin θ  →  7  and  −1          3 sin θ − 4  →  −1  and  −7", XL, Y0 + 0.40 + 3 * LH + NG, CW, 0.40, 16)

# ── a sin + b cos ─────────────────────────────────────────────────────────────
slide = new_slide("Copy", "When sine and cosine are added", TEAL)
add_notes(slide, "Copy. Amplitude is √(a²+b²). They cannot both be +1 at the same θ.")
body(slide, "sin θ and cos θ do not peak together. The combined amplitude is the hypotenuse of a and b.", XL, Y0, CW)
fml(slide, "y  =  a sin θ  +  b cos θ", XL, Y0 + LH + NG, CW, LH, 22)
ncert(slide, "ymax  =  + √(a² + b²)          ymin  =  − √(a² + b²)", XL, Y0 + 2 * LH + 2 * NG, CW, 1.00, 18)
trap_card(slide, "Do not write ymax = a + b. That would need sin = 1 and cos = 1 at the same instant, which never happens.", XL, Y0 + 3 * LH + 2 * NG + 0.90, CW, 0.45)

slide = new_slide("Board", "y  =  3 sin θ  +  4 cos θ", GREEN)
add_notes(slide, "Read. √25=5. min −5.")
qtx(slide, "If  y  =  3 sin θ  +  4 cos θ,  find ymax and ymin.", XL, Y0, CW, 0.40, 17)
fml(slide, "ymax  =  √(3² + 4²)  =  √(9 + 16)  =  √25  =  +5", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "ymin  =  −5", XL, Y0 + 0.50 + LH, CW, LH, 18)
result_card(slide, "ymax  =  +5          ymin  =  −5", XL, Y0 + 0.50 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "Minimum magnitude of a fraction", GREEN)
add_notes(slide, "Read. Bigger denominator, smaller fraction. ymin mag = 4/√13.")
qtx(slide, "If  y  =  4 / (3 sin θ + 2 cos θ),  find the minimum magnitude of y.", XL, Y0, CW, 0.45, 16)
body(slide, "4/10 = 0.4.   4/100 = 0.04.   A larger denominator makes a smaller positive fraction.", XL, Y0 + 0.52, CW)
fml(slide, "ymin  =  4  /  (3 sin θ + 2 cos θ)max  =  4 / √(3² + 2²)  =  4 / √13", XL, Y0 + 0.52 + LH + NG, CW, LH, 18)
result_card(slide, "minimum magnitude  =  4 / √13", XL, Y0 + 0.52 + 2 * LH + 2 * NG, CW, 0.36, 20)

slide = new_slide("Board", "Five max–min — every line", GREEN)
add_notes(slide, "Read. 5, √13, 6 and −4, 3 and −1, 13.")
qtx(slide, "Find maximum and minimum of each.", XL, Y0, CW, 0.32, 16)
grid_table(
    slide,
    ["", "y", "ymax", "ymin"],
    [
        ["(a)", "4 sin θ + 3 cos θ", "√(16+9)  =  5", "−5"],
        ["(b)", "2 sin θ − 3 cos θ", "√(4+9)  =  √13", "−√13"],
        ["(c)", "1 + 5 sin θ", "1+5  =  6", "1−5  =  −4"],
        ["(d)", "1 − 2 sin θ", "1−2(−1)  =  3", "1−2(1)  =  −1"],
        ["(e)", "5 sin θ − 12 cos θ", "√(25+144)  =  13", "−13"],
    ],
    XL, Y0 + 0.40, [0.85, 4.20, 3.50, 3.50], 0.48, 14, 15, scaled=False,
)

# ── PHASE ─────────────────────────────────────────────────────────────────────
slide = new_slide("Copy", "Phase is the angle inside the ratio", TEAL)
add_notes(slide, "Copy. Phase is the whole argument. Used later in SHM, waves, AC, optics, EM, vectors.")
body(slide, "Phase is the angle sitting inside the trigonometric function. Not the 4 in front.", XL, Y0, CW)
fml(slide, "sin θ                         phase  =  θ", XL, Y0 + LH + NG, CW, LH, 20)
fml(slide, "sin (θ + α)                   phase  =  θ + α", XL, Y0 + 2 * LH + NG, CW, LH, 20)
fml(slide, "sin (ωt + α + 90°)            phase  =  ωt + α + 90°", XL, Y0 + 3 * LH + NG, CW, LH, 20)
ncert(slide, "Later:  AC,  SHM,  waves,  wave optics,  EM waves,  vectors.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.85, 16)

slide = new_slide("Copy", "Phase difference — same function first", TEAL)
add_notes(slide, "Copy. Same ratio. Δφ = φ2 − φ1. Plus is lead. Minus is lag.")
body(slide, "Angle difference between two of the same trigonometric function is called phase difference.", XL, Y0, CW)
fml(slide, "y1  =  sin θ                         φ1  =  θ", XL, Y0 + LH + NG, CW, LH, 18)
fml(slide, "y2  =  sin (θ + 60°)                 φ2  =  θ + 60°", XL, Y0 + 2 * LH + NG, CW, LH, 18)
fml(slide, "Δφ  =  φ2 − φ1  =  +60°     (y2 leads y1)          Δφ  =  φ1 − φ2  =  −60°     (lag)", XL, Y0 + 3 * LH + NG, CW, LH, 18)
trap_card(slide, "Do not subtract a cosine phase from a sine phase until both are written as the same ratio.", XL, Y0 + 4 * LH + 2 * NG, CW, 0.45)

slide = new_slide("Board", "y1 leads y2 by 90°", GREEN)
add_notes(slide, "Read. +30 − (−60) = 90. Phasor: +30 and −60 on the sin axis.")
qtx(slide, "Find the phase difference.   y1  =  sin (ωt + 30°)     y2  =  sin (ωt − 60°)", XL, Y0, CW, 0.45, 16)
cx, cy = phasor_axes(slide, 8.55, 1.85, 1.55, "sin (ωt)", "")
draw_line(slide, cx, cy, cx + 1.20, cy - 0.70, GOLD, 2.25)
draw_line(slide, cx, cy, cx + 0.70, cy + 1.20, GOLD, 2.25)
tx(slide, "y1  +30°", cx + 1.10, cy - 0.95, 1.20, 0.22, 12, GOLD, True, spa=0)
tx(slide, "y2  −60°", cx + 0.55, cy + 1.22, 1.20, 0.22, 12, GOLD, True, spa=0)
tx(slide, "90°", cx + 0.55, cy - 0.18, 0.50, 0.22, 14, GREEN, True, spa=0)
fml(slide, "Δφ  =  φ1 − φ2  =  (ωt + 30°)  −  (ωt − 60°)  =  90°", XL, Y0 + 0.55, 7.6, LH, 17)
result_card(slide, "Δφ  =  90°     —  y1 is ahead of y2", XL, Y0 + 0.55 + LH + NG, 7.6, 0.36, 18)

slide = new_slide("Board", "Convert to the same ratio, or the answer is wrong", GREEN)
add_notes(slide, "Read. cos(ωt−90)=sin(ωt). Δφ=0. Naked subtract gives a fake 90°.")
qtx(slide, "Find Δφ.   y1  =  3 cos (ωt − 90°)     y2  =  4 sin (ωt)", XL, Y0, CW, 0.45, 16)
y = warn_row(slide, "Δφ  =  ωt − (ωt − 90°)  =  90°", "first write both as sine (or both as cosine)", XL, Y0 + 0.52, CW)
fml(slide, "cos (ωt − 90°)  =  sin ωt          so          y1  =  3 sin (ωt)", XL, y + NG, CW, LH, 18)
result_card(slide, "Δφ  =  0     —  they are in phase", XL, y + NG + LH + NG, CW, 0.36, 20)
trap_card(slide, "A minus cosine is a sine with a shift. Convert, then subtract.", XL, y + NG + 2 * LH + 2 * NG, CW, 0.40)

slide = new_slide("Board", "cos (ωt + π/2)  points opposite to sine", GREEN)
add_notes(slide, "Read. cos(ωt+π/2)= −sin(ωt). Δφ=180°.")
qtx(slide, "Find Δφ.   y1  =  3 sin (ωt)     y2  =  4 cos (ωt + π/2)", XL, Y0, CW, 0.40, 16)
fml(slide, "cos (ωt + π/2)  =  − sin (ωt)          so          y2  =  − 4 sin (ωt)", XL, Y0 + 0.50, CW, LH, 18)
result_card(slide, "Δφ  =  180°     =     π rad     —  opposite on the phasor", XL, Y0 + 0.50 + LH + NG, CW, 0.36, 18)
cx, cy = phasor_axes(slide, 4.55, 3.55, 1.35, "sin (ωt)", "cos (ωt)")
draw_line(slide, cx, cy, cx + 1.30, cy, GOLD, 2.25)
draw_line(slide, cx, cy, cx - 1.30, cy, GREEN, 2.25)
tx(slide, "y1", cx + 1.15, cy - 0.28, 0.50, 0.22, 14, GOLD, True, spa=0)
tx(slide, "y2", cx - 1.55, cy - 0.28, 0.50, 0.22, 14, GREEN, True, spa=0)

slide = new_slide("Board", "Two more phase differences", GREEN)
add_notes(slide, "Read. 150° and 135°. Then 30° leftover as a one-liner.")
qtx(slide, "(1)  x1 = A cos(ωt − π/2)     y2 = A cos(ωt + π/3)", XL, Y0, CW, 0.36, 16)
fml(slide, "Same function already.   Δφ  =  (ωt + π/3) − (ωt − π/2)  =  5π/6  =  150°", XL, Y0 + 0.42, CW, LH, 16)
result_card(slide, "(1)   Δφ  =  150°", XL, Y0 + 0.42 + LH + NG, CW, 0.32, 18)
qtx(slide, "(2)  y1 = 4 sin(θ + π/4)     y2 = 3 sin(θ − π/2)", XL, Y0 + 0.42 + 2 * LH + 2 * NG, CW, 0.32, 16)
fml(slide, "Δφ  =  (θ + π/4) − (θ − π/2)  =  π/4 + π/2  =  135°", XL, Y0 + 0.42 + 3 * LH + 2 * NG, CW, LH, 16)
result_card(slide, "(2)   Δφ  =  135°          Also:  y1 = 4 sin(θ+π/3), y2 = 3 sin(θ+π/6)  →  30°", XL, Y0 + 0.42 + 4 * LH + 3 * NG, CW, 0.32, 16)

slide = new_slide("Board", "AC current and voltage", GREEN)
add_notes(slide, "Read. Convert V to sine. 90−(−30)=120=2π/3.")
qtx(slide, "I  =  I0 sin(ωt − 30°)     V  =  V0 cos(ωt).   Find the phase difference.", XL, Y0, CW, 0.45, 16)
fml(slide, "V  =  V0 cos(ωt)  =  V0 sin(ωt + 90°)", XL, Y0 + 0.55, CW, LH, 18)
fml(slide, "Δφ  =  90° − (−30°)  =  120°  =  2π/3", XL, Y0 + 0.55 + LH, CW, LH, 18)
result_card(slide, "Δφ  =  2π/3     —  voltage leads current by 120°", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "Negative cosine is a shifted sine", GREEN)
add_notes(slide, "Read. I=sin(ωt−π/2). Δφ=π/3. Option 2.")
qtx(slide, "I  =  − I0 cos(ωt)     V  =  V0 sin(ωt − π/6).   Find Δφ.", XL, Y0, CW, 0.40, 16)
choices(slide, ["2π/3", "π/3", "π/6", "5π/6"], XL, Y0 + 0.48)
fml(slide, "− cos(ωt)  =  sin(ωt − π/2)          Δφ  =  (−π/6) − (−π/2)  =  π/3", XL, Y0 + 2.20, CW, LH, 16)
result_card(slide, "Option 2     Δφ  =  π/3", XL, Y0 + 2.20 + LH + NG, CW, 0.36, 20)

# ── MIXED LIVE ────────────────────────────────────────────────────────────────
slide = new_slide("Board", "Person walks toward a 40 m building", GREEN)
add_notes(slide, "Read. d = 40(√3 − 1/√3)=80/√3. Option 2.")
qtx(slide, "A person sees a 40 m building at elevation 30°, then at 60°. Distance moved toward the building?", XL, Y0, CW, 0.50, 15)
choices(slide, ["40/√3", "80/√3", "60/√3", "120/√3"], XL, Y0 + 0.58, 5.4)
fml(slide, "x1  =  40 cot 30°  =  40√3          x2  =  40 cot 60°  =  40/√3", XL, Y0 + 2.30, CW, LH, 16)
fml(slide, "d  =  x1 − x2  =  40√3 − 40/√3  =  40(3 − 1)/√3  =  80/√3", XL, Y0 + 2.30 + LH, CW, LH, 16)
result_card(slide, "Option 2     d  =  80/√3", XL, Y0 + 2.30 + 2 * LH + NG, CW, 0.36, 18)

slide = new_slide("Board", "2 (sin 15° + sin 75°)²", GREEN)
add_notes(slide, "Read. sin75=cos15. 2(1+sin30)=3.")
qtx(slide, "What is the value of  2 (sin 15° + sin 75°)² ?", XL, Y0, CW, 0.40, 18)
fml(slide, "sin 75°  =  cos 15°          so          2 (sin 15° + cos 15°)²", XL, Y0 + 0.50, CW, LH, 18)
fml(slide, "        =  2 ( sin²15° + cos²15° + 2 sin 15° cos 15° )  =  2 ( 1 + sin 30° )", XL, Y0 + 0.50 + LH, CW, LH, 16)
fml(slide, "        =  2 ( 1 + 1/2 )  =  3", XL, Y0 + 0.50 + 2 * LH, CW, LH, 18)
result_card(slide, "2 (sin 15° + sin 75°)²  =  3", XL, Y0 + 0.50 + 3 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "If sin θ = 1/2, first quadrant", GREEN)
add_notes(slide, "Read. θ=30. √3(√3/2 + 2/√3)=7/2. Option 3.")
qtx(slide, "If  sin θ  =  1/2,  the value of  √3 (cos θ + 2 tan θ)  is:", XL, Y0, CW, 0.40, 16)
choices(slide, ["7", "15/2", "7/2", "none"], XL, Y0 + 0.50)
fml(slide, "First quadrant:  θ  =  30°          cos 30°  =  √3/2          tan 30°  =  1/√3", XL, Y0 + 2.20, CW, LH, 16)
fml(slide, "√3 ( √3/2  +  2/√3 )  =  3/2  +  2  =  7/2", XL, Y0 + 2.20 + LH, CW, LH, 18)
result_card(slide, "Option 3     7/2", XL, Y0 + 2.20 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "If sin θ = √5 / 3", GREEN)
add_notes(slide, "Read. cos=2/3. tan=√5/2. sec=3/2. Answer √5/2 + 6.")
qtx(slide, "If  sin θ  =  √5 / 3,  find  tan θ  +  4 sec θ.   Take θ acute.", XL, Y0, CW, 0.45, 16)
fml(slide, "cos θ  =  √(1 − 5/9)  =  √(4/9)  =  2/3          tan θ  =  (√5/3) / (2/3)  =  √5 / 2", XL, Y0 + 0.55, CW, LH, 16)
fml(slide, "sec θ  =  3/2          tan θ + 4 sec θ  =  √5/2  +  4 × (3/2)  =  √5/2  +  6", XL, Y0 + 0.55 + LH, CW, LH, 16)
result_card(slide, "tan θ  +  4 sec θ  =  6 + √5/2", XL, Y0 + 0.55 + 2 * LH + NG, CW, 0.36, 20)

slide = new_slide("Board", "tan⁻¹(1)  —  principal value", GREEN)
add_notes(slide, "Read. Range of tan⁻¹ is (−π/2, π/2). Only π/4. Not 5π/4.")
qtx(slide, "The value of  tan⁻¹(1)  is", XL, Y0, CW, 0.36, 18)
choices(slide, ["π/4", "5π/4", "both (1) and (2)", "none of these"], XL, Y0 + 0.45)
y = warn_row(slide, "tan(5π/4) = 1, so both", "tan⁻¹ returns only the principal value, in (−π/2, π/2)", XL, Y0 + 2.15, CW)
result_card(slide, "Option 1     π/4", XL, y + NG, CW, 0.36, 20)
trap_card(slide, "5π/4 is an angle whose tangent is 1. It is not the inverse-tangent of 1.", XL, y + NG + 0.42, CW, 0.40)

# ── TONIGHT ───────────────────────────────────────────────────────────────────
slide = new_slide("Tonight", "Phase difference table — copy, do not wait", AMBER)
add_notes(slide, "Do not reveal. Convert to the same ratio first. Answers stay off the board.")
qtx(slide, "Find the phase difference in each row. Convert to the same trigonometric function first.", XL, Y0, CW, 0.45, 15)
grid_table(
    slide,
    ["Equation 1", "Equation 2", "Δφ"],
    [
        ["I0 sin(θ + π/3)", "I0 sin(θ − π/6)", ""],
        ["I0 sin(θ + π/3)", "I0 cos(θ + π/3)", ""],
        ["4 sin(ωt)", "4 cos(ωt − π/3)", ""],
        ["3 sin(ωt + π/2)", "5 cos(ωt)", ""],
        ["4 cos(ωt − π/2)", "4 sin(ωt + π/2)", ""],
        ["3 sin(ωt + π/3)", "3 cos(ωt)", ""],
        ["4 sin(θ − 60°)", "4 cos(θ − 30°)", ""],
    ],
    XL, Y0 + 0.52, [4.15, 4.40, 3.50], 0.42, 13, 14, scaled=False,
)

slide = new_slide("Tonight", "After class", AMBER)
add_notes(slide, "Daily test before next lecture. Compound and ASTC three steps on a rough page.")
body(slide, "Daily test before the next lecture.", XL, Y0, CW)
body(slide, "Write the three ASTC steps ten times on a rough page.", XL, Y0 + LH, CW)
body(slide, "Write  ymax = √(a² + b²)  for  y = a sin θ + b cos θ  ten times.", XL, Y0 + 2 * LH, CW)
body(slide, "Tonight: the seven phase-difference rows. Convert first, then subtract.", XL, Y0 + 3 * LH + NG, CW)

# ── 5-MINUTE SUMMARY ──────────────────────────────────────────────────────────
slide = new_slide("5-minute summary", "ASTC method — now locked", TEAL)
add_notes(slide, "They say the three steps. 90 and 270 change the function.")
fml(slide, "1.  Sign from ASTC.     2.  Function changes at 90° and 270° only.     3.  sin ↔ cos.", XL, Y0, CW, LH, 16)
result_card(slide, "sin 300°  =  −√3/2          cos 480°  =  −1/2          sin 405°  =  1/√2", XL, Y0 + LH + NG, CW, 0.36, 16)
fml(slide, "sin 74°  =  24/25          sin 22.5°  =  √[(√2 − 1)/(2√2)]", XL, Y0 + 2 * LH + 2 * NG, CW, LH, 18)
ncert(slide, "Drop full turns of 360° first.", XL, Y0 + 3 * LH + 3 * NG, CW, 0.80, 18)

slide = new_slide("5-minute summary", "Max–min — the two boxes", TEAL)
add_notes(slide, "They say both boxes. Do not add a and b.")
fml(slide, "y  =  A sin (nθ)          ymax  =  |A|          at  nθ  =  90°", XL, Y0, CW, LH, 18)
ncert(slide, "y  =  a sin θ + b cos θ          ymax  =  √(a² + b²)", XL, Y0 + LH + NG, CW, 0.90, 18)
result_card(slide, "3 sin θ + 4 cos θ  →  ±5          4 / (3 sin θ + 2 cos θ)  →  4/√13", XL, Y0 + 2 * LH + NG + 0.80, CW, 0.36, 16)
trap_card(slide, "3 sin θ + 4 sin θ is 7 sin θ, not √(9+16). Add like ratios first.", XL, Y0 + 3 * LH + 2 * NG + 0.80, CW, 0.40)

slide = new_slide("5-minute summary", "Phase — the one lock-line", TEAL)
add_notes(slide, "Same function first. Then subtract. Fake 90° if you skip the convert.")
body(slide, "Phase is the angle inside the ratio. Phase difference needs the same trigonometric function.", XL, Y0, CW)
fml(slide, "Δφ  =  φ2 − φ1          after both are sine, or both are cosine", XL, Y0 + LH + NG, CW, LH, 18)
result_card(slide, "3 cos(ωt − 90°)  and  4 sin(ωt)   →   Δφ  =  0", XL, Y0 + 2 * LH + 2 * NG, CW, 0.36, 18)
trap_card(slide, "Never subtract a cosine-phase from a sine-phase. Convert, then subtract.", XL, Y0 + 3 * LH + 3 * NG, CW, 0.45)
body(slide, "Daily test before the next lecture. Seven phase rows tonight.", XL, Y0 + 4 * LH + 4 * NG, CW)

prs.save(OUTPUT)
print(f"Created {len(prs.slides)} slides")
print(OUTPUT)
